# StuLink v1.19.47.29 2026-10-10
# 课表服务层：学期管理 / 作息时间（节次与时间） / 视图查询 / 条目CRUD / 冲突检测 / Excel导入导出 / 版本快照
# Copyright (c) 2026 zkxxzf. Apache License 2.0
"""课表核心服务（独立库 timetable.db）。

所有查询均过滤 is_deleted == False；跨库只用逻辑键不 JOIN。
写操作均记录 ScheduleVersion 快照（snapshot_json 为修改前数据）。
"""
import io
import json
import logging
import re
from datetime import datetime, date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, false, or_

# 2026-10-10 架构审计修复：本文件的静默 except 会掩盖节次/配置数据异常。
# 数据损坏类用 warning，用户输入类用 debug；降级行为一律不变。
logger = logging.getLogger(__name__)

from app.extensions import db
from app.models.timetable import (
    TermSchedule, PeriodDef, ScheduleEntry, ScheduleVersion,
    get_default_periods, SCHEDULE_STATUS, WEEKDAY_NAMES, PERIOD_NUMBER_CEILING,
    PERIOD_TYPES, ENTRY_TYPES,
)
# 公共辅助（学期定位 / 节次查询 / 周次解析）：schedule_common 不反向依赖本模块，无循环导入风险
from app.modules.academic.services.schedule_common import (
    get_active_schedule, get_periods, pick_conflict, temp_target_ids,
    week_range_covers as _week_range_covers,
)

# ─── 内部工具 ──────────────────────────────────────────────────────────────

_WEEKDAY_MAP = {'周一': 1, '周二': 2, '周三': 3, '周四': 4, '周五': 5, '周六': 6, '周日': 7}


def _snapshot(entry):
    """生成条目快照 JSON 字符串（委托 schedule_common）"""
    from app.modules.academic.services.schedule_common import snapshot_entry
    return snapshot_entry(entry)


def _record_version(schedule_id, entry_id, action, snapshot_str, operator=None, remark=None):
    """写入版本记录（委托 schedule_common，保持原签名兼容）"""
    from app.modules.academic.services.schedule_common import record_version
    record_version(
        term_schedule_id=schedule_id,
        entry_id=entry_id,
        action=action,
        operator_id=getattr(operator, 'id', None),
        operator_name=getattr(operator, 'real_name', None) or getattr(operator, 'username', None),
        remark=remark or '',
        snapshot_json=snapshot_str,
    )


def _base_entry_query(schedule_id):
    """基础条目查询：过滤软删除 + 排除**临时调课产生的目标条目**。

    临时调课只在调课当天生效（见 swap_service.resolve_effective_entry），
    不应出现在常规周课表/导出/统计里（否则会出现"原课与调课重叠、跨周次也显示"）。
    查课实时课表由 get_day_matrix 单独叠加当天临时调课。
    """
    q = ScheduleEntry.query.filter_by(
        term_schedule_id=schedule_id, is_deleted=False)
    temp_ids = temp_target_ids(schedule_id)
    if temp_ids:
        q = q.filter(~ScheduleEntry.id.in_(temp_ids))
    return q


def _apply_class_scope(query, allowed_grades, allowed_classes):
    """Filter entries by per-grade class allowlists when a user has them."""
    if allowed_classes is None:
        return query
    grades = set(allowed_grades or allowed_classes.keys())
    clauses = []
    for grade in grades:
        classes = allowed_classes.get(grade, set())
        if classes is None:
            clauses.append(ScheduleEntry.grade == grade)
        elif classes:
            clauses.append(and_(ScheduleEntry.grade == grade,
                                ScheduleEntry.class_name.in_(classes)))
    return query.filter(or_(*clauses)) if clauses else query.filter(false())


def _pick_conflict(entries, week_range=None, week=None):
    """从候选条目中挑出真正冲突的一条（周次无交集不算冲突）。

    实现委托 schedule_common.pick_conflict，保持与调课模块同一套判定口径。
    """
    return pick_conflict(entries, week_range=week_range, week=week)


def _filter_by_week(entries, week):
    """按周次过滤条目列表（week=None 时原样返回）"""
    if week is None:
        return list(entries)
    return [e for e in entries if _week_range_covers(e.week_range, week)]


def _week_hint(week_range_str):
    """冲突提示里的周次后缀（默认 1-18 / 全周时不显示，避免噪音）"""
    s = (week_range_str or '').strip()
    if not s or s in ('全周', '全部', '每周', '1-18'):
        return ''
    return f'，周次 {s}'


def _conflict_msg(kind, who, weekday, period_number, conflict):
    """统一的冲突提示文案（带周次，便于判断单双周冲突）"""
    return (f'{kind}冲突：{who} {WEEKDAY_NAMES.get(weekday,"")}第{period_number}节 '
            f'已有「{conflict.subject}」{_week_hint(conflict.week_range)}')


# ═══════════════════════════════════════════════════════════════════════════════
# 学期管理
# ═══════════════════════════════════════════════════════════════════════════════

# get_active_schedule 和 get_periods 已由 schedule_common 提供（见文件顶部导入）


def list_schedules():
    """全部学期，按 created_at 倒序"""
    return TermSchedule.query.order_by(TermSchedule.created_at.desc()).all()


def create_schedule(name, school_year, term, description=None,
                    created_by=None, with_default_periods=True,
                    start_date=None, end_date=None, total_weeks=None,
                    week_start_offset=0, is_current=False):
    """创建学期，with_default_periods 为真时写入 13 条默认节次。

    Task#27 增量追加可选关键字参数（默认值不改变原有行为）：
    start_date / end_date / total_weeks / week_start_offset / is_current。
    """
    ts = TermSchedule(
        name=name, school_year=school_year, term=term,
        status='draft', description=description,
        created_by=getattr(created_by, 'id', None) if created_by else None,
        start_date=start_date, end_date=end_date,
        total_weeks=total_weeks if total_weeks is not None else 20,
        week_start_offset=week_start_offset or 0,
        is_current=bool(is_current),
    )
    db.session.add(ts)
    db.session.flush()  # 获取 id
    if with_default_periods:
        # 2026-10-10：优先套用「全局作息模板」（学校作息一般全校固定，配一次各处复用），
        # 没有全局模板时才用内置的 13 节默认值
        for p in (get_global_periods() or get_default_periods()):
            db.session.add(PeriodDef(term_schedule_id=ts.id, **p))
    db.session.commit()
    return ts


def activate_schedule(schedule_id):
    """把该学期设为 active；其他 active 学期归档并写快照（进入「历史课表」）。

    2026-10-10：原实现直接 `update(status='archived')` 不留快照，与「归档」语义
    不一致（历史课表里看不到归档记录）；改为复用 term_service.archive_term，
    让"自动归档"与"手动归档"完全同源、可追溯。
    """
    target = db.session.get(TermSchedule, schedule_id)
    if not target:
        raise ValueError('学期不存在')
    from app.modules.academic.services import term_service
    for other in TermSchedule.query.filter_by(status='active').all():
        if other.id == schedule_id:
            continue
        try:
            term_service.archive_term(other.id)
        except Exception:  # noqa: BLE001  归档失败不阻断激活，退化为原行为
            other.status = 'archived'
            other.is_current = False
    target.status = 'active'
    db.session.commit()
    return target


def update_schedule(schedule_id, return_warnings=False, **fields):
    """更新学期字段。

    支持 name/school_year/term/description/status 及 Task#27 新增的
    start_date/end_date/total_weeks/week_start_offset/is_current。
    含日期字段时调用 term_service.validate_term_dates 做校验。
    return_warnings=False（默认）保持向后兼容，返回 ts；
    return_warnings=True 时返回 (ts, warnings)。
    """
    ts = db.session.get(TermSchedule, schedule_id)
    if not ts:
        raise ValueError('学期不存在')
    allowed = {'name', 'school_year', 'term', 'description', 'status',
               'start_date', 'end_date', 'total_weeks', 'week_start_offset',
               'is_current'}
    for k, v in fields.items():
        if k in allowed:
            setattr(ts, k, v)
    warnings = []
    if any(k in fields for k in ('start_date', 'end_date', 'total_weeks')):
        from app.modules.academic.services import term_service
        warnings = term_service.validate_term_dates(
            ts.start_date, ts.end_date, ts.total_weeks)
    db.session.commit()
    if return_warnings:
        return ts, warnings
    return ts


def set_current_term(schedule_id):
    """把指定学期设为 is_current=True 并清除其他学期的 is_current（互斥）。

    2026-10-10 语义拆分（与「归档」区分开）：
    - 本操作只管「当前学期」标记，不再隐式激活/归档其他学期；
    - 未启用的学期请先「激活」；历史（归档）学期请先「重新启用」。
    返回该学期对象。
    """
    target = db.session.get(TermSchedule, schedule_id)
    if not target:
        raise ValueError('学期不存在')
    if target.status == 'draft':
        raise ValueError('该学期未启用，不能直接设为当前，请先「激活」')
    if target.status == 'archived':
        raise ValueError('历史学期不能设为当前，请先「重新启用」')
    TermSchedule.query.filter(TermSchedule.id != schedule_id)\
        .update({'is_current': False}, synchronize_session=False)
    target.is_current = True
    db.session.commit()
    return target


def delete_schedule(schedule_id):
    """删除学期（仅限未启用 draft 状态，cascade 清掉节次和条目）。

    2026-10-10：删除是**不可恢复**的物理删除，只对"还没启用过"的学期开放；
    启用中 / 已归档的学期走「归档进历史课表」（留快照、仍可回看），不提供删除。
    """
    ts = db.session.get(TermSchedule, schedule_id)
    if not ts:
        raise ValueError('学期不存在')
    if ts.status != 'draft':
        raise ValueError('只能删除未启用的学期；已启用或已归档的课表请归档进「历史课表」保留')
    # 物理删除条目（含已软删除的）
    ScheduleEntry.query.filter_by(term_schedule_id=schedule_id).delete()
    ScheduleVersion.query.filter_by(term_schedule_id=schedule_id).delete()
    db.session.delete(ts)
    db.session.commit()


# ═══════════════════════════════════════════════════════════════════════════════
# 节次管理
# ═══════════════════════════════════════════════════════════════════════════════

# get_periods 已由 schedule_common 提供（见文件顶部导入）


def save_periods(schedule_id, periods_data, remove_missing=False):
    """批量保存作息时间配置（按 period_number upsert）。

    remove_missing=True 时，未出现在提交列表中的节次会被**删除**
    （修复此前"删掉的行保存后又复活"的问题）；
    若该节次仍被课表条目引用，则保留不删，并通过 kept_in_use 返回提示，
    避免出现"有课但没有节次定义"的孤儿数据。

    返回 {'removed': [节次号...], 'kept_in_use': [节次号...]}
    """
    existing = {p.period_number: p for p in get_periods(schedule_id)}
    submitted = set()
    for item in periods_data:
        pn = item.get('period_number')
        # 2026-10-10：原为 1..13 业务上限；现在只挡非法编号（节次数由学校自定义）
        if not pn or not isinstance(pn, int) or pn < 1 or pn > PERIOD_NUMBER_CEILING:
            continue
        submitted.add(pn)
        if pn in existing:
            p = existing[pn]
            p.period_name = item.get('period_name', p.period_name)
            p.start_time = item.get('start_time', p.start_time)
            p.end_time = item.get('end_time', p.end_time)
            p.period_type = item.get('period_type', p.period_type)
            p.sort_order = item.get('sort_order', p.sort_order)
        else:
            p = PeriodDef(term_schedule_id=schedule_id, **{
                k: item[k] for k in ('period_number', 'period_name', 'start_time',
                                     'end_time', 'period_type', 'sort_order')
                if k in item
            })
            db.session.add(p)
    removed, kept_in_use = [], []
    if remove_missing:
        candidates = sorted(set(existing) - submitted)
        if candidates:
            used = {r[0] for r in db.session.query(ScheduleEntry.period_number)
                    .filter(ScheduleEntry.term_schedule_id == schedule_id,
                            ScheduleEntry.is_deleted.is_(False),
                            ScheduleEntry.period_number.in_(candidates))
                    .distinct().all()}
            for pn in candidates:
                if pn in used:
                    kept_in_use.append(pn)
                    continue
                db.session.delete(existing[pn])
                removed.append(pn)
    db.session.commit()
    return {'removed': removed, 'kept_in_use': kept_in_use}


def reset_default_periods(schedule_id):
    """重置为内置默认模板（早读 + 上午5 + 下午4 + 晚自习3 = 13 节）。已有条目时拒绝，以免节次号改义。"""
    if ScheduleEntry.query.filter_by(term_schedule_id=schedule_id).first():
        raise ValueError(
            '该学期已有课表条目，不能重置默认节次；请手动调整名称和时间，避免已有课程错位')
    PeriodDef.query.filter_by(term_schedule_id=schedule_id).delete()
    for p in get_default_periods():
        db.session.add(PeriodDef(term_schedule_id=schedule_id, **p))
    db.session.commit()


# ═══════════════════════════════════════════════════════════════════════════════
# 全局作息模板（2026-10-10 新增）
#
# 学校的作息时间一般是全校固定的（同一套铃声），但 PeriodDef 是按学期存的（不同学期
# 可能微调，如冬夏作息）。于是提供一份**全局模板**：在任意学期把作息「存为全局作息」，
# 新建学期默认套用、其它学期也能一键「套用全局作息」，避免每建一个学期就重配一遍。
# 存在 system_settings 的 KV 表（key=global_periods，值为 JSON），不新增表/迁移。
# ═══════════════════════════════════════════════════════════════════════════════

GLOBAL_PERIODS_KEY = 'global_periods'
_GLOBAL_PERIOD_KEYS = ('period_number', 'period_name', 'start_time', 'end_time',
                       'period_type', 'sort_order')


def get_global_periods():
    """读全局作息模板 → list[dict]（可直接 PeriodDef(**row)）；没有/坏了返回 None。"""
    from app.models.system_setting import SystemSetting
    raw = SystemSetting.get(GLOBAL_PERIODS_KEY)
    if not raw:
        return None
    try:
        data = json.loads(raw)
    except (ValueError, TypeError):
        return None
    rows = (data.get('periods') if isinstance(data, dict) else data) or []
    out = []
    for i, row in enumerate(rows):
        if not isinstance(row, dict):
            continue
        try:
            pn = int(row.get('period_number'))
        except (TypeError, ValueError):
            continue
        if pn < 1:
            continue
        out.append({
            'period_number': pn,
            'period_name': str(row.get('period_name') or f'第{pn}节')[:20],
            'start_time': row.get('start_time') or None,
            'end_time': row.get('end_time') or None,
            'period_type': row.get('period_type') or 'morning',
            'sort_order': int(row.get('sort_order') or i + 1),
        })
    return sorted(out, key=lambda r: (r['sort_order'], r['period_number'])) or None


def get_global_periods_info():
    """全局作息模板摘要（给页面显示）：{count, updated_at} 或 None。"""
    from app.models.system_setting import SystemSetting
    raw = SystemSetting.get(GLOBAL_PERIODS_KEY)
    if not raw:
        return None
    try:
        data = json.loads(raw)
    except (ValueError, TypeError):
        return None
    rows = get_global_periods() or []
    if not rows:
        return None
    return {'count': len(rows),
            'updated_at': (data.get('updated_at') if isinstance(data, dict) else None)}


def save_global_periods(schedule_id, operator=None):
    """把指定学期的作息存为全局模板 → (success, message)。"""
    periods = get_periods(schedule_id)
    if not periods:
        return False, '该学期还没有节次定义，无法存为全局作息'
    from app.models.system_setting import SystemSetting
    payload = {'updated_at': datetime.now().strftime('%Y-%m-%d %H:%M'),
               'periods': [{k: getattr(p, k) for k in _GLOBAL_PERIOD_KEYS}
                           for p in periods]}
    SystemSetting.set(GLOBAL_PERIODS_KEY,
                      json.dumps(payload, ensure_ascii=False),
                      user_id=getattr(operator, 'id', None),
                      description='全局作息模板（节次名称与时间）：新建学期默认套用，'
                                  '也可在作息时间页套用到任意学期')
    db.session.commit()
    return True, f'已把本学期的 {len(periods)} 节作息存为全局模板（新建学期将默认套用）'


def apply_global_periods(schedule_id):
    """把全局作息模板套用到指定学期 → (success, message, result|None)。

    复用 save_periods：按节次号 upsert + 移除模板里没有的节次；仍被课程引用的节次
    会保留（kept_in_use），避免出现"有课却没有节次定义"的孤儿数据。
    """
    rows = get_global_periods()
    if not rows:
        return False, '还没有全局作息模板：请先在某个学期点「存为全局作息」', None
    result = save_periods(schedule_id, [dict(r) for r in rows], remove_missing=True)
    msg = f'已套用全局作息（{len(rows)} 节）'
    if result['removed']:
        msg += f"；移除了模板里没有的第 {'、'.join(str(n) for n in result['removed'])} 节"
    return True, msg, result


# ═══════════════════════════════════════════════════════════════════════════════
# 视图查询
# ═══════════════════════════════════════════════════════════════════════════════

def _build_grid(entries, periods, weekday=None):
    """构建 {period_number: {weekday: [entry_dict, ...]}} 网格。

    同一格子允许多条（如"单周语文 / 双周数学"交替上课），按周次先后排序，
    便于前端逐条渲染与点击编辑。

    weekday（1-7，可选）：只给这一天调用 to_dict 填格。2026-10-10 优化——全校总
    课表一次只渲染一天（见模板 schedule_overview.html），其余 6 天的 to_dict 属于
    白干（响应体里也渲染不出来）；不传＝全部星期都填，行为与优化前完全一致。
    """
    grid = {}
    for p in periods:
        grid[p.period_number] = {wd: [] for wd in range(1, 8)}
    for e in entries:
        if e.period_number in grid:
            if weekday is not None and e.weekday != weekday:
                continue
            grid[e.period_number][e.weekday].append(e.to_dict())
    return grid


def get_class_view(schedule_id, grade, class_name, weekday=None, week=None,
                   allowed_grades=None, allowed_classes=None):
    """班级课表视图：返回网格 + 节次列表 + 学科课时统计（week 为可选周次过滤）"""
    periods = get_periods(schedule_id)
    q = _base_entry_query(schedule_id).filter_by(grade=grade, class_name=class_name)
    if allowed_grades is not None:
        q = q.filter(ScheduleEntry.grade.in_(allowed_grades))
    q = _apply_class_scope(q, allowed_grades, allowed_classes)
    if weekday:
        q = q.filter_by(weekday=weekday)
    entries = _filter_by_week(q.all(), week)
    grid = _build_grid(entries, periods)
    # 学科课时统计
    stats = {}
    for e in entries:
        stats[e.subject] = stats.get(e.subject, 0) + 1
    return {
        'grid': grid,
        'periods': [p.to_dict() for p in periods],
        'entries': [e.to_dict() for e in entries],
        'stats': stats,
        'total': len(entries),
    }


def get_grade_view(schedule_id, grade, weekday=None, week=None,
                   allowed_grades=None, allowed_classes=None):
    """年级视图：该年级所有班级列表 + 每个班的网格"""
    periods = get_periods(schedule_id)
    q = _base_entry_query(schedule_id).filter_by(grade=grade)
    if allowed_grades is not None:
        q = q.filter(ScheduleEntry.grade.in_(allowed_grades))
    q = _apply_class_scope(q, allowed_grades, allowed_classes)
    if weekday:
        q = q.filter_by(weekday=weekday)
    entries = _filter_by_week(q.all(), week)
    classes = sorted(set(e.class_name for e in entries))
    result = {}
    for cn in classes:
        cls_entries = [e for e in entries if e.class_name == cn]
        result[cn] = _build_grid(cls_entries, periods)
    return {
        'classes': classes,
        'grids': result,
        'periods': [p.to_dict() for p in periods],
        'total': len(entries),
    }


# 2026-10-10：get_master_view()（大课表：按年级分组返回全校网格）已删——随大课表页下线，
# 且早已无调用者（页面后来走 get_grade_class_list + get_class_view 的 AJAX 链路）。


def get_teacher_view(schedule_id, teacher_uid, weekday=None, week=None,
                     allowed_grades=None, allowed_classes=None):
    """教师个人课表 + 学科课时统计"""
    periods = get_periods(schedule_id)
    q = _base_entry_query(schedule_id).filter_by(teacher_uid=teacher_uid)
    if weekday:
        q = q.filter_by(weekday=weekday)
    if allowed_grades is not None:
        q = q.filter(ScheduleEntry.grade.in_(allowed_grades))
    q = _apply_class_scope(q, allowed_grades, allowed_classes)
    entries = _filter_by_week(q.all(), week)
    grid = _build_grid(entries, periods)
    stats = {}
    for e in entries:
        stats[e.subject] = stats.get(e.subject, 0) + 1
    return {
        'grid': grid,
        'periods': [p.to_dict() for p in periods],
        'entries': [e.to_dict() for e in entries],
        'stats': stats,
        'total': len(entries),
    }



def _period_label(p):
    """节次所属的作息分组：早读 / 上午 / 下午 / 晚自习 / 课间（高中作息语义）。"""
    name = p.get('period_name') or ''
    if '早读' in name or '晨读' in name:
        return '早读'
    return {'morning': '上午', 'afternoon': '下午',
            'evening': '晚自习', 'break': '课间'}.get(p.get('period_type'), '其他')


def _period_groups(period_dicts):
    """给每个节次标注作息分组，并返回连续分组表（供总课表表头/导出复用）。

    总课表版式是「行＝班级、列＝节次」，作息分组直接做**二级表头**：
    连续同组（早读/上午/下午/晚自习…）合并成一个跨列表头格；午休等 break
    节次自成「课间」组，渲染成窄灰空档列。
    """
    groups = []
    for p in period_dicts:
        label = _period_label(p)
        p['group_label'] = label
        if groups and groups[-1]['label'] == label:
            groups[-1]['span'] += 1
            p['group_start'] = False
        else:
            groups.append({'label': label, 'span': 1})
            p['group_start'] = True
    return groups


def _build_vertical_matrix(blocks, period_dicts, day):
    """把「行＝班级、列＝节次」的分块转置成**一张竖版表**（2026-10-10 改版）。

        ┌──────┬────────┬────────┬────────┬────────┐
        │ 作息 │  节次  │ 高三1班 │ 高三2班 │ 高二1班 │  ← 表头＝班级（按年级分组）
        ├──────┼────────┼────────┼────────┼────────┤
        │ 早读 │  早读  │ 语文   │ 英语   │ 英语   │
        │ 上午 │  第1节 │ 数学   │ 语文   │ 数学   │
        └──────┴────────┴────────┴────────┴────────┘

    行＝节次（左列作息分组、次列节次名），列＝班级。三个年级的班并排放在**同一张表**里，
    节次行只出现一次 —— 比"一个年级一张表"省掉两张重复表头，也才能一屏容下全校班级。
    返回 {grade_groups, columns, rows, grade_starts}；cells 与 columns 按下标对齐。
    """
    usable = [b for b in blocks if b.get('classes')]
    columns = []
    grade_groups = []
    grade_starts = []
    for b in usable:
        grade_groups.append({'grade': b['grade'],
                             'label': b['grade_label'] or b['grade'],
                             'span': len(b['classes']),
                             'total': b['total']})
        grade_starts.append(len(columns))
        short = (b['grade_label'] or b['grade']).split('(')[0]   # 高三(2024级) → 高三
        for cn in b['classes']:
            columns.append({
                'grade': b['grade'],
                'grade_label': b['grade_label'] or b['grade'],
                'grade_short': short,
                'class_name': cn,
                'meta': b['class_meta'].get(cn, {}),
                'total': b['class_totals'].get(cn, 0),
            })

    rows = []
    for p in period_dicts:
        pn = p['period_number']
        cells = []
        for b in usable:
            for cn in b['classes']:
                cells.append(((b['grids'].get(cn) or {}).get(pn) or {}).get(day) or [])
        rows.append({'period': p,
                     'group_label': p.get('group_label') or _period_label(p),
                     'group_start': False, 'group_span': 1, 'cells': cells})

    # 作息分组（早读/上午/下午/晚自习）：连续同组只在首行出一格，用 rowspan 跨行
    i = 0
    while i < len(rows):
        label = rows[i]['group_label']
        span = 1
        while i + span < len(rows) and rows[i + span]['group_label'] == label:
            span += 1
        rows[i]['group_start'] = True
        rows[i]['group_span'] = span
        i += span

    return {'grade_groups': grade_groups, 'columns': columns, 'rows': rows,
            'grade_starts': grade_starts}


def get_overview_view(schedule_id, week=None, grades=None, include_break=True,
                      max_weekday=5, direction=None, class_type=None,
                      allowed_grades=None, allowed_classes=None, weekday=None,
                      prefer_weekday=None):
    """全校总课表（2026-09-25 新增；2026-10-09 定版为「行=班级、列=节次」矩阵）。

    对齐"一纸打印全校课表"的形态：每个年级一张大表，第一列＝该年级各班，
    其余列＝一天的各节次（早读 + 上午 5 节 + 下午 4 节 + 晚自习 3 节），
    格内显示"学科 + 教师"（可带教室）；星期由页面顶部标签切换，每次只渲染一天。
    参数：
    - grades: 指定年级列表（None=全部）
    - include_break: 是否包含「课间/午休」节次（打印时常常去掉）
    - max_weekday: 显示到星期几（1-7，默认 5=周五；数据里有周六日的自动放宽）
    - direction/class_type: 按**选科方向**（物理/历史，新高考 3+1+2）与**班型**
      （强基班/卓越班…）筛选，读 ClassProfile；高中场景下教务常按这两维度看课表。
    - weekday: 只构建/只填充这一天（1-7）的网格。2026-10-10 优化——页面一次只渲染
      一天，其余 6 天不调用 to_dict；传 None 时构建全部 7 天，行为与优化前完全一致
      （导出路径 _export_overview_sheet 仍走全周，故全周能力保留）。
    """
    from app.modules.academic.services.grade_utils import (class_profile_map,
                                                           grade_labels,
                                                           grade_sort_key)

    periods = get_periods(schedule_id)
    if not include_break:
        periods = [p for p in periods if p.period_type != 'break']
    entries = _filter_by_week(_base_entry_query(schedule_id).all(), week)
    if allowed_grades is not None:
        entries = [e for e in entries if e.grade in allowed_grades]
    if allowed_classes is not None:
        entries = [e for e in entries
                   if e.class_name in allowed_classes.get(e.grade, set())
                   or allowed_classes.get(e.grade) is None]
    all_grades = sorted({e.grade for e in entries if e.grade})
    profiles = class_profile_map()
    labels = grade_labels(all_grades)
    visible_pairs = {(e.grade, e.class_name) for e in entries}
    visible_profiles = [profiles.get(pair, {}) for pair in visible_pairs]
    if grades:
        wanted = {g for g in grades if g}
        entries = [e for e in entries if e.grade in wanted]
    if direction:
        entries = [e for e in entries
                   if profiles.get((e.grade, e.class_name), {}).get('direction') == direction]
    if class_type:
        entries = [e for e in entries
                   if profiles.get((e.grade, e.class_name), {}).get('class_type') == class_type]
    if entries:
        max_weekday = max(max_weekday, max((e.weekday or 1) for e in entries))

    # 2026-10-10：缺省入口的"今天超出课表范围就退回周一"在这里一次决定 —— 原来路由要
    # 先构建一遍才知道 max_weekday，周末（今天不在课表范围内）会白跑一整趟取数+建格。
    build_weekday = weekday
    if prefer_weekday is not None:
        build_weekday = prefer_weekday if prefer_weekday <= max_weekday else 1

    by_grade = {}
    for e in entries:
        by_grade.setdefault(e.grade, {}).setdefault(e.class_name, []).append(e)

    blocks = []
    # 高中习惯：毕业年级在前（高三 → 高二 → 高一），而不是按名称升序
    for grade in sorted(by_grade, key=lambda g: (-grade_sort_key(g, labels), g)):
        class_map = by_grade[grade]
        classes = sorted(class_map)
        # 2026-10-10：weekday 有值时只为这一天建格（其余天的 to_dict 白干）
        grids = {cn: _build_grid(class_map[cn], periods, build_weekday) for cn in classes}
        blocks.append({
            'grade': grade,
            'grade_label': labels.get(grade, grade),   # 如「高三(2024级)」
            'classes': classes,
            'class_meta': {cn: profiles.get((grade, cn), {}) for cn in classes},
            'grids': grids,
            'class_totals': {cn: len(class_map[cn]) for cn in classes},  # 每班节数（卡片角标）
            'total': sum(len(v) for v in class_map.values()),
        })
    period_dicts = [p.to_dict() for p in periods]
    period_groups = _period_groups(period_dicts)   # 给每个节次标 group_label
    # 竖版（2026-10-10）：行＝节次、列＝班级，三个年级并到一张表；blocks 保留给
    # Excel 导出等旧消费方，页面走 ovv
    ovv = _build_vertical_matrix(blocks, period_dicts, build_weekday)
    return {
        'blocks': blocks,
        'ovv': ovv,
        'periods': period_dicts,
        'period_groups': period_groups,
        'grades': [b['grade'] for b in blocks],
        'all_grades': all_grades,   # 未过滤的年级全集（供筛选控件渲染，避免过滤后无法切回）
        'labels': labels,
        'directions': sorted({v.get('direction') for v in visible_profiles
                              if v.get('direction')}),
        'class_types': sorted({v.get('class_type') for v in visible_profiles
                               if v.get('class_type')}),
        'max_weekday': max_weekday,
        'day': build_weekday,       # 实际渲染的星期几（缺省入口由 prefer_weekday 决定）
        'total': len(entries),
    }


def _grade_filter_values(grade):
    if grade is None:
        return None
    if isinstance(grade, str):
        return {grade} if grade else set()
    return {g for g in grade if g}


def _empty_day_matrix(target_date=None):
    """空矩阵（未建学期/无权限范围时的兜底结构，字段与 get_day_matrix 对齐）。"""
    return {'entries': [], 'current_period_number': None, 'weekday': 0,
            'weekday_text': '', 'periods': [], 'period_groups': [],
            'blocks': [], 'labels': {}, 'max_weekday': 0, 'total': 0,
            'date': (target_date or date.today()).isoformat(),
            'schedule_id': None, 'week': None}


def get_day_matrix(schedule_id, target_date, week=None, grade=None, class_name=None,
                   allowed_classes=None, highlight_now=True):
    """任意一天的「行=班级、列=节次」矩阵（全校总课表 / 查课实时课表共用）。

    blocks/grids/periods/period_groups 可直接喂 `_schedule_matrix.html` 宏；
    并叠加该日期的临时调课（调入的显示、被调走的不显示）。

    :param target_date: 要看的日期（date 对象）——查课要按日期回查/补录历史，
                        因此不再写死 now()。
    :param highlight_now: 只有值为 True 且 target_date 就是今天时才给出
                        current_period_number（"当前节次整列高亮"），否则为 None。
    :param allowed_classes: {年级: {班级}} 数据范围（查课页复用同一套范围过滤）。
    """
    if not schedule_id:
        return _empty_day_matrix(target_date)
    wd = target_date.isoweekday()  # 1=周一 ... 7=周日
    periods = get_periods(schedule_id)
    grade_values = _grade_filter_values(grade)
    q = _base_entry_query(schedule_id).filter_by(weekday=wd)
    if grade_values is not None:
        q = q.filter(ScheduleEntry.grade.in_(grade_values))
    q = _apply_class_scope(q, grade_values, allowed_classes)
    if class_name:
        q = q.filter_by(class_name=class_name)
    entries = _filter_by_week(
        q.order_by(ScheduleEntry.grade, ScheduleEntry.class_name,
                   ScheduleEntry.period_number).all(), week)

    # 叠加"当天临时调课"：调入的显示、被调走的不显示（常规周课表已排除临时条目）
    moved_away, temp_dicts = set(), []
    try:
        from app.modules.academic.services import swap_service
        for sw in swap_service.get_temp_swaps_by_date(schedule_id, target_date):
            oe = sw.get('original_entry')
            te = sw.get('target_entry')
            if oe:
                moved_away.add(oe['id'])
            if te and te.get('weekday') == wd:
                if grade_values is not None and te.get('grade') not in grade_values:
                    continue
                if (allowed_classes is not None and
                        (te.get('class_name') not in allowed_classes.get(te.get('grade'), set())
                         and allowed_classes.get(te.get('grade')) is not None)):
                    continue
                if class_name and te.get('class_name') != class_name:
                    continue
                d = dict(te)
                d['is_temp_swap'] = True
                temp_dicts.append(d)
    except Exception as _e:  # noqa: BLE001
        # 2026-10-11 审计：调课叠加失败退化为基础课表，但需留痕（原为静默 pass）
        logger.warning('叠加当天临时调课失败，退化为基础课表: %s', _e)
    entry_dicts = [e.to_dict() for e in entries if e.id not in moved_away]
    has_ids = {d['id'] for d in entry_dicts}
    entry_dicts += [d for d in temp_dicts if d['id'] not in has_ids]
    entry_dicts.sort(key=lambda d: (d.get('grade', ''), d.get('class_name', ''),
                                    d.get('period_number', 0)))

    # 确定当前节次（只在"看今天"时有意义）
    current_period = None
    if highlight_now and target_date == date.today():
        now = datetime.now()
        cur_minutes = now.hour * 60 + now.minute
        for p in periods:
            if p.start_time and p.end_time:
                try:
                    sh, sm = map(int, p.start_time.split(':'))
                    eh, em = map(int, p.end_time.split(':'))
                    if sh * 60 + sm <= cur_minutes <= eh * 60 + em:
                        current_period = p.period_number
                        break
                except (ValueError, AttributeError) as _e:
                    # 节次时间字段格式异常（脏数据）→ 该节次不参与"当前节次"判定。
                    # 属配置数据问题，用 warning 便于发现后修数据。
                    logger.warning('节次时间格式异常，已跳过该节次(id=%s, start=%r, end=%r): %s',
                                   getattr(p, 'id', '?'),
                                   getattr(p, 'start_time', None),
                                   getattr(p, 'end_time', None), _e)

    # ── 按年级分块：与「全校总课表」同款矩阵（行=班级、列=节次）──
    # 这一天的格子只填 weekday=wd 那一列；临时调课（当天调入）已并入 entry_dicts。
    from app.modules.academic.services.grade_utils import (class_profile_map,
                                                           grade_labels,
                                                           grade_sort_key)
    profiles = class_profile_map()
    labels = grade_labels(sorted({d.get('grade') for d in entry_dicts if d.get('grade')}))
    by_grade = {}
    for d in entry_dicts:
        by_grade.setdefault(d.get('grade'), {}).setdefault(d.get('class_name'), []).append(d)
    blocks = []
    for g in sorted(by_grade, key=lambda x: (-grade_sort_key(x, labels), x)):
        cmap = by_grade[g]
        classes = sorted(cmap)
        grids = {}
        for cn in classes:
            grid = {p.period_number: {w: [] for w in range(1, 8)} for p in periods}
            for d in cmap[cn]:
                pn = d.get('period_number')
                if pn in grid:
                    grid[pn][wd].append(d)
            grids[cn] = grid
        blocks.append({
            'grade': g,
            'grade_label': labels.get(g, g),
            'classes': classes,
            'class_meta': {cn: profiles.get((g, cn), {}) for cn in classes},
            'grids': grids,
            'class_totals': {cn: len(cmap[cn]) for cn in classes},
            'total': sum(len(v) for v in cmap.values()),
        })
    period_dicts = [p.to_dict() for p in periods]
    return {
        'entries': entry_dicts,
        'current_period_number': current_period,
        'weekday': wd,
        'weekday_text': WEEKDAY_NAMES.get(wd, ''),
        'periods': period_dicts,
        'period_groups': _period_groups(period_dicts),
        'blocks': blocks,
        'labels': labels,
        'max_weekday': wd,
        'date': target_date.isoformat(),
        'total': len(entry_dicts),
        'schedule_id': schedule_id,
        'week': week,
    }


def build_matrix_from_entries(entry_dicts, periods):
    """把扁平条目列表整理成「行=班级、列=节次」矩阵（导入预览用，不查库）。

    与 get_day_matrix() 的 blocks/grids 同构，可直接喂 `_schedule_matrix.html`
    宏 —— 导入前就能看到"导进去是什么样"，与全校总课表/查课核对页同版式。
    entry_dicts 需含 grade/class_name/weekday/period_number/subject 等字段。
    返回 {'blocks','periods','period_groups','days','total'}，days = 出现过的星期几。
    """
    period_dicts = []
    for p in (periods or []):
        period_dicts.append(p.to_dict() if hasattr(p, 'to_dict') else dict(p))
    numbers = [d.get('period_number') for d in period_dicts if d.get('period_number')]
    by_grade = {}
    for e in (entry_dicts or []):
        g = e.get('grade') or '未分年级'
        cn = e.get('class_name') or '未分班'
        by_grade.setdefault(g, {}).setdefault(cn, []).append(e)
    blocks = []
    for g in sorted(by_grade):
        cmap = by_grade[g]
        classes = sorted(cmap)
        grids = {}
        for cn in classes:
            grid = {pn: {w: [] for w in range(1, 8)} for pn in numbers}
            for e in cmap[cn]:
                pn, w = e.get('period_number'), e.get('weekday')
                if pn in grid and 1 <= (w or 0) <= 7:
                    grid[pn][w].append(e)
            grids[cn] = grid
        blocks.append({'grade': g, 'grade_label': g, 'classes': classes,
                       'class_meta': {}, 'grids': grids,
                       'class_totals': {cn: len(cmap[cn]) for cn in classes},
                       'total': sum(len(v) for v in cmap.values())})
    days = sorted({e.get('weekday') for e in (entry_dicts or []) if e.get('weekday')})
    return {'blocks': blocks, 'periods': period_dicts,
            'period_groups': _period_groups(period_dicts),
            'days': days, 'total': len(entry_dicts or [])}



def get_entry_detail(entry_id):
    """单条目详情"""
    e = db.session.get(ScheduleEntry, entry_id)
    if not e or e.is_deleted:
        return None
    return e.to_dict()


# ═══════════════════════════════════════════════════════════════════════════════
# 条目 CRUD
# ═══════════════════════════════════════════════════════════════════════════════

def check_class_conflict(schedule_id, grade, class_name, weekday, period_number,
                         exclude_entry_id=None, week_range=None, week=None,
                         teaching_class=None):
    """检测同班同时段冲突，返回冲突 entry 或 None。

    week_range / week 用于周次维度过滤：周次无交集的条目不算冲突
    （同一格子可以是"单周语文 / 双周数学"）。

    teaching_class（教学班，2026-09-26）：
    - 待排的条目若填了教学班，只与"同一教学班"的课判冲突。不同教学班在同一时段
      并行开课是走班形态（物化生1 与 物化生2 同时上），不能算班级冲突。
      注：走班教学班课表页已删（2026-10-10），但该列仍由排课编辑使用。
    - 未填教学班的条目视为行政班课，与全班所有课判冲突（含填了教学班的课）。
    """
    q = _base_entry_query(schedule_id).filter_by(
        grade=grade, class_name=class_name,
        weekday=weekday, period_number=period_number)
    if exclude_entry_id:
        q = q.filter(ScheduleEntry.id != exclude_entry_id)
    if teaching_class:
        q = q.filter(db.or_(ScheduleEntry.teaching_class == teaching_class,
                            ScheduleEntry.teaching_class.is_(None),
                            ScheduleEntry.teaching_class == ''))
    return _pick_conflict(q.all(), week_range=week_range, week=week)


def check_teacher_conflict(schedule_id, teacher_uid, weekday, period_number,
                           exclude_entry_id=None, week_range=None, week=None,
                           teaching_class=None):
    """检测同教师同时段冲突，返回冲突 entry 或 None（周次维度同上）。

    走班课（2026-10-12）：同一个**教学班**（teaching_class 同名）的课在别的班课表里
    同时出现，其实是"同一次走班课"（混合班两拨学生一起上），不算教师冲突 ——
    如陈丽莉带的政治走班班同时在 2025级03班 / 04 班课表里各占一行。
    """
    if not teacher_uid:
        return None
    q = _base_entry_query(schedule_id).filter_by(
        teacher_uid=teacher_uid, weekday=weekday, period_number=period_number)
    if exclude_entry_id:
        q = q.filter(ScheduleEntry.id != exclude_entry_id)
    rows = q.all()
    if teaching_class:
        rows = [e for e in rows if (e.teaching_class or '') != teaching_class]
    return _pick_conflict(rows, week_range=week_range, week=week)


def add_entry(schedule_id, grade, class_name, weekday, period_number, subject,
              teacher_uid=None, teacher_name=None, room=None,
              week_range='1-18', note=None, operator=None, teaching_class=None):
    """添加课条目 → (success, message_or_entry)"""
    # 校验节次：2026-10-10 去掉"一天最多 13 节"的业务上限，只挡非法编号；
    # 真正的约束是下一句的"该节次必须在本学期定义过"
    if not (1 <= period_number <= PERIOD_NUMBER_CEILING):
        return False, '节次编号无效'
    pd = PeriodDef.query.filter_by(term_schedule_id=schedule_id,
                                   period_number=period_number).first()
    if not pd:
        return False, f'该学期未定义第 {period_number} 节'
    week_range = week_range or '1-18'
    # 班级冲突（按周次交集判定）
    conflict = check_class_conflict(schedule_id, grade, class_name, weekday,
                                    period_number, week_range=week_range,
                                    teaching_class=teaching_class)
    if conflict:
        return False, _conflict_msg('班级', f'{grade}{class_name}', weekday,
                                    period_number, conflict)
    # 教师冲突（走班课按教学班合并：同名教学班跨班同格不算冲突）
    if teacher_uid:
        tc = check_teacher_conflict(schedule_id, teacher_uid, weekday, period_number,
                                    week_range=week_range,
                                    teaching_class=teaching_class)
        if tc:
            return False, (f'教师冲突：{teacher_name or teacher_uid} '
                           f'{WEEKDAY_NAMES.get(weekday,"")}第{period_number}节 '
                           f'已有「{tc.subject}」({tc.grade}{tc.class_name}'
                           f'{_week_hint(tc.week_range)})')

    entry = ScheduleEntry(
        term_schedule_id=schedule_id, grade=grade, class_name=class_name,
        weekday=weekday, period_number=period_number, subject=subject,
        teacher_uid=teacher_uid, teacher_name=teacher_name,
        room=room, week_range=week_range or '1-18', note=note,
        teaching_class=(teaching_class or '').strip() or None,
        entry_type='normal', is_deleted=False,
    )
    db.session.add(entry)
    db.session.flush()
    _record_version(schedule_id, entry.id, 'create', _snapshot(entry),
                    operator, f'新增：{grade}{class_name} {WEEKDAY_NAMES.get(weekday,"")}第{period_number}节 {subject}')
    db.session.commit()
    return True, entry


def edit_entry(entry_id, operator=None, **fields):
    """编辑课条目（冲突校验排除自身）→ (success, message_or_entry)"""
    entry = db.session.get(ScheduleEntry, entry_id)
    if not entry or entry.is_deleted:
        return False, '条目不存在'

    new_weekday = fields.get('weekday', entry.weekday)
    new_period = fields.get('period_number', entry.period_number)
    new_grade = fields.get('grade', entry.grade)
    new_class = fields.get('class_name', entry.class_name)
    new_teacher_uid = fields.get('teacher_uid', entry.teacher_uid)
    new_teacher_name = fields.get('teacher_name', entry.teacher_name)
    new_week_range = fields.get('week_range', entry.week_range)

    # 冲突校验（按周次交集判定：单周课与双周课同格不冲突）
    conflict = check_class_conflict(entry.term_schedule_id, new_grade, new_class,
                                    new_weekday, new_period, exclude_entry_id=entry_id,
                                    week_range=new_week_range,
                                    teaching_class=fields.get('teaching_class',
                                                              entry.teaching_class))
    if conflict:
        return False, _conflict_msg('班级', f'{new_grade}{new_class}', new_weekday,
                                    new_period, conflict)
    if new_teacher_uid:
        tc = check_teacher_conflict(entry.term_schedule_id, new_teacher_uid,
                                    new_weekday, new_period, exclude_entry_id=entry_id,
                                    week_range=new_week_range,
                                    teaching_class=fields.get('teaching_class',
                                                              entry.teaching_class))
        if tc:
            return False, (f'教师冲突：{new_teacher_name or new_teacher_uid} '
                           f'{WEEKDAY_NAMES.get(new_weekday,"")}第{new_period}节 '
                           f'已有课({tc.grade}{tc.class_name}'
                           f'{_week_hint(tc.week_range)})')

    # 快照旧数据
    old_snapshot = _snapshot(entry)
    allowed = {'grade', 'class_name', 'weekday', 'period_number', 'subject',
               'teacher_uid', 'teacher_name', 'room', 'week_range', 'note',
               'teaching_class'}
    for k, v in fields.items():
        if k in allowed:
            setattr(entry, k, v)
    entry.updated_at = datetime.now()
    _record_version(entry.term_schedule_id, entry_id, 'update', old_snapshot,
                    operator, f'编辑：{entry.grade}{entry.class_name} {entry.subject}')
    db.session.commit()
    return True, entry


def delete_entry(entry_id, operator=None):
    """软删除 → (success, message)"""
    entry = db.session.get(ScheduleEntry, entry_id)
    if not entry or entry.is_deleted:
        return False, '条目不存在'
    old_snapshot = _snapshot(entry)
    entry.is_deleted = True
    entry.updated_at = datetime.now()
    _record_version(entry.term_schedule_id, entry_id, 'delete', old_snapshot,
                    operator, f'删除：{entry.grade}{entry.class_name} {WEEKDAY_NAMES.get(entry.weekday,"")}第{entry.period_number}节 {entry.subject}')
    db.session.commit()
    return True, '已删除'


def clear_entries(schedule_id, operator=None):
    """清空学期**全部课表条目**（软删除 + 汇总版本留痕）→ {'deleted': n}

    - 与 `scripts/import_school_timetable.py --replace` 同口径：只软删除课表条目，
      **不动学期本身与作息时间配置**（要删整个学期用 delete_schedule）；
    - 已归档学期的只读拦截由路由层 `_readonly_guard` 负责，这里不做状态判断；
    - 写一条汇总版本（action='clear'，entry_id 为空），在页面「变更历史」可追溯，
      但**不逐条记版本**（整学期动辄上千条，逐条会把版本表撑爆）。
    """
    ts = db.session.get(TermSchedule, schedule_id)
    if not ts:
        raise ValueError('学期不存在')
    entries = ScheduleEntry.query.filter_by(term_schedule_id=schedule_id,
                                            is_deleted=False).all()
    if not entries:
        return {'deleted': 0}
    now = datetime.now()
    for e in entries:
        e.is_deleted = True
        e.updated_at = now
    _record_version(schedule_id, None, 'clear', '{}', operator,
                    f'清空课表：{len(entries)} 条课程条目')
    db.session.commit()
    return {'deleted': len(entries)}


def check_teacher_mapping(entries, subjects=None):
    """按**任课教师映射**（TeacherSubjectLink）核对导入条目 → {'checked', 'issues', 'counts'}

    用户 2026-10-12 要求：课表导入时，"这个老师不是教这个班这一科"要给提示，
    但**绝对不准修改任课教师映射** —— 本函数只读不写（不新增/不更新/不删除任何 link），
    页面上的处理建议也只是"去任课教师映射页改"。

    口径：
    - 只核对映射体系覆盖的 9 个学科（SUBJECTS：语数外物历化生政地）；
      体育/美术/信息技术/心理健康/自习这些不在映射体系里，跳过以免整屏噪音；
    - 课表里的教师先经 `Teacher.user_id` 换算到 users.id 再与映射比对（两库主键不同，
      直接比会把同一个人全报成不一致）；没绑账号的（如学校表里只写姓氏「陈」「邓」）
      退化为姓名比对；
    - issues 已按「年级+班级+科目+课表教师+映射教师+结论」聚合（带 count 与首个位置），
      避免上千条明细糊在预览页上。
    kind：
      no_link  — 该班该科还没有任课映射（先去映射页补，再核对）
      mismatch — 该班该科映射的任课教师不是课表里的这位
    """
    from app.models import User
    from app.models.academic import Teacher
    from app.models.grades import TeacherSubjectLink, SUBJECTS

    want = set(subjects or SUBJECTS)
    rows = [e for e in (entries or []) if (e.get('subject') or '') in want]
    if not rows:
        return {'checked': 0, 'issues': [], 'counts': {}}

    grades = {e.get('grade') for e in rows if e.get('grade')}
    links = (TeacherSubjectLink.query
             .filter(TeacherSubjectLink.grade.in_(grades)).all()
             if grades else [])
    link_map = {(l.grade, l.class_name, l.subject): l for l in links if l.active}

    # 关键：课表条目的 teacher_uid 是 academic.Teacher.id，映射的 user_id 是 users.id，
    # 两套库主键**不能直接比**（直接比会把"同一个人"全报成不一致）。经 Teacher.user_id
    # 换算到 users.id 后再比；没绑账号的老师（Teacher.user_id 为空）退化为姓名比对。
    tids = {e.get('teacher_uid') for e in rows if e.get('teacher_uid')}
    teacher_to_uid = {}
    if tids:
        teacher_to_uid = {t.id: t.user_id
                          for t in Teacher.query.filter(Teacher.id.in_(tids)).all()}

    uids = {l.user_id for l in link_map.values()}
    uids |= {u for u in teacher_to_uid.values() if u}
    names = {}
    if uids:
        names = {u.id: (u.real_name or u.username)
                 for u in User.query.filter(User.id.in_(uids)).all()}

    agg = {}
    checked = 0
    for e in rows:
        tname = (e.get('teacher_name') or '').strip()
        uid = e.get('teacher_uid')
        if not tname and not uid:
            continue
        checked += 1
        link = link_map.get((e.get('grade'), e.get('class_name'), e.get('subject')))
        expected_uid = link.user_id if link else None
        expected = names.get(expected_uid, '') if expected_uid else ''
        kind = None
        if link is None:
            kind = 'no_link'
        else:
            actual_uid = teacher_to_uid.get(uid) if uid else None
            if actual_uid:
                if actual_uid != expected_uid:
                    kind = 'mismatch'
            elif tname and expected and expected != tname:
                # 未绑账号的老师只能比姓名；学校表常写姓氏缩写（「陈」「邓」），
                # 能对上姓氏时只标「疑似同一人」，不硬判成"映射是别人"
                if len(tname) == 1 and (expected.startswith(tname) or expected.endswith(tname)):
                    kind = 'abbrev'
                else:
                    kind = 'mismatch'
        if not kind:
            continue
        akey = (e.get('grade'), e.get('class_name'), e.get('subject'), tname,
                expected, kind)
        item = agg.get(akey)
        if item is None:
            try:
                wd = '一二三四五六日'[int(e.get('weekday') or 1) - 1]
            except (TypeError, ValueError, IndexError):
                wd = '?'
            item = agg[akey] = {
                'grade': e.get('grade'), 'class_name': e.get('class_name'),
                'subject': e.get('subject'), 'teacher_name': tname,
                'expected': expected, 'kind': kind, 'count': 0,
                'first': f'周{wd}第{e.get("period_number")}节',
            }
        item['count'] += 1
    issues = sorted(agg.values(),
                    key=lambda x: (x['grade'] or '', x['class_name'] or '', x['subject'] or ''))
    counts = {}
    for it in issues:
        counts[it['kind']] = counts.get(it['kind'], 0) + 1
    return {'checked': checked, 'issues': issues, 'counts': counts}


def mapping_target(grade, class_name, subject):
    """取「该班该科」任课映射里的那位教师 → {'found', 'teacher_name', 'teacher_uid', 'user_id'}

    只读。`teacher_uid` 取 academic.Teacher.teacher_uid（课表条目用的就是这个编号，
    不是 users.id），该映射账号在教师名单里没对应记录时留空。
    """
    from app.models import User
    from app.models.academic import Teacher
    from app.models.grades import TeacherSubjectLink

    link = (TeacherSubjectLink.query
            .filter_by(grade=grade, class_name=class_name, subject=subject, active=True)
            .first())
    if not link:
        return {'found': False}
    user = db.session.get(User, link.user_id)
    teacher = Teacher.query.filter_by(user_id=link.user_id).first()
    return {'found': True, 'user_id': link.user_id,
            'teacher_name': ((user.real_name or user.username) if user else '') or '',
            'teacher_uid': (teacher.teacher_uid if teacher else '') or ''}


def apply_mapping_fix(data, grade, class_name, subject):
    """把导入草稿里**该班该科**的课表教师改成任课映射里的那位 → {'fixed': n, ...}

    只改本次导入的 `data`（预览/写入用的是它），**绝不改 TeacherSubjectLink**
    （用户 2026-10-12 明确要求"不准修改任课教师映射"）。
    教师名与教师编号一起换 —— 否则入库后教师课表/教师冲突检测对不上人。
    """
    target = mapping_target(grade, class_name, subject)
    if not target.get('found'):
        return {'fixed': 0, 'reason': 'no_link', 'target': target}
    fixed = 0
    for sh in (data.get('sheets') or []):
        if (sh.get('grade') or '') != grade or (sh.get('class_name') or '') != class_name:
            continue
        for e in (sh.get('entries') or []):
            if (e.get('subject') or '') != subject:
                continue
            if target.get('teacher_uid'):
                e['teacher_uid'] = target['teacher_uid']
            if target.get('teacher_name'):
                e['teacher_name'] = target['teacher_name']
            fixed += 1
    return {'fixed': fixed, 'reason': '', 'target': target}


def apply_all_mapping_fixes(data, issues):
    """按核对结果批量更正所有「疑似缩写 / 映射是别人」组 → 更正条数

    同样是**只改导入草稿**。`no_link`（该班该科无映射）不动 —— 那要去映射页设置，
    不能凭空造映射。
    """
    total = 0
    for it in (issues or []):
        if it.get('kind') not in ('abbrev', 'mismatch'):
            continue
        res = apply_mapping_fix(data, it.get('grade'), it.get('class_name'),
                                it.get('subject'))
        total += res.get('fixed', 0)
    return total


def batch_add_entries(schedule_id, entries_list, operator=None):
    """批量添加（跳过冲突项）→ {success: n, failed: n, errors: [...]}"""
    success = 0
    failed = 0
    errors = []
    for idx, item in enumerate(entries_list):
        ok, result = add_entry(
            schedule_id=schedule_id,
            grade=item.get('grade', ''),
            class_name=item.get('class_name', ''),
            weekday=item.get('weekday', 1),
            period_number=item.get('period_number', 1),
            subject=item.get('subject', ''),
            teacher_uid=item.get('teacher_uid'),
            teacher_name=item.get('teacher_name'),
            room=item.get('room'),
            week_range=item.get('week_range', '1-18'),
            note=item.get('note'),
            teaching_class=item.get('teaching_class'),
            operator=operator,
        )
        if ok:
            success += 1
        else:
            failed += 1
            errors.append({'row': idx + 1, 'message': result})
    return {'success': success, 'failed': failed, 'errors': errors}


# ═══════════════════════════════════════════════════════════════════════════════
# 版本查询
# ═══════════════════════════════════════════════════════════════════════════════

def get_entry_versions(entry_id, limit=50):
    """变更历史"""
    return ScheduleVersion.query.filter_by(entry_id=entry_id)\
        .order_by(ScheduleVersion.operated_at.desc()).limit(limit).all()


def get_schedule_versions(schedule_id, page=1, per_page=50):
    """学期变更历史分页"""
    return ScheduleVersion.query.filter_by(term_schedule_id=schedule_id)\
        .order_by(ScheduleVersion.operated_at.desc())\
        .paginate(page=page, per_page=per_page, error_out=False)


# ═══════════════════════════════════════════════════════════════════════════════
# Excel 导出
# ═══════════════════════════════════════════════════════════════════════════════

def export_schedule(schedule_id, view_type='class', grade=None, class_name=None,
                    teacher_uid=None, week=None, allowed_grades=None,
                    allowed_classes=None):
    """导出 Excel → BytesIO（week 为可选周次过滤，与视图口径一致）"""
    periods = get_periods(schedule_id)
    ts = db.session.get(TermSchedule, schedule_id)
    title = ts.name if ts else '课表'

    wb = Workbook()
    wb.remove(wb.active)

    if view_type == 'class' and grade and class_name:
        _export_class_sheet(wb, schedule_id, grade, class_name, periods, title,
                            week, allowed_grades, allowed_classes)
    elif view_type == 'grade' and grade:
        data = get_grade_view(schedule_id, grade, week=week,
                              allowed_grades=allowed_grades,
                              allowed_classes=allowed_classes)
        for cn in data['classes']:
            _export_class_sheet(wb, schedule_id, grade, cn, periods,
                                f'{title} {grade}{cn}', week,
                                allowed_grades, allowed_classes)
    elif view_type == 'teacher' and teacher_uid:
        _export_teacher_sheet(wb, schedule_id, teacher_uid, periods, title, week,
                              allowed_grades, allowed_classes)
    elif view_type == 'all':
        # 整校：一个班一张 sheet（2026-10-10 由原 `else`/master 分支改名而来）。
        # 大课表页已下线，但"整校分班导出"这个能力仍被「学期管理」卡片的导出按钮使用，
        # 所以只把名字与页面解耦，不删功能。
        grade_classes = get_grade_class_list(
            schedule_id, allowed_grades=allowed_grades,
            allowed_classes=allowed_classes)
        for g, cls_list in grade_classes.items():
            for cn in cls_list:
                _export_class_sheet(wb, schedule_id, g, cn, periods, f'{g}{cn}',
                                    week, allowed_grades, allowed_classes)
    else:
        # 'overview' 及任何未知 view_type（含旧分享短链里的 'master'）→ 全校总课表一张表
        _export_overview_sheet(wb, schedule_id, periods, title, week,
                               allowed_grades, allowed_classes)

    if not wb.sheetnames:
        ws = wb.create_sheet('空')
        ws.append(['无数据'])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _export_overview_sheet(wb, schedule_id, periods, title, week=None,
                           allowed_grades=None, allowed_classes=None):
    """全校总课表工作表（2026-09-25）：按年级分块、班级并列，一张表看全校。

    行＝星期 × 节次（星期列纵向合并），列＝年级内各班，格内「学科 + 教师」，
    与页面 /schedule/<sid>/overview 的形态一致，方便直接打印或二次编辑。
    """
    view = get_overview_view(schedule_id, week=week,
                             allowed_grades=allowed_grades,
                             allowed_classes=allowed_classes)
    ws = wb.create_sheet('全校总课表'[:31])

    hf = Font(bold=True, color='FFFFFF', size=10)
    hfl = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    gf = Font(bold=True, size=12)
    tb = Border(left=Side('thin'), right=Side('thin'),
                top=Side('thin'), bottom=Side('thin'))
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    def put(row, col, value, font=None, fill=None, align=center):
        c = ws.cell(row=row, column=col, value=value)
        c.border = tb
        c.alignment = align
        if font:
            c.font = font
        if fill:
            c.fill = fill
        return c

    ri = 1
    put(ri, 1, title, font=gf, align=Alignment(horizontal='left', vertical='center'))
    ri += 2
    for b in view['blocks']:
        classes = b['classes']
        put(ri, 1, f"{b['grade']}（{len(classes)} 个班 · {b['total']} 节/周）", font=gf,
            align=Alignment(horizontal='left', vertical='center'))
        ri += 1
        put(ri, 1, '星期', font=hf, fill=hfl)
        put(ri, 2, '节次', font=hf, fill=hfl)
        for ci, cn in enumerate(classes, 3):
            put(ri, ci, f'{b["grade"]}{cn}', font=hf, fill=hfl)
        ri += 1
        for wd in range(1, view['max_weekday'] + 1):
            first = ri
            for p in periods:
                time_label = (f'{p.period_name}\n{p.start_time}-{p.end_time}'
                              if p.start_time else p.period_name)
                put(ri, 1, WEEKDAY_NAMES.get(wd, ''))
                put(ri, 2, time_label, font=Font(bold=True, size=9))
                for ci, cn in enumerate(classes, 3):
                    items = (b['grids'][cn].get(p.period_number) or {}).get(wd) or []
                    put(ri, ci, _cell_lines(items))
                ri += 1
            if ri - 1 > first:   # 星期列纵向合并
                ws.merge_cells(start_row=first, start_column=1, end_row=ri - 1, end_column=1)
            ws.cell(row=first, column=1).alignment = center
        ri += 1   # 年级之间空一行

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 15
    max_cols = 3 + max((len(b['classes']) for b in view['blocks']), default=1) - 1
    for i in range(3, max_cols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 13


def _cell_lines(items, with_class=False):
    """把格子里的多条条目整理成 Excel 单元格文本（单双周交替课分行显示）

    末尾做公式注入转义：科目/教师名由用户维护，理论上可被写成 =cmd|... 之类的文本。
    """
    lines = []
    for e in items or []:
        who = (f"{e.get('grade','')}{e.get('class_name','')}" if with_class
               else (e.get('teacher_name') or ''))
        head = e.get('subject', '')
        badge = e.get('week_badge') or ''
        if badge:
            head += f'（{badge}）'
        lines.append(f'{head}\n{who}' if who else head)
    from app.utils.export_helpers import xl_safe
    return xl_safe('\n'.join(lines))


def _export_class_sheet(wb, schedule_id, grade, class_name, periods, sheet_title,
                        week=None, allowed_grades=None, allowed_classes=None):
    """导出单个班级课表工作表"""
    ws = wb.create_sheet(sheet_title[:31])  # Excel sheet名最长31字符
    view = get_class_view(schedule_id, grade, class_name, week=week,
                          allowed_grades=allowed_grades,
                          allowed_classes=allowed_classes)
    grid = view['grid']

    # 表头样式
    hf = Font(bold=True, color='FFFFFF', size=11)
    hfl = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    tb = Border(left=Side('thin'), right=Side('thin'),
                top=Side('thin'), bottom=Side('thin'))
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # 表头行
    headers = ['节次'] + [WEEKDAY_NAMES.get(wd, f'周{wd}') for wd in range(1, 8)]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = hf
        c.fill = hfl
        c.alignment = center
        c.border = tb

    # 数据行
    for ri, p in enumerate(periods, 2):
        time_label = f'{p.period_name}\n{p.start_time}-{p.end_time}' if p.start_time else p.period_name
        c = ws.cell(row=ri, column=1, value=time_label)
        c.alignment = center
        c.border = tb
        c.font = Font(bold=True, size=10)
        for wd in range(1, 8):
            items = grid.get(p.period_number, {}).get(wd) or []
            c = ws.cell(row=ri, column=wd + 1, value=_cell_lines(items))
            c.alignment = center
            c.border = tb

    ws.column_dimensions['A'].width = 16
    for i in range(2, 9):
        ws.column_dimensions[get_column_letter(i)].width = 14


def _export_teacher_sheet(wb, schedule_id, teacher_uid, periods, sheet_title,
                          week=None, allowed_grades=None, allowed_classes=None):
    """导出教师个人课表工作表"""
    ws = wb.create_sheet(sheet_title[:31])
    view = get_teacher_view(schedule_id, teacher_uid, week=week,
                            allowed_grades=allowed_grades,
                            allowed_classes=allowed_classes)
    grid = view['grid']

    hf = Font(bold=True, color='FFFFFF', size=11)
    hfl = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    tb = Border(left=Side('thin'), right=Side('thin'),
                top=Side('thin'), bottom=Side('thin'))
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    headers = ['节次'] + [WEEKDAY_NAMES.get(wd, '') for wd in range(1, 8)]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = hf
        c.fill = hfl
        c.alignment = center
        c.border = tb

    for ri, p in enumerate(periods, 2):
        time_label = f'{p.period_name}\n{p.start_time}-{p.end_time}' if p.start_time else p.period_name
        c = ws.cell(row=ri, column=1, value=time_label)
        c.alignment = center
        c.border = tb
        c.font = Font(bold=True, size=10)
        for wd in range(1, 8):
            items = grid.get(p.period_number, {}).get(wd) or []
            c = ws.cell(row=ri, column=wd + 1,
                        value=_cell_lines(items, with_class=True))
            c.alignment = center
            c.border = tb

    ws.column_dimensions['A'].width = 16
    for i in range(2, 9):
        ws.column_dimensions[get_column_letter(i)].width = 14


# ═══════════════════════════════════════════════════════════════════════════════
# 辅助查询（供路由/API 使用）
# ═══════════════════════════════════════════════════════════════════════════════

def get_grade_class_list(schedule_id=None, allowed_grades=None,
                         allowed_classes=None):
    """获取年级+班级列表（从课表条目中提取 distinct，供级联下拉）"""
    if schedule_id:
        q = _base_entry_query(schedule_id)
    else:
        ts = get_active_schedule()
        if not ts:
            return {}
        q = _base_entry_query(ts.id)
    if allowed_grades is not None:
        q = q.filter(ScheduleEntry.grade.in_(allowed_grades))
    q = _apply_class_scope(q, allowed_grades, allowed_classes)
    rows = q.with_entities(ScheduleEntry.grade, ScheduleEntry.class_name).distinct().all()
    result = {}
    for g, cn in rows:
        result.setdefault(g, set()).add(cn)
    return {g: sorted(cs) for g, cs in sorted(result.items())}


def grade_class_map(schedule_ids=None):
    """一次查询返回 {schedule_id: {grade: [class_name, ...]}}。

    供学期管理卡片（年级课表入口）等场景使用，避免逐学期查一次。
    """
    q = ScheduleEntry.query.filter(ScheduleEntry.is_deleted.is_(False))
    if schedule_ids is not None:
        ids = list(schedule_ids)
        if not ids:
            return {}
        q = q.filter(ScheduleEntry.term_schedule_id.in_(ids))
    rows = (q.with_entities(ScheduleEntry.term_schedule_id,
                            ScheduleEntry.grade, ScheduleEntry.class_name)
            .distinct().all())
    out = {}
    for sid, g, cn in rows:
        out.setdefault(sid, {}).setdefault(g, set()).add(cn)
    return {sid: {g: sorted(cs) for g, cs in sorted(gm.items())}
            for sid, gm in out.items()}


def get_all_teachers():
    """获取教师列表（从 academic.db Teacher 表，跨库逻辑键查询）"""
    from app.models.academic import Teacher
    teachers = Teacher.query.filter_by(status='active')\
        .order_by(Teacher.teacher_uid).all()
    return [{'uid': t.teacher_uid, 'name': t.name, 'subject': t.subject or ''}
            for t in teachers]


# ═══════════════════════════════════════════════════════════════════════════════
# 组课表查询（2026-10-12 新增）：教师课表 / 备课组课表 / 教研组课表
#
# 成员口径：**按本学期课表条目聚合** —— 教这个年级/学科的老师就是该组成员。
# 备课组（年级 × 学科）与教研组（学科，跨年级）都不另建组织关系表，与课表天然一致，
# 不会出现「组里有人但课表查不到」的口径分叉；组长标注取自 SubjectLeader
# （无法从课表推导，故落库在 academic.db），只做标签展示，不参与成员筛选。
# ═══════════════════════════════════════════════════════════════════════════════

GROUP_SCOPES = {'lesson': '备课组', 'subject': '教研组'}


def _subject_key(subject):
    """学科排序键（SUBJECT_ORDER 优先，名单外的学科排最后）"""
    from app.models.academic import SUBJECT_ORDER
    try:
        return (SUBJECT_ORDER.index(subject), subject or '')
    except ValueError:
        return (len(SUBJECT_ORDER), subject or '')


def list_grade_options(schedule_id):
    """该学期有排课的年级（低年级在前）"""
    from app.modules.academic.services.grade_utils import grade_sort_key
    q = _base_entry_query(schedule_id).with_entities(ScheduleEntry.grade).distinct()
    grades = {r[0] for r in q.all() if r[0]}
    return sorted(grades, key=grade_sort_key)


def list_subject_options(schedule_id, grade=None):
    """该学期出现过的学科（按 SUBJECT_ORDER 排序；给 grade 时只列该年级的）"""
    q = _base_entry_query(schedule_id)
    if grade:
        q = q.filter_by(grade=grade)
    subjects = {r[0] for r in q.with_entities(ScheduleEntry.subject).distinct().all() if r[0]}
    return sorted(subjects, key=_subject_key)


def leader_roles(schedule_id, subject, grade=None):
    """组长标注 {teacher_uid: [角色标签]}（来自 SubjectLeader 落库数据）。

    - grade=None（教研组页）：全校级登记（grade=''）标「教研组长」，各年级登记标「高一备课组长」；
    - 给了 grade（备课组页）：只取该年级的备课组长。
    学期口径：登记为空（''＝整学年）或与本学期一致才算数。
    """
    from app.models.academic import SubjectLeader
    ts = db.session.get(TermSchedule, schedule_id)
    roles = {}
    if not ts or not subject:
        return roles
    for ld in SubjectLeader.query.filter_by(subject=subject,
                                            school_year=ts.school_year).all():
        if (ld.term or '') not in ('', ts.term or ''):
            continue
        if grade is not None and (ld.grade or '') != grade:
            continue
        if not ld.leader_uid:
            continue
        if (ld.grade or ''):
            label = f'{ld.grade}备课组长'
        else:
            label = '教研组长' if grade is None else '备课组长'
        labels = roles.setdefault(ld.leader_uid, [])
        if label not in labels:
            labels.append(label)
    return roles


def _group_members(buckets, roles=None):
    """把 {teacher_uid: [ScheduleEntry, ...]} 聚合成成员列表。

    展示名 label：组内重名时补教师编号（行标不能撞车）；排序＝组长优先、课时多的在前。
    """
    roles = roles or {}
    members = []
    for uid, items in buckets.items():
        counts = {}
        for e in items:
            nm = (e.teacher_name or '').strip() or uid
            counts[nm] = counts.get(nm, 0) + 1
        name = sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))[0][0]
        subs = {}
        grades, classes = set(), set()
        for e in items:
            if e.subject:
                subs[e.subject] = subs.get(e.subject, 0) + 1
            if e.grade:
                grades.add(e.grade)
            classes.add(f'{e.grade or ""}{e.class_name or ""}')
        sub_list = [s for s, _c in sorted(subs.items(), key=lambda kv: (-kv[1],
                                                                       _subject_key(kv[0])))]
        members.append({
            'uid': uid, 'name': name, 'total': len(items),
            'subjects': sub_list, 'primary_subject': sub_list[0] if sub_list else '',
            'grades': sorted(grades), 'classes': sorted(classes),
            'roles': roles.get(uid, []), 'is_leader': bool(roles.get(uid)),
        })
    # 重名消歧（同一组内出现同名教师时才补编号，正常情况行标就是姓名）
    name_count = {}
    for m in members:
        name_count[m['name']] = name_count.get(m['name'], 0) + 1
    for m in members:
        m['label'] = (f"{m['name']}({m['uid']})"
                      if name_count.get(m['name'], 0) > 1 else m['name'])
    members.sort(key=lambda m: (0 if m['is_leader'] else 1, -m['total'], m['name']))
    return members


def get_group_view(schedule_id, scope='lesson', grade=None, subject=None,
                   day=None, week=None):
    """备课组 / 教研组课表视图。

    - scope='lesson'： 年级 × 学科（备课组），两者都要给；
    - scope='subject'：学科（教研组，跨年级）；
    - day：渲染星期（1-7，默认今天）；week：周次过滤（单双周课）。
    返回 dict：members / grids（键＝成员 label）/ metas / periods / period_groups /
    free_periods（该天全组都没课的节次，排除课间空档）/ group_name / 选项列表。
    """
    periods = get_periods(schedule_id)
    period_dicts = [p.to_dict() for p in periods]
    period_groups = _period_groups(period_dicts)

    # 参数不完整（备课组缺年级/学科）时按**空组**处理 —— 页面给"该组合本学期没排课"的
    # 空态，绝不能退化成"整个年级全部学科"这种口径不符的结果。
    if scope == 'lesson':
        q = (_base_entry_query(schedule_id).filter_by(grade=grade, subject=subject)
             if (grade and subject) else None)
    else:
        q = _base_entry_query(schedule_id).filter_by(subject=subject) if subject else None
    entries = _filter_by_week(q.all(), week) if q is not None else []

    buckets = {}
    for e in entries:
        uid = (e.teacher_uid or '').strip()
        if uid:
            buckets.setdefault(uid, []).append(e)
    roles = leader_roles(schedule_id, subject,
                         grade if scope == 'lesson' else None)
    members = _group_members(buckets, roles)
    grids = {m['label']: _build_grid(buckets[m['uid']], periods) for m in members}
    metas = {m['label']: {'class_type': ' · '.join(m['roles'])}
             for m in members if m['roles']}

    day = int(day) if str(day or '').strip().isdigit() else date.today().isoweekday()
    day = max(1, min(7, day))
    free = []
    for p in period_dicts:
        if p.get('period_type') == 'break':
            continue
        busy = any((grids.get(m['label'], {}).get(p['period_number']) or {}).get(day)
                   for m in members)
        if not busy:
            free.append({'period_number': p['period_number'],
                         'period_name': p.get('period_name') or f"第{p['period_number']}节",
                         'start_time': p.get('start_time') or ''})

    if scope == 'lesson':
        group_name = f'{grade or ""}{subject or ""}备课组'
    else:
        group_name = f'{subject or ""}教研组'
    return {
        'scope': scope, 'scope_label': GROUP_SCOPES.get(scope, '组'),
        'grade': grade, 'subject': subject, 'day': day, 'week': week,
        'group_name': group_name,
        'periods': period_dicts, 'period_groups': period_groups,
        'members': members, 'grids': grids, 'metas': metas,
        'member_count': len(members),
        'course_total': len(entries),
        'free_periods': free,
        'grades': list_grade_options(schedule_id),
        'subjects': list_subject_options(schedule_id,
                                        grade if scope == 'lesson' else None),
    }


def list_teacher_digest(schedule_id, week=None):
    """本学期有排课的教师（按学科分组），供「教师课表」选择态展示。

    返回 [{'subject': 学科, 'teachers': [{'uid','name','total','grades','classes'}]}]；
    教师的归属学科取**课时最多的那门**（个人课表里能看到他的全部课程）。
    """
    entries = _filter_by_week(_base_entry_query(schedule_id).all(), week)
    buckets = {}
    for e in entries:
        uid = (e.teacher_uid or '').strip()
        if uid:
            buckets.setdefault(uid, []).append(e)
    members = _group_members(buckets)
    grouped = {}
    for m in members:
        grouped.setdefault(m['primary_subject'] or '其他', []).append({
            'uid': m['uid'], 'name': m['name'], 'total': m['total'],
            'grades': m['grades'], 'classes': m['classes'],
        })
    return [{'subject': subj, 'teachers': sorted(rows, key=lambda r: (-r['total'], r['name']))}
            for subj, rows in sorted(grouped.items(), key=lambda kv: _subject_key(kv[0]))]
