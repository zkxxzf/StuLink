# StuLink v1.18.9.1 2026-10-10
# Copyright (c) 2026 zkxxzf. Apache License 2.0
"""晚自习值班服务（2026-09-26，高中教务）。

高中晚自习不排学科课，而是排教师**值班看班/巡楼**。这里负责：
- 取晚自习节次（PeriodDef 里 period_type='evening'）；
- 手工指定 / 清空单个班次；
- 今日值班（给门卫/巡楼用）与教师值班次数统计；
- 导出 Excel（贴墙用）。

值班只读课表，不写课表条目 —— 值班不是课，不该出现在班级课表网格里。

2026-10-10：系统不再提供「一键均衡排班」与「同天只值一节 / 周次数上限」检查 ——
值班由教务自行安排，系统只做登记、统计与导出；旧地址 /night-duty/auto 已下线。
"""
from collections import defaultdict
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.extensions import db
from app.models.timetable import (WEEKDAY_NAMES, NightDuty, PeriodDef,
                                  ScheduleEntry)
from app.modules.academic.services.schedule_service import _base_entry_query
from app.utils.export_helpers import xl_safe

WEEKDAYS = (1, 2, 3, 4, 5)


def evening_period_numbers(schedule_id):
    """晚自习节次号（升序）。没有配置 evening 节次时返回空表，页面会给出引导。"""
    rows = (PeriodDef.query.filter_by(term_schedule_id=schedule_id)
            .order_by(PeriodDef.period_number).all())
    return [p.period_number for p in rows if (p.period_type or '') == 'evening']


def grades_of(schedule_id):
    """该学期有课的年级（升序）。"""
    rows = (db.session.query(ScheduleEntry.grade)
            .filter_by(term_schedule_id=schedule_id, is_deleted=False)
            .distinct().all())
    return sorted({g[0] for g in rows if g[0]})


def teacher_pool():
    """可值班教师池：在职教师（含行政，学科可能为空）。"""
    try:
        from app.models.academic import Teacher
        rows = Teacher.query.filter_by(status='active').order_by(Teacher.id).all()
    except Exception:  # noqa: BLE001  教师表不可用时退化为空池
        return []
    return [{'uid': t.teacher_uid, 'name': t.name, 'subject': t.subject or ''}
            for t in rows if t.teacher_uid]


def _roster_from(rows, grades, periods):
    """（内部）由已取出的值班行构造值班表数据：年级分块 + {weekday: {period: duty}}。"""
    grid = {g: {wd: {} for wd in WEEKDAYS} for g in grades}
    for d in rows:
        if d.grade in grid:
            grid[d.grade].setdefault(d.weekday, {})[d.period_number] = d.to_dict()
    blocks = [{'grade': g, 'grid': grid[g]} for g in grades]
    return {'blocks': blocks, 'periods': periods, 'grades': grades,
            'total': len(rows)}


def _today_from(rows, weekday):
    """（内部）今日值班清单（按年级、节次升序）。"""
    picked = [d for d in rows if d.weekday == weekday]
    picked.sort(key=lambda d: ((d.grade or ''), (d.period_number or 0)))
    return [d.to_dict() for d in picked]


def _stats_from(rows):
    """（内部）教师值班次数统计（按次数降序）。"""
    stat = {}
    for d in rows:
        if not d.teacher_uid:
            continue
        item = stat.setdefault(d.teacher_uid, {
            'uid': d.teacher_uid, 'name': d.teacher_name or d.teacher_uid,
            'count': 0, 'days': []})
        item['count'] += 1
        wd = WEEKDAY_NAMES.get(d.weekday, '')
        if wd not in item['days']:
            item['days'].append(wd)
    return sorted(stat.values(), key=lambda x: (-x['count'], x['name']))


def get_roster(schedule_id, grades=None):
    """值班表数据：按年级分块，每块 {weekday: {period: duty}}。"""
    grades = grades or grades_of(schedule_id)
    periods = evening_period_numbers(schedule_id)
    rows = NightDuty.query.filter_by(term_schedule_id=schedule_id).all()
    return _roster_from(rows, grades, periods)


def set_duty(schedule_id, grade, weekday, period_number, teacher_uid=None,
             teacher_name=None, note=None, operator=None):
    """手工指定某个班次的值班教师；teacher_uid 为空＝清空该班次。"""
    slot = NightDuty.query.filter_by(term_schedule_id=schedule_id, grade=grade,
                                     weekday=weekday,
                                     period_number=period_number).first()
    if not teacher_uid:
        if slot:
            db.session.delete(slot)
            db.session.commit()
        return True, '已清空'
    if slot is None:
        slot = NightDuty(term_schedule_id=schedule_id, grade=grade, weekday=weekday,
                         period_number=period_number)
        db.session.add(slot)
    slot.teacher_uid = teacher_uid
    slot.teacher_name = teacher_name
    slot.note = note
    db.session.commit()
    if operator is not None:
        try:
            from app.utils.helpers import log_operation
            log_operation(operator, '更新', '晚自习值班', schedule_id,
                          f'{grade} 周{weekday} 第{period_number}节 {teacher_name}',
                          module='academic')
        except Exception:  # noqa: BLE001
            pass
    return True, '已保存'


def get_page_data(schedule_id, grades=None, weekday=None):
    """值班页一次取数 → (roster, today, stats)。

    2026-10-10 优化：值班页原先依次调用 get_roster / today_duties / teacher_stats，
    同一张 night_duties 表被全查 3 遍；改为一次全查 + 内存分组，页面输出完全不变。
    """
    grades = grades or grades_of(schedule_id)
    periods = evening_period_numbers(schedule_id)
    rows = NightDuty.query.filter_by(term_schedule_id=schedule_id).all()
    wd = weekday or date.today().isoweekday()
    return (_roster_from(rows, grades, periods),
            _today_from(rows, wd), _stats_from(rows))


def today_duties(schedule_id, weekday=None):
    """今日值班清单（默认取系统当天星期；周末返回空表）。"""
    wd = weekday or date.today().isoweekday()
    rows = NightDuty.query.filter_by(term_schedule_id=schedule_id).all()
    return _today_from(rows, wd)


def teacher_stats(schedule_id):
    """教师值班次数统计（按次数降序）。"""
    rows = NightDuty.query.filter_by(term_schedule_id=schedule_id).all()
    return _stats_from(rows)


def export_workbook(schedule_id, grades=None):
    """导出值班表 Excel（每行一个班次，按年级分块，贴墙/发年级组用）。"""
    data = get_roster(schedule_id, grades=grades)
    wb = Workbook()
    ws = wb.active
    ws.title = '晚自习值班表'
    thin = Side(style='thin', color='BBBBBB')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    title = ws.cell(row=1, column=1, value='晚自习值班表')
    title.font = Font(bold=True, size=14)
    periods = data['periods']
    head = ws.cell(row=2, column=1, value='年级')
    head.font = Font(bold=True)
    for i, wd in enumerate(WEEKDAYS):
        ws.cell(row=2, column=2 + i, value=WEEKDAY_NAMES.get(wd, ''))
        ws.cell(row=2, column=2 + i).font = Font(bold=True)
    fill = PatternFill('solid', fgColor='EEF2FF')
    for c in ws[2]:
        c.fill = fill
        c.alignment = Alignment(horizontal='center')
        c.border = border
    ri = 3
    for b in data['blocks']:
        for pn in periods:
            ws.cell(row=ri, column=1,
                    value=xl_safe(f"{b['grade']} 第{pn}节")).border = border
            for i, wd in enumerate(WEEKDAYS):
                cell = ws.cell(row=ri, column=2 + i,
                               value=xl_safe((b['grid'].get(wd, {}) or {})
                                             .get(pn, {}).get('teacher_name') or ''))
                cell.alignment = Alignment(horizontal='center')
                cell.border = border
            ri += 1
    ws.column_dimensions['A'].width = 18
    for i in range(len(WEEKDAYS)):
        ws.column_dimensions[chr(ord('B') + i)].width = 14
    return wb
