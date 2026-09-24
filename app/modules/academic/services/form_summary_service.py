# StuLink v1.18.2.1 2026-09-24
# 表单收集「汇总服务层」：应交名单 / 提交统计 / 横向汇总矩阵 / Excel 导出
#                        / 材料清单与打包下载 / 未交催交
# Copyright (c) 2026 zkxxzf. Apache License 2.0
"""表单收集之后的能力（只读汇总，不改动 form_service 已有函数）。

跨库说明：
- FormTemplate/FormQuestion/FormSubmission/FormAnswer/Teacher 绑定 academic.db；
- Student 无 bind_key，落默认主库 system.db；Notification 绑定 system.db。
  Flask-SQLAlchemy 会按 bind 自动路由，故分别用各自的 .query 取数后在 Python 侧
  按 uid（学生=student_number / 教师=teacher_uid）匹配，绝不跨库 JOIN。

应交名单口径：
- target_type in ('students','all') 或 grade*  → 学生名单（system.db Student）
- target_type == 'teachers'                    → 教师名单（academic.db Teacher）
  说明：'all'（全体）在收集场景按「全体学生」计算应交名单（班级拆分需要学生维度），
  这一口径在报告中已注明。
"""
import io
import json
import os
import re
import csv
import shutil
import tempfile
import zipfile
from datetime import datetime

from flask import current_app, url_for
from sqlalchemy import or_, func

from app.extensions import db
from app.models.academic import (
    FormTemplate, FormQuestion, FormSubmission, FormAnswer, Teacher,
)
from app.models.student import Student
from app.utils.export_helpers import xl_row, xl_safe


SUBMISSION_STATUS_TEXT = {'submitted': '待审核', 'approved': '已通过', 'rejected': '已驳回'}
QUESTION_TYPE_TEXT = {
    'text': '单行文本', 'textarea': '多行文本', 'single_choice': '单选',
    'multi_choice': '多选', 'file': '文件上传', 'date': '日期', 'number': '数字',
}

# 打包总大小超过该阈值时改为落盘（避免 BytesIO 撑爆内存）
PACKAGE_DISK_THRESHOLD = 500 * 1024 * 1024  # 500MB
_ILLEGAL_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


# ── 通用工具 ──────────────────────────────────────────────────

def _parse_target_scope(raw):
    """解析 FormTemplate.target_scope，兼容两种写法：
    - 旧：JSON 数组 ["2026级"]            → (grades, [])
    - 新：JSON 对象 {"grades":[..],"classes":[..]} → (grades, classes)
    返回 (grades:list[str], classes:list[str])。
    """
    if not raw:
        return [], []
    try:
        data = json.loads(raw)
    except (json.JSONDecodeError, TypeError):
        return [], []
    if isinstance(data, list):
        return [str(x) for x in data], []
    if isinstance(data, dict):
        grades = [str(x) for x in (data.get('grades') or [])]
        classes = [str(x) for x in (data.get('classes') or [])]
        return grades, classes
    return [], []


def _load_template(form_id):
    tpl = db.session.get(FormTemplate, form_id)
    if not tpl:
        raise ValueError('表单不存在')
    return tpl


def _load_questions(form_id):
    """按 sort_order 返回题目列表"""
    qs = FormQuestion.query.filter_by(template_id=form_id).all()
    return sorted(qs, key=lambda q: (q.sort_order if q.sort_order is not None else 0, q.id))


def _question_options(q):
    """解析选择题选项（list[str]）"""
    if not q.options_json:
        return []
    try:
        opts = json.loads(q.options_json)
        if isinstance(opts, list):
            return [str(o) for o in opts]
    except (json.JSONDecodeError, TypeError):
        pass
    return []


def _sanitize(name, fallback='file'):
    """清洗文件/目录名中的非法字符（Windows + zip 安全）"""
    name = (name or '').strip()
    name = _ILLEGAL_CHARS.sub('_', name)
    name = name.replace('\n', ' ').replace('\r', ' ')
    name = name.strip('. ')  # Windows 不允许结尾的点/空格
    return name or fallback


def _human_size(n):
    n = n or 0
    for unit in ('B', 'KB', 'MB', 'GB'):
        if n < 1024:
            return f'{n:.0f}{unit}' if unit == 'B' else f'{n:.1f}{unit}'
        n /= 1024
    return f'{n:.1f}TB'


def _static_abspath(rel_path):
    """把 FormAnswer.file_path（相对 static）转为磁盘绝对路径"""
    return os.path.join(current_app.static_folder, rel_path.replace('/', os.sep))


def _file_url(rel_path):
    """附件下载 URL —— 必须走带鉴权的 academic.form_file 路由。

    历史上这里返回 /static/uploads/forms/... 的静态直链，任何拿到链接的人
    （含未登录用户）都能下载学生上传的材料。详情见
    app/modules/academic/routes/forms.py 的 form_file 路由。
    """
    parts = (rel_path or '').split('/')
    form_id = None
    if len(parts) >= 4 and parts[0] == 'uploads' and parts[1] == 'forms':
        if parts[2].isdigit():
            form_id = int(parts[2])
    if form_id:
        try:
            return url_for('academic.form_file', form_id=form_id, rel_path=rel_path)
        except Exception:
            pass
    return '#'


def _query_submissions(form_id, status=None, grade=None, class_name=None,
                       keyword=None, include_rejected=True):
    """构造提交查询（可按审核状态/年级/班级/关键词筛选）"""
    q = FormSubmission.query.filter_by(template_id=form_id)
    if status and status in SUBMISSION_STATUS_TEXT:
        q = q.filter(FormSubmission.status == status)
    elif not include_rejected:
        q = q.filter(FormSubmission.status != 'rejected')
    if grade:
        q = q.filter(FormSubmission.submitter_grade == grade)
    if class_name:
        q = q.filter(FormSubmission.submitter_class == class_name)
    if keyword:
        kw = f'%{keyword}%'
        q = q.filter(or_(FormSubmission.submitter_name.like(kw),
                         FormSubmission.submitter_uid.like(kw)))
    return q


def _paginate(items, page, per_page):
    """内存分页 → (page_items, pagination_dict)"""
    total = len(items)
    per_page = max(1, int(per_page or 50))
    page = max(1, int(page or 1))
    pages = (total + per_page - 1) // per_page if total else 0
    start = (page - 1) * per_page
    page_items = items[start:start + per_page]
    return page_items, {
        'page': page, 'per_page': per_page, 'total': total, 'pages': pages,
        'has_prev': page > 1, 'has_next': page < pages,
        'prev_num': max(1, page - 1), 'next_num': min(pages or 1, page + 1),
    }


# ── A. 应交名单与提交情况 ─────────────────────────────────────

def get_expected_submitters(form_id, scope=None):
    """计算「应交名单」。

    scope: 可选 dict(grade=?, class_name=?) 用于收窄范围。
    返回 list[dict(uid, name, grade, class_name, role)]，按 年级→班级→uid 排序。
    """
    tpl = _load_template(form_id)
    tt = (tpl.target_type or 'all').strip()
    scope = scope or {}
    s_grade = scope.get('grade')
    s_class = scope.get('class_name')

    result = []
    if tt == 'teachers':
        q = Teacher.query.filter(Teacher.status == 'active')
        for t in q.all():
            result.append({'uid': t.teacher_uid, 'name': t.name,
                           'grade': '', 'class_name': '', 'role': 'teacher'})
    else:
        # students / all / grade*  → 学生名单
        q = Student.query
        if tt.startswith('grade'):
            grades, classes = _parse_target_scope(tpl.target_scope)
            if grades:
                q = q.filter(Student.grade.in_(grades))
            if classes:
                q = q.filter(Student.class_name.in_(classes))
        if s_grade:
            q = q.filter(Student.grade == s_grade)
        if s_class:
            q = q.filter(Student.class_name == s_class)
        rows = q.order_by(Student.grade, Student.class_name, Student.student_number).all()
        for s in rows:
            result.append({'uid': s.student_number, 'name': s.name,
                           'grade': s.grade or '', 'class_name': s.class_name or '',
                           'role': 'student'})
    return result


def _submitted_uid_set(form_id):
    """该表单所有提交的 submitter_uid 集合（去重）"""
    rows = (FormSubmission.query.filter_by(template_id=form_id)
            .with_entities(FormSubmission.submitter_uid).distinct().all())
    return {r[0] for r in rows if r[0]}


def get_submission_stats(form_id):
    """汇总统计（expected == submitted + not_submitted 自洽）"""
    expected_list = get_expected_submitters(form_id)
    expected_uids = {e['uid'] for e in expected_list if e['uid']}
    submitted_uids = _submitted_uid_set(form_id)

    expected = len(expected_uids)
    submitted = len(expected_uids & submitted_uids)
    not_submitted = expected - submitted
    rate = round(submitted * 100.0 / expected, 1) if expected else 0.0

    subs = FormSubmission.query.filter_by(template_id=form_id).all()
    sub_ids = [s.id for s in subs]
    approved = sum(1 for s in subs if s.status == 'approved')
    pending = sum(1 for s in subs if s.status == 'submitted')
    rejected = sum(1 for s in subs if s.status == 'rejected')

    total_answers = 0
    total_files = 0
    total_file_size = 0
    if sub_ids:
        answers = (FormAnswer.query.filter(FormAnswer.submission_id.in_(sub_ids)).all())
        total_answers = len(answers)
        for a in answers:
            if a.file_path:
                total_files += 1
                total_file_size += (a.file_size or 0)

    return {
        'expected': expected, 'submitted': submitted, 'not_submitted': not_submitted,
        'rate': rate, 'approved': approved, 'pending': pending, 'rejected': rejected,
        'total_submissions': len(subs),
        'total_answers': total_answers, 'total_files': total_files,
        'total_file_size': total_file_size,
        'total_file_size_text': _human_size(total_file_size),
    }


def get_class_breakdown(form_id):
    """按班级拆分提交情况（按 年级→班级 排序）。

    返回 list[dict(grade, class_name, expected, submitted, not_submitted, rate)]。
    """
    expected_list = get_expected_submitters(form_id)
    submitted_uids = _submitted_uid_set(form_id)

    groups = {}
    for e in expected_list:
        key = (e.get('grade') or '—', e.get('class_name') or '—')
        g = groups.setdefault(key, {'grade': key[0], 'class_name': key[1],
                                    'expected': 0, 'submitted': 0})
        g['expected'] += 1
        if e.get('uid') in submitted_uids:
            g['submitted'] += 1

    rows = []
    for key in sorted(groups.keys(), key=lambda k: (k[0], k[1])):
        g = groups[key]
        g['not_submitted'] = g['expected'] - g['submitted']
        g['rate'] = round(g['submitted'] * 100.0 / g['expected'], 1) if g['expected'] else 0.0
        rows.append(g)
    return rows


def class_breakdown_totals(rows):
    """由班级行汇总各年级小计与全校合计"""
    grade_totals = {}
    for r in rows:
        gt = grade_totals.setdefault(r['grade'], {'grade': r['grade'], 'expected': 0,
                                                 'submitted': 0, 'not_submitted': 0})
        gt['expected'] += r['expected']
        gt['submitted'] += r['submitted']
        gt['not_submitted'] += r['not_submitted']
    for gt in grade_totals.values():
        gt['rate'] = round(gt['submitted'] * 100.0 / gt['expected'], 1) if gt['expected'] else 0.0
    overall = {'expected': 0, 'submitted': 0, 'not_submitted': 0}
    for gt in grade_totals.values():
        overall['expected'] += gt['expected']
        overall['submitted'] += gt['submitted']
        overall['not_submitted'] += gt['not_submitted']
    overall['rate'] = round(overall['submitted'] * 100.0 / overall['expected'], 1) if overall['expected'] else 0.0
    return {'grade_totals': [grade_totals[k] for k in sorted(grade_totals.keys())],
            'overall': overall}


def get_submission_trend(form_id, days=14):
    """每日提交趋势（近 days 天，按 submitted_at 日期聚合）"""
    subs = (FormSubmission.query.filter_by(template_id=form_id)
            .with_entities(FormSubmission.submitted_at).all())
    buckets = {}
    for (dt,) in subs:
        if not dt:
            continue
        buckets[dt.strftime('%Y-%m-%d')] = buckets.get(dt.strftime('%Y-%m-%d'), 0) + 1
    return [{'date': k, 'count': v} for k, v in sorted(buckets.items())][-days:]


def get_missing_submitters(form_id, grade=None, class_name=None, page=1, per_page=50):
    """未提交名单（应交但没交），分页 → (items, pagination)"""
    scope = {}
    if grade:
        scope['grade'] = grade
    if class_name:
        scope['class_name'] = class_name
    expected_list = get_expected_submitters(form_id, scope=scope or None)
    submitted_uids = _submitted_uid_set(form_id)
    missing = [e for e in expected_list if e.get('uid') not in submitted_uids]
    missing.sort(key=lambda x: (x.get('grade') or '', x.get('class_name') or '', str(x.get('uid') or '')))
    return _paginate(missing, page, per_page)


def get_all_missing(form_id, grade=None, class_name=None):
    """未提交名单全量（催交用，不分页）"""
    scope = {}
    if grade:
        scope['grade'] = grade
    if class_name:
        scope['class_name'] = class_name
    expected_list = get_expected_submitters(form_id, scope=scope or None)
    submitted_uids = _submitted_uid_set(form_id)
    return [e for e in expected_list if e.get('uid') not in submitted_uids]


def get_submitter_list(form_id, status=None, grade=None, class_name=None,
                       keyword=None, page=1, per_page=50):
    """已提交名单（按审核状态/班级筛选、姓名或学号搜索、分页）→ (items, pagination)"""
    q = _query_submissions(form_id, status=status, grade=grade,
                           class_name=class_name, keyword=keyword)
    subs = q.order_by(FormSubmission.submitter_grade, FormSubmission.submitter_class,
                      FormSubmission.submitter_uid).all()
    items = []
    for s in subs:
        items.append({
            'submission_id': s.id, 'uid': s.submitter_uid, 'name': s.submitter_name,
            'grade': s.submitter_grade or '', 'class_name': s.submitter_class or '',
            'submitted_at': s.submitted_at.strftime('%Y-%m-%d %H:%M') if s.submitted_at else '',
            'status': s.status, 'status_text': SUBMISSION_STATUS_TEXT.get(s.status, s.status),
            'review_note': s.review_note or '',
        })
    return _paginate(items, page, per_page)


def get_scope_options(form_id):
    """返回该表单应交范围内的 年级→班级 级联选项（供前端筛选下拉）"""
    expected_list = get_expected_submitters(form_id)
    tree = {}
    for e in expected_list:
        g = e.get('grade') or '—'
        c = e.get('class_name') or '—'
        tree.setdefault(g, set()).add(c)
    return [{'grade': g, 'classes': sorted(tree[g])} for g in sorted(tree.keys())]


# ── B. 横向汇总矩阵 ───────────────────────────────────────────

def _answers_by_question(sub):
    """把一条提交的 answers 按 question_id 分组（一题可能多条：多选/多文件）"""
    m = {}
    for a in sub.answers:
        m.setdefault(a.question_id, []).append(a)
    return m


def _build_cell(question, answers):
    """构造单元格 {display, raw, files}"""
    qtype = question.question_type
    if qtype == 'file':
        files = []
        for a in answers:
            if a.file_path:
                files.append({
                    'filename': a.file_name or os.path.basename(a.file_path),
                    'path': a.file_path, 'size': a.file_size or 0,
                    'size_text': _human_size(a.file_size or 0),
                    'url': _file_url(a.file_path),
                })
        return {'display': '、'.join(f['filename'] for f in files), 'raw': files, 'files': files}
    if qtype == 'multi_choice':
        selected = []
        for a in answers:
            if a.answer_json:
                try:
                    v = json.loads(a.answer_json)
                    if isinstance(v, list):
                        selected.extend(str(x) for x in v)
                    else:
                        selected.append(str(v))
                except (json.JSONDecodeError, TypeError):
                    selected.append(a.answer_json)
            elif a.answer_text:
                selected.append(a.answer_text)
        return {'display': '、'.join(selected), 'raw': selected, 'files': []}
    # text/textarea/single_choice/date/number
    texts = [a.answer_text for a in answers if a.answer_text not in (None, '')]
    display = texts[0] if texts else ''
    return {'display': display, 'raw': display, 'files': []}


def build_summary_matrix(form_id, status=None, grade=None, class_name=None,
                         keyword=None, include_rejected=False):
    """横向汇总：一行一人，一列一题。

    返回 dict(columns, rows, question_stats)。
    """
    tpl = _load_template(form_id)
    questions = _load_questions(form_id)

    columns = [{
        'key': f'q{q.id}', 'question_id': q.id, 'title': q.title,
        'question_type': q.question_type,
        'type_text': QUESTION_TYPE_TEXT.get(q.question_type, q.question_type),
        'is_file': q.question_type == 'file', 'sort_order': q.sort_order or 0,
    } for q in questions]

    subs = _query_submissions(form_id, status=status, grade=grade, class_name=class_name,
                              keyword=keyword, include_rejected=include_rejected) \
        .order_by(FormSubmission.submitter_grade, FormSubmission.submitter_class,
                  FormSubmission.submitter_uid).all()

    rows = []
    # 每题统计累加器
    stats_acc = {}
    for q in questions:
        stats_acc[q.id] = {
            'question_id': q.id, 'title': q.title, 'type': q.question_type,
            'type_text': QUESTION_TYPE_TEXT.get(q.question_type, q.question_type),
            'answered_count': 0, 'blank_count': 0,
            'options': [], 'option_counter': {},
            'numbers': [], 'text_len_total': 0,
            'file_count': 0, 'file_size_total': 0, 'file_missing': 0,
        }
        if q.question_type in ('single_choice', 'multi_choice'):
            stats_acc[q.id]['options'] = _question_options(q)

    for idx, sub in enumerate(subs, 1):
        amap = _answers_by_question(sub)
        cells = {}
        for q in questions:
            answers = amap.get(q.id, [])
            cell = _build_cell(q, answers)
            cells[str(q.id)] = cell
            acc = stats_acc[q.id]
            _accumulate_question_stat(acc, q, cell, answers)
        rows.append({
            'index': idx, 'submission_id': sub.id,
            'uid': sub.submitter_uid, 'name': sub.submitter_name,
            'grade': sub.submitter_grade or '', 'class_name': sub.submitter_class or '',
            'submitted_at': sub.submitted_at.strftime('%Y-%m-%d %H:%M') if sub.submitted_at else '',
            'status': sub.status, 'status_text': SUBMISSION_STATUS_TEXT.get(sub.status, sub.status),
            'answers': cells,
        })

    total_rows = len(rows)
    question_stats = []
    for q in questions:
        acc = stats_acc[q.id]
        stat = {
            'question_id': q.id, 'title': q.title, 'type': q.question_type,
            'type_text': acc['type_text'],
            'answered_count': acc['answered_count'], 'blank_count': acc['blank_count'],
        }
        if q.question_type in ('single_choice', 'multi_choice'):
            opts = acc['options']
            # 保证出现过的答案即便不在预设选项里也统计
            labels = list(opts)
            for lab in acc['option_counter'].keys():
                if lab not in labels:
                    labels.append(lab)
            answered = acc['answered_count'] or 0
            stat['options'] = [{
                'label': lab,
                'count': acc['option_counter'].get(lab, 0),
                'percent': round(acc['option_counter'].get(lab, 0) * 100.0 / answered, 1) if answered else 0.0,
            } for lab in labels]
        elif q.question_type == 'number':
            nums = acc['numbers']
            stat['avg'] = round(sum(nums) / len(nums), 2) if nums else None
            stat['min'] = min(nums) if nums else None
            stat['max'] = max(nums) if nums else None
            stat['count'] = len(nums)
        elif q.question_type == 'file':
            stat['file_count'] = acc['file_count']
            stat['file_size_total'] = acc['file_size_total']
            stat['file_size_text'] = _human_size(acc['file_size_total'])
            stat['file_missing'] = acc['file_missing']
        else:
            stat['avg_len'] = round(acc['text_len_total'] / acc['answered_count'], 1) if acc['answered_count'] else 0
        question_stats.append(stat)

    return {'columns': columns, 'rows': rows, 'question_stats': question_stats,
            'total_rows': total_rows, 'form_title': tpl.title}


def _accumulate_question_stat(acc, q, cell, answers):
    """把一个人的一题答案累加进题目标题统计"""
    qtype = q.question_type
    display = cell.get('display') or ''
    if qtype == 'file':
        nfiles = len(cell.get('files') or [])
        acc['file_count'] += nfiles
        for f in cell.get('files') or []:
            acc['file_size_total'] += f.get('size') or 0
        if nfiles > 0:
            acc['answered_count'] += 1
        else:
            acc['file_missing'] += 1
            acc['blank_count'] += 1
        return

    if qtype == 'multi_choice':
        selected = cell.get('raw') or []
        if selected:
            acc['answered_count'] += 1
            for lab in selected:
                acc['option_counter'][lab] = acc['option_counter'].get(lab, 0) + 1
        else:
            acc['blank_count'] += 1
        return

    if qtype == 'single_choice':
        if display:
            acc['answered_count'] += 1
            acc['option_counter'][display] = acc['option_counter'].get(display, 0) + 1
        else:
            acc['blank_count'] += 1
        return

    if qtype == 'number':
        if display:
            acc['answered_count'] += 1
            try:
                acc['numbers'].append(float(display))
            except ValueError:
                pass
        else:
            acc['blank_count'] += 1
        return

    # text / textarea / date
    if display:
        acc['answered_count'] += 1
        acc['text_len_total'] += len(display)
    else:
        acc['blank_count'] += 1


# ── B2. Excel 导出 ────────────────────────────────────────────

def _xlsx_headers():
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    return openpyxl, Font, Alignment, PatternFill, Border, Side


def export_summary_excel(form_id, status=None, grade=None, class_name=None,
                         keyword=None, include_rejected=False):
    """导出多工作表汇总 Excel → (BytesIO, 下载文件名)"""
    openpyxl, Font, Alignment, PatternFill, Border, Side = _xlsx_headers()
    from openpyxl.utils import get_column_letter

    tpl = _load_template(form_id)
    matrix = build_summary_matrix(form_id, status=status, grade=grade,
                                  class_name=class_name, keyword=keyword,
                                  include_rejected=include_rejected)
    stats = get_submission_stats(form_id)
    breakdown = get_class_breakdown(form_id)
    missing = get_all_missing(form_id)

    hf = Font(bold=True, color='FFFFFF')
    hfill = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
    sub_fill = PatternFill(start_color='8EAADB', end_color='8EAADB', fill_type='solid')
    tb = Border(left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin'))
    link_font = Font(color='0563C1', underline='single')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    wb = openpyxl.Workbook()

    # ── Sheet1 提交汇总 ──
    ws = wb.active
    ws.title = '提交汇总'
    fixed = ['序号', '学号', '姓名', '年级', '班级', '提交时间', '审核状态']
    questions = _load_questions(form_id)
    # 表头两行
    for i, h in enumerate(fixed, 1):
        ws.merge_cells(start_row=1, start_column=i, end_row=2, end_column=i)
        c = ws.cell(row=1, column=i, value=h)
        c.font = hf; c.fill = hfill; c.alignment = center; c.border = tb
        ws.cell(row=2, column=i).border = tb; ws.cell(row=2, column=i).fill = hfill
    col = len(fixed) + 1
    qcol = {}
    for q in questions:
        c1 = ws.cell(row=1, column=col, value=xl_safe(q.title))
        c1.font = hf; c1.fill = hfill; c1.alignment = center; c1.border = tb
        c2 = ws.cell(row=2, column=col,
                     value=QUESTION_TYPE_TEXT.get(q.question_type, q.question_type))
        c2.font = Font(bold=True, color='1F3864'); c2.fill = sub_fill
        c2.alignment = center; c2.border = tb
        qcol[q.id] = col
        col += 1

    r = 3
    for row in matrix['rows']:
        ws.cell(row=r, column=1, value=row['index']).border = tb
        ws.cell(row=r, column=2, value=xl_safe(row['uid'] or '')).border = tb
        ws.cell(row=r, column=3, value=xl_safe(row['name'] or '')).border = tb
        ws.cell(row=r, column=4, value=xl_safe(row['grade'])).border = tb
        ws.cell(row=r, column=5, value=xl_safe(row['class_name'])).border = tb
        ws.cell(row=r, column=6, value=row['submitted_at']).border = tb
        sc = ws.cell(row=r, column=7, value=row['status_text']); sc.border = tb
        if row['status'] == 'rejected':
            sc.font = Font(color='C00000')
        elif row['status'] == 'approved':
            sc.font = Font(color='2E7D32')
        for q in questions:
            cell = row['answers'].get(str(q.id), {})
            cc = ws.cell(row=r, column=qcol[q.id])
            cc.border = tb
            cc.alignment = Alignment(vertical='center', wrap_text=True)
            if q.question_type == 'file':
                files = cell.get('files') or []
                names = '、'.join(f['filename'] for f in files)
                cc.value = xl_safe(names or '-')
                if files:
                    cc.hyperlink = files[0]['url']
                    cc.font = link_font
            else:
                cc.value = xl_safe(cell.get('display') or '')
        r += 1

    for i in range(1, len(fixed) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 12
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['F'].width = 18
    for q in questions:
        ws.column_dimensions[get_column_letter(qcol[q.id])].width = 22
    ws.freeze_panes = 'D3'

    # ── Sheet2 提交统计 ──
    ws2 = wb.create_sheet('提交统计')
    ws2.append(['统计项', '数值'])
    for c in ws2[1]:
        c.font = hf; c.fill = hfill; c.alignment = center; c.border = tb
    stat_rows = [
        ('应交人数', stats['expected']), ('已交人数', stats['submitted']),
        ('未交人数', stats['not_submitted']), ('提交率(%)', stats['rate']),
        ('已通过', stats['approved']), ('待审核', stats['pending']),
        ('已驳回', stats['rejected']), ('答案条数', stats['total_answers']),
        ('附件总数', stats['total_files']), ('附件总大小', stats['total_file_size_text']),
    ]
    for name, val in stat_rows:
        ws2.append([name, val])
    ws2.append([])
    ws2.append(['年级', '班级', '应交', '已交', '未交', '提交率(%)'])
    hdr_row = ws2.max_row
    for c in ws2[hdr_row]:
        c.font = hf; c.fill = hfill; c.alignment = center; c.border = tb
    for b in breakdown:
        ws2.append(xl_row([b['grade'], b['class_name'], b['expected'],
                           b['submitted'], b['not_submitted'], b['rate']]))
    for col_i in range(1, 7):
        ws2.column_dimensions[get_column_letter(col_i)].width = 14

    # ── Sheet3 题目分析 ──
    ws3 = wb.create_sheet('题目分析')
    ws3.append(['题号', '题目', '题型', '已答', '未答', '统计明细'])
    for c in ws3[1]:
        c.font = hf; c.fill = hfill; c.alignment = center; c.border = tb
    for i, qs in enumerate(matrix['question_stats'], 1):
        detail = ''
        if qs['type'] in ('single_choice', 'multi_choice'):
            detail = '；'.join(f"{o['label']}: {o['count']}({o['percent']}%)"
                               for o in qs.get('options', []))
        elif qs['type'] == 'number':
            detail = f"均值 {qs.get('avg')} / 最小 {qs.get('min')} / 最大 {qs.get('max')}"
        elif qs['type'] == 'file':
            detail = f"文件 {qs.get('file_count',0)} 个 / {qs.get('file_size_text','0B')} / 缺失 {qs.get('file_missing',0)} 人"
        else:
            detail = f"平均字数 {qs.get('avg_len',0)}"
        ws3.append(xl_row([i, qs['title'], qs['type_text'], qs['answered_count'],
                           qs['blank_count'], detail]))
    ws3.column_dimensions['B'].width = 30
    ws3.column_dimensions['F'].width = 60
    for col_i in ('A', 'C', 'D', 'E'):
        ws3.column_dimensions[col_i].width = 10

    # ── Sheet4 未提交名单 ──
    ws4 = wb.create_sheet('未提交名单')
    ws4.append(['年级', '班级', '学号', '姓名'])
    for c in ws4[1]:
        c.font = hf; c.fill = hfill; c.alignment = center; c.border = tb
    for m in missing:
        ws4.append(xl_row([m.get('grade', ''), m.get('class_name', ''),
                           m.get('uid', ''), m.get('name', '')]))
    for col_i in ('A', 'B', 'C', 'D'):
        ws4.column_dimensions[col_i].width = 14

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    fname = f"材料汇总_{_sanitize(tpl.title)}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return out, fname


def export_missing_list_excel(form_id, grade=None, class_name=None):
    """单独导出未提交名单 Excel → (BytesIO, 下载文件名)"""
    openpyxl, Font, Alignment, PatternFill, Border, Side = _xlsx_headers()
    from openpyxl.utils import get_column_letter
    tpl = _load_template(form_id)
    missing = get_all_missing(form_id, grade=grade, class_name=class_name)
    hf = Font(bold=True, color='FFFFFF')
    hfill = PatternFill(start_color='C00000', end_color='C00000', fill_type='solid')
    center = Alignment(horizontal='center', vertical='center')
    tb = Border(left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin'))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '未提交名单'
    ws.append(['序号', '年级', '班级', '学号', '姓名'])
    for c in ws[1]:
        c.font = hf; c.fill = hfill; c.alignment = center; c.border = tb
    for i, m in enumerate(missing, 1):
        ws.append(xl_row([i, m.get('grade', ''), m.get('class_name', ''),
                          m.get('uid', ''), m.get('name', '')]))
        for c in ws[ws.max_row]:
            c.border = tb
    for col_i in range(1, 6):
        ws.column_dimensions[get_column_letter(col_i)].width = 14
    ws.freeze_panes = 'A2'
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    fname = f"未交名单_{_sanitize(tpl.title)}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return out, fname


# ── C. 材料清单与打包 ─────────────────────────────────────────

def get_file_inventory(form_id, grade=None, class_name=None, question_id=None,
                       submitter_uid=None, approved_only=False):
    """附件清单。返回 dict(items, total, total_size, total_size_text, missing_count)。

    items 每项：submission_id/uid/name/grade/class_name/question_id/question_title/
               filename/stored_path/rel_path/size/size_text/ext/exists/download_url
    排序：年级→班级→学号→题目。
    """
    questions = {q.id: q for q in _load_questions(form_id)}
    q = _query_submissions(form_id, grade=grade, class_name=class_name,
                           include_rejected=not approved_only)
    if approved_only:
        q = q.filter(FormSubmission.status == 'approved')
    subs = q.all()
    if submitter_uid:
        subs = [s for s in subs if s.submitter_uid == submitter_uid]

    items = []
    for s in subs:
        for a in s.answers:
            if not a.file_path:
                continue
            if question_id and a.question_id != int(question_id):
                continue
            qobj = questions.get(a.question_id)
            apath = _static_abspath(a.file_path)
            exists = os.path.exists(apath)
            size = a.file_size or (os.path.getsize(apath) if exists else 0)
            items.append({
                'submission_id': s.id, 'uid': s.submitter_uid, 'name': s.submitter_name,
                'grade': s.submitter_grade or '', 'class_name': s.submitter_class or '',
                'question_id': a.question_id,
                'question_title': qobj.title if qobj else f'题目{a.question_id}',
                'question_sort': (qobj.sort_order if qobj and qobj.sort_order is not None else 0),
                'filename': a.file_name or os.path.basename(a.file_path),
                'rel_path': a.file_path, 'stored_path': apath,
                'size': size, 'size_text': _human_size(size),
                'ext': os.path.splitext(a.file_path)[1].lower().lstrip('.'),
                'exists': exists, 'download_url': _file_url(a.file_path),
            })
    items.sort(key=lambda x: (x['grade'], x['class_name'], str(x['uid'] or ''),
                              x['question_sort'], x['filename']))
    total_size = sum(i['size'] for i in items if i['exists'])
    missing_count = sum(1 for i in items if not i['exists'])
    return {'items': items, 'total': len(items), 'total_size': total_size,
            'total_size_text': _human_size(total_size), 'missing_count': missing_count}


def _dedup_arcname(used, arcname):
    """重名追加 (1)(2)，绝不覆盖"""
    if arcname not in used:
        used.add(arcname)
        return arcname
    base, ext = os.path.splitext(arcname)
    n = 1
    while True:
        cand = f'{base}({n}){ext}'
        if cand not in used:
            used.add(cand)
            return cand
        n += 1


def _build_arcname(item, form_title, structure):
    """按目录结构生成 zip 内路径（各段做非法字符清洗）"""
    ft = _sanitize(form_title, 'form')
    grade = _sanitize(item['grade'] or '未分年级', '未分年级')
    cls = _sanitize(item['class_name'] or '未分班', '未分班')
    stu = _sanitize(f"{item['uid'] or ''}_{item['name'] or ''}".strip('_'), '未知')
    qtitle = _sanitize(item['question_title'] or '附件', '附件')
    fname = _sanitize(item['filename'] or 'file')
    if structure == 'question_class':
        return f'{ft}/{qtitle}/{grade}/{cls}/{stu}_{fname}'
    if structure == 'flat':
        return f'{ft}/{grade}_{cls}_{stu}_{qtitle}_{fname}'
    # class_student（默认）
    return f'{ft}/{grade}/{cls}/{stu}/{qtitle}/{fname}'


def build_material_package(form_id, scope='all', grade=None, class_name=None,
                           question_id=None, submitter_uid=None,
                           structure='class_student', approved_only=False,
                           include_manifest=True):
    """打包下载。返回 (buffer_or_path, download_name, stats, is_disk)。

    - 总大小 ≤ 阈值：写入 BytesIO（is_disk=False）；否则落盘临时文件（is_disk=True）。
    - zip 内中文用 UTF-8 flag（ZipInfo.flag_bits |= 0x800），流式写入避免大文件爆内存。
    """
    tpl = _load_template(form_id)
    inv = get_file_inventory(form_id, grade=grade, class_name=class_name,
                             question_id=question_id, submitter_uid=submitter_uid,
                             approved_only=approved_only)
    items = inv['items']
    existing = [i for i in items if i['exists']]
    total_size = inv['total_size']
    use_disk = total_size > PACKAGE_DISK_THRESHOLD

    if use_disk:
        tmp_dir = tempfile.mkdtemp(prefix='stulink_pkg_')
        tmp_path = os.path.join(tmp_dir, 'package.zip')
        target = open(tmp_path, 'wb')
    else:
        target = io.BytesIO()

    used = set()
    file_count = 0
    written_size = 0
    zf = zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED)
    try:
        for it in existing:
            arc = _build_arcname(it, tpl.title, structure)
            arc = _dedup_arcname(used, arc)
            zi = zipfile.ZipInfo(arc, date_time=datetime.now().timetuple()[:6])
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.flag_bits |= 0x800  # UTF-8 文件名标志
            zi.external_attr = 0o600 << 16
            with zf.open(zi, 'w') as dest, open(it['stored_path'], 'rb') as src:
                shutil.copyfileobj(src, dest, 1024 * 256)
            file_count += 1
            written_size += it['size']

        if include_manifest:
            manifest = _build_manifest_csv(items)
            mi = zipfile.ZipInfo('_材料清单.csv',
                                 date_time=datetime.now().timetuple()[:6])
            mi.compress_type = zipfile.ZIP_DEFLATED
            mi.flag_bits |= 0x800
            zf.writestr(mi, manifest)
    finally:
        zf.close()

    if use_disk:
        target.close()
        buffer_or_path = target.name
    else:
        target.seek(0)
        buffer_or_path = target

    scope_label = {'all': '全部'}.get(scope, scope or '全部')
    if grade and not class_name:
        scope_label = _sanitize(grade)
    elif grade and class_name:
        scope_label = _sanitize(f'{grade}{class_name}')
    dname = (f"材料包_{_sanitize(tpl.title)}_{scope_label}_"
             f"{datetime.now().strftime('%Y%m%d_%H%M')}.zip")
    stats = {'file_count': file_count, 'total_size': written_size,
             'total_size_text': _human_size(written_size),
             'skipped': len(items) - file_count, 'manifest': bool(include_manifest)}
    return buffer_or_path, dname, stats, use_disk


def _build_manifest_csv(items):
    """清单 CSV（utf-8-sig，Excel 直接可读）"""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(['年级', '班级', '学号', '姓名', '题目', '文件名', '大小', '状态'])
    for i in items:
        w.writerow([i['grade'], i['class_name'], i['uid'] or '', i['name'] or '',
                    i['question_title'], i['filename'], i['size_text'],
                    '正常' if i['exists'] else '文件缺失'])
    return ('\ufeff' + buf.getvalue()).encode('utf-8')


def get_package_preview(form_id, grade=None, class_name=None, question_id=None,
                        submitter_uid=None, structure='class_student',
                        approved_only=False):
    """打包前预览：文件数 / 总大小 / 目录树前 3 层示例 / 缺失警告"""
    tpl = _load_template(form_id)
    inv = get_file_inventory(form_id, grade=grade, class_name=class_name,
                             question_id=question_id, submitter_uid=submitter_uid,
                             approved_only=approved_only)
    items = inv['items']
    existing = [i for i in items if i['exists']]

    # 目录树前 3 层（取前若干文件构造示例）
    tree = {}
    for it in existing[:60]:
        arc = _build_arcname(it, tpl.title, structure)
        parts = arc.split('/')
        node = tree
        for p in parts[:3]:
            node = node.setdefault(p, {})
    tree_lines = _render_tree(tree, '', 0, 3)

    return {
        'form_title': tpl.title, 'structure': structure,
        'total_files': len(items), 'packable_files': len(existing),
        'total_size': inv['total_size'], 'total_size_text': inv['total_size_text'],
        'missing_count': inv['missing_count'],
        'use_disk': inv['total_size'] > PACKAGE_DISK_THRESHOLD,
        'tree': tree_lines,
        'warning': (f"有 {inv['missing_count']} 个附件物理文件缺失，将被跳过"
                    if inv['missing_count'] else ''),
    }


def _render_tree(node, prefix, depth, max_depth):
    lines = []
    if depth >= max_depth:
        return lines
    keys = list(node.keys())
    for i, k in enumerate(keys[:12]):
        last = (i == len(keys) - 1) or i == 11
        lines.append(f"{prefix}{'└─ ' if last else '├─ '}{k}")
        child_prefix = prefix + ('   ' if last else '│  ')
        lines.extend(_render_tree(node[k], child_prefix, depth + 1, max_depth))
    if len(keys) > 12:
        lines.append(f"{prefix}└─ …（共 {len(keys)} 项）")
    return lines


# ── D. 催交 ───────────────────────────────────────────────────

def get_remind_rounds(form_id):
    """已催交轮次：统计该表单已发出的催交通知条数。

    v2.0：优先按 biz_type='form_remind' AND biz_id=form_id 精确统计；
    兜底按标题匹配（兼容改造前历史数据）。
    """
    try:
        from app.models.notification import Notification
        # 精确路径：biz_type + biz_id
        cnt = Notification.query.filter_by(
            biz_type='form_remind', biz_id=form_id).count()
        if cnt:
            return cnt
        # 兜底：旧数据按标题匹配
        tpl = db.session.get(FormTemplate, form_id)
        if not tpl:
            return 0
        title = f'【催交】{tpl.title}'
        return Notification.query.filter_by(title=title).count()
    except Exception:
        return 0


def remind_submitters(form_id, uids=None, all_missing=False, operator=None,
                      grade=None, class_name=None):
    """对未提交者发提醒（v2.0：精确到人推送，已交者不再被误伤）。

    使用 notify_users(target_type='users') 精确推送到未交者的 uid。
    返回 (success, message, {notified, notifications_created, rounds})。
    notified = 实际收到通知的人数（精确等于 recipients 行数）。
    """
    from app.modules.notifications.services import notification_service

    tpl = _load_template(form_id)
    missing = get_all_missing(form_id, grade=grade, class_name=class_name)

    if uids:
        uid_set = {str(u) for u in uids}
        targets = [m for m in missing if str(m.get('uid')) in uid_set]
    elif all_missing:
        targets = missing
    else:
        return False, '请指定要催交的人员或选择催交全部未交', {'notified': 0,
                                                          'notifications_created': 0}

    if not targets:
        return False, '没有可催交的未提交人员', {'notified': 0, 'notifications_created': 0}

    # 填写链接
    try:
        fill_url = url_for('academic.form_fill', form_id=form_id, _external=True)
    except Exception:
        fill_url = f'/academic/forms/{form_id}/fill'
    deadline_txt = tpl.deadline.strftime('%Y-%m-%d %H:%M') if tpl.deadline else '无'

    # 收集未交者 uid 列表
    target_uids = [t['uid'] for t in targets if t.get('uid')]
    names = '、'.join(f"{m.get('class_name','')}{m.get('name','')}"
                      for m in targets[:40])
    if len(targets) > 40:
        names += f' 等 {len(targets)} 人'

    title = f'【催交】{tpl.title}'
    content = (f'尚有 {len(targets)} 人未提交「{tpl.title}」。\n'
               f'未交名单：{names}\n'
               f'截止时间：{deadline_txt}\n'
               f'请尽快完成提交。')

    success, msg, info = notification_service.notify_users(
        uids=target_uids,
        title=title,
        content=content,
        category='reminder',
        biz_type='form_remind',
        biz_id=form_id,
        link_url=fill_url,
        creator_id=(operator.id if operator else None),
        priority='urgent',
    )
    notified = info.get('notified', 0) if info else 0
    rounds = get_remind_rounds(form_id)
    if not success:
        return False, msg, {'notified': 0, 'notifications_created': 0, 'rounds': rounds}
    return True, f'已向 {notified} 名未提交者发出催交通知', {
        'notified': notified, 'notifications_created': 1, 'rounds': rounds}
