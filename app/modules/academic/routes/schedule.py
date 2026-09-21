# StuLink v1.17.0 2026-09-20
# 教务 · 学期课表（timetable.db）：管理 / 视图 / 条目CRUD / 导入导出 / 查课联动 / JSON API
# Copyright (c) 2026 zkxxzf. Apache License 2.0
"""新课表系统路由。

- 管理端（/academic/schedule/...）：需 academic.timetable 权限
- 查看端（my-schedule / today / inspection 联动 / api）：需 academic.view 权限
- 所有 JSON 接口统一返回 {success, message, data}
"""
from datetime import date, datetime

import io

from flask import (render_template, request, redirect, url_for, flash,
                   send_file, abort, jsonify)
from flask_login import login_required, current_user
from sqlalchemy import func

from app.extensions import db
from app.models.timetable import (TermSchedule, ScheduleEntry, WEEKDAY_NAMES,
                                  MAX_PERIOD, PERIOD_TYPES)
from app.modules.academic import bp
from app.modules.academic.services import schedule_service as svc
from app.modules.academic.services import term_service as tsvc
from app.utils.decorators import perm_required
from app.utils.export_helpers import xl_row, xl_safe
from app.utils.helpers import log_operation

_XLSX_MIME = ('application/vnd.openxmlformats-officedocument'
              '.spreadsheetml.sheet')


# ─── 内部工具 ──────────────────────────────────────────────────────────────

def _json_ok(data=None, message='ok'):
    return jsonify({'success': True, 'message': message, 'data': data})


def _json_err(message, code=400):
    return jsonify({'success': False, 'message': message, 'data': None}), code


def _payload():
    """兼容 JSON 与表单两种提交方式"""
    if request.is_json:
        return request.get_json(silent=True) or {}
    return request.form.to_dict()


def _resolve_teacher(teacher_uid, teacher_name):
    """按教师编号补全姓名（跨库查 academic.db Teacher，失败静默）"""
    if teacher_uid and not teacher_name:
        try:
            from app.models.academic import Teacher
            t = Teacher.query.filter_by(teacher_uid=teacher_uid).first()
            if t:
                teacher_name = t.name
        except Exception:
            pass
    return teacher_uid or None, teacher_name or None


def _get_schedule_or_404(sid):
    ts = db.session.get(TermSchedule, sid)
    if not ts:
        abort(404)
    return ts


def _can_edit():
    return current_user.has_perm('academic.timetable')


def _week_param():
    """解析 ?week=N 周次参数（预留钩子，透传给 service 视图函数）"""
    w = request.args.get('week', type=int)
    return w if w and w > 0 else None


def _editable(ts):
    """页面编辑态：有课表管理权限且学期未归档（archived 只读回看）"""
    return _can_edit() and ts.status != 'archived'


def _readonly_guard(ts, back_endpoint, **back_kwargs):
    """写操作拦截：归档学期只读，返回重定向响应或 None"""
    if ts.status == 'archived':
        flash('历史学期课表为只读，不能修改', 'warning')
        return redirect(url_for(back_endpoint, **back_kwargs))
    return None


# ═══════════════════════════════════════════════════════════════════════════
# 管理端页面
# ═══════════════════════════════════════════════════════════════════════════

@bp.route('/schedule')
@login_required
@perm_required('academic.timetable')
def schedule_manage():
    """学期课表管理首页：学期列表 + 创建/激活/归档/删除入口"""
    schedules = svc.list_schedules()
    ids = [ts.id for ts in schedules]
    # 条目数一次聚合（原为逐学期 count，N 次查询）
    counts = {}
    if ids:
        counts = dict(
            db.session.query(ScheduleEntry.term_schedule_id,
                             func.count(ScheduleEntry.id))
            .filter(ScheduleEntry.is_deleted.is_(False),
                    ScheduleEntry.term_schedule_id.in_(ids))
            .group_by(ScheduleEntry.term_schedule_id).all())
    # 各学期的年级/班级（用于「年级课表」入口，不再硬编码高一）
    grade_map = svc.grade_class_map(ids)
    return render_template('academic/schedule_manage.html',
                           schedules=schedules, counts=counts,
                           grade_map=grade_map)


@bp.route('/schedule/create', methods=['POST'])
@login_required
@perm_required('academic.timetable')
def schedule_create():
    """创建学期（自动带 13 节默认节次）"""
    name = (request.form.get('name') or '').strip()
    school_year = (request.form.get('school_year') or '').strip()
    term = (request.form.get('term') or '').strip()
    description = (request.form.get('description') or '').strip()
    if not name or not school_year or not term:
        flash('名称、学年、学期为必填项', 'danger')
        return redirect(url_for('academic.schedule_manage'))
    try:
        ts = svc.create_schedule(name, school_year, term,
                                 description=description or None,
                                 created_by=current_user)
    except Exception:
        db.session.rollback()
        flash('创建失败，请重试', 'danger')
        return redirect(url_for('academic.schedule_manage'))
    log_operation(current_user, '新增', '学期课表', ts.id, ts.name, module='academic')
    flash(f'学期「{ts.name}」已创建（草稿状态，含 13 节默认作息）', 'success')
    return redirect(url_for('academic.schedule_manage'))


@bp.route('/schedule/<int:sid>/activate', methods=['POST'])
@login_required
@perm_required('academic.timetable')
def schedule_activate(sid):
    """激活学期（其他 active 自动归档）"""
    ts = _get_schedule_or_404(sid)
    try:
        svc.activate_schedule(sid)
    except ValueError as e:
        flash(str(e), 'danger')
        return redirect(url_for('academic.schedule_manage'))
    log_operation(current_user, '激活', '学期课表', sid, ts.name, module='academic')
    flash(f'学期「{ts.name}」已激活', 'success')
    return redirect(url_for('academic.schedule_manage'))


@bp.route('/schedule/<int:sid>/archive', methods=['POST'])
@login_required
@perm_required('academic.timetable')
def schedule_archive(sid):
    """归档学期（与 archive-term 同一实现：写归档快照，避免两套语义）"""
    ts = _get_schedule_or_404(sid)
    try:
        ok, msg = tsvc.archive_term(sid, operator=current_user)
    except Exception:
        db.session.rollback()
        flash('归档失败，请重试', 'danger')
        return redirect(url_for('academic.schedule_manage'))
    if not ok:
        flash(msg, 'danger')
        return redirect(url_for('academic.schedule_manage'))
    log_operation(current_user, '归档', '学期课表', sid, ts.name, module='academic')
    flash(msg + '（已生成归档快照，可追溯）', 'success')
    return redirect(url_for('academic.schedule_manage'))


@bp.route('/schedule/<int:sid>/delete', methods=['POST'])
@login_required
@perm_required('academic.timetable')
def schedule_delete(sid):
    """删除学期（仅 draft）"""
    ts = _get_schedule_or_404(sid)
    name = ts.name
    try:
        svc.delete_schedule(sid)
    except ValueError as e:
        flash(str(e), 'danger')
        return redirect(url_for('academic.schedule_manage'))
    log_operation(current_user, '删除', '学期课表', sid, name, module='academic')
    flash(f'学期「{name}」已删除', 'success')
    return redirect(url_for('academic.schedule_manage'))


@bp.route('/schedule/<int:sid>/master')
@login_required
@perm_required('academic.timetable')
def schedule_master(sid):
    """大课表（全校总览）：年级标签 + AJAX 按需加载班级网格"""
    ts = _get_schedule_or_404(sid)
    grade_classes = svc.get_grade_class_list(sid)
    init_grade = (request.args.get('grade') or '').strip()
    if init_grade not in grade_classes:
        init_grade = next(iter(grade_classes), '')
    return render_template('academic/schedule_master.html',
                           ts=ts, grade_classes=grade_classes,
                           init_grade=init_grade, week=_week_param(),
                           schedules=svc.list_schedules(),
                           can_edit=_editable(ts))


@bp.route('/schedule/<int:sid>/grade/<grade>')
@login_required
@perm_required('academic.timetable')
def schedule_grade(sid, grade):
    """年级课表：年级内班级标签切换（?class= 指定班级）"""
    ts = _get_schedule_or_404(sid)
    week = _week_param()
    data = svc.get_grade_view(sid, grade, week=week)
    classes = data['classes']
    cur_class = (request.args.get('class') or '').strip()
    if cur_class not in classes:
        cur_class = classes[0] if classes else ''
    view = svc.get_class_view(sid, grade, cur_class, week=week) if cur_class else \
        {'grid': {}, 'periods': data['periods'], 'stats': {}, 'total': 0}
    return render_template('academic/schedule_grade.html',
                           ts=ts, grade=grade, classes=classes,
                           cur_class=cur_class, view=view, week=week,
                           schedules=svc.list_schedules(),
                           can_edit=_editable(ts))


@bp.route('/schedule/<int:sid>/class/<grade>/<class_name>')
@login_required
@perm_required('academic.timetable')
def schedule_class(sid, grade, class_name):
    """班级课表：13x7 网格 + 学科课时统计"""
    ts = _get_schedule_or_404(sid)
    week = _week_param()
    view = svc.get_class_view(sid, grade, class_name, week=week)
    return render_template('academic/schedule_class.html',
                           ts=ts, grade=grade, class_name=class_name,
                           view=view, week=week,
                           schedules=svc.list_schedules(),
                           can_edit=_editable(ts))


@bp.route('/schedule/<int:sid>/teacher/<uid>')
@login_required
@perm_required('academic.timetable')
def schedule_teacher(sid, uid):
    """教师个人课表：搜索选择 + 网格 + 课时统计"""
    ts = _get_schedule_or_404(sid)
    view = svc.get_teacher_view(sid, uid, week=_week_param())
    teacher = None
    try:
        from app.models.academic import Teacher
        teacher = Teacher.query.filter_by(teacher_uid=uid).first()
    except Exception:
        pass
    teachers = svc.get_all_teachers()
    return render_template('academic/schedule_teacher.html',
                           ts=ts, uid=uid, teacher=teacher, view=view,
                           teachers=teachers, week=_week_param(),
                           schedules=svc.list_schedules(),
                           can_edit=_editable(ts), my_mode=False)


@bp.route('/schedule/<int:sid>/periods', methods=['GET', 'POST'])
@login_required
@perm_required('academic.timetable')
def schedule_periods(sid):
    """节次时间配置（GET 表单 / POST 保存）"""
    ts = _get_schedule_or_404(sid)
    if request.method == 'POST':
        guard = _readonly_guard(ts, 'academic.schedule_periods', sid=sid)
        if guard:
            return guard
        nums = request.form.getlist('period_number')
        names = request.form.getlist('period_name')
        starts = request.form.getlist('start_time')
        ends = request.form.getlist('end_time')
        types = request.form.getlist('period_type')
        periods_data = []
        for i, n in enumerate(nums):
            try:
                pn = int(n)
            except (ValueError, TypeError):
                continue
            if not (1 <= pn <= MAX_PERIOD):
                continue
            periods_data.append({
                'period_number': pn,
                'period_name': (names[i] if i < len(names) else '').strip() or f'第{pn}节',
                'start_time': (starts[i] if i < len(starts) else '').strip() or None,
                'end_time': (ends[i] if i < len(ends) else '').strip() or None,
                'period_type': (types[i] if i < len(types) else '').strip() or 'morning',
                'sort_order': i + 1,
            })
        try:
            result = svc.save_periods(sid, periods_data, remove_missing=True)
        except Exception:
            db.session.rollback()
            flash('保存失败，请检查时间格式', 'danger')
            return redirect(url_for('academic.schedule_periods', sid=sid))
        log_operation(current_user, '更新', '节次配置', sid,
                      f'{ts.name}：{len(periods_data)} 节'
                      f"{'，移除 ' + '、'.join(str(n) for n in result['removed']) if result['removed'] else ''}"
                      f"{'，保留在用 ' + '、'.join(str(n) for n in result['kept_in_use']) if result['kept_in_use'] else ''}",
                      module='academic')
        msg = '节次配置已保存'
        if result['removed']:
            msg += f"；已移除第 {'、'.join(str(n) for n in result['removed'])} 节"
        flash(msg, 'success')
        if result['kept_in_use']:
            nums = '、'.join(str(n) for n in result['kept_in_use'])
            flash(f'第 {nums} 节仍有课程安排，未移除（请先调整这些课，再删除节次）', 'warning')
        return redirect(url_for('academic.schedule_periods', sid=sid))

    periods = svc.get_periods(sid)
    return render_template('academic/schedule_periods.html',
                           ts=ts, periods=periods, period_types=PERIOD_TYPES,
                           schedules=svc.list_schedules(),
                           can_edit=_editable(ts))


@bp.route('/schedule/<int:sid>/periods/reset', methods=['POST'])
@login_required
@perm_required('academic.timetable')
def schedule_periods_reset(sid):
    """重置为默认 13 节"""
    ts = _get_schedule_or_404(sid)
    guard = _readonly_guard(ts, 'academic.schedule_periods', sid=sid)
    if guard:
        return guard
    svc.reset_default_periods(sid)
    log_operation(current_user, '重置', '节次配置', sid, ts.name, module='academic')
    flash('已重置为默认 13 节作息模板', 'success')
    return redirect(url_for('academic.schedule_periods', sid=sid))


@bp.route('/schedule/<int:sid>/versions')
@login_required
@perm_required('academic.timetable')
def schedule_versions(sid):
    """课表变更历史（分页）"""
    ts = _get_schedule_or_404(sid)
    page = request.args.get('page', 1, type=int)
    pagination = svc.get_schedule_versions(sid, page=page, per_page=50)
    return render_template('academic/schedule_versions.html',
                           ts=ts, pagination=pagination,
                           versions=pagination.items)


@bp.route('/schedule/<int:sid>/import', methods=['GET', 'POST'])
@login_required
@perm_required('academic.timetable')
def schedule_import(sid):
    """Excel 批量导入页面 + 上传处理"""
    ts = _get_schedule_or_404(sid)
    result = None
    if request.method == 'POST':
        guard = _readonly_guard(ts, 'academic.schedule_import', sid=sid)
        if guard:
            return guard
        file = request.files.get('file')
        if not file or not file.filename:
            flash('请选择 Excel 文件', 'danger')
            return redirect(url_for('academic.schedule_import', sid=sid))
        try:
            result = svc.import_from_excel(sid, file, operator=current_user)
        except Exception:
            db.session.rollback()
            flash('导入失败：文件解析异常，请使用标准模板', 'danger')
            return redirect(url_for('academic.schedule_import', sid=sid))
        log_operation(current_user, '导入', '学期课表', sid,
                      f'{ts.name}：成功 {result["success"]} 条，失败 {result["failed"]} 条',
                      module='academic')
        flash(f'导入完成：成功 {result["success"]} 条，失败 {result["failed"]} 条',
              'success' if result['success'] else 'warning')
    return render_template('academic/schedule_import.html', ts=ts, result=result)


@bp.route('/schedule/<int:sid>/export')
@login_required
@perm_required('academic.timetable')
def schedule_export(sid):
    """导出 Excel（query: view_type/grade/class/teacher）"""
    ts = _get_schedule_or_404(sid)
    view_type = (request.args.get('view_type') or 'master').strip()
    grade = (request.args.get('grade') or '').strip() or None
    class_name = (request.args.get('class') or '').strip() or None
    teacher_uid = (request.args.get('teacher') or '').strip() or None
    try:
        buf = svc.export_schedule(sid, view_type=view_type, grade=grade,
                                  class_name=class_name, teacher_uid=teacher_uid,
                                  week=_week_param())
    except Exception:
        db.session.rollback()
        flash('导出失败，请重试', 'danger')
        return redirect(url_for('academic.schedule_master', sid=sid))
    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    suffix = {'class': f'_{grade or ""}{class_name or ""}',
              'grade': f'_{grade or ""}',
              'teacher': f'_{teacher_uid or ""}'}.get(view_type, '')
    return send_file(buf, as_attachment=True,
                     download_name=f'{ts.name}{suffix}_{stamp}.xlsx',
                     mimetype=_XLSX_MIME)


@bp.route('/schedule/<int:sid>/template')
@login_required
@perm_required('academic.timetable')
def schedule_template(sid):
    """下载导入模板"""
    _get_schedule_or_404(sid)
    buf = svc.generate_import_template(sid)
    return send_file(buf, as_attachment=True,
                     download_name='课表导入模板.xlsx', mimetype=_XLSX_MIME)


# ═══════════════════════════════════════════════════════════════════════════
# 条目 CRUD（JSON）
# ═══════════════════════════════════════════════════════════════════════════

@bp.route('/schedule/<int:sid>/entry/add', methods=['POST'])
@login_required
@perm_required('academic.timetable')
def schedule_entry_add(sid):
    """添加课条目"""
    ts = _get_schedule_or_404(sid)
    if ts.status == 'archived':
        return _json_err('历史学期课表为只读，不能修改', 403)
    d = _payload()
    grade = (d.get('grade') or '').strip()
    class_name = (d.get('class_name') or '').strip()
    subject = (d.get('subject') or '').strip()
    try:
        weekday = int(d.get('weekday') or 0)
        period_number = int(d.get('period_number') or 0)
    except (ValueError, TypeError):
        return _json_err('星期/节次格式无效')
    if not grade or not class_name or not subject:
        return _json_err('年级、班级、学科为必填项')
    if not (1 <= weekday <= 7):
        return _json_err('星期超出范围（1-7）')

    teacher_uid, teacher_name = _resolve_teacher(
        (d.get('teacher_uid') or '').strip(), (d.get('teacher_name') or '').strip())
    try:
        ok, result = svc.add_entry(
            sid, grade, class_name, weekday, period_number, subject,
            teacher_uid=teacher_uid, teacher_name=teacher_name,
            room=(d.get('room') or '').strip() or None,
            week_range=(d.get('week_range') or '').strip() or '1-18',
            note=(d.get('note') or '').strip() or None,
            operator=current_user)
    except Exception:
        db.session.rollback()
        return _json_err('添加失败，请重试', 500)
    if not ok:
        return _json_err(result, 409)
    log_operation(current_user, '新增', '课表条目', result.id,
                  f'{ts.name}：{grade}{class_name} {WEEKDAY_NAMES.get(weekday,"")}第{period_number}节 {subject}',
                  module='academic')
    return _json_ok(result.to_dict(), '课表条目已添加')


@bp.route('/schedule/<int:sid>/entry/<int:eid>/edit', methods=['POST'])
@login_required
@perm_required('academic.timetable')
def schedule_entry_edit(sid, eid):
    """编辑课条目"""
    ts = _get_schedule_or_404(sid)
    if ts.status == 'archived':
        return _json_err('历史学期课表为只读，不能修改', 403)
    entry = db.session.get(ScheduleEntry, eid)
    if not entry or entry.term_schedule_id != sid or entry.is_deleted:
        return _json_err('条目不存在', 404)
    d = _payload()
    fields = {}
    for key in ('grade', 'class_name', 'subject', 'room', 'week_range', 'note'):
        if key in d:
            fields[key] = (d.get(key) or '').strip() or None
    for key in ('weekday', 'period_number'):
        if d.get(key):
            try:
                fields[key] = int(d[key])
            except (ValueError, TypeError):
                return _json_err(f'{key} 格式无效')
    if 'weekday' in fields and not (1 <= fields['weekday'] <= 7):
        return _json_err('星期超出范围（1-7）')
    if 'teacher_uid' in d:
        fields['teacher_uid'], fields['teacher_name'] = _resolve_teacher(
            (d.get('teacher_uid') or '').strip(), (d.get('teacher_name') or '').strip())
    try:
        ok, result = svc.edit_entry(eid, operator=current_user, **fields)
    except Exception:
        db.session.rollback()
        return _json_err('保存失败，请重试', 500)
    if not ok:
        return _json_err(result, 409)
    log_operation(current_user, '更新', '课表条目', eid, ts.name, module='academic')
    return _json_ok(result.to_dict(), '课表条目已更新')


@bp.route('/schedule/<int:sid>/entry/<int:eid>/delete', methods=['POST'])
@login_required
@perm_required('academic.timetable')
def schedule_entry_delete(sid, eid):
    """删除课条目（软删除）"""
    ts = _get_schedule_or_404(sid)
    if ts.status == 'archived':
        return _json_err('历史学期课表为只读，不能修改', 403)
    entry = db.session.get(ScheduleEntry, eid)
    if not entry or entry.term_schedule_id != sid or entry.is_deleted:
        return _json_err('条目不存在', 404)
    detail = f'{entry.grade}{entry.class_name} {entry.subject}'
    try:
        ok, msg = svc.delete_entry(eid, operator=current_user)
    except Exception:
        db.session.rollback()
        return _json_err('删除失败，请重试', 500)
    if not ok:
        return _json_err(msg, 409)
    log_operation(current_user, '删除', '课表条目', eid, f'{ts.name}：{detail}',
                  module='academic')
    return _json_ok(None, '课表条目已删除')


# ═══════════════════════════════════════════════════════════════════════════
# 查看端（教师/班主任）
# ═══════════════════════════════════════════════════════════════════════════

@bp.route('/my-schedule')
@login_required
@perm_required('academic.view')
def my_schedule():
    """我的课表（教师视角，取当前登录用户关联的教师编号；?sid= 可回看历史学期）"""
    from app.modules.academic.services.teacher_service import teacher_of_user
    sid = request.args.get('sid', type=int)
    ts = db.session.get(TermSchedule, sid) if sid else svc.get_active_schedule()
    teacher = teacher_of_user(current_user) if ts else None
    view = None
    if ts and teacher:
        view = svc.get_teacher_view(ts.id, teacher.teacher_uid,
                                    week=_week_param())
    return render_template('academic/schedule_teacher.html',
                           ts=ts, uid=teacher.teacher_uid if teacher else '',
                           teacher=teacher, view=view, teachers=[],
                           week=_week_param(), schedules=svc.list_schedules(),
                           can_edit=False, my_mode=True)


@bp.route('/today')
@login_required
@perm_required('academic.view')
def today_schedule():
    """今日课表（全校当天，支持 grade/class/week 过滤，标记当前节次）"""
    grade = (request.args.get('grade') or '').strip() or None
    class_name = (request.args.get('class') or '').strip() or None
    sid = request.args.get('sid', type=int)
    ts = db.session.get(TermSchedule, sid) if sid else svc.get_active_schedule()
    data = svc.get_today_schedule(schedule_id=ts.id if ts else None,
                                  grade=grade, class_name=class_name,
                                  week=_week_param())
    # 按 年级→班级 分组供模板渲染
    grouped = {}
    for e in data['entries']:
        grouped.setdefault(e['grade'], {}).setdefault(e['class_name'], []).append(e)
    grade_opts = sorted(grouped.keys())
    return render_template('academic/schedule_today.html',
                           ts=ts, data=data, grouped=grouped,
                           grade_opts=grade_opts, week=_week_param(),
                           schedules=svc.list_schedules(),
                           f_grade=grade or '', f_class=class_name or '')


# ═══════════════════════════════════════════════════════════════════════════
# 查课联动
# ═══════════════════════════════════════════════════════════════════════════

@bp.route('/inspection/schedule')
@login_required
@perm_required('academic.view')
def inspection_schedule():
    """查课联动：已统一到 inspection_live，本路由保留并重定向。"""
    return redirect(url_for('academic.inspection_live_schedule',
                            date=request.args.get('date', ''),
                            grade=request.args.get('grade', '')))


@bp.route('/inspection/schedule/period/<int:n>')
@login_required
@perm_required('academic.view')
def inspection_schedule_period(n):
    """查课联动：指定节次 → 重定向到实时课表（带 period 参数）。"""
    if not (1 <= n <= MAX_PERIOD):
        abort(404)
    d = (request.args.get('date') or '').strip()
    grade = (request.args.get('grade') or '').strip()
    return redirect(url_for('academic.inspection_live_schedule',
                            date=d, period=n, grade=grade))


# ═══════════════════════════════════════════════════════════════════════════
# AJAX 数据接口（JSON）
# ═══════════════════════════════════════════════════════════════════════════

@bp.route('/api/schedule/<int:sid>/class-data')
@login_required
@perm_required('academic.view')
def api_schedule_class_data(sid):
    """班级网格 JSON（query: grade/class_name/weekday）"""
    _get_schedule_or_404(sid)
    grade = (request.args.get('grade') or '').strip()
    class_name = (request.args.get('class_name') or '').strip()
    weekday = request.args.get('weekday', type=int)
    if not grade or not class_name:
        return _json_err('缺少 grade / class_name 参数')
    view = svc.get_class_view(sid, grade, class_name, weekday=weekday,
                              week=_week_param())
    return _json_ok(view)


@bp.route('/api/schedule/<int:sid>/grade-data')
@login_required
@perm_required('academic.view')
def api_schedule_grade_data(sid):
    """年级各班数据 JSON（大课表切换年级时按需加载）"""
    _get_schedule_or_404(sid)
    grade = (request.args.get('grade') or '').strip()
    if not grade:
        return _json_err('缺少 grade 参数')
    data = svc.get_grade_view(sid, grade, week=_week_param())
    return _json_ok(data)


@bp.route('/api/schedule/today-data')
@login_required
@perm_required('academic.view')
def api_schedule_today_data():
    """今日课表 JSON（query: grade/sid/week）"""
    grade = (request.args.get('grade') or '').strip() or None
    sid = request.args.get('sid', type=int)
    data = svc.get_today_schedule(schedule_id=sid, grade=grade,
                                  week=_week_param())
    return _json_ok(data)


@bp.route('/api/schedule/period-data/<int:n>')
@login_required
@perm_required('academic.view')
def api_schedule_period_data(n):
    """指定节次全校 JSON（query: date/grade）"""
    if not (1 <= n <= MAX_PERIOD):
        return _json_err('节次超出范围')
    grade = (request.args.get('grade') or '').strip() or None
    d = (request.args.get('date') or '').strip()
    sid = request.args.get('sid', type=int)
    data = svc.get_period_schedule(n, schedule_id=sid, grade=grade,
                                   date_str=d or None, week=_week_param())
    return _json_ok(data)


@bp.route('/api/schedule/entry/<int:eid>')
@login_required
@perm_required('academic.view')
def api_schedule_entry_detail(eid):
    """条目详情 JSON（含变更历史）"""
    detail = svc.get_entry_detail(eid)
    if not detail:
        return _json_err('条目不存在', 404)
    versions = [v.to_dict() for v in svc.get_entry_versions(eid, limit=20)]
    return _json_ok({'entry': detail, 'versions': versions})


@bp.route('/api/schedule/teachers')
@login_required
@perm_required('academic.view')
def api_schedule_teachers():
    """教师列表 JSON（供下拉选择，academic.db Teacher）"""
    try:
        teachers = svc.get_all_teachers()
    except Exception:
        teachers = []
    kw = (request.args.get('q') or '').strip()
    if kw:
        teachers = [t for t in teachers
                    if kw in t['name'] or kw in t['uid'] or kw in t['subject']]
    return _json_ok(teachers)


@bp.route('/api/schedule/classes')
@login_required
@perm_required('academic.view')
def api_schedule_classes():
    """年级+班级列表 JSON（学籍库 distinct + 课表已有条目合并，供级联下拉）"""
    merged = {}
    try:
        from app.models.student import Student
        rows = db.session.query(Student.grade, Student.class_name).distinct().all()
        skip = {'已转出', '离校', '不分班', '已毕业', ''}
        for g, cn in rows:
            if g and cn and cn not in skip:
                merged.setdefault(g, set()).add(cn)
    except Exception:
        pass
    sid = request.args.get('sid', type=int)
    for g, cs in svc.get_grade_class_list(sid).items():
        merged.setdefault(g, set()).update(cs)
    data = {g: sorted(cs) for g, cs in sorted(merged.items())}
    return _json_ok(data)


@bp.route('/api/schedule/check-conflict')
@login_required
@perm_required('academic.view')
def api_schedule_check_conflict():
    """冲突检测 JSON（query: sid/teacher_uid/grade/class_name/weekday/period_number/week_range）

    week_range 参与判定：周次无交集的条目不算冲突（单周课与双周课可同格）。
    """
    sid = request.args.get('sid', type=int)
    if not sid:
        return _json_err('缺少 sid 参数')
    teacher_uid = (request.args.get('teacher_uid') or '').strip()
    grade = (request.args.get('grade') or '').strip()
    class_name = (request.args.get('class_name') or '').strip()
    weekday = request.args.get('weekday', type=int)
    period_number = request.args.get('period_number', type=int)
    exclude = request.args.get('exclude_entry_id', type=int)
    week_range = (request.args.get('week_range') or '').strip() or None
    if not weekday or not period_number:
        return _json_err('缺少 weekday / period_number 参数')

    data = {'class_conflict': None, 'teacher_conflict': None}
    if grade and class_name:
        c = svc.check_class_conflict(sid, grade, class_name, weekday,
                                     period_number, exclude_entry_id=exclude,
                                     week_range=week_range)
        if c:
            data['class_conflict'] = c.to_dict()
    if teacher_uid:
        t = svc.check_teacher_conflict(sid, teacher_uid, weekday, period_number,
                                       exclude_entry_id=exclude, week_range=week_range)
        if t:
            data['teacher_conflict'] = t.to_dict()
    data['has_conflict'] = bool(data['class_conflict'] or data['teacher_conflict'])
    return _json_ok(data)


# ═══════════════════════════════════════════════════════════════════════════
# 学期周期维度 / 历史课表追溯（Task#27 增量追加，不改动以上已有路由）
# ═══════════════════════════════════════════════════════════════════════════

def _parse_date(raw):
    """解析 YYYY-MM-DD 字符串为 date，空/非法返回 None"""
    raw = (raw or '').strip()
    if not raw:
        return None
    try:
        return datetime.strptime(raw[:10], '%Y-%m-%d').date()
    except (ValueError, TypeError):
        return None


@bp.route('/schedule/<int:sid>/dates', methods=['GET', 'POST'])
@login_required
@perm_required('academic.timetable')
def schedule_dates(sid):
    """学期起止日期与教学周配置（GET 表单 / POST 保存 + 校验警告展示）"""
    ts = _get_schedule_or_404(sid)
    warnings = []
    if request.method == 'POST':
        guard = _readonly_guard(ts, 'academic.schedule_dates', sid=sid)
        if guard:
            return guard
        start_date = _parse_date(request.form.get('start_date'))
        end_date = _parse_date(request.form.get('end_date'))
        try:
            total_weeks = int(request.form.get('total_weeks') or 0)
        except (ValueError, TypeError):
            total_weeks = 0
        try:
            week_start_offset = int(request.form.get('week_start_offset') or 0)
        except (ValueError, TypeError):
            week_start_offset = 0
        is_current = request.form.get('is_current') in ('1', 'on', 'true', 'True')
        warnings = tsvc.validate_term_dates(start_date, end_date, total_weeks)
        # 硬错误拦截：结束<=开始 或 周数<=0 时不保存
        if (start_date and end_date and end_date <= start_date) or total_weeks <= 0:
            flash('保存失败：' + '；'.join(warnings), 'danger')
            return redirect(url_for('academic.schedule_dates', sid=sid))
        try:
            if is_current:
                TermSchedule.query.filter(TermSchedule.id != sid)\
                    .update({'is_current': False}, synchronize_session=False)
            svc.update_schedule(sid, start_date=start_date, end_date=end_date,
                                total_weeks=total_weeks,
                                week_start_offset=week_start_offset,
                                is_current=is_current)
            db.session.commit()
        except Exception:
            db.session.rollback()
            flash('保存失败，请重试', 'danger')
            return redirect(url_for('academic.schedule_dates', sid=sid))
        log_operation(current_user, '更新', '学期日期配置', sid, ts.name, module='academic')
        if warnings:
            flash('已保存，但有提醒：' + '；'.join(warnings), 'warning')
        else:
            flash('学期日期与周次配置已保存', 'success')
        return redirect(url_for('academic.schedule_dates', sid=sid))
    return render_template('academic/schedule_dates.html',
                           ts=ts, warnings=warnings,
                           schedules=svc.list_schedules(),
                           can_edit=_editable(ts))


@bp.route('/schedule/<int:sid>/calendar')
@login_required
@perm_required('academic.view')
def schedule_calendar(sid):
    """学期校历页（周次 x 日期矩阵，标记当前周/周末）"""
    ts = _get_schedule_or_404(sid)
    cal = tsvc.get_school_calendar(sid)
    return render_template('academic/schedule_calendar.html',
                           ts=ts, cal=cal,
                           schedules=svc.list_schedules(),
                           can_edit=_editable(ts))


@bp.route('/schedule/<int:sid>/copy', methods=['GET', 'POST'])
@login_required
@perm_required('academic.timetable')
def schedule_copy(sid):
    """学期交接：以该学期为底版创建新学期草稿（GET 表单预填 / POST 执行）"""
    ts = _get_schedule_or_404(sid)
    if request.method == 'POST':
        new_name = (request.form.get('new_name') or '').strip()
        new_year = (request.form.get('new_school_year') or '').strip()
        new_term = (request.form.get('new_term') or '').strip()
        start_date = _parse_date(request.form.get('start_date'))
        end_date = _parse_date(request.form.get('end_date'))
        try:
            total_weeks = int(request.form.get('total_weeks') or 0) or None
        except (ValueError, TypeError):
            total_weeks = None
        copy_entries = request.form.get('copy_entries') in ('1', 'on', 'true', 'True')
        copy_periods = request.form.get('copy_periods') in ('1', 'on', 'true', 'True')
        if not new_name or not new_year or not new_term:
            flash('新学期名称、学年、学期为必填项', 'danger')
            return redirect(url_for('academic.schedule_copy', sid=sid))
        warnings = tsvc.validate_term_dates(start_date, end_date, total_weeks or 20)
        try:
            ok, msg, new, stats = tsvc.copy_term_as_new_draft(
                sid, new_name, new_year, new_term,
                start_date=start_date, end_date=end_date, total_weeks=total_weeks,
                copy_entries=copy_entries, copy_periods=copy_periods,
                created_by=current_user)
        except Exception:
            db.session.rollback()
            flash('复制失败，请重试', 'danger')
            return redirect(url_for('academic.schedule_copy', sid=sid))
        if not ok:
            flash(msg, 'danger')
            return redirect(url_for('academic.schedule_copy', sid=sid))
        log_operation(current_user, '新增', '学期交接', new.id,
                      f'{ts.name} → {new_name}（节次 {stats["periods"]}，条目 {stats["entries"]}）',
                      module='academic')
        flash(f'{msg}：复制节次 {stats["periods"]} 条、课表条目 {stats["entries"]} 条'
              + ('（提醒：' + '；'.join(warnings) + '）' if warnings else ''),
              'warning' if warnings else 'success')
        return redirect(url_for('academic.schedule_manage'))
    # GET：预填建议
    suggest_year = ts.school_year
    suggest_term = '第二学期' if '第一' in (ts.term or '') else '第一学期'
    return render_template('academic/schedule_copy.html',
                           ts=ts, suggest_year=suggest_year, suggest_term=suggest_term)


@bp.route('/schedule/<int:sid>/set-current', methods=['POST'])
@login_required
@perm_required('academic.timetable')
def schedule_set_current(sid):
    """设为当前学期（互斥；draft 自动转 active）"""
    ts = _get_schedule_or_404(sid)
    try:
        svc.set_current_term(sid)
    except ValueError as e:
        flash(str(e), 'danger')
        return redirect(url_for('academic.schedule_manage'))
    log_operation(current_user, '更新', '当前学期', sid, ts.name, module='academic')
    flash(f'已将「{ts.name}」设为当前学期', 'success')
    return redirect(request.referrer or url_for('academic.schedule_manage'))


@bp.route('/schedule/<int:sid>/archive-term', methods=['POST'])
@login_required
@perm_required('academic.timetable')
def schedule_archive_term(sid):
    """归档学期（含快照记录，写一条 ScheduleVersion）"""
    ts = _get_schedule_or_404(sid)
    try:
        ok, msg = tsvc.archive_term(sid, operator=current_user)
    except Exception:
        db.session.rollback()
        flash('归档失败，请重试', 'danger')
        return redirect(url_for('academic.schedule_manage'))
    if not ok:
        flash(msg, 'danger')
        return redirect(url_for('academic.schedule_manage'))
    log_operation(current_user, '归档', '学期课表', sid, ts.name, module='academic')
    flash(msg + '（已生成归档快照，可追溯）', 'success')
    return redirect(url_for('academic.schedule_manage'))


@bp.route('/schedule/<int:sid>/usage')
@login_required
@perm_required('academic.timetable')
def schedule_usage(sid):
    """学期课表使用情况报告页"""
    ts = _get_schedule_or_404(sid)
    report = tsvc.get_term_usage_report(sid)
    return render_template('academic/schedule_usage.html',
                           ts=ts, report=report,
                           schedules=svc.list_schedules(),
                           can_edit=_editable(ts))


@bp.route('/schedule/<int:sid>/usage/export')
@login_required
@perm_required('academic.timetable')
def schedule_usage_export(sid):
    """导出学期使用情况报告 Excel"""
    ts = _get_schedule_or_404(sid)
    report = tsvc.get_term_usage_report(sid)
    try:
        buf = _build_usage_workbook(ts, report)
    except Exception:
        db.session.rollback()
        flash('导出失败，请重试', 'danger')
        return redirect(url_for('academic.schedule_usage', sid=sid))
    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    return send_file(buf, as_attachment=True,
                     download_name=f'{ts.name}_使用情况报告_{stamp}.xlsx',
                     mimetype=_XLSX_MIME)


@bp.route('/schedule/history')
@login_required
@perm_required('academic.view')
def schedule_history():
    """历史课表总览：按学年分组列出全部学期（含 archived）"""
    terms = tsvc.list_terms_with_stats()
    grouped = {}
    for t in terms:
        grouped.setdefault(t['school_year'] or '未分学年', []).append(t)
    # 学年倒序
    years = sorted(grouped.keys(), reverse=True)
    return render_template('academic/schedule_history.html',
                           grouped=grouped, years=years, total=len(terms))


@bp.route('/schedule/history/<int:sid>')
@login_required
@perm_required('academic.view')
def schedule_history_detail(sid):
    """历史学期课表入口页（大课表/年级/班级/教师 视图导航，只读）"""
    ts = _get_schedule_or_404(sid)
    periods = svc.get_periods(sid)
    grade_classes = svc.get_grade_class_list(sid)
    versions = tsvc.compare_term_versions(sid, limit=20)
    return render_template('academic/schedule_history_detail.html',
                           ts=ts, periods=periods, grade_classes=grade_classes,
                           versions=versions, schedules=svc.list_schedules())


@bp.route('/schedule/compare')
@login_required
@perm_required('academic.timetable')
def schedule_compare():
    """跨学期对比页（选择学期 A/B + 可选年级班级）"""
    schedules = svc.list_schedules()
    sid_a = request.args.get('sid_a', type=int)
    sid_b = request.args.get('sid_b', type=int)
    grade = (request.args.get('grade') or '').strip() or None
    class_name = (request.args.get('class') or '').strip() or None
    result = None
    if sid_a and sid_b:
        result = tsvc.compare_terms(sid_a, sid_b, grade=grade, class_name=class_name)
    return render_template('academic/schedule_compare.html',
                           schedules=schedules, sid_a=sid_a, sid_b=sid_b,
                           f_grade=grade or '', f_class=class_name or '',
                           result=result)


@bp.route('/schedule/compare/export')
@login_required
@perm_required('academic.timetable')
def schedule_compare_export():
    """导出跨学期对比结果 Excel"""
    sid_a = request.args.get('sid_a', type=int)
    sid_b = request.args.get('sid_b', type=int)
    if not sid_a or not sid_b:
        abort(400)
    grade = (request.args.get('grade') or '').strip() or None
    class_name = (request.args.get('class') or '').strip() or None
    result = tsvc.compare_terms(sid_a, sid_b, grade=grade, class_name=class_name)
    try:
        buf = _build_compare_workbook(result)
    except Exception:
        flash('导出失败，请重试', 'danger')
        return redirect(url_for('academic.schedule_compare',
                                sid_a=sid_a, sid_b=sid_b))
    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    return send_file(buf, as_attachment=True,
                     download_name=f'课表对比_{stamp}.xlsx', mimetype=_XLSX_MIME)


# ── 学期周期 JSON API ──────────────────────────────────────────────────────

@bp.route('/api/schedule/compare-data')
@login_required
@perm_required('academic.view')
def api_schedule_compare_data():
    """跨学期对比结果 JSON（query: sid_a/sid_b/grade/class_name）"""
    sid_a = request.args.get('sid_a', type=int)
    sid_b = request.args.get('sid_b', type=int)
    if not sid_a or not sid_b:
        return _json_err('缺少 sid_a / sid_b 参数')
    grade = (request.args.get('grade') or '').strip() or None
    class_name = (request.args.get('class_name') or '').strip() or None
    data = tsvc.compare_terms(sid_a, sid_b, grade=grade, class_name=class_name)
    return _json_ok(data)


@bp.route('/api/schedule/<int:sid>/week-calendar')
@login_required
@perm_required('academic.view')
def api_schedule_week_calendar(sid):
    """周次日历 JSON（供前端周次选择器）"""
    _get_schedule_or_404(sid)
    return _json_ok(tsvc.get_week_calendar(sid))


@bp.route('/api/schedule/current-context')
@login_required
@perm_required('academic.view')
def api_schedule_current_context():
    """当前学期上下文 JSON（schedule_id/week/weekday/date/period_text）"""
    ctx = tsvc.get_current_term_context()
    sched = ctx.get('schedule')
    d = ctx.get('date')
    return _json_ok({
        'schedule_id': ctx.get('schedule_id'),
        'schedule_name': sched.name if sched else None,
        'week': ctx.get('week'),
        'weekday': ctx.get('weekday'),
        'weekday_text': WEEKDAY_NAMES.get(ctx.get('weekday'), ''),
        'date': d.strftime('%Y-%m-%d') if d else None,
        'period_text': ctx.get('period_text'),
        'term': sched.to_dict() if sched else None,
    })


@bp.route('/api/schedule/terms-stats')
@login_required
@perm_required('academic.view')
def api_schedule_terms_stats():
    """学期列表 + 统计 JSON"""
    return _json_ok(tsvc.list_terms_with_stats())


# ── Excel 导出辅助（对比 / 使用报告） ────────────────────────────────────

def _build_compare_workbook(result):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    wb.remove(wb.active)
    hf = Font(bold=True, color='FFFFFF')
    center = Alignment(horizontal='center', vertical='center')

    def _sheet(name, rows, headers, color):
        ws = wb.create_sheet(name[:31])
        fill = PatternFill(start_color=color, end_color=color, fill_type='solid')
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=ci, value=h)
            c.font = hf; c.fill = fill; c.alignment = center
        for ri, row in enumerate(rows, 2):
            for ci, v in enumerate(row, 1):
                ws.cell(row=ri, column=ci, value=xl_safe(v))
        return ws

    ta = (result.get('term_a') or {}).get('name', 'A')
    tb = (result.get('term_b') or {}).get('name', 'B')
    base_h = ['年级', '班级', '星期', '节次', '学科', '教师', '教室', '周次']
    _sheet(f'新增({tb}有{ta}无)',
           [[r['grade'], r['class_name'], r['weekday_text'], r['period_number'],
             r['subject'], r['teacher_name'] or '', r['room'] or '', r['week_range'] or '']
            for r in result['added']], base_h, '2E7D32')
    _sheet(f'减少({ta}有{tb}无)',
           [[r['grade'], r['class_name'], r['weekday_text'], r['period_number'],
             r['subject'], r['teacher_name'] or '', r['room'] or '', r['week_range'] or '']
            for r in result['removed']], base_h, 'C62828')
    _sheet('变更',
           [[r['grade'], r['class_name'], r['weekday_text'], r['period_number'],
             r['field'], r['old'] or '', r['new'] or '']
            for r in result['changed']],
           ['年级', '班级', '星期', '节次', '字段', '原值', '新值'], 'F9A825')
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _build_usage_workbook(ts, report):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    hf = Font(bold=True, color='FFFFFF')
    fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    center = Alignment(horizontal='center', vertical='center')

    def _sheet(name, headers, rows):
        ws = wb.create_sheet(name[:31])
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=ci, value=h)
            c.font = hf; c.fill = fill; c.alignment = center
        for ri, row in enumerate(rows, 2):
            for ci, v in enumerate(row, 1):
                ws.cell(row=ri, column=ci, value=xl_safe(v))

    ws0 = wb.create_sheet('概览')
    ws0.append(xl_row(['学期', ts.name]))
    ws0.append(['日期区间', ts.period_text()])
    ws0.append(['条目总数', report['total_entries']])
    ws0.append(['班级数', report['total_classes']])
    ws0.append(['学科数', report['total_subjects']])
    ws0.append(['教师数', report['total_teachers']])
    for c in ws0['A']:
        c.font = Font(bold=True)
    _sheet('班级周课时', ['班级', '周课时数'],
           [[r['class'], r['hours']] for r in report['class_hours']])
    _sheet('学科节数', ['学科', '总节数'],
           [[r['subject'], r['count']] for r in report['subject_counts']])
    _sheet('教师课时Top', ['教师', '总节数'],
           [[r['teacher'], r['count']] for r in report['teacher_top']])
    _sheet('节次类型分布', ['类型', '节数'],
           [[k, v] for k, v in report['type_dist'].items()])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
