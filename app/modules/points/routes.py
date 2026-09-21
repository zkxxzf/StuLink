# StuLink v1.17.0 2026-09-20
# 积分管理：单页记录（独立库 points.db）+ 范围权限过滤 + 批量导入导出 + 可视化
# Copyright (c) 2026 zkxxzf. Apache License 2.0
import io
from datetime import date, datetime

from flask import Blueprint, render_template, request, jsonify, abort, send_file
from flask_login import login_required, current_user
from sqlalchemy import func

from app.extensions import db
from app.models import Student, UserClassLink, PointRecord, PointRuleTemplate
from app.utils.decorators import perm_required
from app.utils.helpers import log_operation, get_graduated_grades

bp = Blueprint('points', __name__, url_prefix='/points')

CATEGORIES = ['纪律', '学习', '卫生', '活动', '其他']


# ==================== 数据范围（复用权限组 scope_type 体系） ====================

def _scope():
    """返回 (scope_type, grade, 范围明细 or None)

    范围明细：scope='class' 时为班级白名单 [(grade, class_name)]；
    scope='user_grades' 时为授权年级清单（v1.9.2 用户级数据范围，优先于组级规则）。
    """
    from app.modules.grades.services.scope import user_grade_scope
    ug = user_grade_scope(current_user)
    if ug is not None:
        return 'user_grades', '', ug
    pg = current_user.permission_group
    scope = pg.scope_type if pg else 'none'
    grade = current_user.grade or ''
    classes = None
    if scope == 'class':
        links = UserClassLink.query.filter_by(user_id=current_user.id).all()
        classes = [(l.grade, l.class_name) for l in links]
    return scope, grade, classes


def _apply_scope(q, scope, grade, classes):
    """把范围条件应用到 PointRecord 查询（classes：班级白名单或授权年级清单）"""
    if scope == 'school':
        return q
    if scope == 'user_grades':
        return q.filter(PointRecord.grade.in_(classes or []))
    if scope == 'grade' and grade:
        return q.filter_by(grade=grade)
    if scope == 'class' and classes:
        from sqlalchemy import or_, and_
        conds = [and_(PointRecord.grade == g, PointRecord.class_name == c)
                 for g, c in classes]
        return q.filter(or_(*conds))
    # none / 无范围：什么都看不到
    return q.filter(PointRecord.id == -1)


def _in_scope(student, scope, grade, classes):
    if scope == 'school':
        return True
    if scope == 'user_grades':
        return student.grade in (classes or [])
    if scope == 'grade':
        return student.grade == grade
    if scope == 'class':
        return (student.grade, student.class_name) in (classes or [])
    return False


# ==================== 页面 ====================

@bp.route('/')
@login_required
@perm_required('points.view')
def index():
    return render_template('points/index.html')


@bp.route('/api/options')
@login_required
@perm_required('points.view')
def api_options():
    scope, grade, classes = _scope()
    gds = set(get_graduated_grades())
    if scope == 'school':
        rows = (Student.query.with_entities(Student.grade).distinct().all())
        grades = sorted({r[0] for r in rows if r[0] and r[0] not in gds},
                        key=lambda g: int(''.join(filter(str.isdigit, g)) or 0), reverse=True)
        class_map = {}
        for g in grades:
            cls = Student.query.filter_by(grade=g).with_entities(Student.class_name).distinct().all()
            class_map[g] = sorted([r[0] for r in cls if r[0] and r[0].endswith('班')
                                   and r[0][:-1].isdigit()],
                                  key=lambda c: int(c[:-1]))
    elif scope == 'grade' and grade:
        grades = [grade]
        cls = Student.query.filter_by(grade=grade).with_entities(Student.class_name).distinct().all()
        class_map = {grade: sorted([r[0] for r in cls if r[0] and r[0].endswith('班')
                                    and r[0][:-1].isdigit()], key=lambda c: int(c[:-1]))}
    elif scope == 'class':
        grades = sorted({g for g, _c in (classes or [])}, reverse=True)
        class_map = {}
        for g, c in (classes or []):
            class_map.setdefault(g, [])
            if c not in class_map[g]:
                class_map[g].append(c)
        for g in class_map:
            class_map[g].sort(key=lambda c: int(c[:-1]) if c[:-1].isdigit() else 99)
    else:
        grades, class_map = [], {}
    # v1.17.0：导入/导出按钮按各自权限显示（路由已分别要求 points.import / points.export，
    # 否则会出现"按钮可见、点进去 403"）
    return jsonify(success=True, data={
        'scope': scope, 'grades': grades, 'classes': class_map,
        'categories': CATEGORIES,
        'can_edit': current_user.has_perm('points.edit'),
        'can_import': current_user.has_perm('points.import'),
        'can_export': current_user.has_perm('points.export'),
        'today': date.today().strftime('%Y-%m-%d'),
    })


# ==================== 列表与汇总 ====================

def _filter_query():
    scope, grade, classes = _scope()
    q = PointRecord.query
    q = _apply_scope(q, scope, grade, classes)
    g = request.args.get('grade', '').strip()
    c = request.args.get('class_name', '').strip()
    kw = request.args.get('q', '').strip()
    d1 = request.args.get('date_from', '').strip()
    d2 = request.args.get('date_to', '').strip()
    if g:
        q = q.filter_by(grade=g)
    if c:
        q = q.filter_by(class_name=c)
    if kw:
        q = q.filter(db.or_(PointRecord.student_no.contains(kw),
                            PointRecord.student_name.contains(kw)))
    if d1:
        q = q.filter(PointRecord.recorded_at >= date.fromisoformat(d1))
    if d2:
        q = q.filter(PointRecord.recorded_at <= date.fromisoformat(d2))
    return q


@bp.route('/api/list')
@login_required
@perm_required('points.view')
def api_list():
    q = _filter_query()
    rows = q.order_by(PointRecord.recorded_at.desc(), PointRecord.id.desc()) \
            .limit(500).all()
    total_points = sum(r.points for r in rows)
    students = len({r.student_no for r in rows})
    return jsonify(success=True, data={
        'records': [r.to_dict() for r in rows],
        'stat': {'count': len(rows), 'students': students, 'total_points': total_points},
    })


@bp.route('/api/summary')
@login_required
@perm_required('points.view')
def api_summary():
    q = _filter_query()
    rows = (q.with_entities(PointRecord.student_no, PointRecord.student_name,
                            PointRecord.grade, PointRecord.class_name,
                            func.sum(PointRecord.points).label('pts'),
                            func.count(PointRecord.id).label('cnt'))
            .group_by(PointRecord.student_no)
            .order_by(func.sum(PointRecord.points).desc())
            .limit(500).all())
    return jsonify(success=True, data=[{
        'student_no': r[0], 'student_name': r[1], 'grade': r[2],
        'class_name': r[3], 'points': int(r[4] or 0), 'count': int(r[5] or 0),
    } for r in rows])


@bp.route('/api/students')
@login_required
@perm_required('points.view')
def api_students():
    """录入时按学号/姓名搜索学生（限当前用户范围）"""
    kw = request.args.get('q', '').strip()
    if len(kw) < 1:
        return jsonify(success=True, data=[])
    scope, grade, classes = _scope()
    q = Student.query
    if scope == 'school':
        pass
    elif scope == 'grade' and grade:
        q = q.filter_by(grade=grade)
    elif scope == 'class' and classes:
        from sqlalchemy import or_, and_
        conds = [and_(Student.grade == g, Student.class_name == c) for g, c in classes]
        q = q.filter(or_(*conds))
    else:
        return jsonify(success=True, data=[])
    rows = q.filter(db.or_(Student.student_number.contains(kw),
                           Student.name.contains(kw))) \
            .order_by(Student.grade, Student.class_name, Student.student_number) \
            .limit(20).all()
    return jsonify(success=True, data=[{
        'student_no': str(s.student_number), 'name': s.name,
        'grade': s.grade, 'class_name': s.class_name,
    } for s in rows])


# ==================== 录入 / 编辑 / 删除 ====================

def _parse_payload():
    data = request.get_json(silent=True) or {}
    student_no = str(data.get('student_no') or '').strip()
    try:
        points = int(data.get('points'))
    except (TypeError, ValueError):
        return None, '分值必须是整数'
    reason = (data.get('reason') or '').strip()
    category = (data.get('category') or '').strip()[:20]
    remark = (data.get('remark') or '').strip()[:200]
    rd = (data.get('recorded_at') or '').strip() or date.today().isoformat()
    if not student_no:
        return None, '请选择学生'
    if points == 0 or abs(points) > 100:
        return None, '分值需在 -100 ~ 100 之间且不为 0'
    if not reason:
        return None, '请填写事由'
    try:
        recorded_at = date.fromisoformat(rd)
    except ValueError:
        return None, '日期格式不正确'
    return {'student_no': student_no, 'points': points, 'reason': reason,
            'category': category, 'remark': remark, 'recorded_at': recorded_at}, None


@bp.route('/api/record', methods=['POST'])
@login_required
@perm_required('points.edit')
def api_record_create():
    payload, err = _parse_payload()
    if err:
        return jsonify(success=False, message=err), 400
    scope, grade, classes = _scope()
    student = Student.query.filter_by(student_number=payload['student_no']).first()
    if not student:
        return jsonify(success=False, message='学号未匹配到学生'), 400
    if not _in_scope(student, scope, grade, classes):
        return jsonify(success=False, message='该学生不在您的管理范围内'), 403
    rec = PointRecord(
        student_no=str(student.student_number), student_name=student.name,
        grade=student.grade, class_name=student.class_name,
        points=payload['points'], category=payload['category'],
        reason=payload['reason'], remark=payload['remark'],
        recorded_at=payload['recorded_at'],
        operator_id=current_user.id, operator_name=current_user.real_name)
    db.session.add(rec)
    db.session.commit()
    log_operation(current_user, '录入', '积分', rec.id,
                  f"{student.name} {payload['points']:+d}分 {payload['reason']}",
                  module='points')
    return jsonify(success=True, message='已记录', data=rec.to_dict())


@bp.route('/api/record/<int:rid>', methods=['POST'])
@login_required
@perm_required('points.edit')
def api_record_update(rid):
    rec = PointRecord.query.get_or_404(rid)
    scope, grade, classes = _scope()
    student = Student.query.filter_by(student_number=rec.student_no).first()
    if student and not _in_scope(student, scope, grade, classes):
        return jsonify(success=False, message='该记录不在您的管理范围内'), 403
    payload, err = _parse_payload()
    if err:
        return jsonify(success=False, message=err), 400
    rec.points = payload['points']
    rec.category = payload['category']
    rec.reason = payload['reason']
    rec.remark = payload['remark']
    rec.recorded_at = payload['recorded_at']
    db.session.commit()
    return jsonify(success=True, message='已更新', data=rec.to_dict())


@bp.route('/api/record/<int:rid>/delete', methods=['POST'])
@login_required
@perm_required('points.edit')
def api_record_delete(rid):
    rec = PointRecord.query.get_or_404(rid)
    scope, grade, classes = _scope()
    student = Student.query.filter_by(student_number=rec.student_no).first()
    if student and not _in_scope(student, scope, grade, classes):
        return jsonify(success=False, message='该记录不在您的管理范围内'), 403
    db.session.delete(rec)
    db.session.commit()
    log_operation(current_user, '删除', '积分', rid, f'{rec.student_name} {rec.points:+d}', module='points')
    return jsonify(success=True, message='已删除')


# ==================== 批量导入 ====================

@bp.route('/import')
@login_required
@perm_required('points.import')
def import_page():
    return render_template('points/import.html')


@bp.route('/template')
@login_required
@perm_required('points.import')
def download_template():
    from app.modules.points.services.import_service import generate_template
    out = generate_template()
    return send_file(out, as_attachment=True,
                     download_name=f'积分导入模板_{date.today().strftime("%Y%m%d")}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@bp.route('/import/upload', methods=['POST'])
@login_required
@perm_required('points.import')
def import_upload():
    from app.modules.points.services.import_service import parse_points_excel, validate_records
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify(success=False, message='请选择文件'), 400
    if not f.filename.endswith(('.xlsx', '.xls')):
        return jsonify(success=False, message='仅支持 .xlsx / .xls 格式'), 400
    try:
        result = parse_points_excel(f.stream)
    except Exception as e:
        return jsonify(success=False, message=f'解析失败：{str(e)}'), 400
    valid, invalid = validate_records(result['rows'])
    errors = result['errors'] + [{'line': r.get('line'), 'student_no': r.get('student_no'),
                                  'reason': r.get('error')} for r in invalid]
    return jsonify(success=True, data={
        'valid': valid,
        'errors': errors,
        'stats': {
            'total': result['stats']['total_rows'],
            'ok': len(valid),
            'error': len(errors),
        }
    })


@bp.route('/import/confirm', methods=['POST'])
@login_required
@perm_required('points.import')
def import_confirm():
    """确认导入

    v1.17.0（PR#5 安全审查 M2）：前端提交的 rows 一律不信任，
    写入前用 sanitize_rows 重校验分值/类别/学号与数据范围。
    """
    from app.modules.points.services.import_service import import_records, sanitize_rows
    payload = request.get_json(silent=True) or {}
    rows = payload.get('rows', [])
    if not rows:
        return jsonify(success=False, message='没有可导入的数据'), 400

    scope, grade, classes = _scope()
    clean, errors, skipped = sanitize_rows(
        rows, in_scope=lambda s: _in_scope(s, scope, grade, classes))
    if not clean:
        return jsonify(success=False,
                       message='没有通过校验的合法数据',
                       data={'errors': errors, 'skipped': skipped}), 400

    count = import_records(clean, current_user.id, current_user.real_name)
    log_operation(current_user, '批量导入', '积分', None,
                  f'导入 {count} 条积分记录（跳过越权 {skipped} 行、'
                  f'校验失败 {len(errors)} 行）', module='points')
    msg = f'成功导入 {count} 条记录'
    if errors:
        msg += f'；{len(errors)} 行未通过校验'
    if skipped:
        msg += f'；{skipped} 行不在您的管理范围内已跳过'
    return jsonify(success=True, message=msg,
                   data={'errors': errors, 'skipped': skipped})


# ==================== 导出 ====================

@bp.route('/export')
@login_required
@perm_required('points.export')
def export_excel():
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from app.utils.export_helpers import xl_safe
    q = _filter_query()
    rows = q.order_by(PointRecord.recorded_at.desc(), PointRecord.id.desc()).limit(5000).all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '积分记录'
    headers = ['日期', '学号', '姓名', '年级', '班级', '分值', '类别', '事由', '备注', '录入人']
    hf = Font(bold=True, color='FFFFFF')
    hfl = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    tb = Border(left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin'))
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = hf
        c.fill = hfl
        c.alignment = Alignment(horizontal='center')
        c.border = tb
    for ri, r in enumerate(rows, 2):
        vals = [r.recorded_at.strftime('%Y-%m-%d') if r.recorded_at else '',
                r.student_no, r.student_name, r.grade, r.class_name,
                r.points, r.category or '', r.reason, r.remark or '', r.operator_name or '']
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=ri, column=ci, value=xl_safe(v))
            c.border = tb
    widths = [12, 15, 10, 8, 10, 8, 8, 25, 20, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    return send_file(out, as_attachment=True,
                     download_name=f'积分记录_{timestamp}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ==================== 规则模板 ====================

@bp.route('/rules')
@login_required
@perm_required('points.rules')
def rules_page():
    return render_template('points/rules.html')


@bp.route('/api/rules')
@login_required
@perm_required('points.rules')
def api_rules_list():
    rules = PointRuleTemplate.query.order_by(PointRuleTemplate.category,
                                              PointRuleTemplate.default_points.desc()).all()
    return jsonify(success=True, data=[r.to_dict() for r in rules])


@bp.route('/api/rules', methods=['POST'])
@login_required
@perm_required('points.rules')
def api_rules_create():
    data = request.get_json(silent=True) or {}
    name = (data.get('name') or '').strip()[:50]
    category = (data.get('category') or '').strip()[:20]
    if not name or not category:
        return jsonify(success=False, message='名称和类别不能为空'), 400
    try:
        points = int(data.get('default_points', 0))
    except (TypeError, ValueError):
        return jsonify(success=False, message='分值必须是整数'), 400
    if points == 0 or abs(points) > 100:
        return jsonify(success=False, message='分值需在 -100 ~ 100 之间且不为 0'), 400
    rule = PointRuleTemplate(name=name, category=category, default_points=points,
                              description=(data.get('description') or '').strip()[:200])
    db.session.add(rule)
    db.session.commit()
    return jsonify(success=True, message='已创建', data=rule.to_dict())


@bp.route('/api/rules/<int:rule_id>/edit', methods=['POST'])
@login_required
@perm_required('points.rules')
def api_rules_edit(rule_id):
    rule = PointRuleTemplate.query.get_or_404(rule_id)
    data = request.get_json(silent=True) or {}
    name = (data.get('name') or '').strip()[:50]
    category = (data.get('category') or '').strip()[:20]
    if not name or not category:
        return jsonify(success=False, message='名称和类别不能为空'), 400
    try:
        points = int(data.get('default_points', 0))
    except (TypeError, ValueError):
        return jsonify(success=False, message='分值必须是整数'), 400
    if points == 0 or abs(points) > 100:
        return jsonify(success=False, message='分值需在 -100 ~ 100 之间且不为 0'), 400
    rule.name = name
    rule.category = category
    rule.default_points = points
    rule.description = (data.get('description') or '').strip()[:200]
    db.session.commit()
    return jsonify(success=True, message='已更新', data=rule.to_dict())


@bp.route('/api/rules/<int:rule_id>/toggle', methods=['POST'])
@login_required
@perm_required('points.rules')
def api_rules_toggle(rule_id):
    rule = PointRuleTemplate.query.get_or_404(rule_id)
    rule.is_active = not rule.is_active
    db.session.commit()
    return jsonify(success=True, message='已' + ('启用' if rule.is_active else '禁用'),
                   data=rule.to_dict())


@bp.route('/api/rules/<int:rule_id>/delete', methods=['POST'])
@login_required
@perm_required('points.rules')
def api_rules_delete(rule_id):
    rule = PointRuleTemplate.query.get_or_404(rule_id)
    db.session.delete(rule)
    db.session.commit()
    return jsonify(success=True, message='已删除')


# ==================== 可视化 API ====================

@bp.route('/api/trend')
@login_required
@perm_required('points.view')
def api_trend():
    """积分趋势数据：按日期聚合"""
    q = _filter_query()
    rows = q.with_entities(PointRecord.recorded_at,
                           func.sum(PointRecord.points).label('total'),
                           func.count(PointRecord.id).label('cnt')) \
            .group_by(PointRecord.recorded_at) \
            .order_by(PointRecord.recorded_at) \
            .limit(90).all()
    return jsonify(success=True, data=[{
        'date': r[0].strftime('%Y-%m-%d') if r[0] else '',
        'total': int(r[1] or 0),
        'count': int(r[2] or 0),
    } for r in rows])


@bp.route('/api/category-distribution')
@login_required
@perm_required('points.view')
def api_category_distribution():
    """类别分布数据：按 category 聚合"""
    q = _filter_query()
    rows = q.with_entities(PointRecord.category,
                           func.sum(PointRecord.points).label('total'),
                           func.count(PointRecord.id).label('cnt')) \
            .group_by(PointRecord.category) \
            .order_by(func.count(PointRecord.id).desc()).all()
    return jsonify(success=True, data=[{
        'category': r[0] or '未分类',
        'total': int(r[1] or 0),
        'count': int(r[2] or 0),
    } for r in rows])


@bp.route('/api/class-ranking')
@login_required
@perm_required('points.view')
def api_class_ranking():
    """班级积分排名：按班级聚合"""
    q = _filter_query()
    rows = q.with_entities(PointRecord.grade, PointRecord.class_name,
                           func.sum(PointRecord.points).label('total'),
                           func.count(PointRecord.id).label('cnt')) \
            .group_by(PointRecord.grade, PointRecord.class_name) \
            .order_by(func.sum(PointRecord.points).desc()) \
            .limit(20).all()
    return jsonify(success=True, data=[{
        'grade': r[0] or '',
        'class_name': r[1] or '',
        'total': int(r[2] or 0),
        'count': int(r[3] or 0),
    } for r in rows])
