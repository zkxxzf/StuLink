# StuLink v1.19.49.44 2026-10-20
# 教育事业统计基表导出（Excel）
# Copyright (c) 2026 zkxxzf. Apache License 2.0
"""把计算结果写成与基表同结构的 Excel（多 sheet）。

只做"把算出来的数填进对应格子"：行列含义与基表一致，便于人工对照后誊抄/上报。
未填写的属性写空串（不是 0），避免把"没填"当成"确实是 0"。
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.utils.export_helpers import xl_safe

_HEAD_FILL = PatternFill(start_color='4472C4', end_color='4472C4',
                         fill_type='solid')
_HEAD_FONT = Font(bold=True, color='FFFFFF')


def _title(ws, text, sub, ncols):
    c = ws.cell(row=1, column=1, value=text)
    c.font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.cell(row=2, column=1, value=sub)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)


def _header(ws, row_idx, headers):
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=row_idx, column=ci, value=h)
        c.font = _HEAD_FONT
        c.fill = _HEAD_FILL
        c.alignment = Alignment(horizontal='center', vertical='center',
                                wrap_text=True)


def _put(ws, r, values):
    for ci, v in enumerate(values, 1):
        cell = ws.cell(row=r, column=ci,
                       value=xl_safe(v) if isinstance(v, str) else v)
        cell.alignment = Alignment(horizontal='center', vertical='center',
                                   wrap_text=True)


def _n(v):
    """None → 空串（未填 ≠ 0）。"""
    return '' if v is None else v


def _widths(ws, first, last, w=11):
    for ci in range(first, last + 1):
        ws.column_dimensions[ws.cell(row=3, column=ci).column_letter].width = w


def build_workbook(school_year, stat):
    """stat 见 edu_stat_service.collect()。"""
    wb = Workbook()
    wb.remove(wb.active)

    # ── 教基2108 班额班数（列＝合计/一年级/二年级/三年级）──────────────
    cs = stat['class_size']
    ws = wb.create_sheet('教基2108_班额班数')
    _title(ws, '普通高中班额班数情况',
           f'表号：教基2108 | 统计时点：{school_year}学年', 6)
    _header(ws, 3, ['指标名称', '代码', '合计', '一年级', '二年级', '三年级'])
    _put(ws, 4, ['总计', '01', cs['class_count'],
                 *[cs['total'][g] for g in cs['grades']]])
    for i, (name, cells) in enumerate(cs['rows'], 1):
        _put(ws, 4 + i, [name, '%02d' % (i + 1), sum(cells.values()),
                         *[cells[g] for g in cs['grades']]])
    ws.cell(row=4 + len(cs['rows']) + 2, column=1,
            value='注：一年级=%s，二年级=%s，三年级=%s（共 %d 名学生）'
                  % (*cs['grades'], cs['student_count']))
    ws.column_dimensions['A'].width = 18
    _widths(ws, 2, 6, 13)

    # ── 教基3118 分类型学生数 ──────────────────────────────────────────
    st = stat['student_types']
    gs = st['grades']
    ws = wb.create_sheet('教基3118_分类型学生数')
    _title(ws, '普通高中分类型学生数',
           f'表号：教基3118 | 统计时点：{school_year}学年', 9)
    _header(ws, 3, ['指标名称', '代码', '毕业生数', '招生数', '在校生数',
                    '一年级', '二年级', '三年级', '预计毕业生数'])
    rows = [
        ('总计', '01', st['graduated'], st['new_students'],
         st['enrolled_total']['total'],
         st['enrolled'][gs['g1']]['total'], st['enrolled'][gs['g2']]['total'],
         st['enrolled'][gs['g3']]['total'], st['graduates_expected']),
        ('#女', '02', st['graduated_female'], '', st['enrolled_total']['female'],
         st['enrolled'][gs['g1']]['female'], st['enrolled'][gs['g2']]['female'],
         st['enrolled'][gs['g3']]['female'], ''),
        ('#少数民族', '03', '', '', st['enrolled_total']['minority'],
         st['enrolled'][gs['g1']]['minority'], st['enrolled'][gs['g2']]['minority'],
         st['enrolled'][gs['g3']]['minority'], ''),
        ('#寄宿生', '04', '', '', st['enrolled_total']['boarding'],
         st['enrolled'][gs['g1']]['boarding'], st['enrolled'][gs['g2']]['boarding'],
         st['enrolled'][gs['g3']]['boarding'], ''),
    ]
    for i, vals in enumerate(rows):
        _put(ws, 4 + i, list(vals))
    ws.cell(row=9, column=1,
            value='注：一年级=%s，二年级=%s，三年级=%s，毕业生=%s'
                  % (gs['g1'], gs['g2'], gs['g3'], gs['graduated']))
    ws.column_dimensions['A'].width = 14
    _widths(ws, 2, 9, 13)

    # ── 教基3119 分年龄学生数 ──────────────────────────────────────────
    sa = stat['student_ages']
    ws = wb.create_sheet('教基3119_分年龄学生数')
    _title(ws, '高中阶段分年龄学生数',
           f'表号：教基3119 | 统计时点：{school_year}学年（年龄按 {sa["ref_date"]} 计算）', 4)
    _header(ws, 3, ['指标名称', '代码', '招生数', '在校生数'])
    _put(ws, 4, ['总计', '01', sum(r[1]['new'] for r in sa['rows']),
                 sum(r[1]['enrolled'] for r in sa['rows'])])
    for i, (name, cell) in enumerate(sa['rows'], 1):
        _put(ws, 4 + i, [name, '%02d' % (i + 1), cell['new'], cell['enrolled']])
    if sa['unknown']:
        ws.cell(row=5 + len(sa['rows']), column=1,
                value='注：%d 名学生缺身份证，未计入分年龄统计' % sa['unknown'])
    ws.column_dimensions['A'].width = 18
    _widths(ws, 2, 4, 13)

    # ── 教基3045 在校生中其他情况 ──────────────────────────────────────
    so = stat['student_other']
    ws = wb.create_sheet('教基3045_在校生其他情况')
    _title(ws, '在校生中其他情况',
           f'表号：教基3045 | 统计时点：{school_year}学年', 11)
    _header(ws, 3, ['指标名称', '代码', '中共党员', '共青团员', '民主党派',
                    '香港', '澳门', '台湾', '华侨', '少数民族', '残疾人'])
    _put(ws, 4, ['普通高中', '07', '', '', '', '', '', '', '',
                 so['minority'], ''])
    _put(ws, 5, ['#女', '08', '', '', '', '', '', '', '',
                 so['female_minority'], ''])
    ws.cell(row=7, column=1,
            value='注：系统无党员/团员/民主党派/港澳台/华侨/残疾人字段，'
                  '这些列需人工补填；少数民族由学生档案自动统计。')
    ws.column_dimensions['A'].width = 14
    _widths(ws, 2, 11, 10)

    # ── 教基4149 教职工情况 ────────────────────────────────────────────
    tm = stat['teacher_summary']
    ws = wb.create_sheet('教基4149_教职工情况')
    _title(ws, '中小学教职工情况',
           f'表号：教基4149 | 统计时点：{school_year}学年', 9)
    _header(ws, 3, ['指标名称', '代码', '教职工数', '专任教师', '行政人员',
                    '教辅人员', '工勤人员', '其他', '校外教师'])
    _put(ws, 4, ['总计', '01', tm['total'],
                 *[tm['by_type'][k] for k in
                   ('专任教师', '行政人员', '教辅人员', '工勤人员', '其他')], ''])
    _put(ws, 5, ['#女', '02', tm['female'], '', '', '', '', '', ''])
    _put(ws, 6, ['#少数民族', '03', tm['minority'], '', '', '', '', '', ''])
    _put(ws, 7, ['#在编人员', '04', tm['official'], '', '', '', '', '', ''])
    _put(ws, 8, ['#外籍人员', '05', tm['foreign'], '', '', '', '', '', ''])
    if tm['untagged_staff_type']:
        ws.cell(row=10, column=1,
                value='注：%d 名教师未填「人员类别」，已按有无学科推断'
                      '（有学科=专任教师）；建议在教师名单页补齐。'
                      % tm['untagged_staff_type'])
    ws.column_dimensions['A'].width = 16
    _widths(ws, 2, 9, 12)

    # ── 教基4153 专任教师分年龄（列＝年龄档，行＝总计/女/少数民族/职称）──
    ta = stat['teacher_ages']
    bands = [r[0] for r in ta['rows']]
    by_band = dict(ta['rows'])
    ws = wb.create_sheet('教基4153_专任教师分年龄')
    _title(ws, '基础教育专任教师分年龄情况',
           f'表号：教基4153 | 统计时点：{school_year}学年（年龄按 {ta["ref_date"]} 计算）',
           len(bands) + 3)
    _header(ws, 3, ['指标名称', '代码', '合计'] + bands)
    _put(ws, 4, ['普通高中', '28', sum(c['total'] for c in by_band.values())]
         + [by_band[b]['total'] for b in bands])
    _put(ws, 5, ['#女', '29', sum(c['female'] for c in by_band.values())]
         + [by_band[b]['female'] for b in bands])
    _put(ws, 6, ['#少数民族', '30', sum(c['minority'] for c in by_band.values())]
         + [by_band[b]['minority'] for b in bands])
    from app.modules.edustat.services.edu_stat_service import TITLE_LEVELS
    r = 7
    for title in TITLE_LEVELS:
        vals = [sum(c['titles'].get(title, 0) for c in by_band.values())]
        vals += [by_band[b]['titles'].get(title, 0) for b in bands]
        _put(ws, r, [title, '%02d' % (30 + r - 6), *vals])
        r += 1
    if ta['unknown']:
        ws.cell(row=r + 1, column=1,
                value='注：%d 名教师缺身份证，无法计算年龄，未计入本表'
                      % ta['unknown'])
    ws.column_dimensions['A'].width = 16
    _widths(ws, 2, len(bands) + 3, 12)

    # ── 教基4156 专任教师分课程分学历 ──────────────────────────────────
    tsub = stat['teacher_by_subject']
    ws = wb.create_sheet('教基4156_分课程分学历')
    _title(ws, '普通高中专任教师分课程、分学历情况',
           f'表号：教基4156 | 统计时点：{school_year}学年', len(tsub['columns']) + 2)
    _header(ws, 3, ['指标名称', '代码'] + tsub['columns'])
    for i, (lv, cells) in enumerate(tsub['rows'], 1):
        _put(ws, 3 + i, [lv, '%02d' % i] + [cells[c] for c in tsub['columns']])
    note_row = 4 + len(tsub['rows'])
    notes = []
    if tsub['unknown_education']:
        notes.append('%d 人未填学历（已单列「未填学历」行）' % tsub['unknown_education'])
    if tsub['unmapped_subject']:
        notes.append('%d 人学科无法映射到基表课程列（已归入「其他」）'
                     % tsub['unmapped_subject'])
    if notes:
        ws.cell(row=note_row, column=1, value='注：' + '；'.join(notes))
    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 8
    _widths(ws, 3, len(tsub['columns']) + 2, 9)

    # ── 校舍 / 资产（5170 + 5176 + 1102 达标）──────────────────────────
    asset = stat.get('asset')
    if asset:
        ws = wb.create_sheet('教基5170_5176_校舍资产')
        _title(ws, '校舍情况 / 资产等办学条件',
               f'表号：教基5170、教基5176 | 统计时点：{school_year}学年', 3)
        _header(ws, 3, ['指标名称', '单位', '本学年数量'])
        rows = [
            ('本学年校舍建筑面积（总计）', '平方米', asset.building_total),
            ('#C级危房', '平方米', asset.building_c_danger),
            ('#D级危房', '平方米', asset.building_d_danger),
            ('#租用外单位', '平方米', asset.building_rented),
            ('一、教学及辅助用房', '平方米', asset.room_teaching),
            ('  普通教室', '平方米', asset.room_classroom),
            ('  专用教室', '平方米', asset.room_special),
            ('    理化生实验室', '平方米', asset.room_lab),
            ('  公共教学用房', '平方米', asset.room_public),
            ('    图书阅览室', '平方米', asset.room_library),
            ('    室内体育用房', '平方米', asset.room_sports_indoor),
            ('    心理辅导室', '平方米', asset.room_psych),
            ('二、行政办公用房', '平方米', asset.room_office),
            ('  教师办公室', '平方米', asset.room_teacher_office),
            ('三、生活用房', '平方米', asset.room_living),
            ('  教工值班宿舍', '平方米', asset.room_staff_dorm),
            ('  教师周转宿舍', '平方米', asset.room_teacher_turnover),
            ('  学生宿舍', '平方米', asset.room_student_dorm),
            ('  学生餐厅', '平方米', asset.room_canteen),
            ('  厕所', '平方米', asset.room_toilet),
            ('四、其他用房', '平方米', asset.room_other),
            ('#被外单位租（借）用', '平方米', asset.room_lent_out),
            ('占地面积', '平方米', asset.land_area),
            ('#绿化用地面积', '平方米', asset.green_area),
            ('#运动场地面积', '平方米', asset.sports_field_area),
            ('校园足球场（11人制）', '个', asset.football_11),
            ('校园足球场（7人制）', '个', asset.football_7),
            ('校园足球场（5人制）', '个', asset.football_5),
            ('图书', '册', asset.books),
            ('数字终端数', '台', asset.terminal_total),
            ('#教师终端数', '台', asset.terminal_teacher),
            ('#学生终端数', '台', asset.terminal_student),
            ('教室', '间', asset.classroom_total),
            ('#网络多媒体教室', '间', asset.classroom_multimedia),
            ('固定资产总值', '万元', asset.asset_fixed),
            ('#教学仪器设备资产值', '万元', asset.asset_teaching_equipment),
        ]
        for i, (label, unit, val) in enumerate(rows, 4):
            _put(ws, i, [label, unit, _n(val)])
        base = 4 + len(rows) + 1
        _put(ws, base, ['【教基1102 达标项】', '', ''])
        extra = [
            ('体育运动场（馆）面积是否达标', asset.ok_sports_venue),
            ('体育器械配备是否达标', asset.ok_sports_equipment),
            ('音乐器材配备是否达标', asset.ok_music_equipment),
            ('美术器材配备是否达标', asset.ok_art_equipment),
            ('理科实验仪器是否达标（中学）', asset.ok_science_lab),
        ]
        for i, (label, val) in enumerate(extra, base + 1):
            _put(ws, i, [label, '', '' if val is None else ('是' if val else '否')])
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 10
        ws.column_dimensions['C'].width = 14

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
