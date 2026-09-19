"""考勤管理服务层"""
import io
from datetime import date, datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from sqlalchemy import func

from app.extensions import db
from app.models.academic import AttendanceRecord, ATTENDANCE_STATUS
from app.models.student import Student
from app.models.user_class_link import UserClassLink


def record_attendance(records_list, recorded_by):
    """批量记录考勤
    records_list: [{student_no, student_name, grade, class_name, attend_date, period, status, remark}, ...]
    """
    created = []
    for r in records_list:
        attend_date = r.get('attend_date')
        if isinstance(attend_date, str):
            try:
                attend_date = date.fromisoformat(attend_date)
            except ValueError:
                continue
        rec = AttendanceRecord(
            student_no=r.get('student_no', ''),
            student_name=r.get('student_name', ''),
            grade=r.get('grade', ''),
            class_name=r.get('class_name', ''),
            attend_date=attend_date,
            period=r.get('period') or None,
            status=r.get('status', 'present'),
            recorded_by=recorded_by,
            remark=(r.get('remark') or '').strip() or None,
        )
        db.session.add(rec)
        created.append(rec)
    db.session.commit()
    return created


def get_attendance(class_name=None, grade=None, date_from=None, date_to=None,
                   student_no=None, page=1, per_page=30):
    """查询考勤记录（分页+筛选）"""
    q = AttendanceRecord.query
    if class_name and grade:
        q = q.filter_by(class_name=class_name, grade=grade)
    elif class_name:
        q = q.filter_by(class_name=class_name)
    if grade:
        q = q.filter_by(grade=grade)
    if date_from:
        q = q.filter(AttendanceRecord.attend_date >= date_from)
    if date_to:
        q = q.filter(AttendanceRecord.attend_date <= date_to)
    if student_no:
        q = q.filter(AttendanceRecord.student_no == student_no)
    q = q.order_by(AttendanceRecord.attend_date.desc(),
                   AttendanceRecord.class_name, AttendanceRecord.student_no)
    pagination = q.paginate(page=page, per_page=per_page, error_out=False)
    return pagination


def get_attendance_stats(class_name, grade=None, month=None):
    """月度考勤统计
    返回：{total_students, present_days, absent_days, late_days, leave_days, rate}
    """
    if month is None:
        month = date.today().strftime('%Y-%m')
    year, mon = int(month[:4]), int(month[5:7])
    start = date(year, mon, 1)
    if mon == 12:
        end = date(year + 1, 1, 1)
    else:
        end = date(year, mon + 1, 1)
    end = end - timedelta(days=1)

    q = db.session.query(
        AttendanceRecord.status, func.count(AttendanceRecord.id)
    ).filter(
        AttendanceRecord.class_name == class_name,
        AttendanceRecord.attend_date >= start,
        AttendanceRecord.attend_date <= end,
    )
    if grade:
        q = q.filter(AttendanceRecord.grade == grade)
    # v1.16.0 性能改造：旧版 q.all() 把所有考勤记录载入内存再用 Python sum 计数，
    # 改为一次 GROUP BY status 在数据库端聚合（避免大量行传输），口径一致。
    status_map = {st: cnt for st, cnt in q.group_by(AttendanceRecord.status).all()}

    present = status_map.get('present', 0)
    absent = status_map.get('absent', 0)
    late = status_map.get('late', 0)
    leave = status_map.get('leave', 0)
    total = sum(status_map.values()) or 1

    sq = Student.query.filter_by(class_name=class_name)
    if grade:
        sq = sq.filter_by(grade=grade)
    total_students = sq.count()

    rate = round(present / total * 100, 1) if total else 0
    return {
        'total_students': total_students,
        'present_days': present,
        'absent_days': absent,
        'late_days': late,
        'leave_days': leave,
        'rate': rate,
    }


def get_student_attendance(student_no, month=None):
    """单个学生考勤记录"""
    if month is None:
        month = date.today().strftime('%Y-%m')
    year, mon = int(month[:4]), int(month[5:7])
    start = date(year, mon, 1)
    if mon == 12:
        end = date(year + 1, 1, 1)
    else:
        end = date(year, mon + 1, 1)
    end = end - timedelta(days=1)

    records = (AttendanceRecord.query
               .filter_by(student_no=student_no)
               .filter(AttendanceRecord.attend_date >= start,
                       AttendanceRecord.attend_date <= end)
               .order_by(AttendanceRecord.attend_date.desc())
               .all())
    return records


def get_class_students(class_name, grade=None):
    """获取管辖班级的学生列表"""
    q = Student.query.filter_by(class_name=class_name)
    if grade:
        q = q.filter_by(grade=grade)
    return q.order_by(Student.student_number).all()


def get_teacher_classes(user_id):
    """获取教师管辖的班级列表"""
    links = UserClassLink.query.filter_by(user_id=user_id).all()
    return [(l.grade, l.class_name) for l in links]


def export_attendance(class_name, date_from, date_to, grade=None):
    """导出考勤 Excel"""
    q = AttendanceRecord.query.filter_by(class_name=class_name)
    if grade:
        q = q.filter_by(grade=grade)
    if date_from:
        q = q.filter(AttendanceRecord.attend_date >= date_from)
    if date_to:
        q = q.filter(AttendanceRecord.attend_date <= date_to)
    records = q.order_by(AttendanceRecord.attend_date, AttendanceRecord.student_no).all()

    status_map = dict(ATTENDANCE_STATUS)

    wb = Workbook()
    ws = wb.active
    ws.title = '考勤记录'

    hf = Font(bold=True, color='FFFFFF')
    hfl = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    tb = Border(left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin'))

    headers = ['日期', '学号', '姓名', '年级', '班级', '节次', '状态', '备注']
    widths = [12, 14, 10, 8, 8, 6, 8, 20]

    for ci, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = hf
        c.fill = hfl
        c.alignment = Alignment(horizontal='center')
        c.border = tb
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(ci)].width = w

    for ri, r in enumerate(records, 2):
        row = [
            r.attend_date.strftime('%Y-%m-%d') if r.attend_date else '',
            r.student_no, r.student_name, r.grade, r.class_name,
            f'第{r.period}节' if r.period else '全天',
            status_map.get(r.status, r.status),
            r.remark or '',
        ]
        for ci, v in enumerate(row, 1):
            c = ws.cell(row=ri, column=ci, value=v)
            c.border = tb

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
