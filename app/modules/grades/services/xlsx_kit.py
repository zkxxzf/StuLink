# -*- coding: utf-8 -*-
"""成绩名单类 Excel 导出的公共件（2026-10-19）。

表彰名单 / 临界生名单两个导出共用：文件名清洗、样式常量、口径注释行、
表体写入（含公式注入防护）。本模块只负责「写表」，不含任何业务口径。
"""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.utils.export_helpers import xl_safe
from app.utils.text_guard import safe_download_name

DASH = '—'

HEAD_FONT = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill(start_color='1F4E79', end_color='1F4E79',
                        fill_type='solid')
NOTE_FONT = Font(size=9, color='595959')
NOTE_ALIGN = Alignment(horizontal='left', vertical='center', wrap_text=True)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
RIGHT = Alignment(horizontal='right', vertical='center')
_SIDE = Side(style='thin', color='B7C3D0')
BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)


def safe_fname(name):
    """下载文件名清洗：统一走 `text_guard.safe_download_name`（L-9 头注入防护）。"""
    return safe_download_name(name, '未命名')


def write_note(ws, text, ncol):
    """首行写口径说明：跨全表合并、灰字小号、自动换行（表头从第 2 行起）。"""
    cell = ws.cell(row=1, column=1, value=xl_safe(text))
    cell.font, cell.alignment = NOTE_FONT, NOTE_ALIGN
    ws.merge_cells(start_row=1, start_column=1, end_row=1,
                   end_column=max(ncol, 1))
    ws.row_dimensions[1].height = 30


def write_head(ws, row, heads):
    """写表头行：深蓝底白字加粗、细边框、居中。"""
    for ci, value in enumerate(heads, start=1):
        cell = ws.cell(row=row, column=ci, value=xl_safe(value))
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment, cell.border = CENTER, BORDER


def write_rows(ws, start_row, rows, numeric=()):
    """写表体：整行转义 + 细边框；numeric 为需右对齐的数字列（1 基序号）。"""
    for ri, values in enumerate(rows, start=start_row):
        for ci, value in enumerate(values, start=1):
            cell = ws.cell(row=ri, column=ci, value=xl_safe(value))
            cell.border = BORDER
            cell.alignment = RIGHT if ci in numeric else CENTER


def write_empty(ws, row, ncol, text='（当前筛选无匹配）'):
    """明细为空时写一行提示（跨列合并），避免用户以为导出失败。"""
    cell = ws.cell(row=row, column=1, value=text)
    cell.font, cell.alignment = NOTE_FONT, NOTE_ALIGN
    ws.merge_cells(start_row=row, start_column=1, end_row=row,
                   end_column=max(ncol, 1))


def set_widths(ws, widths):
    """按序号设置列宽。"""
    for ci, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = width


def cell_value(value):
    """空值统一成「—」，与页面表格的显示口径一致。"""
    return DASH if value is None or value == '' else value
