# StuLink v1.19.49.49 2026-10-20
# 限时练单科导入：考号优先匹配 + 缺考文字摒弃
# Copyright (c) 2026 zkxxzf. Apache License 2.0
"""限时练单科成绩解析（只解析，不写库；学生匹配见 ``timed_match``）。

与普通考试导入的差异（2026-10-20 用户口径）：

1. **考号优先**：先按文件里的【考号】列匹配 ``students.zhixue_no``
   （智学网固定考号），未命中再按【学号】匹配 ``student_number``；
   姓名/班级 **不用于猜测**，避免同名误配。
2. **缺考文字摒弃**：``缺考 / 缺 / 未考 / 免考 / — / - / 无`` 等一律
   视为"无成绩"，不报错、不落库，页面上也不再显示"缺考"二字。
3. **一科一文件**：科目由页面选择（同一次导入只处理一个科目）；
   重复导入同科 = 该科整科覆盖（见 ``routes/timed_import_routes.py``）。
"""
import io
import re

import openpyxl

from app.models.grades import DEFAULT_FULL_MARKS

HEADER_SCAN_MAX_ROW = 15
MAX_IMPORT_ROWS = 6000

ZHIXUE_KEYS = ('考号', '智学网考号', '智学考号', '学生考号', '考生号',
               '准考证号', '考试号')
NO_KEYS = ('学号', '学生学号')
NAME_KEYS = ('姓名', '学生姓名')
CLASS_KEYS = ('班级', '行政班', '班别')
SCORE_KEYS = ('分数', '成绩', '得分', '分值', '卷面分')
# 缺考 / 无成绩的文字口径（去空白与括号说明后比对）
ABSENT_WORDS = {'缺考', '缺', '未考', '未参加', '未参加考试', '免考', '弃考',
                '无', '无成绩', '未录入', '空', '—', '–', '−', '-', '－',
                '--', '/', '\\'}


class TimedParseError(Exception):
    """整表级解析错误（表头不可识别 / 无有效列等）"""


def _text(value):
    """单元格 → 去空白文本（数字 12345.0 → '12345'）"""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _code(value):
    """考号 / 学号归一化：去空白、去 Excel 千分位、去尾随 .0、去前导 '"""
    s = _text(value).replace(',', '').replace('，', '').lstrip("'")
    s = re.sub(r'^(\d+)\.0+$', r'\1', s)
    return s


def _score(value):
    """分数单元格 → (分数 or None, 是否缺考)

    「缺考（病假）」「免考(伤病)」这类带说明的写法也要认作缺考，
    不能落进 float() 失败分支变成错误行。
    """
    s = _text(value)
    if s == '':
        return None, True
    key = re.sub(r'[（(].*?[)）]', '', s).strip().lower()
    if key in ABSENT_WORDS:
        return None, True
    s = s.replace('分', '').replace(',', '').strip()
    try:
        return float(s), False
    except ValueError as e:
        raise TimedParseError(f'分数“{_text(value)}”不是数值') from e


def _col_index(headers, keys):
    """列定位：精确名优先，其次包含匹配（不依赖模板列的先后顺序）"""
    for i, h in enumerate(headers):
        if h in keys:
            return i
    for i, h in enumerate(headers):
        if any(k in h for k in keys):
            return i
    return None


def _locate_header(ws):
    """扫描前 N 行找表头：必须同时有（考号或学号）与分数列"""
    for row_no, cells in enumerate(
            ws.iter_rows(min_row=1, max_row=HEADER_SCAN_MAX_ROW, values_only=True),
            start=1):
        headers = [_text(c) for c in cells]
        idx = {
            'zhixue': _col_index(headers, ZHIXUE_KEYS),
            'no': _col_index(headers, NO_KEYS),
            'name': _col_index(headers, NAME_KEYS),
            'class': _col_index(headers, CLASS_KEYS),
            'score': _col_index(headers, SCORE_KEYS),
        }
        if idx['score'] is None:
            continue
        if idx['zhixue'] is None and idx['no'] is None:
            continue
        return row_no, headers, idx
    raise TimedParseError(
        '未找到表头：至少需要【考号】或【学号】列，以及【分数】列')


def _iter_rows(ws, row_no, idx):
    """逐行产出原始行（跳过整行空白）"""
    for line, cells in enumerate(
            ws.iter_rows(min_row=row_no + 1, values_only=True), start=row_no + 1):
        values = [_text(c) for c in cells]
        if not any(values):
            continue
        yield line, values


def parse_sheet(stream, subject, grade=''):
    """解析限时练成绩文件（不写库）。

    入参：stream（文件字节流）、subject（本次导入科目）、grade（限年级过滤，
    空 = 不做年级限制）。返回 ``{'rows','errors','stats','headers'}``。
    """
    if subject not in DEFAULT_FULL_MARKS:
        raise TimedParseError(f'不支持的科目：{subject}')
    try:
        wb = openpyxl.load_workbook(io.BytesIO(stream), read_only=True,
                                    data_only=True)
    except Exception as e:                              # noqa: BLE001
        raise TimedParseError(f'无法读取 Excel：{e}') from e
    ws = wb[wb.sheetnames[0]]
    row_no, headers, idx = _locate_header(ws)

    rows, errors = [], []
    total = absent = 0
    seen_codes = {}
    for line, values in _iter_rows(ws, row_no, idx):
        total += 1
        if total > MAX_IMPORT_ROWS:
            raise TimedParseError(f'数据超过 {MAX_IMPORT_ROWS} 行，请分批导入')
        zhixue = _code(values[idx['zhixue']]) if idx['zhixue'] is not None else ''
        no = _code(values[idx['no']]) if idx['no'] is not None else ''
        name = values[idx['name']] if idx['name'] is not None else ''
        class_name = values[idx['class']] if idx['class'] is not None else ''
        raw_score = values[idx['score']] if idx['score'] is not None else ''
        key = zhixue or no
        if not key:
            errors.append({'line': line, 'code': '', 'reason': '缺少考号/学号'})
            continue
        if key in seen_codes:
            errors.append({'line': line, 'code': key,
                           'reason': f'与第 {seen_codes[key]} 行重复'})
            continue
        seen_codes[key] = line
        try:
            score, is_absent = _score(raw_score)
        except TimedParseError as e:
            errors.append({'line': line, 'code': key, 'reason': str(e)})
            continue
        if is_absent:
            absent += 1          # 缺考：不落库、不报错（用户口径：摒弃文字）
            continue
        full = DEFAULT_FULL_MARKS.get(subject, 100)
        if score < 0 or score > full + 0.5:
            errors.append({'line': line, 'code': key,
                           'reason': f'分数 {score:g} 超出 0~{full:g}'})
            continue
        rows.append({'line': line, 'zhixue_no': zhixue, 'student_no': no,
                     'name': name, 'class_name': class_name,
                     'score': score, 'grade': grade})

    if not rows:
        raise TimedParseError(
            f'没有可用成绩（共 {total} 行，缺考 {absent} 行，'
            f'错误 {len(errors)} 行），请检查文件')
    return {'rows': rows, 'errors': errors,
            'headers': {k: v is not None for k, v in idx.items()},
            'stats': {'total_lines': total, 'ok_rows': len(rows),
                      'error_rows': len(errors), 'absent_rows': absent}}
