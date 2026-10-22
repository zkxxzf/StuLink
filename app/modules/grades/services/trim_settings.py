# StuLink v1.19.63.64 2026-10-21
# 成绩 · 去差人数设置（按班覆盖，导入页展示 / 手工改 / Excel / AI 导入共用）
# Copyright (c) 2026 zkxxzf. Apache License 2.0
"""去差人数设置。

用户口径（2026-10-21）：
- 数据导入页展示「班级人数 / 参考人数 / 去差人数 / 去差后人数」，只允许改去差人数；
- 默认卓越班去差 2 人、强基班 0 人（其余班型 0），可在页面上按班修改；
- 支持 Excel 导入（标准上传 + AI 智能导入识别同一张表）。

落库：``Exam.config_json['class_trim']`` = ``{班级: 去差人数}``（无需加列迁移）。
计算侧由 ``report_service.trim_quota`` 统一读取（本场设置 > 班型规则）。
"""
from __future__ import annotations

import io

from sqlalchemy import func

from app.extensions import db
from app.models import ClassProfile, Student
from app.models.grades import ExamScore, TOTAL_SUBJECT
from app.modules.grades.services.report_service import (
    DEFAULT_TRIM, TRIM_RULES, trim_quota)
from app.modules.grades.utils import normalize_class_name

# 这些伪班级不参与去差设置（与教学班统计口径一致）
_SKIP_CLASSES = {'', '不分班', '已转出', '离校', '—'}


def _class_sizes(grade):
    """{班级: 班级人数}（主库学籍名册；库不可用时返回空 dict）"""
    try:
        rows = (db.session.query(Student.class_name, func.count(Student.id))
                .filter(Student.grade == grade)
                .group_by(Student.class_name).all())
    except Exception:                                   # noqa: BLE001
        return {}
    return {normalize_class_name(c): n for c, n in rows if c}


def _attend_counts(exam_id):
    """{班级: 参考人数}：优先总分行的去重人数，没有总分行时回落全部科目去重人数"""
    q = (db.session.query(ExamScore.class_name,
                          func.count(func.distinct(ExamScore.student_no)))
         .filter(ExamScore.exam_id == exam_id, ExamScore.subject == TOTAL_SUBJECT)
         .group_by(ExamScore.class_name))
    rows = q.all()
    if not rows:
        rows = (db.session.query(ExamScore.class_name,
                                 func.count(func.distinct(ExamScore.student_no)))
                .filter(ExamScore.exam_id == exam_id)
                .group_by(ExamScore.class_name).all())
    return {normalize_class_name(c): n for c, n in rows if c}


def _class_types(exam):
    """{班级: 班型}：成绩行快照优先，历史数据回落当前班级档案"""
    out = {}
    for cls, ct in (db.session.query(ExamScore.class_name, ExamScore.class_type)
                    .filter(ExamScore.exam_id == exam.id,
                            ExamScore.class_type.isnot(None))
                    .distinct().all()):
        c = normalize_class_name(cls)
        if c and ct and c not in out:
            out[c] = ct
    try:
        for p in ClassProfile.query.filter_by(grade=exam.grade).all():
            c = normalize_class_name(p.class_name)
            if c and c not in out:
                out[c] = p.class_type or ''
    except Exception:                                   # noqa: BLE001
        pass
    return out


def build_rows(exam):
    """去差设置表格行（按班级名排序）

    每行：class_name / class_type / class_size / attend / trim / after / max_trim
    - class_size：学籍名册人数（取不到时回落参考人数）
    - attend：本场参考人数
    - trim：本场按班设置优先，否则班型默认（卓越班 2、强基班 0）
    - after：参考人数 − 去差人数（至少保留 1 人，去差上限 = attend-1）
    """
    sizes = _class_sizes(exam.grade)
    attend = _attend_counts(exam.id)
    types = _class_types(exam)
    overrides = exam.class_trim()
    names = set(sizes) | set(attend) | set(types) | set(overrides)
    rows = []
    for cls in sorted(n for n in names if n not in _SKIP_CLASSES):
        ct = types.get(cls, '')
        ref = int(attend.get(cls, 0) or 0)
        max_trim = max(ref - 1, 0)
        if cls in overrides:
            trim = overrides[cls]
        else:
            trim = TRIM_RULES.get(ct, DEFAULT_TRIM)
        trim = max(0, min(int(trim or 0), max_trim))
        rows.append({
            'class_name': cls,
            'class_type': ct,
            'class_size': int(sizes.get(cls, ref) or 0),
            'attend': ref,
            'trim': trim,
            'after': max(ref - trim, 0),
            'max_trim': max_trim,
        })
    return rows


def save_trims(exam, mapping, dry=False):
    """保存按班去差人数（只保留本场存在的班级，值收敛到 0..参考人数-1）

    dry=True 时只做校验与收敛，不写库。
    返回实际写入的 {班级: 人数}。
    """
    allowed = {r['class_name']: r for r in build_rows(exam)}
    clean = {}
    for cls, val in (mapping or {}).items():
        cls = normalize_class_name(cls)
        if cls not in allowed:
            continue
        try:
            n = int(val)
        except (TypeError, ValueError):
            continue
        clean[cls] = max(0, min(n, allowed[cls]['max_trim']))
    if not dry:
        exam.set_class_trim(clean)
    return clean


# ==================== Excel 导入 / 模板 ====================

_TEMPLATE_HEADERS = ('班级', '去差人数', '班级人数', '参考人数', '去差后人数')


def build_template(exam):
    """下载用模板：班级 / 去差人数（其余三列仅供对照，导入时忽略）"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = '去差人数设置'
    ws.append(list(_TEMPLATE_HEADERS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='E7ECF3')
    for r in build_rows(exam):
        ws.append([r['class_name'], r['trim'], r['class_size'],
                   r['attend'], r['after']])
    for i, width in enumerate([12, 12, 12, 12, 14], start=1):
        ws.column_dimensions[chr(64 + i)].width = width
    ws.freeze_panes = 'A2'
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _find_col(headers, *keywords):
    for i, h in enumerate(headers):
        t = str(h or '').strip()
        if not t:
            continue
        if all(k in t for k in keywords):
            return i
    return None


def parse_excel(stream):
    """解析标准模板：返回 ({班级: 去差人数}, errors)

    识别规则：表头（前 10 行内）需同时含「班级」与「去差人数」；
    「去差后人数」列不会被误当成「去差人数」。
    """
    from openpyxl import load_workbook
    try:
        wb = load_workbook(stream, data_only=True, read_only=True)
    except Exception as e:                              # noqa: BLE001
        return {}, ['文件无法解析：%s' % str(e)[:80]]
    mapping, errors = {}, []
    try:
        for ws in wb.worksheets:
            header_row, idx_cls, idx_trim = None, None, None
            for ri, row in enumerate(ws.iter_rows(min_row=1, max_row=10,
                                                  values_only=True), start=1):
                headers = [str(c or '').strip() for c in (row or [])]
                ci = _find_col(headers, '班级')
                ti = _find_col(headers, '去差人数')
                if ti is None:
                    ti = next((i for i, h in enumerate(headers)
                               if '去差' in h and '后' not in h), None)
                if ci is not None and ti is not None:
                    header_row, idx_cls, idx_trim = ri, ci, ti
                    break
            if header_row is None:
                continue
            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                cells = list(row or [])
                cls = normalize_class_name(cells[idx_cls] if idx_cls < len(cells) else '')
                if not cls:
                    continue
                raw = cells[idx_trim] if idx_trim < len(cells) else None
                try:
                    n = int(float(str(raw).strip()))
                except (TypeError, ValueError):
                    errors.append('%s：去差人数「%s」不是整数，已跳过' % (cls, raw))
                    continue
                mapping[cls] = max(0, n)
            break
    finally:
        try:
            wb.close()
        except Exception:                               # noqa: BLE001
            pass
    if not mapping and not errors:
        errors.append('未找到「班级 + 去差人数」表头（前 10 行内）')
    return mapping, errors


def parse_sheet(sheet):
    """解析 AI 导入草稿里的一张工作表 → ({班级: 去差人数}, errors)

    sheet: {headers: [...], rows: [(行号, [单元格...]), ...]}
    """
    headers = [str(h or '').strip() for h in (sheet.get('headers') or [])]
    ci = _find_col(headers, '班级')
    ti = _find_col(headers, '去差人数')
    if ti is None:
        ti = next((i for i, h in enumerate(headers)
                   if '去差' in h and '后' not in h), None)
    if ci is None or ti is None:
        return {}, [{'line': sheet.get('header_row') or 1, 'no': '',
                     'reason': '去差设置表需包含「班级」与「去差人数」两列'}]
    mapping, errors = {}, []
    for line_no, cells in (sheet.get('rows') or []):
        cells = list(cells or [])
        cls = normalize_class_name(cells[ci] if ci < len(cells) else '')
        if not cls:
            continue
        raw = cells[ti] if ti < len(cells) else ''
        try:
            n = int(float(str(raw).strip()))
        except (TypeError, ValueError):
            errors.append({'line': line_no, 'no': cls,
                           'reason': '去差人数「%s」不是整数' % raw})
            continue
        mapping[cls] = max(0, n)
    return mapping, errors
