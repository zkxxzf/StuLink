# StuLink v1.19.72.64 2026-10-20
# 限时练 AI 智能导入：只让模型做「列映射」，逐行解析仍走确定性代码
# Copyright (c) 2026 zkxxzf. Apache License 2.0
"""限时练 AI 导入的列映射。

设计沿用普通考试 AI 导入的既有原则（`ai_import_service`）：
**AI 不解析数据，只回答"哪一列是什么"**；全量逐行解析、分数校验、
缺考摒弃、学生匹配都由 `timed_import` 的确定性代码完成。
"""
import io
import json

import openpyxl

from app.modules.grades.services import timed_import as imp

SAMPLE_ROWS = 3
FIELDS = ('zhixue', 'no', 'name', 'class', 'score')
FIELD_NAMES = {'zhixue': '考号（智学网）', 'no': '学号', 'name': '姓名',
               'class': '班级', 'score': '分数'}
_KEYS = {'zhixue': imp.ZHIXUE_KEYS, 'no': imp.NO_KEYS,
         'name': imp.NAME_KEYS, 'class': imp.CLASS_KEYS,
         'score': imp.SCORE_KEYS}


def auto_mapping(headers):
    """先按内置别名自动映射（AI 不可用时的兜底，也是默认值）"""
    return {k: imp._col_index(headers, keys) for k, keys in _KEYS.items()}


def probe(stream):
    """读取第一张表的表头与前 3 行样本（只用于展示与提示词）"""
    wb = openpyxl.load_workbook(io.BytesIO(stream), read_only=True,
                               data_only=True)
    ws = wb[wb.sheetnames[0]]
    row_no, headers, _idx = imp._locate_header(ws)
    sample = []
    for cells in ws.iter_rows(min_row=row_no + 1, max_row=row_no + SAMPLE_ROWS,
                              values_only=True):
        sample.append([imp._text(c) for c in cells])
    return {'sheet': wb.sheetnames[0], 'header_row': row_no,
            'headers': headers, 'sample': sample,
            'auto': auto_mapping(headers)}


def build_messages(probe_data, subject):
    """构造列映射提示词（要求只回 JSON）"""
    lines = [f'这是一份「{subject}」单科成绩表，表头如下（从左到右）：']
    for i, h in enumerate(probe_data['headers']):
        lines.append(f'第{i + 1}列：{h or "（空）"}')
    lines.append('\n前几行样本：')
    for r in probe_data['sample']:
        lines.append(' | '.join(r))
    lines.append('\n请判断每个字段对应第几列（从 1 开始；没有该字段填 0）。'
                 '字段：考号/学号/姓名/班级/分数。'
                 '只输出 JSON，形如 '
                 '{"zhixue":1,"no":0,"name":2,"class":3,"score":4}。')
    return [{'role': 'system',
             'content': '你是成绩表列识别助手，只输出严格 JSON，不要解释。'},
            {'role': 'user', 'content': '\n'.join(lines)}]


def parse_ai_json(raw):
    """从模型回复里容错取出 {字段: 列号}（1 起；0/空 = 无）"""
    if not isinstance(raw, str):
        return None
    text = raw.strip()
    if '```' in text:
        text = text.split('```')[1] if text.count('```') >= 2 else text
        text = text.replace('json', '', 1).strip()
    start, end = text.find('{'), text.rfind('}')
    if start < 0 or end <= start:
        return None
    try:
        data = json.loads(text[start:end + 1])
    except (ValueError, TypeError):
        return None
    out = {}
    for key in FIELDS:
        try:
            n = int(data.get(key) or 0)
        except (TypeError, ValueError):
            n = 0
        out[key] = n - 1 if n > 0 else None
    return out


def mapping_from_index(data):
    """列号映射归一化（前端回传的就是 0 起索引；负数/非法值一律当"不指定"）"""
    if not data:
        return {}
    out = {}
    for key in FIELDS:
        v = data.get(key)
        try:
            n = int(v)
        except (TypeError, ValueError):
            n = -1
        out[key] = n if n >= 0 else None
    return out


def _header_row_of(ws, mapping):
    """表头行定位：别名识别优先，失败时按"映射列有值"的首次命中兜底。

    智学网表的分数列常直接叫「语文」「数学」，不在别名表里 → 自动识别会失败；
    此时若还从第 2 行开始读，第 2 行表头会被当数据。用映射列是否有值反推表头行。
    """
    try:
        return imp._locate_header(ws)[0]
    except imp.TimedParseError:
        pass
    want = [i for i in mapping.values() if i is not None]
    if not want:
        return 1
    for row_no, cells in enumerate(
            ws.iter_rows(min_row=1, max_row=imp.HEADER_SCAN_MAX_ROW,
                         values_only=True), start=1):
        vals = [imp._text(c) for c in cells]
        hits = sum(1 for i in want if i < len(vals) and vals[i])
        if hits >= max(2, len(want) - 1):
            return row_no
    return 1


def apply_mapping(stream, mapping, subject, grade=''):
    """按给定列映射解析全表（结构与 ``timed_import.parse_sheet`` 一致）"""
    if subject not in imp.DEFAULT_FULL_MARKS:
        raise imp.TimedParseError(f'不支持的科目：{subject}')
    wb = openpyxl.load_workbook(io.BytesIO(stream), read_only=True,
                               data_only=True)
    ws = wb[wb.sheetnames[0]]
    mapping = mapping_from_index(mapping)
    if mapping.get('score') is None:
        raise imp.TimedParseError('未指定分数列')
    if mapping.get('zhixue') is None and mapping.get('no') is None:
        raise imp.TimedParseError('未指定考号列或学号列')

    # 表头行：优先用自动识别；识别不到就按"映射列名有值的首行"兜底
    data_start = _header_row_of(ws, mapping)

    rows, errors, total, absent = [], [], 0, 0
    for line, cells in enumerate(
            ws.iter_rows(min_row=data_start + 1, values_only=True),
            start=data_start + 1):
        values = [_text_of(c) for c in cells]
        if not any(values):
            continue
        total += 1
        if total > imp.MAX_IMPORT_ROWS:
            raise imp.TimedParseError(f'数据超过 {imp.MAX_IMPORT_ROWS} 行，请分批导入')

        def _get(key):
            idx = mapping.get(key)
            if idx is None or idx < 0 or idx >= len(values):
                return ''
            return values[idx]

        zhixue = imp._code(_get('zhixue'))
        no = imp._code(_get('no'))
        key = zhixue or no
        if not key:
            errors.append({'line': line, 'code': '', 'reason': '缺少考号/学号'})
            continue
        try:
            score, is_absent = imp._score(_get('score'))
        except imp.TimedParseError as e:
            errors.append({'line': line, 'code': key, 'reason': str(e)})
            continue
        if is_absent:
            absent += 1
            continue
        full = imp.DEFAULT_FULL_MARKS.get(subject, 100)
        if score < 0 or score > full + 0.5:
            errors.append({'line': line, 'code': key,
                           'reason': f'分数 {score:g} 超出 0~{full:g}'})
            continue
        rows.append({'line': line, 'zhixue_no': zhixue, 'student_no': no,
                     'name': _get('name'), 'class_name': _get('class'),
                     'score': score, 'grade': grade})
    if not rows:
        raise imp.TimedParseError('没有可用成绩，请检查列映射与文件内容')
    return {'rows': rows, 'errors': errors,
            'headers': {k: mapping.get(k) is not None for k in FIELDS},
            'stats': {'total_lines': total, 'ok_rows': len(rows),
                      'error_rows': len(errors), 'absent_rows': absent}}


def _text_of(value):
    """单元格文本（与 timed_import._text 同口径的薄别名）"""
    return imp._text(value)
