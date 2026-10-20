# -*- coding: utf-8 -*-
"""临界生名单 Excel 导出（2026-10-19）。

数据由 `near_line_view.load` 提供（与页面同一份计算、同一套数据范围），
本模块只负责排版：明细一张表、班级对比一张表。
"""
from openpyxl import Workbook

from app.modules.grades.services import xlsx_kit as kit


def _note(ctx):
    """首行口径：考试 / 方向 / 层级线 / 上下分区间 / 班级筛选。"""
    data = ctx['data']
    exam = data['exam']
    parts = [f"临界生名单｜{exam['name']}"]
    parts.append(f"方向 {data['direction']}，层级 {data['layer']}")
    if isinstance(data['line'], (int, float)):
        parts.append(f"线 {data['line']} 分")
    else:
        lines = '、'.join(f'{k} {v}' for k, v in (data.get('lines') or {}).items())
        parts.append(f"线按选科组合（{lines}）" if lines else '线按选科组合')
    parts.append(f"线上 {data['above']} 分内 / 线下 {data['below']} 分内")
    parts.append(f"班级：{data['class_name'] or '全部'}")
    parts.append('线下组按总分降序、线上组按总分升序')
    return '口径：' + '；'.join(parts) + '。'


def _detail_sheet(ws, ctx):
    """临界生明细：状态 / 班级 / 班主任 / 姓名 / 总分 / 距线 / 排名 / 各科。"""
    data = ctx['data']
    subjects = data['subjects']
    heads = (['状态', '班级', '班主任', '姓名', '总分', '距线', '方向排名']
             + [s['label'] for s in subjects])
    kit.write_note(ws, _note(ctx), len(heads))
    kit.write_head(ws, 2, heads)
    rows = []
    for row in data['rows']:
        rows.append([row['status'], row['class_name'],
                     row.get('headteacher') or '', row.get('name') or '',
                     row.get('score'), row.get('gap'),
                     row.get('rank_dir') or '']
                    + [kit.cell_value(row.get(s['key'])) for s in subjects])
    if rows:
        kit.write_rows(ws, 3, rows,
                       numeric={5, 6, 7} | set(range(8, 8 + len(subjects))))
    else:
        kit.write_empty(ws, 3, len(heads), '（该范围内没有临界生）')
    kit.set_widths(ws, [7, 11, 10, 11, 8, 7, 9] + [7] * len(subjects))
    ws.freeze_panes = 'D3'


def _class_sheet(ws, ctx):
    """班级对比：各班线下 / 线上人数与合计（末尾给总计行）。"""
    data = ctx['data']
    heads = ['班级', '班主任', '线下临界', '线上临界', '合计']
    note = (f"班级对比｜{data['exam']['name']}｜{data['direction']} · "
            f"{data['layer']}｜按临界生人数降序")
    kit.write_note(ws, note, len(heads))
    kit.write_head(ws, 2, heads)
    rows = [[c['class_name'], c.get('headteacher') or '',
             c['below'], c['above'], c['above'] + c['below']]
            for c in ctx['by_class']]
    rows.append(['合计', '', data['below_n'], data['above_n'],
                 data['below_n'] + data['above_n']])
    kit.write_rows(ws, 3, rows, numeric={3, 4, 5})
    last = ws.max_row
    for ci in range(1, len(heads) + 1):
        ws.cell(row=last, column=ci).font = kit.NOTE_FONT
    kit.set_widths(ws, [12, 10, 10, 10, 8])
    ws.freeze_panes = 'A3'


def build(ctx):
    """构建临界生导出工作簿（明细 + 班级对比）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = '临界生名单'
    _detail_sheet(ws, ctx)
    _class_sheet(wb.create_sheet('班级对比'), ctx)
    return wb
