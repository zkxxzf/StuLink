# StuLink v1.19.50.53 2026-10-20
# 一键导出的取数与工作簿组装（路由只负责权限与参数）
# Copyright (c) 2026 zkxxzf. Apache License 2.0
"""成绩分析一键导出：各模块取数 + 工作簿组装。

取数全部复用页面同一份 `tab_service` / `compare_service`，
所以导出内容与「页面上看到的那张表」逐格一致（口径改动只改一处）。
权限判定不在这里做 —— 路由先判定再调用，避免服务层依赖 current_user。
"""
from datetime import datetime

import openpyxl

from app.models.grades import SUBJECTS
from app.modules.grades.services import compare_service, tab_service
from app.utils.xlsx_style import write_table


def sheets_of(payload):
    """payload['tables'] → [(工作表名, table)]，表名用表格自带标题"""
    return [(t.get('title') or k, t)
            for k, t in (payload.get('tables') or {}).items()]


def grade_sheets(exam_id, p):
    return sheets_of(tab_service.grade_tab(
        exam_id, p['direction'], p['selection'], p['denom']))


def subject_sheets(exam_id, p):
    """9 科的「各班对比」合并成一张（多一列科目），无数据的科目跳过"""
    rows, columns = [], None
    for subject in SUBJECTS:
        payload = tab_service.subject_tab(
            exam_id, subject, p['direction'], p['selection'], p['denom'])
        tab = (payload.get('tables') or {}).get('class_compare') or {}
        if not tab.get('rows'):
            continue
        columns = columns or tab['columns']
        for row in tab['rows']:
            item = {'subject': subject}
            item.update(row)
            rows.append(item)
    if not rows:
        return []
    table = {
        'title': '学科-各班对比',
        'rows': rows,
        'columns': [{'key': 'subject', 'label': '科目'}] + list(columns),
    }
    return [('学科-各班对比', table)]


def class_sheets(exam_id, class_name, p):
    return sheets_of(tab_service.class_tab(
        exam_id, class_name, p['direction'], p['selection'], p['denom']))


def compare_sheets(exam_id, classes, p):
    return sheets_of(compare_service.compare_tab(
        exam_id, classes, p['direction'], p['denom']))


def _note_rows(exam, p, skipped, user):
    denom = ('在籍人数（缺考按 0 分计入）' if p['denom'] == 'enrolled'
             else '参考人数（有成绩的学生）')
    rows = [
        ['成绩分析一键导出'],
        ['考试', '%s %s（%s）' % (exam.grade, exam.name, exam.exam_date)],
        ['统计口径', denom],
    ]
    if p['direction']:
        rows.append(['方向', p['direction']])
    if p['selection']:
        rows.append(['选科组合', p['selection']])
    if p['class_name']:
        rows.append(['班级明细', p['class_name']])
    if p['classes']:
        rows.append(['对比班级', '、'.join(p['classes'])])
    if skipped:
        rows.append(['未导出', '、'.join(skipped) + '（无权限或未选择）'])
    rows += [
        ['逐生明细', '仅当前选中班；其它班请到「班级分析」单选后导出'],
        ['导出人', getattr(user, 'real_name', '') or ''],
        ['导出时间', datetime.now().strftime('%Y-%m-%d %H:%M:%S')],
    ]
    return rows


def build_workbook(exam, p, sheets, skipped, user):
    """说明页 + 每个模块一张表；标题重名/超 31 字按 Excel 规则收口"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '说明'
    for row in _note_rows(exam, p, skipped, user):
        ws.append(row)
    if not sheets:
        ws.append(['提示', '本场考试暂无成绩数据（未导入或已清空）'])
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 70
    for title, table in sheets:
        write_table(wb.create_sheet(title=_title(wb, title)), table)
    return wb


def _title(wb, base):
    base = base or 'Sheet'
    title = base[:31]
    n = 2
    while title in wb.sheetnames:
        suffix = '(%d)' % n
        title = base[:31 - len(suffix)] + suffix
        n += 1
    return title
