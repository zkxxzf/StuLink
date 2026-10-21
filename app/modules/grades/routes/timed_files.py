# StuLink v1.19.50.54 2026-10-20
# 限时练：导入模板下载 / 错误明细下载
# Copyright (c) 2026 zkxxzf. Apache License 2.0
"""限时练的文件下载路由（模板与错误明细），与导入逻辑分开维护。"""
import io

import openpyxl
from flask import abort, send_file
from flask_login import login_required

from app.modules.grades import bp
from app.modules.grades.services import timed_service
from app.utils.decorators import perm_required
from app.utils.text_guard import safe_download_name


@bp.route('/timed/template.xlsx')
@login_required
@perm_required('grades.view')
def timed_template_xlsx():
    """限时练单科导入模板（考号 / 学号 / 姓名 / 班级 / 分数）"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '限时练单科成绩'
    ws.append(['考号', '学号', '姓名', '班级', '分数'])
    ws.append(['59468662', '20250001', '张三', '01班', 108])
    notes = wb.create_sheet('填表说明')
    for line in (
            '1. 一次只导入一个科目；科目在导入页面选择。',
            '2. 【考号】填智学网考号（学生档案上的固定 8 位号），优先按考号匹配。',
            '3. 没有考号可填【学号】；两者至少有一个。',
            '4. 缺考不要写「缺考」，留空即可（写「缺考」也按空值处理）。',
            '5. 同一学生重复出现会报错，请先合并。',
            '6. 重复导入同一科目 = 该科整科覆盖（上次有、这次没有的学生会被移除）。'):
        notes.append([line])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True,
                     download_name=safe_download_name('限时练导入模板.xlsx'))


@bp.route('/timed/<int:tid>/import/errors')
@login_required
@perm_required('grades.view')
def timed_import_errors(tid):
    """下载错误明细（行号 / 考号 / 原因）"""
    exam = timed_service.assert_visible(tid)
    draft = timed_service.report_of(exam)
    if not draft:
        abort(404)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '错误明细'
    ws.append(['行号', '考号/学号', '原因'])
    for e in draft.get('errors') or []:
        ws.append([e.get('line'), e.get('code'), e.get('reason')])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True,
                     download_name=safe_download_name(
                         f'限时练导入错误_{exam.name[:20]}.xlsx'))
