# StuLink v1.19.69.64 2026-10-20
# 限时练：成绩汇总表导出（xlsx）
# Copyright (c) 2026 zkxxzf. Apache License 2.0
"""限时练汇总表导出。

2026-10-20 用户口径：学生成绩页是"各科 + 总分 + 年级/班级名次"的汇总表，
导出也要把这张表原样带走（Excel 老师可直接排序、透视）。
"""
from io import BytesIO

from flask import request, send_file
from flask_login import current_user, login_required
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.modules.grades import bp
from app.modules.grades.services import timed_analysis, timed_service
from app.utils.decorators import perm_required

HEAD_FILL = PatternFill('solid', fgColor='EEF2F7')


@bp.route('/timed/<int:tid>/export.xlsx')
@login_required
@perm_required('grades.view')
def timed_export_summary(tid):
    """导出汇总表（按当前班级筛选口径，与页面一致）"""
    exam = timed_service.assert_visible(tid)
    class_name = (request.args.get('class') or '').strip()
    q = (request.args.get('q') or '').strip()
    # 导出不分页：把当前「班级 + 搜索」口径下的**全部**学生都写进 Excel
    # 2026-10-21：按当前账号可见范围裁剪（与页面同一份口径）
    data = timed_analysis.summary_matrix(exam, class_name, q=q,
                                         user=current_user)

    wb = Workbook()
    ws = wb.active
    ws.title = '成绩汇总'
    head = ['学号', '姓名', '班级'] + data['subjects']
    # 总分 / 名次列按可见范围决定是否输出（任课教师只给各科分 + 年级名次）
    if data['total_visible']:
        head += ['总分', 'T分数']
    if data['rank_g_visible']:
        head.append('年级名次')
    if data['rank_c_visible']:
        head.append('班级名次')
    ws.append(head)
    for r in data['rows']:
        # 「不分班/未分班」的行不参与班级排名（summary_matrix 不给它 rank_class），
        # 用 get 兜底 —— 之前用下标会 KeyError → 整场导出 500（2026-10-21 审计）
        line = ([r['student_no'], r['name'], r['class_name']]
                + [r['scores'].get(s) for s in data['subjects']])
        if data['total_visible']:
            line += [r.get('total'), r.get('t')]
        if data['rank_g_visible']:
            line.append(r.get('rank_grade'))
        if data['rank_c_visible']:
            line.append(r.get('rank_class'))
        ws.append(line)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal='center')
    for i in range(1, len(head) + 1):
        width = 14 if i <= 3 else 10
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = 'A2'

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f'{exam.grade}_{exam.name}_成绩汇总'
    return send_file(
        buf, as_attachment=True, download_name=f'{name}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument'
                 '.spreadsheetml.sheet')
