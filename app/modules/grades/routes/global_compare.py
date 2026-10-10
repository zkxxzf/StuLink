# StuLink v1.18.9.1 2026-10-09
# 全局对比：班级 × 学科宽表（去差均分 / 特优线 / 本科线），全科与分科自适配
#  - 页面/接口门槛与「成绩汇报区」一致（_guard_report：管理员/校领导/年级长）
#  - Excel 导出与页面同口径（复用 global_compare_service 的同一份缓存数据）
# Copyright (c) 2026 zkxxzf. Apache License 2.0
import io
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from flask import jsonify, render_template, request, send_file
from flask_login import current_user, login_required

from app.modules.grades import bp
from app.modules.grades.routes.report import _cached, _guard_report, _visible_exam
from app.modules.grades.services import global_compare_service as ld
from app.utils.decorators import perm_required
from app.utils.export_helpers import xl_row   # M-4：公式注入防护
from app.utils.helpers import log_operation

# 固定列（6 列，两级表头下 rowSpan=2）
_BASE_HEADS = ['年级', '班级', '班级人数', '方向', '属性', '班主任']
_PER_SUBJECT = 3   # 去差均分 / 特优线 / 本科线


@bp.route('/global-compare')
@login_required
@perm_required('grades.view')
def global_compare_page():
    """全局对比页（从成绩分析页或侧边栏进入，可携 exam_id）"""
    _guard_report()
    return render_template('grades/global_compare.html')


@bp.route('/api/global-compare/report')
@login_required
@perm_required('grades.view')
def api_global_compare_report():
    """全局对比数据 JSON（direction 为空 = 全部方向）"""
    _guard_report()
    exam_id = request.args.get('exam_id', type=int)
    direction = (request.args.get('direction') or '').strip()
    _visible_exam(exam_id)
    data = _cached(f'grades_report_ld_{exam_id}_{direction}',
                   lambda: ld.global_compare_report(exam_id, direction))
    return jsonify(success='error' not in data, data=data,
                   message=data.get('error') if 'error' in data else '')


# ============ Excel 导出（两行表头 + 小计/总计行，与页面所见一致） ============

_HF = Font(bold=True, color='FFFFFF')
_HFL = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
_HF2 = Font(bold=True, color='1F4E79')
_HFL2 = PatternFill(start_color='DCE6F1', end_color='DCE6F1', fill_type='solid')
_SUB_FILL = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
_GRAND_FILL = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')
_TB = Border(left=Side(style='thin'), right=Side(style='thin'),
             top=Side(style='thin'), bottom=Side(style='thin'))
_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _fmt_avg(v):
    return v if v is not None else '—'


def _fmt_online(n, rate):
    """上线单元格：'12 / 32%'；无该层单科线时 '—'"""
    if n is None:
        return '—'
    return f'{n} / {rate if rate is not None else 0}%'


def _row_values(row, subjects):
    vals = [row.get('grade', ''), row.get('class_name', ''), row.get('count', ''),
            row.get('direction', ''), row.get('class_type', ''), row.get('headteacher', '')]
    for sub in subjects:
        c = (row.get('cells') or {}).get(sub) or {}
        vals.append(_fmt_avg(c.get('trim_avg')))
        vals.append(_fmt_online(c.get('l1_n'), c.get('l1_rate')))
        vals.append(_fmt_online(c.get('l2_n'), c.get('l2_rate')))
    return vals


def _build_workbook(data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '全局对比'
    subjects = data['subjects']
    l1, l2 = data['l1_name'], data['l2_name']

    # 第一行：固定列 rowSpan=2 + 科目组 colSpan=3
    heads1 = list(_BASE_HEADS)
    for sub in subjects:
        heads1.extend([sub] + [''] * (_PER_SUBJECT - 1))
    ws.append(xl_row(heads1))
    # 第二行：指标名（固定列位置留空）
    heads2 = [''] * len(_BASE_HEADS)
    for _sub in subjects:
        heads2.extend(['去差均分', ld.line_label(l1), ld.line_label(l2)])
    ws.append(xl_row(heads2))

    # 合并固定列表头（竖向）+ 科目组表头（横向）
    for ci in range(1, len(_BASE_HEADS) + 1):
        ws.merge_cells(start_row=1, start_column=ci, end_row=2, end_column=ci)
    for gi, _sub in enumerate(subjects):
        c0 = len(_BASE_HEADS) + gi * _PER_SUBJECT + 1
        ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + _PER_SUBJECT - 1)

    # 数据行：组内班级 → 小计 → 末尾总计
    body = []      # (values, kind)  kind: normal/subtotal/grand
    for g in data['groups']:
        for r in g['rows']:
            body.append((r, 'normal'))
        if g.get('subtotal'):
            body.append((g['subtotal'], 'subtotal'))
    if data.get('grand'):
        body.append((data['grand'], 'grand'))
    for r, kind in body:
        ws.append(xl_row(_row_values(r, subjects)))

    # 样式
    ncol = len(_BASE_HEADS) + len(subjects) * _PER_SUBJECT
    for ci in range(1, ncol + 1):
        c = ws.cell(row=1, column=ci)
        c.font, c.fill, c.alignment, c.border = _HF, _HFL, _CENTER, _TB
        c2 = ws.cell(row=2, column=ci)
        c2.font, c2.fill, c2.alignment, c2.border = _HF2, _HFL2, _CENTER, _TB
    for ri, (r, kind) in enumerate(body, start=3):
        fill = _SUB_FILL if kind == 'subtotal' else (_GRAND_FILL if kind == 'grand' else None)
        for ci in range(1, ncol + 1):
            c = ws.cell(row=ri, column=ci)
            c.border, c.alignment = _TB, _CENTER
            if fill is not None:
                c.fill = fill
                c.font = Font(bold=True)
    ws.freeze_panes = 'G3'
    widths = [8, 12, 10, 8, 10, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for i in range(len(_BASE_HEADS) + 1, ncol + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 11
    return wb


@bp.route('/api/global-compare/export')
@login_required
@perm_required('grades.view')
def api_global_compare_export():
    """导出全局对比 xlsx（与页面同口径；复用同一份缓存数据）"""
    _guard_report()
    exam_id = request.args.get('exam_id', type=int)
    direction = (request.args.get('direction') or '').strip()
    exam = _visible_exam(exam_id)
    data = _cached(f'grades_report_ld_{exam_id}_{direction}',
                   lambda: ld.global_compare_report(exam_id, direction))
    if 'error' in data:
        return jsonify(success=False, message=data['error']), 400
    buf = io.BytesIO()
    _build_workbook(data).save(buf)
    buf.seek(0)
    log_operation(current_user, '导出', '全局对比', exam_id,
                  f'{exam.name} {direction or "全部"}', module='grades')
    fname = f'全局对比_{exam.name}.xlsx'.replace('/', '_')
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype='application/vnd.openxmlformats-officedocument'
                              '.spreadsheetml.sheet')
