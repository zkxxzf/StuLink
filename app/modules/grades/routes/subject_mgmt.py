# StuLink v1.19.65.64 2026-10-10
# 选科维护：按班级快速查看/批量修改学生选科组合
# Copyright (c) 2026 zkxxzf. Apache License 2.0
from io import BytesIO

from flask import (abort, flash, redirect, render_template, request, send_file,
                   url_for)
from flask_login import login_required, current_user
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.extensions import db
from app.models import Student
from app.modules.grades import bp
from app.modules.grades.services.exam_guard import assert_grade_visible
from app.utils.decorators import perm_required
from app.utils.helpers import (get_active_grades, get_dict_values,
                               get_graduated_grades, log_operation)


@bp.route('/subjects')
@login_required
@perm_required('grades.subject_mgmt')
def subject_mgmt():
    """选科维护页面：按班级列出学生，批量修改选科组合"""
    graduated = get_graduated_grades()
    grades = [g for g in get_active_grades() if g not in graduated]
    classes = get_dict_values('class')
    subjects = get_dict_values('subject')

    grade = request.args.get('grade', '')
    class_name = request.args.get('class_name', '')
    students = []
    if grade and class_name:
        # 2026-10-12 安全审计 H-2：页面按年级/班级取数，须先过可见性守卫，
        # 否则持权限者可借 URL 参数浏览任意年级班级的学生名单。
        assert_grade_visible(grade)
        students = Student.query.filter(
            Student.grade == grade, Student.class_name == class_name,
            Student.class_name != '不分班'
        ).order_by(Student.student_number).all()

    return render_template('grades/subject_mgmt.html', grades=grades, classes=classes,
                           subjects=subjects, grade=grade, class_name=class_name,
                           students=students)


@bp.route('/subjects/save', methods=['POST'])
@login_required
@perm_required('grades.subject_mgmt')
def subject_mgmt_save():
    """批量保存选科修改（仅更新有变化的）

    2026-10-12 安全审计 H-2：原先 `sid = int(key[4:])` 后直接 get(Student) 落库，
    参数里的 grade/class_name 已传入却**未用于校验** → 持 grades.subject_mgmt
    权限者可改任意年级任意班级学生的选科；而选科直接决定方向判定，
    进而影响排名/分层/汇报统计。现逐条校验学号归属本次提交的 (年级, 班级)。
    """
    grade = request.form.get('grade', '')
    class_name = request.form.get('class_name', '')
    assert_grade_visible(grade)
    changed = 0
    for key in list(request.form.keys()):
        if not key.startswith('sel_'):
            continue
        try:
            sid = int(key[4:])
        except ValueError:
            continue
        val = request.form.get(key, '').strip()
        s = db.session.get(Student, sid)
        if s and (s.subject_selection or '') != val:
            if s.grade != grade or s.class_name != class_name:
                abort(403, description='学号不属于所选年级/班级')
            s.subject_selection = val or None
            changed += 1
    if changed:
        db.session.commit()
        log_operation(current_user, '选科维护', '学生', None,
                      f'{grade} {class_name} 修改 {changed} 人选科')
        flash(f'已更新 {changed} 名学生的选科', 'success')
    else:
        flash('没有需要修改的选科', 'info')
    return redirect(url_for('grades.subject_mgmt', grade=grade, class_name=class_name))


# ---------------------------------------------------------------- 模板 / 导入
# 2026-10-22 用户口径：「选科维护要有模板导入 —— 把学生信息下载下来，只导入
# 选科的数据」。所以模板里带学号/姓名/班级/性别做对照，但导入时**只认
# 「选科组合」这一列**，其它列一律忽略，避免顺手改坏学生档案。
_SEL_HEADERS = ['学号', '姓名', '班级', '性别', '选科组合']


@bp.route('/subjects/template.xlsx')
@login_required
@perm_required('grades.subject_mgmt')
def subject_mgmt_template():
    """下载选科导入模板（学生名单 + 现有选科）。未选班级时导整个年级。"""
    grade = (request.args.get('grade') or '').strip()
    class_name = (request.args.get('class_name') or '').strip()
    if not grade:
        flash('请先选择年级，再下载模板', 'warning')
        return redirect(url_for('grades.subject_mgmt'))
    assert_grade_visible(grade)
    q = Student.query.filter(Student.grade == grade,
                             Student.class_name != '不分班')
    if class_name:
        q = q.filter(Student.class_name == class_name)
    students = q.order_by(Student.class_name, Student.student_number).all()

    wb = Workbook()
    ws = wb.active
    ws.title = '选科导入'
    ws.append(_SEL_HEADERS)
    for i in range(1, len(_SEL_HEADERS) + 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='4F81BD')
        c.alignment = Alignment(horizontal='center')
    for s in students:
        ws.append([s.student_number, s.name, s.class_name, s.gender or '',
                   s.subject_selection or ''])
    for i, w in enumerate([16, 12, 10, 6, 20], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'

    tip = wb.create_sheet('填写说明')
    for line in [
        '1. 只填「选科组合」列，其余列不要动（学号用来匹配学生，改了会对不上）。',
        '2. 选科组合请用系统字典里的名称，例如：物化生 / 史地政（数据字典 → 科目）。',
        '3. 留空的行会被跳过，不会清空该生原有选科。',
        '4. 表头行不要删，「学号」「选科组合」两列必须有。',
        f'5. 本表 {len(students)} 人（{grade}{class_name or " 全年级"}）。',
    ]:
        tip.append([line])
    tip.column_dimensions['A'].width = 88

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf, as_attachment=True,
        download_name=f'选科维护_{grade}{class_name}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument'
                 '.spreadsheetml.sheet')


@bp.route('/subjects/import', methods=['POST'])
@login_required
@perm_required('grades.subject_mgmt')
def subject_mgmt_import():
    """导入选科（只读「选科组合」列，按学号匹配，空值跳过）。

    2026-12 安全口径与 `subject_mgmt_save` 一致：先过年级可见性，
    再逐行校验学生属于本次导入的年级（跨班可以，跨年级不行），
    选科值必须落在数据字典里，避免脏值把方向判定带偏。
    """
    grade = (request.form.get('grade') or '').strip()
    class_name = (request.form.get('class_name') or '').strip()
    back = url_for('grades.subject_mgmt', grade=grade, class_name=class_name)
    assert_grade_visible(grade)

    f = request.files.get('file')
    if not f or not (f.filename or '').strip():
        flash('请选择要导入的 Excel 文件', 'warning')
        return redirect(back)
    if not f.filename.lower().endswith(('.xlsx', '.xlsm')):
        flash('只支持 .xlsx 文件，请用「下载模板」导出的格式填写', 'warning')
        return redirect(back)
    try:
        wb = load_workbook(f, data_only=True)
    except Exception:                                  # noqa: BLE001
        flash('文件无法解析，请确认是有效的 .xlsx（不是改后缀的 .xls）', 'danger')
        return redirect(back)
    ws = wb['选科导入'] if '选科导入' in wb.sheetnames else wb.worksheets[0]

    header = [str(c.value or '').strip() for c in ws[1]]

    def _col(*names):
        for n in names:
            if n in header:
                return header.index(n)
        return -1

    c_no = _col('学号', 'student_number')
    c_sel = _col('选科组合', '选科')
    if c_no < 0 or c_sel < 0:
        flash('表头里找不到「学号」或「选科组合」列，请用「下载模板」重新填写',
              'danger')
        return redirect(back)

    valid = set(get_dict_values('subject') or [])
    updated = unchanged = 0
    errors = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row:
            continue
        no = str(row[c_no] or '').strip() if c_no < len(row) else ''
        if not no:
            continue
        val = str(row[c_sel] or '').strip() if c_sel < len(row) else ''
        if not val:
            unchanged += 1          # 留空 = 不动这个学生
            continue
        if valid and val not in valid:
            errors.append(f'第 {i} 行 {no}：选科「{val}」不在字典里')
            continue
        s = Student.query.filter_by(student_number=no).first()
        if s is None:
            errors.append(f'第 {i} 行 {no}：学号不存在')
            continue
        if s.grade != grade or s.class_name == '不分班':
            errors.append(f'第 {i} 行 {no}：不属于 {grade}')
            continue
        if (s.subject_selection or '') != val:
            s.subject_selection = val
            updated += 1
        else:
            unchanged += 1

    if updated:
        db.session.commit()
        log_operation(current_user, '选科维护', '学生', None,
                      f'{grade}{class_name} 导入选科，更新 {updated} 人')
    msg = f'导入完成：更新 {updated} 人，未变化/留空 {unchanged} 人'
    if errors:
        msg += f'，{len(errors)} 行有问题'
        flash(msg, 'warning')
        flash('；'.join(errors[:10]) + ('…' if len(errors) > 10 else ''),
              'danger')
    else:
        flash(msg, 'success')
    return redirect(back)
