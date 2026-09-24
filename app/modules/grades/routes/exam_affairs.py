# StuLink v1.18.1.0 2026-09-24
# 考务管理：完整考务流程（对应 Excel 宏工作簿 2025考场学生考号与考场信息编排v1.2）
#   步骤：① 学生名单（学生信息表） → ② 考场设置（考场信息表）
#        → ③ 编排与考号生成（三种模式，镜像宏 编排考场考号2）
#        → ④ 导出（按班级/按考场/竖版桌签） → ⑤ 成绩关联（绑定考试、按考场看成绩）
# Copyright (c) 2026 zkxxzf. Apache License 2.0
import io
import json  # v1.12.1 选科前后缀配置解析
import random
from datetime import date, datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from flask import (render_template, request, jsonify, abort, redirect,
                   url_for, send_file, flash, current_app)
from flask_login import login_required, current_user

from app.extensions import db
from app.models import Student
from app.models.grades import (ExamAffair, AffairRoom, AffairRoomLib, AffairStudent,
                               Exam)  # v1.12.1 加考场房间库；v1.12.2 移除成绩关联查询依赖
from app.modules.grades import bp
from app.utils.decorators import perm_required
from app.utils.helpers import log_operation
from app.utils.text_guard import (sanitize_label, sanitize_prefix,
                                  safe_download_name)
from app.utils.upload_guard import validate_upload   # L-8：导入文件类型校验
from app.utils.export_helpers import xl_row   # M-4：公式注入防护

MODE_LABEL = {0: '模式0：前缀+考场号+座号', 1: '模式1：学号即考号', 2: '模式2：自定义考号'}

_HF = Font(bold=True, color='FFFFFF')
_HFL = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
_TB = Border(left=Side(style='thin'), right=Side(style='thin'),
             top=Side(style='thin'), bottom=Side(style='thin'))
_CENTER = Alignment(horizontal='center', vertical='center')


def _get_affair_checked(aid):
    """H-4：取考务批次并校验年级范围（越界 403）。

    此前 20 余处直接 `ExamAffair.query.get_or_404(aid)`，只验存在不验归属，
    任何有 grades.edit 的人都能操作其它年级的考务批次。
    """
    from app.modules.grades.services.exam_guard import assert_grade_visible
    affair = ExamAffair.query.get_or_404(aid)
    assert_grade_visible(affair.grade)
    return affair


def _pad2(v):
    s = str(v)
    return s.zfill(2) if s.isdigit() else s


# v1.12.2 考场编号规范：纯数字编号统一补零为两位（01、02…），保证排序与考号拼接一致
def _norm_room_no(no):
    s = str(no).strip()
    return s.zfill(2) if s.isdigit() else s


def _grade_options():
    from app.modules.grades.services.scope import visible_grades
    return visible_grades(current_user)


# ==================== 考务批次列表 / 新建 / 删除 ====================

@bp.route('/affairs')
@login_required
@perm_required('grades.edit')
def affairs_list():
    grade = request.args.get('grade', '')
    q = ExamAffair.query
    if grade:
        q = q.filter_by(grade=grade)
    affairs = q.order_by(ExamAffair.created_at.desc(), ExamAffair.id.desc()).all()
    stats = {}
    for a in affairs:
        stus = AffairStudent.query.filter_by(affair_id=a.id).all()
        stats[a.id] = {
            'students': len(stus),
            'attend': sum(1 for s in stus if s.is_attend),
            'rooms': AffairRoom.query.filter_by(affair_id=a.id).count(),
            'arranged': sum(1 for s in stus if s.room_no),
            'linked': bool(a.exam_id),
        }
    return render_template('grades/affairs/list.html', affairs=affairs, stats=stats,
                           grade=grade, grade_options=_grade_options())


@bp.route('/affairs/create', methods=['POST'])
@login_required
@perm_required('grades.edit')
def affair_create():
    name = (request.form.get('name') or '').strip()
    grade = (request.form.get('grade') or '').strip()
    if not name or not grade:
        flash('批次名与年级必填', 'danger')
        return redirect(url_for('grades.affairs_list'))
    # H-4：创建时同样校验年级范围（不能为其它年级建考务批次）
    from app.modules.grades.services.exam_guard import assert_grade_visible
    assert_grade_visible(grade)
    exam_date = None
    try:
        exam_date = date.fromisoformat(request.form.get('exam_date'))
    except (TypeError, ValueError):
        exam_date = None
    # v1.12.1 新建批次即定编排模式：selected 按选科分组 / plain 不选科全体一组
    selection_mode = request.form.get('selection_mode')
    if selection_mode not in ('selected', 'plain'):
        selection_mode = 'selected'
    affair = ExamAffair(name=name, grade=grade, exam_date=exam_date,
                        default_prefix=sanitize_prefix(request.form.get('default_prefix')) or '1701',
                        selection_mode=selection_mode,
                        operator_id=current_user.id)
    db.session.add(affair)
    db.session.commit()
    log_operation(current_user, '新建', '考务批次', affair.id, f'{grade} {name}', module='grades')
    flash('考务批次已创建', 'success')
    return redirect(url_for('grades.affair_detail', aid=affair.id))


@bp.route('/exams/<int:exam_id>/affair')
@login_required
@perm_required('grades.edit')
def exam_affair_go(exam_id):
    """考试列表「考务安排」入口：已有批次→进详情；无→按原流程创建（预填考试信息）后进向导"""
    # H-4：统一走考试年级范围校验（原先只比对下拉选项，不校验数据范围）
    from app.modules.grades.services.exam_guard import assert_exam_visible
    exam = assert_exam_visible(exam_id)
    affair = (ExamAffair.query.filter_by(exam_id=exam_id)
              .order_by(ExamAffair.id.desc()).first())
    if affair:
        return redirect(url_for('grades.affair_detail', aid=affair.id))
    affair = ExamAffair(name=f'{exam.name} 考务',
                        grade=exam.grade,
                        exam_date=exam.exam_date,
                        exam_id=exam_id,
                        default_prefix='1701',
                        selection_mode='selected',
                        operator_id=current_user.id)
    db.session.add(affair)
    db.session.commit()
    log_operation(current_user, '新建', '考务批次', affair.id,
                  f'{exam.grade} {affair.name}（自考试列表）', module='grades')
    flash('已为该考试创建考务安排，请按原流程编排', 'success')
    return redirect(url_for('grades.affair_detail', aid=affair.id))


@bp.route('/affairs/<int:aid>/delete', methods=['POST'])
@login_required
@perm_required('grades.edit')
def affair_delete(aid):
    affair = _get_affair_checked(aid)
    db.session.delete(affair)
    db.session.commit()
    log_operation(current_user, '删除', '考务批次', aid, affair.name, module='grades')
    flash('考务批次已删除', 'success')
    return redirect(url_for('grades.affairs_list'))


# ==================== 考务向导详情（含 5 个步骤 + 成绩关联） ====================

@bp.route('/affairs/<int:aid>')
@login_required
@perm_required('grades.edit')
def affair_detail(aid):
    affair = _get_affair_checked(aid)
    students = (AffairStudent.query.filter_by(affair_id=aid)
                .order_by(AffairStudent.class_name, AffairStudent.student_no).all())
    rooms = (AffairRoom.query.filter_by(affair_id=aid)
             .order_by(AffairRoom.room_no).all())
    scope_students = [s for s in students if s.is_attend]
    arranged = [s for s in scope_students if s.room_no]

    # v1.12.2 步骤⑤成绩关联 UI 已移除，不再查询考试与成绩（省去每次访问的无用查询）

    # v1.12.1 编排容量统计：学生总数/参考数、考场总数/总容量、各选科人数与可用容量差额
    stats = {
        'students': len(students), 'attend': len(scope_students),
        'rooms': len(rooms), 'capacity': sum(r.capacity or 0 for r in rooms),
        'per_sel': {},
    }
    for s in scope_students:
        key = (s.subject_selection or '默认') or '默认'
        d = stats['per_sel'].setdefault(key, {'students': 0, 'capacity': 0})
        d['students'] += 1
    for r in rooms:
        if r.is_universal:  # 通用考场容量计入每个有参考学生的选科
            for d in stats['per_sel'].values():
                d['capacity'] += r.capacity or 0
            if not stats['per_sel']:
                stats['per_sel']['默认'] = {'students': 0, 'capacity': r.capacity or 0}
        elif r.subject in stats['per_sel']:
            stats['per_sel'][r.subject]['capacity'] += r.capacity or 0
        else:
            stats['per_sel'][r.subject] = {'students': 0, 'capacity': r.capacity or 0}
    # v1.12.1 选科选项（批次设置弹窗中按选科设前后缀）
    sel_options = sorted({s.subject_selection for s in students if s.subject_selection})
    return render_template('grades/affairs/detail.html',
                           affair=affair, students=students, rooms=rooms,
                           scope_students=scope_students, arranged=arranged,
                           mode_label=MODE_LABEL,
                           stats=stats, sel_options=sel_options)


# ==================== 步骤① 学生名单：导入 / 模板 / 增删改 ====================

@bp.route('/affairs/<int:aid>/students/template')
@login_required
@perm_required('grades.edit')
def affair_students_template(aid):
    affair = _get_affair_checked(aid)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '学生信息表'
    headers = ['学号', '姓名', '班级', '选科组合', '是否参与考试', '固定考场', '固定座号',
               '考试科目', '自定义考号']
    ws.append(headers)
    for ci in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=ci)
        c.font = _HF
        c.fill = _HFL
        c.alignment = _CENTER
        c.border = _TB
    # 预填系统该年级学生（与 Excel「学生信息表」口径一致，考务人员可直接编辑）
    sys_stus = (Student.query.filter_by(grade=affair.grade)
                .order_by(Student.class_name, Student.student_number).all())
    for st in sys_stus:
        ws.append([st.student_number, st.name, st.class_name,
                   st.subject_selection or '', '是', '', '', '', ''])
    for col, w in zip('ABCDEFGHI', [12, 10, 8, 10, 12, 10, 10, 16, 12]):
        ws.column_dimensions[col].width = w
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    fn = f'学生信息表模板_{affair.grade}.xlsx'
    return send_file(out, as_attachment=True,
                     download_name=fn, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def _header_index(ws, names):
    first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    idx = {}
    for i, cell in enumerate(first or []):
        name = str(cell).strip() if cell is not None else ''
        if name in names:
            idx[name] = i
    return idx


@bp.route('/affairs/<int:aid>/students/import', methods=['POST'])
@login_required
@perm_required('grades.edit')
def affair_students_import(aid):
    affair = _get_affair_checked(aid)
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'ok': False, 'msg': '请选择 Excel 文件'})
    # L-8：统一上传校验（白名单 + 危险类型 + magic），不再只看扩展名
    ok, msg = validate_upload(f.filename, allowed_exts=['xlsx', 'xls'], stream=f.stream)
    if not ok:
        return jsonify({'ok': False, 'msg': msg})
    try:
        wb = openpyxl.load_workbook(f.stream, read_only=True, data_only=True)
    except Exception as e:
        return jsonify({'ok': False, 'msg': f'Excel 解析失败：{e}'})
    ws = wb.active
    idx = _header_index(ws, {'学号', '姓名', '班级', '选科组合', '是否参与考试',
                             '固定考场', '固定座号', '考试科目', '自定义考号'})
    if '学号' not in idx:
        return jsonify({'ok': False, 'msg': '缺少「学号」列，请使用模板'})
    added = 0
    skipped = 0
    sys_map = {s.student_number: s for s in Student.query.filter_by(grade=affair.grade).all()}
    for row in ws.iter_rows(min_row=2, values_only=True):
        no = str(row[idx['学号']]).strip() if idx['学号'] < len(row) and row[idx['学号']] is not None else ''
        if not no:
            continue
        exist = AffairStudent.query.filter_by(affair_id=aid, student_no=no).first()
        if exist:
            skipped += 1
            continue
        name = (str(row[idx['姓名']]).strip() if '姓名' in idx and idx['姓名'] < len(row) and row[idx['姓名']] else '')
        cls = (str(row[idx['班级']]).strip() if '班级' in idx and idx['班级'] < len(row) and row[idx['班级']] else '')
        sel = (str(row[idx['选科组合']]).strip() if '选科组合' in idx and idx['选科组合'] < len(row) and row[idx['选科组合']] else '')
        attend = str(row[idx['是否参与考试']]).strip() if '是否参与考试' in idx and idx['是否参与考试'] < len(row) and row[idx['是否参与考试']] else '是'
        fixed_room = (str(row[idx['固定考场']]).strip() if '固定考场' in idx and idx['固定考场'] < len(row) and row[idx['固定考场']] else '') or None
        fixed_seat = None
        if '固定座号' in idx and idx['固定座号'] < len(row) and row[idx['固定座号']]:
            try:
                fixed_seat = int(float(row[idx['固定座号']]))
            except (TypeError, ValueError):
                fixed_seat = None
        custom = (str(row[idx['自定义考号']]).strip() if '自定义考号' in idx and idx['自定义考号'] < len(row) and row[idx['自定义考号']] else '') or None
        # 用系统学生补全空白字段，保证与学籍一致
        sys_stu = sys_map.get(no)
        if sys_stu:
            name = name or sys_stu.name
            cls = cls or sys_stu.class_name
            sel = sel or (sys_stu.subject_selection or '')
        st = AffairStudent(affair_id=aid, student_no=no, name=name, class_name=cls,
                           subject_selection=sel, subject=sel or '默认',
                           is_attend=(attend != '否'), fixed_room=fixed_room,
                           fixed_seat=fixed_seat, custom_number=custom)
        db.session.add(st)
        added += 1
    db.session.commit()
    wb.close()
    log_operation(current_user, '导入', '考务学生', aid, f'新增{added}条/跳过{skipped}条', module='grades')
    return jsonify({'ok': True, 'added': added, 'skipped': skipped})


@bp.route('/affairs/<int:aid>/students/sync', methods=['POST'])
@login_required
@perm_required('grades.edit')
def affair_students_sync(aid):
    """v1.12.1 从学生学籍库同步本年级学生到考务名单（增量：新增+补全信息，可选拨除已删学籍）"""
    affair = _get_affair_checked(aid)
    sys_stus = (Student.query.filter_by(grade=affair.grade)
                .order_by(Student.class_name, Student.student_number).all())
    exist = {s.student_no: s for s in AffairStudent.query.filter_by(affair_id=aid).all()}
    added = removed = 0
    for st in sys_stus:
        row = exist.get(st.student_number)
        if not row:
            row = AffairStudent(affair_id=aid, student_no=st.student_number)
            db.session.add(row)
            added += 1
        # 刷新基本信息，保留考务特有字段（固定考场/座号/自定义考号/是否参考）
        row.name = st.name
        row.class_name = st.class_name
        row.subject_selection = st.subject_selection or ''
        row.subject = row.subject_selection or '默认'
        if row.is_attend is None:
            row.is_attend = True
    if request.form.get('remove_missing') == '1':  # 可选：学籍已删除的学生同步移除
        sys_nos = {st.student_number for st in sys_stus}
        for no, row in exist.items():
            if no not in sys_nos:
                db.session.delete(row)
                removed += 1
    db.session.commit()
    log_operation(current_user, '同步', '考务学生', aid,
                  f'新增{added}人/移除{removed}人/共{len(sys_stus)}人', module='grades')
    return jsonify({'ok': True, 'added': added, 'removed': removed,
                    'total': len(sys_stus)})


@bp.route('/affairs/<int:aid>/students/pick')
@login_required
@perm_required('grades.edit')
def affair_students_pick(aid):
    """v1.12.2 「加入学生」选择器数据源：返回本年级学籍学生，
    带 in_batch 标记（是否已在当前考务名单），供前端筛选/排序/搜索。"""
    affair = _get_affair_checked(aid)
    have = {s.student_no for s in AffairStudent.query.filter_by(affair_id=aid).all()}
    rows = (Student.query.filter_by(grade=affair.grade)
            .order_by(Student.class_name, Student.student_number).all())
    return jsonify({'ok': True, 'students': [
        {'student_no': st.student_number, 'name': st.name, 'class_name': st.class_name,
         'subject_selection': st.subject_selection or '', 'gender': st.gender or '',
         'in_batch': st.student_number in have}
        for st in rows]})


@bp.route('/affairs/<int:aid>/students/batch-add', methods=['POST'])
@login_required
@perm_required('grades.edit')
def affair_students_batch_add(aid):
    """v1.12.2 从选择器批量加入：按学号从学籍库带出信息写入考务名单"""
    affair = _get_affair_checked(aid)
    nos = [n for n in (request.form.get('student_nos') or '').split(',') if n.strip()]
    if not nos:
        return jsonify({'ok': False, 'msg': '未选择学生'})
    exist = {s.student_no for s in AffairStudent.query.filter_by(affair_id=aid).all()}
    added = 0
    for st in Student.query.filter(Student.student_number.in_(nos)).all():
        if st.student_number in exist:
            continue
        db.session.add(AffairStudent(
            affair_id=aid, student_no=st.student_number, name=st.name,
            class_name=st.class_name, subject_selection=st.subject_selection or '',
            subject=st.subject_selection or '默认', is_attend=True))
        added += 1
    db.session.commit()
    log_operation(current_user, '加入', '考务学生', aid, f'选择器加入{added}人', module='grades')
    return jsonify({'ok': True, 'added': added})


@bp.route('/affairs/<int:aid>/students/<int:sid>/fixed', methods=['POST'])
@login_required
@perm_required('grades.edit')
def affair_student_fixed(aid, sid):
    """v1.12.2 表格内直接编辑固定考场/固定座号（失焦即时保存）"""
    st = AffairStudent.query.get_or_404(sid)
    if st.affair_id != aid:
        abort(403)
    room = (request.form.get('fixed_room') or '').strip() or None
    seat_raw = (request.form.get('fixed_seat') or '').strip()
    try:
        seat = int(seat_raw) if seat_raw else None
    except ValueError:
        seat = None
    st.fixed_room, st.fixed_seat = room, seat
    db.session.commit()
    return jsonify({'ok': True})


@bp.route('/affairs/<int:aid>/students/<int:sid>')
@login_required
@perm_required('grades.edit')
def affair_student_detail(aid, sid):
    st = AffairStudent.query.get_or_404(sid)
    if st.affair_id != aid:
        abort(403)
    return jsonify({
        'student_no': st.student_no, 'name': st.name, 'class_name': st.class_name,
        'subject_selection': st.subject_selection, 'is_attend': st.is_attend,
        'fixed_room': st.fixed_room, 'fixed_seat': st.fixed_seat, 'custom_number': st.custom_number,
    })


@bp.route('/affairs/<int:aid>/students/save', methods=['POST'])
@login_required
@perm_required('grades.edit')
def affair_student_save(aid):
    affair = _get_affair_checked(aid)
    sid = request.form.get('sid')
    no = (request.form.get('student_no') or '').strip()
    if not no:
        return jsonify({'ok': False, 'msg': '学号必填'})
    if sid:
        st = AffairStudent.query.get_or_404(int(sid))
        if st.affair_id != aid:
            abort(403)
    else:
        st = AffairStudent.query.filter_by(affair_id=aid, student_no=no).first()
        if st:
            return jsonify({'ok': False, 'msg': f'学号 {no} 已存在'})
        st = AffairStudent(affair_id=aid, student_no=no)
        db.session.add(st)
    st.name = (request.form.get('name') or '').strip()
    st.class_name = (request.form.get('class_name') or '').strip()
    st.subject_selection = (request.form.get('subject_selection') or '').strip()
    st.subject = st.subject_selection or '默认'
    st.is_attend = (request.form.get('is_attend') == '1')
    st.fixed_room = (request.form.get('fixed_room') or '').strip() or None
    try:
        st.fixed_seat = int(request.form.get('fixed_seat')) if request.form.get('fixed_seat') else None
    except (TypeError, ValueError):
        st.fixed_seat = None
    st.custom_number = (request.form.get('custom_number') or '').strip() or None
    # v1.12.1 勾选后同步写入/更新学生学籍库（新增学生不再依赖 Excel 回导）
    if request.form.get('save_to_db') == '1':
        sys_stu = Student.query.filter_by(student_number=no).first()
        if not sys_stu:
            sys_stu = Student(student_number=no, name=st.name or '未命名',
                              gender=(request.form.get('gender') or '男'),
                              grade=affair.grade, class_name=st.class_name or '')
            db.session.add(sys_stu)
        else:
            sys_stu.name = st.name or sys_stu.name
            sys_stu.class_name = st.class_name or sys_stu.class_name
            sys_stu.subject_selection = st.subject_selection or sys_stu.subject_selection
    db.session.commit()
    return jsonify({'ok': True})


@bp.route('/affairs/<int:aid>/students/<int:sid>/delete', methods=['POST'])
@login_required
@perm_required('grades.edit')
def affair_student_delete(aid, sid):
    st = AffairStudent.query.get_or_404(sid)
    if st.affair_id != aid:
        abort(403)
    db.session.delete(st)
    db.session.commit()
    return jsonify({'ok': True})


@bp.route('/affairs/<int:aid>/students/clear', methods=['POST'])
@login_required
@perm_required('grades.edit')
def affair_students_clear(aid):
    AffairStudent.query.filter_by(affair_id=aid).delete()
    db.session.commit()
    return jsonify({'ok': True})


# ==================== 步骤② 考场设置：导入 / 模板 / 增删改 ====================

@bp.route('/affairs/<int:aid>/rooms/template')
@login_required
@perm_required('grades.edit')
def affair_rooms_template(aid):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '考场信息表'
    headers = ['考场编号', '考场位置', '考场科目', '考场容量', '备注', '考号前缀']
    ws.append(headers)
    for ci in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=ci)
        c.font = _HF
        c.fill = _HFL
        c.alignment = _CENTER
        c.border = _TB
    for col, w in zip('ABCDEF', [10, 16, 12, 10, 12, 10]):
        ws.column_dimensions[col].width = w
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, as_attachment=True,
                     download_name='考场信息表模板.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@bp.route('/affairs/<int:aid>/rooms/import', methods=['POST'])
@login_required
@perm_required('grades.edit')
def affair_rooms_import(aid):
    affair = _get_affair_checked(aid)
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'ok': False, 'msg': '请选择 Excel 文件'})
    # L-8：统一上传校验（白名单 + 危险类型 + magic），不再只看扩展名
    ok, msg = validate_upload(f.filename, allowed_exts=['xlsx', 'xls'], stream=f.stream)
    if not ok:
        return jsonify({'ok': False, 'msg': msg})
    try:
        wb = openpyxl.load_workbook(f.stream, read_only=True, data_only=True)
    except Exception as e:
        return jsonify({'ok': False, 'msg': f'Excel 解析失败：{e}'})
    ws = wb.active
    idx = _header_index(ws, {'考场编号', '考场位置', '考场科目', '考场容量', '备注', '考号前缀'})
    if '考场编号' not in idx:
        return jsonify({'ok': False, 'msg': '缺少「考场编号」列'})
    added = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        no = _norm_room_no(str(row[idx['考场编号']]) if idx['考场编号'] < len(row) and row[idx['考场编号']] is not None else '')  # v1.12.2 编号补零
        if not no:
            continue
        room = AffairRoom.query.filter_by(affair_id=aid, room_no=no).first()
        if not room:
            room = AffairRoom(affair_id=aid, room_no=no)
            db.session.add(room)
            added += 1
        room.location = (str(row[idx['考场位置']]).strip() if '考场位置' in idx and idx['考场位置'] < len(row) and row[idx['考场位置']] else '') or ''
        room.subject = (str(row[idx['考场科目']]).strip() if '考场科目' in idx and idx['考场科目'] < len(row) and row[idx['考场科目']] else '') or '默认'
        cap = 0
        if '考场容量' in idx and idx['考场容量'] < len(row) and row[idx['考场容量']]:
            m = ''.join(ch for ch in str(row[idx['考场容量']]) if ch.isdigit())
            cap = int(m) if m else 0
        room.capacity = cap or 30
        # H-5：导入路径同样清洗（Excel 内容可被任意构造）
        room.prefix = (sanitize_prefix(str(row[idx['考号前缀']])) if '考号前缀' in idx and idx['考号前缀'] < len(row) and row[idx['考号前缀']] else '') or None
        room.note = (sanitize_label(str(row[idx['备注']]), 100) if '备注' in idx and idx['备注'] < len(row) and row[idx['备注']] else '') or ''
    db.session.commit()
    wb.close()
    return jsonify({'ok': True, 'added': added})


@bp.route('/affairs/<int:aid>/rooms/<int:rid>')
@login_required
@perm_required('grades.edit')
def affair_room_detail(aid, rid):
    room = AffairRoom.query.get_or_404(rid)
    if room.affair_id != aid:
        abort(403)
    return jsonify({
        'room_no': room.room_no, 'location': room.location, 'subject': room.subject,
        'capacity': room.capacity, 'prefix': room.prefix, 'note': room.note,
    })


@bp.route('/affairs/<int:aid>/rooms/save', methods=['POST'])
@login_required
@perm_required('grades.edit')
def affair_room_save(aid):
    affair = _get_affair_checked(aid)
    rid = request.form.get('rid')
    no = _norm_room_no(request.form.get('room_no') or '')  # v1.12.2 编号补零
    if rid:
        room = AffairRoom.query.get_or_404(int(rid))
        if room.affair_id != aid:
            abort(403)
        # v1.12.2 编辑时编号与容量锁定不可改：编号由选取时自动生成、容量由考场库带出，
        # 改动会破坏已生成考号与编排一致性
    else:
        if not no:
            return jsonify({'ok': False, 'msg': '考场编号必填'})
        room = AffairRoom.query.filter_by(affair_id=aid, room_no=no).first()
        if room:
            return jsonify({'ok': False, 'msg': f'考场 {no} 已存在'})
        room = AffairRoom(affair_id=aid, room_no=no)
        db.session.add(room)
        try:
            room.capacity = int(request.form.get('capacity') or 30)
        except (TypeError, ValueError):
            room.capacity = 30
    # H-5：服务端字符白名单（前端转义之外的第二道防线）
    room.location = sanitize_label(request.form.get('location'), 50)
    room.subject = sanitize_label(request.form.get('subject') or '默认', 20) or '默认'
    room.prefix = sanitize_prefix(request.form.get('prefix')) or None
    room.note = sanitize_label(request.form.get('note'), 100)
    # v1.12.1 可选：把该房间位置/容量存入考场库，供后续批次直接选取
    if request.form.get('save_to_lib') == '1' and room.location:
        lib_row = AffairRoomLib.query.filter_by(location=room.location).first()
        if not lib_row:
            lib_row = AffairRoomLib(location=room.location)
            db.session.add(lib_row)
        lib_row.capacity = room.capacity
        lib_row.note = room.note
    db.session.commit()
    return jsonify({'ok': True})


@bp.route('/affairs/<int:aid>/rooms/<int:rid>/delete', methods=['POST'])
@login_required
@perm_required('grades.edit')
def affair_room_delete(aid, rid):
    room = AffairRoom.query.get_or_404(rid)
    if room.affair_id != aid:
        abort(403)
    db.session.delete(room)
    db.session.commit()
    return jsonify({'ok': True})


# ==================== v1.12.1 考场房间库（跨批次共享，选取后自动带出位置/容量） ====================

@bp.route('/affairs/room-lib/list')
@login_required
@perm_required('grades.edit')
def affair_room_lib_list():
    rows = AffairRoomLib.query.order_by(AffairRoomLib.location).all()
    return jsonify({'ok': True, 'rooms': [
        {'id': r.id, 'location': r.location, 'capacity': r.capacity, 'note': r.note}
        for r in rows]})


@bp.route('/affairs/room-lib/save', methods=['POST'])
@login_required
@perm_required('grades.edit')
def affair_room_lib_save():
    """库中新增/编辑房间；同位置自动去重（唯一约束兜底）"""
    lid = request.form.get('lid')
    location = sanitize_label(request.form.get('location'), 50)   # H-5
    if not location:
        return jsonify({'ok': False, 'msg': '房间位置必填'})
    if lid:
        row = AffairRoomLib.query.get_or_404(int(lid))
        row.location = location
    else:
        row = AffairRoomLib.query.filter_by(location=location).first()
        if not row:
            row = AffairRoomLib(location=location)
            db.session.add(row)
    try:
        row.capacity = int(request.form.get('capacity') or 30)
    except (TypeError, ValueError):
        row.capacity = 30
    row.note = sanitize_label(request.form.get('note'), 100)      # H-5
    db.session.commit()
    return jsonify({'ok': True, 'id': row.id, 'location': row.location,
                    'capacity': row.capacity})


@bp.route('/affairs/room-lib/<int:lid>/delete', methods=['POST'])
@login_required
@perm_required('grades.edit')
def affair_room_lib_delete(lid):
    row = AffairRoomLib.query.get_or_404(lid)
    db.session.delete(row)
    db.session.commit()
    return jsonify({'ok': True})


@bp.route('/affairs/<int:aid>/rooms/from-lib', methods=['POST'])
@login_required
@perm_required('grades.edit')
def affair_room_from_lib(aid):
    """v1.12.1 从考场房间库选取房间生成考场：
    自动生成考场编号（现有最大数字+1）、确定考场科目（默认，可再编辑）、带出默认容量"""
    AffairRoom.query.get_or_404 and ExamAffair.query.get_or_404(aid)  # 批次存在性校验
    lid = request.form.get('lib_id')
    lib_row = AffairRoomLib.query.get_or_404(int(lid))
    exist_rooms = AffairRoom.query.filter_by(affair_id=aid).all()
    if lib_row.location in {r.location for r in exist_rooms}:
        return jsonify({'ok': False, 'msg': f'考场位置 {lib_row.location} 已在本批次中'})
    nums = []
    for r in exist_rooms:
        digits = ''.join(ch for ch in r.room_no if ch.isdigit())
        if digits:
            nums.append(int(digits))
    # v1.12.2 考场编号统一补零为两位（01、02…）
    new_no = _norm_room_no((max(nums) + 1) if nums else 1)
    room = AffairRoom(affair_id=aid, room_no=new_no, location=lib_row.location,
                      subject='默认', capacity=lib_row.capacity or 30)
    db.session.add(room)
    db.session.commit()
    log_operation(current_user, '选取', '考务考场', aid,
                  f'从库选房间 {lib_row.location} → 考场{new_no}', module='grades')
    return jsonify({'ok': True, 'id': room.id, 'room_no': room.room_no,
                    'location': room.location, 'capacity': room.capacity, 'subject': room.subject})


# ==================== v1.12.1 批次设置：编排模式 / 非参考处理 / 选科前后缀 ====================

@bp.route('/affairs/<int:aid>/settings', methods=['POST'])
@login_required
@perm_required('grades.edit')
def affair_settings(aid):
    affair = _get_affair_checked(aid)
    sm = request.form.get('selection_mode')
    if sm in ('selected', 'plain'):
        affair.selection_mode = sm
    nm = request.form.get('non_attend_mode')
    if nm in ('skip', 'tail'):
        affair.non_attend_mode = nm
    # 选科考号前后缀：前端以 prefix_<选科> / suffix_<选科> 成对提交，sels 为选科列表 JSON
    try:
        sel_list = json.loads(request.form.get('sels') or '[]')
    except (TypeError, ValueError, json.JSONDecodeError):
        sel_list = []
    if isinstance(sel_list, list):
        cfg = {}
        for sel in sel_list:
            sel = str(sel).strip()
            if not sel:
                continue
            p = sanitize_prefix(request.form.get(f'prefix_{sel}'), 16)    # H-5
            sfx = sanitize_prefix(request.form.get(f'suffix_{sel}'), 16)  # H-5
            if p or sfx:
                cfg[sel] = {'prefix': p, 'suffix': sfx}
        affair.subject_prefixes = json.dumps(cfg, ensure_ascii=False) if cfg else None
    db.session.commit()
    log_operation(current_user, '设置', '考务批次', aid,
                  f'模式={affair.selection_mode}/非参考={affair.non_attend_mode}', module='grades')
    return jsonify({'ok': True})


# ==================== 步骤③ 编排与考号生成（镜像宏 编排考场考号2） ====================

def _arrange(affair, students, rooms, mode, log):
    """镜像 Excel 宏的编排逻辑：固定座位优先 → 同班邻座限制 → 随机分配。
    v1.12.1：支持不选科模式、非参考学生尾场处理、选科考号前后缀。
    返回 (assigned_count, failed_count)"""
    room_by_no = {r.room_no: r for r in rooms}
    # 座位占用表
    seats = {r.room_no: [None] * r.capacity for r in rooms}
    used = {r.room_no: set() for r in rooms}
    # v1.12.1 不选科模式：全体一组，所有考场通用
    plain = (affair.selection_mode == 'plain')
    # v1.12.1 选科考号前后缀配置：{"物化生":{"prefix":"1701","suffix":""}}
    sel_cfg = affair.get_subject_prefixes()

    def neighbors(i, cap):
        n = []
        if i > 0:
            n.append(i - 1)
        if i < cap - 1:
            n.append(i + 1)
        return n

    def exam_number_for(student, room_no, seat):
        if mode == 0:
            # v1.12.1 前后缀优先级：选科配置 > 考场前缀 > 批次默认前缀（满足全校同前缀或文理分别前缀）
            cfg = sel_cfg.get((student.subject_selection or '') or '')
            prefix = (cfg or {}).get('prefix') or (room_by_no[room_no].prefix or affair.default_prefix or '1701')
            suffix = (cfg or {}).get('suffix') or ''
            return f'{prefix}{_pad2(room_no)}{_pad2(seat)}{suffix}'
        if mode == 1:
            return student.student_no
        return student.custom_number or f'未填考号_{student.name}'

    assigned = 0
    failed = 0

    # 1) 固定座位优先
    fixed = [s for s in students if s.is_attend and s.fixed_room and s.fixed_seat]
    for s in fixed:
        room = room_by_no.get(s.fixed_room)
        if not room:
            log.append(f'【固定考场】[{s.name}] 指定考场 {s.fixed_room} 不存在，跳过')
            continue
        if s.fixed_seat > room.capacity or s.fixed_seat in used[room.room_no]:
            log.append(f'【固定座位】[{s.name}] 座位超出容量或已占用，跳过')
            continue
        seats[room.room_no][s.fixed_seat - 1] = s
        used[room.room_no].add(s.fixed_seat)
        s.room_no = room.room_no
        s.seat_no = s.fixed_seat
        s.exam_number = exam_number_for(s, room.room_no, s.fixed_seat)
        assigned += 1

    # 2) 按选科分组，分配无固定座位学生
    groups = {}
    for s in students:
        if not s.is_attend or s.room_no:
            continue
        key = (s.subject_selection or '默认') or '默认'
        groups.setdefault(key, []).append(s)

    MAX_SAME_CLASS = 1
    for key, group in groups.items():
        # v1.12.1 不选科模式：全体一组，考场全部通用；选科模式按科目匹配
        avail = [(r.room_no, r) for r in rooms
                 if len(used[r.room_no]) < r.capacity
                 and (plain or r.is_universal or r.subject == key)]
        for s in group:
            candidates = []
            for rno, r in avail:
                if len(used[rno]) >= r.capacity:
                    continue
                for i in range(r.capacity):
                    if (i + 1) in used[rno]:
                        continue
                    same = sum(1 for nb in neighbors(i, r.capacity)
                               if seats[rno][nb] and seats[rno][nb].class_name == s.class_name)
                    if same <= MAX_SAME_CLASS:
                        candidates.append((rno, i + 1))
            if not candidates:
                # 放宽：任意空位
                for rno, r in avail:
                    if len(used[rno]) >= r.capacity:
                        continue
                    for i in range(r.capacity):
                        if (i + 1) not in used[rno]:
                            candidates.append((rno, i + 1))
            if not candidates:
                log.append(f'【分配失败】[{s.name}] 无可用座位（选科 {key}）')
                failed += 1
                continue
            rno, seat = random.choice(candidates)
            seats[rno][seat - 1] = s
            used[rno].add(seat)
            s.room_no = rno
            s.seat_no = seat
            s.exam_number = exam_number_for(s, rno, seat)
            assigned += 1
    # v1.12.1 非参考学生尾场处理：tail 模式下安排到同选科（或通用）考场靠后座号；
    # skip 模式（默认）不安排考场，保持 room_no 为空
    if affair.non_attend_mode == 'tail':
        for s in [x for x in students if not x.is_attend and not x.room_no and not x.fixed_room]:
            key = (s.subject_selection or '默认') or '默认'
            avail = [r for r in rooms if len(used[r.room_no]) < r.capacity
                     and (plain or r.is_universal or r.subject == key)]
            placed = False
            for r in sorted(avail, key=lambda x: x.room_no):
                seat = next((i + 1 for i in range(r.capacity - 1, -1, -1)
                             if (i + 1) not in used[r.room_no]), None)  # 从尾座往前找
                if seat:
                    seats[r.room_no][seat - 1] = s
                    used[r.room_no].add(seat)
                    s.room_no = r.room_no
                    s.seat_no = seat
                    s.exam_number = exam_number_for(s, r.room_no, seat)
                    assigned += 1
                    placed = True
                    break
            if not placed:
                log.append(f'【尾场】[{s.name}] 无可用空位，未安排')
    return assigned, failed


@bp.route('/affairs/<int:aid>/arrange', methods=['POST'])
@login_required
@perm_required('grades.edit')
def affair_arrange(aid):
    affair = _get_affair_checked(aid)
    try:
        mode = int(request.form.get('mode', affair.mode))
    except (TypeError, ValueError):
        mode = 0
    if mode not in (0, 1, 2):
        mode = 0
    # 重新编排前清空旧结果
    AffairStudent.query.filter_by(affair_id=aid).update(
        {AffairStudent.room_no: None, AffairStudent.seat_no: None, AffairStudent.exam_number: None})
    affair.mode = mode
    students = AffairStudent.query.filter_by(affair_id=aid).all()
    rooms = AffairRoom.query.filter_by(affair_id=aid).all()
    if not rooms:
        return jsonify({'ok': False, 'msg': '请先在「考场设置」中添加考场'})
    scope = [s for s in students if s.is_attend]
    log = [f'当前模式：{MODE_LABEL[mode]}']
    assigned, failed = _arrange(affair, students, rooms, mode, log)
    affair.status = 'arranged' if assigned else 'draft'
    db.session.commit()
    log_operation(current_user, '编排', '考务批次', aid,
                  f'{MODE_LABEL[mode]}，分配{assigned}人/失败{failed}人', module='grades')
    return jsonify({'ok': True, 'assigned': assigned, 'failed': failed,
                    'scope': len(scope), 'log': log})


# ==================== 步骤④ 导出：按班级 / 按考场 / 竖版桌签 ====================

def _arranged_sorted(affair):
    return [s for s in AffairStudent.query.filter_by(affair_id=affair.id).all()
            if s.room_no]


def _write_sheet(ws, headers, rows):
    # M-4：学生姓名/班级/考场位置来自导入数据，统一做公式注入转义
    rows = [xl_row(r) for r in rows]
    ws.append(headers)
    for ci in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=ci)
        c.font = _HF
        c.fill = _HFL
        c.alignment = _CENTER
        c.border = _TB
    for r in rows:
        ws.append(xl_row(r))   # M-4
        for ci in range(1, len(headers) + 1):
            ws.cell(row=ws.max_row, column=ci).border = _TB


@bp.route('/affairs/<int:aid>/export/class')
@login_required
@perm_required('grades.edit')
def affair_export_class(aid):
    affair = _get_affair_checked(aid)
    stus = sorted(_arranged_sorted(affair),
                  key=lambda s: (s.class_name or '', s.name or '', s.seat_no or 0))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '按班级学生信息'
    _write_sheet(ws, ['班级', '姓名', '考场位置', '考号', '考场编号', '座号'],
                 [[s.class_name, s.name, _room_loc(affair, s.room_no),
                   s.exam_number, s.room_no, _pad2(s.seat_no)] for s in stus])
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, as_attachment=True,
                     download_name=safe_download_name(   # L-9
                         f'按班级学生信息_{affair.name}.xlsx', '按班级学生信息.xlsx'),
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@bp.route('/affairs/<int:aid>/export/room')
@login_required
@perm_required('grades.edit')
def affair_export_room(aid):
    affair = _get_affair_checked(aid)
    stus = sorted(_arranged_sorted(affair),
                  key=lambda s: (s.room_no or '', s.seat_no or 0))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '按考场学生信息'
    _write_sheet(ws, ['考场编号', '座号', '考号', '姓名', '班级', '考场位置'],
                 [[s.room_no, _pad2(s.seat_no), s.exam_number, s.name,
                   s.class_name, _room_loc(affair, s.room_no)] for s in stus])
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, as_attachment=True,
                     download_name=safe_download_name(   # L-9
                         f'按考场学生信息_{affair.name}.xlsx', '按考场学生信息.xlsx'),
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def _room_loc(affair, room_no):
    room = AffairRoom.query.filter_by(affair_id=affair.id, room_no=room_no).first()
    return room.location if room else ''


# ==================== v1.12.2 A4 打印版 PDF 导出（reportlab） ====================

def _pdf_resp(buf, filename):
    return send_file(buf, as_attachment=True, download_name=filename, mimetype='application/pdf')


def _room_loc_map(affair):
    """一次性取 {考场编号: 位置}，避免桌签/名单逐条查库（N+1）"""
    return {r.room_no: r.location for r in
            AffairRoom.query.filter_by(affair_id=affair.id).all()}


@bp.route('/affairs/<int:aid>/cards/pdf')
@login_required
@perm_required('grades.edit')
def affair_cards_pdf(aid):
    """考试桌签 PDF：一页 A4 排 3 栏 × 15 行 = 45 张小标签，裁开贴桌"""
    from app.utils.affair_pdf import desk_cards_pdf
    affair = _get_affair_checked(aid)
    stus = sorted(_arranged_sorted(affair),
                  key=lambda s: (s.room_no or '', s.seat_no or 0))
    if not stus:
        abort(400, description='尚未编排，无桌签可导出')
    loc_map = _room_loc_map(affair)
    buf = desk_cards_pdf(stus, lambda rn: loc_map.get(rn, ''))
    log_operation(current_user, '导出', '桌签PDF', aid, f'{len(stus)}人', module='grades')
    return _pdf_resp(buf, f'桌签_{affair.name}.pdf')


@bp.route('/affairs/<int:aid>/export/class/pdf')
@login_required
@perm_required('grades.edit')
def affair_export_class_pdf(aid):
    """按班级学生名单 A4 打印版 PDF：v1.12.2 一个班从新的一页开始"""
    from app.utils.affair_pdf import seating_list_sections
    affair = _get_affair_checked(aid)
    stus = sorted(_arranged_sorted(affair),
                  key=lambda s: (s.class_name or '', s.name or '', s.seat_no or 0))
    loc_map = _room_loc_map(affair)
    # 按班级分组（保持排序后的班级顺序）
    groups = {}
    for s in stus:
        groups.setdefault(s.class_name or '未分班', []).append(
            [s.name, s.exam_number, s.room_no, _pad2(s.seat_no),
             loc_map.get(s.room_no, '')])
    sections = [(f'{cls}（共 {len(rows)} 人）', rows) for cls, rows in groups.items()]
    buf = seating_list_sections(
        f'{affair.name}　按班级学生名单（共 {len(stus)} 人 {len(sections)} 个班）',
        ['姓名', '考号', '考场编号', '座号', '考场位置'],
        sections, widths_mm=[26, 36, 18, 14, 50])
    return _pdf_resp(buf, f'按班级学生名单_{affair.name}.pdf')


@bp.route('/affairs/<int:aid>/export/room/pdf')
@login_required
@perm_required('grades.edit')
def affair_export_room_pdf(aid):
    """按考场学生名单 A4 打印版 PDF（监考教师用）：v1.12.2 一个考场从新的一页开始"""
    from app.utils.affair_pdf import seating_list_sections
    affair = _get_affair_checked(aid)
    stus = sorted(_arranged_sorted(affair),
                  key=lambda s: (s.room_no or '', s.seat_no or 0))
    loc_map = _room_loc_map(affair)
    # 按考场分组（保持排序后的考场顺序）
    groups = {}
    for s in stus:
        groups.setdefault(s.room_no or '未分配', []).append(
            [_pad2(s.seat_no), s.exam_number, s.name, s.class_name])
    sections = [(f'考场 {rn}（{loc_map.get(rn, "")}，共 {len(rows)} 人）', rows)
                for rn, rows in groups.items()]
    buf = seating_list_sections(
        f'{affair.name}　按考场学生名单（共 {len(stus)} 人 {len(sections)} 个考场）',
        ['座号', '考号', '姓名', '班级'],
        sections, widths_mm=[14, 40, 26, 20])
    return _pdf_resp(buf, f'按考场学生名单_{affair.name}.pdf')


# v1.12.2 网页版桌签预览路由已删除：千人级渲染会卡死浏览器，统一改走 /cards/pdf


@bp.route('/affairs/<int:aid>/cards/export')
@login_required
@perm_required('grades.edit')
def affair_cards_export(aid):
    """竖版桌签导出（三栏布局，对应 Excel 竖版桌签（可打印））"""
    affair = _get_affair_checked(aid)
    stus = sorted(_arranged_sorted(affair),
                  key=lambda s: (s.room_no or '', s.seat_no or 0))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '竖版桌签'
    # 表头（每栏 4 列：考号/姓名/班级/考场-座号）
    per = 3
    headers = []
    for _ in range(per):
        headers += ['学号/考号', '姓名', '班级', '考场-座号']
    ws.append(headers)
    for ci in range(1, len(headers) + 1):
        ws.cell(row=1, column=ci).font = _HF
        ws.cell(row=1, column=ci).fill = _HFL
        ws.cell(row=1, column=ci).alignment = _CENTER
    # 按三栏分块写入
    rows = []
    for i in range(0, len(stus), per):
        chunk = stus[i:i + per]
        row = []
        for s in chunk:
            row += [s.exam_number, s.name, s.class_name, f'{s.room_no}-{_pad2(s.seat_no)}']
        while len(row) < per * 4:
            row += ['', '', '', '']
        rows.append(row)
    for r in rows:
        ws.append(xl_row(r))   # M-4
    # 列宽
    for ci in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = 14
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, as_attachment=True,
                     download_name=safe_download_name(   # L-9
                         f'竖版桌签_{affair.name}.xlsx', '竖版桌签.xlsx'),
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ==================== 步骤⑤ 成绩关联：绑定考试 ====================

@bp.route('/affairs/<int:aid>/link-exam', methods=['POST'])
@login_required
@perm_required('grades.edit')
def affair_link_exam(aid):
    affair = _get_affair_checked(aid)
    exam_id = request.form.get('exam_id')
    if exam_id:
        exam = Exam.query.get(int(exam_id))
        if not exam or exam.grade != affair.grade:
            return jsonify({'ok': False, 'msg': '考试与批次年级不一致'})
        affair.exam_id = exam.id
        affair.status = 'linked' if affair.status != 'draft' else affair.status
    else:
        affair.exam_id = None
    db.session.commit()
    log_operation(current_user, '关联', '考务批次', aid, f'绑定考试 {affair.exam_id}', module='grades')
    return jsonify({'ok': True})
