# StuLink v1.18.9.2 2026-10-10
# Copyright (c) 2026 zkxxzf. Apache License 2.0
"""openpyxl 读取兼容层（架构审计修复：消除重复的 monkeypatch）。

背景：部分 Excel 文件（多由 WPS / 旧版 Excel 生成）的 DataValidation 带有
openpyxl 不接受的 `id` 参数，直接 `load_workbook()` 会抛 TypeError。
此前 `system/routes/users.py` 与 `system/routes/students.py` 各自贴了一份
完全相同的 patch 代码 —— 升级 openpyxl 时须同步改两处，漏一处则该导入静默崩。

现在统一由本模块提供：
    - `patch_openpyxl_datavalidation()`：幂等地打上补丁（可重复调用）
    - `load_sheet(file)`：patch + 载入 + 返回活动工作表
    - `read_header_map(ws)`：读首行表头 → {列名: 列号}

注意：补丁是**进程级全局**的（改的是 openpyxl 类的 __init__），与原先各调用点
的行为完全一致，不影响功能与 UI。
"""
_patched = False


def patch_openpyxl_datavalidation():
    """幂等打补丁：剥掉 DataValidation 的非法 id 参数（重复调用无副作用）"""
    global _patched
    if _patched:
        return
    import openpyxl.worksheet.datavalidation as dv

    _orig_init = dv.DataValidation.__init__

    def _patched_init(self, *a, **kw):
        kw.pop('id', None)
        _orig_init(self, *a, **kw)

    dv.DataValidation.__init__ = _patched_init
    _patched = True


def load_sheet(file, data_only=True):
    """载入 Excel 并返回活动工作表（已打好兼容性补丁）"""
    import openpyxl
    patch_openpyxl_datavalidation()
    wb = openpyxl.load_workbook(file, data_only=data_only)
    return wb.active


def read_header_map(ws):
    """读工作表首行 → {表头文本: 列号(1-based)}；跳过空单元格"""
    header_map = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=1, column=col).value
        if val:
            header_map[str(val).strip()] = col
    return header_map
