# StuLink v1.19.49.45 2026-10-20
# Excel 导出统一样式：表头蓝底白字、细边框、冻结首行
# Copyright (c) 2026 zkxxzf. Apache License 2.0
"""Excel 导出统一样式（单页签导出与一键导出共用，避免两份样式漂移）。"""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.utils.export_helpers import xl_row

HEAD_FONT = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill(start_color='4472C4', end_color='4472C4',
                        fill_type='solid')
THIN = Border(left=Side(style='thin'), right=Side(style='thin'),
              top=Side(style='thin'), bottom=Side(style='thin'))
CENTER = Alignment(horizontal='center', vertical='center')


def write_table(ws, table):
    """把 {columns:[{key,label}], rows:[{...}]} 写进工作表（首行=表头）"""
    headers = [c['label'] for c in table['columns']]
    ws.append(headers)
    for ci in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=ci)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment, cell.border = CENTER, THIN
    for row in table.get('rows') or []:
        ws.append(xl_row([row.get(c['key']) for c in table['columns']]))
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            cell.border = THIN
    ws.freeze_panes = 'A2'
