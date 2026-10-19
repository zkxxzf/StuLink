# -*- coding: utf-8 -*-
"""AI 导入「考试时间表 / 监考表 → 划分场次」冒烟（脚本级）

覆盖：
- 服务层：parse_schedule_sheet 解析多级表头监考表（场次 + 考场 + 监考老师）
- 服务层：parse_schedule_columns 解析两列式时间表（日期/开始/结束/科目）
- 服务层：parse_schedule_json 容错解析识图 JSON（```json 包裹 / 前后有解释）
- 识别：looks_like_schedule_sheet / guess_sheet_kind(rows)
- HTTP：analyze（上传监考表）→ apply（kind=exam_schedule）→ 落库校验

用法：python scripts/smoke/smoke_ai_schedule.py
"""
import os as _os
import sys as _sys
_HERE = _os.path.dirname(_os.path.abspath(__file__))
_SCRIPTS = _os.path.dirname(_HERE)
_ROOT = _os.path.dirname(_SCRIPTS)
for _p in (_ROOT, _SCRIPTS):
    if _p not in _sys.path:
        _sys.path.insert(0, _p)

import io
import os
import re
import sys
import tempfile

import openpyxl

BASE = _ROOT
sys.path.insert(0, BASE)
_TMP = tempfile.mkdtemp(prefix='stulink_ai_sched_')

import config as config_mod  # noqa: E402
config_mod.BASE_DIR = _TMP
config_mod.Config.SQLALCHEMY_DATABASE_URI = 'sqlite:///' + os.path.join(_TMP, 'system.db')
config_mod.Config.SQLALCHEMY_BINDS = {
    k: 'sqlite:///' + os.path.join(_TMP, k + '.db')
    for k in ('dormitory', 'history', 'grades', 'points', 'academic', 'inspection',
              'achievement', 'forms', 'system', 'portrait', 'timetable', 'morality')
}
config_mod.Config.SECRET_KEY = 'ai-sched-check'

from app import create_app  # noqa: E402
from app.extensions import db  # noqa: E402
from app.models import User  # noqa: E402
from app.models.grades import (Exam, ExamAffair, AffairSchedule, AffairRoom,  # noqa: E402
                               AffairInvigilation)
from app.modules.grades.services import ai_import_service as aisvc  # noqa: E402

app = create_app()
ok = fail = 0


def check(name, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
        print(f'[PASS] {name}')
    else:
        fail += 1
        print(f'[FAIL] {name}  {extra}')


def build_invig_xlsx():
    """合成「监考表」：多级表头（日期行/上午下午行/时间行/科目行）+ 数据行（考场+老师）"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws.append(['郑州市郑开学校2027届高三第一学期10月学情评估监考表（2026年10月8日）'])
    ws.append(['主考：王书欣 张常立       巡视：孙燕梅 雷坤'])
    ws.append(['考场位置', '组合', '考场号', '人数', '10月8日（周四）', None, None,
               '10月9日（周五）', None])
    ws.append([None, None, None, None, '上午', None, '下午', '上午', None])
    ws.append([None, None, None, None, '8:05-10:35', '10:55-12:10', '14:30-15:45',
               '7:50-9:50', '10:10-12:10'])
    ws.append([None, None, None, None, '语文', '物理/历史', None, '数学', '英语'])
    ws.append([None, None, None, None, None, None, None, None, None])       # 空行
    ws.append(['527（1班）', '史地政', '01', 30, '徐卓', '祝平平', '徐卓', '祝平平', '祝平平'])
    ws.append(['525（2班）', '史地政', '02', 29, '吴若澜', '郝文哲', '吴若澜', '郝文哲', '郝文哲'])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_time_xlsx():
    """合成「两列式考试时间表」：日期 / 开始时间 / 结束时间 / 考试科目 / 选科"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '考试时间表'
    ws.append(['考试日期', '开始时间', '结束时间', '考试科目', '选科'])
    ws.append(['2026-10-08', '07:55', '10:25', '语文', ''])
    ws.append(['2026-10-08', '10:45', '12:00', '物理', '物理方向'])
    ws.append(['2026-10-08', '10:45', '12:00', '历史', '历史方向'])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_parallel_xlsx():
    """合成「并行多科监考表」：**同一时段（10:45）两列**分别考 物理 与 历史，
    用于验证监考格不会错挂到另一场次（审查修复 E）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws.append(['某校监考表（2026年10月8日）'])
    ws.append(['考场位置', '考场号', '人数', '10月8日（周四）', None])
    ws.append([None, None, None, '10:45-12:00', '10:45-12:00'])
    ws.append([None, None, None, '物理', '历史'])
    ws.append(['527（1班）', '01', 30, '徐卓', '吴若澜'])     # 物理@01 / 历史@01
    ws.append(['525（2班）', '02', 29, '祝平平', '郝文哲'])   # 物理@02 / 历史@02
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ==================== 1. 服务层：监考表解析 ====================
raw_xlsx = build_invig_xlsx()
pr = aisvc.probe(raw_xlsx)
sh = pr['sheets'][0]
rows = aisvc.read_sheet_rows(raw_xlsx, sh['name'], sh['header_row'])

check('监考表被识别为 exam_schedule',
      aisvc.guess_sheet_kind(sh['headers'], rows) == aisvc.SCHEDULE_KIND,
      str(aisvc.guess_sheet_kind(sh['headers'], rows)))
check('looks_like_schedule_sheet=True', aisvc.looks_like_schedule_sheet(rows))

r = aisvc.parse_schedule_sheet(rows, default_year=2026)
sess = r['sessions']
check('监考表解析出 5 个场次', len(sess) == 5, str(len(sess)))
check('场次1 = 10-08 08:05-10:35 语文',
      sess[0] == {'date': '2026-10-08', 'start': '08:05', 'end': '10:35',
                  'subject': '语文', 'selection': ''}, str(sess[0]))
check('日期右填充：场次2 仍是 10-08', sess[1]['date'] == '2026-10-08', str(sess[1]))
check('合并科目归一为「、」（便于环节⑥ 拆分回避）',
      sess[1]['subject'] == '物理、历史', str(sess[1]))
check('科目行缺失时该场次科目为空（14:30 场）', sess[2]['subject'] == '' and
      sess[2]['start'] == '14:30', str(sess[2]))
check('第二天日期正确：场次4 = 10-09 数学',
      sess[3]['date'] == '2026-10-09' and sess[3]['subject'] == '数学', str(sess[3]))
check('英语按系统口径归一为「外语」', sess[4]['subject'] == '外语', str(sess[4]))
check('解析出 2 个考场', [x['room'] for x in r['rooms']] == ['01', '02'],
      str(r['rooms']))
check('考场容量取自「人数」列', r['rooms'][0]['capacity'] == 30, str(r['rooms'][0]))
check('解析出 10 格监考（2 考场 × 5 场次）', len(r['invigs']) == 10, str(len(r['invigs'])))
check('监考格绑定到场次序号', all(0 <= iv['session_index'] < 5 for iv in r['invigs']))

# ==================== 2. 服务层：两列式时间表 ====================
raw_t = build_time_xlsx()
prt = aisvc.probe(raw_t)
sht = prt['sheets'][0]
sessions2, errs2 = aisvc.parse_schedule_columns(
    {'headers': sht['headers'], 'header_row': sht['header_row'],
     'rows': aisvc.read_sheet_rows(raw_t, sht['name'], sht['header_row'])})
check('两列式时间表解析 3 场', len(sessions2) == 3, str(sessions2))
check('两列式：日期/时间/科目均归一',
      sessions2[0]['date'] == '2026-10-08' and sessions2[0]['start'] == '07:55'
      and sessions2[0]['subject'] == '语文', str(sessions2[0]))
_rows_t = aisvc.read_sheet_rows(raw_t, sht['name'], sht['header_row'])
check('两列式时间表被自动识别为 exam_schedule',
      aisvc.guess_sheet_kind(sht['headers'], _rows_t) == aisvc.SCHEDULE_KIND,
      str(aisvc.guess_sheet_kind(sht['headers'], _rows_t)))
# 反向：含学号的成绩表不得被判成时间表
_score_headers = ['学号', '姓名', '班级', '语文', '数学', '总分']
_score_rows = [(1, _score_headers), (2, ['2501', '张某', '01班', '100', '110', '210'])]
check('含学号的成绩表不误判为 exam_schedule',
      aisvc.guess_sheet_kind(_score_headers, _score_rows) != aisvc.SCHEDULE_KIND)

# ==================== 3. 服务层：识图 JSON 容错解析 ====================
js1 = '```json\n{"sessions":[{"date":"2026-10-08","start":"8:05","end":"10:35",' \
      '"subject":"语文"},{"date":"10月8日","start":"14:30","end":"15:45",' \
      '"subject":"化学/地理/政治/生物"}],"notes":"四选二"}\n```'
s1, note1 = aisvc.parse_schedule_json(js1)
check('识图 JSON（```json 包裹）解析 2 场', len(s1) == 2, str(s1))
check('识图：时间补零为 08:05', s1[0]['start'] == '08:05', str(s1[0]))
check('识图：10月8日 → 当年 YYYY-MM-DD', re.match(r'^\d{4}-10-08$', s1[1]['date'] or '') is not None,
      str(s1[1]))
check('识图：notes 提取', note1 == '四选二', note1)

js2 = '好的，识别结果如下：\n[{"date":"2026-10-09","start":"07:40","end":"09:40",' \
      '"subject":"数学"}]\n以上。'
s2, _n2 = aisvc.parse_schedule_json(js2)
check('识图 JSON（前后有解释文字 + 数组）解析 1 场', len(s2) == 1 and s2[0]['subject'] == '数学',
      str(s2))
check('识图 JSON 空输入 → 空', aisvc.parse_schedule_json('') == ([], ''))

# ==================== 4. HTTP：analyze → apply 落库 ====================
with app.app_context():
    adm = User(username='schedadm', real_name='考务管理员', role='admin', must_change_pwd=False)
    adm.set_password('SchedAdm#2026')
    db.session.add(adm)
    from datetime import date as _d
    ex = Exam(name='冒烟AI场次考试', exam_kind='本校考', exam_type='月考', grade='2025级',
              exam_date=_d(2026, 10, 8), term='2026-2027', status='draft')
    db.session.add(ex)
    db.session.commit()
    exam_id = ex.id

with app.test_client() as c:
    html = c.get('/login').get_data(as_text=True)
    m = re.search(r'name="csrf_token"[^>]*value="([^"]*)"', html)
    tok = m.group(1) if m else ''
    r = c.post('/login', data={'csrf_token': tok, 'username': 'schedadm',
                               'password': 'SchedAdm#2026'}, follow_redirects=False)
    check('管理员登录成功', r.status_code in (301, 302), str(r.status_code))
    m2 = re.search(r'var\s+csrfToken\s*=\s*"([^"]+)"', c.get('/').get_data(as_text=True))
    csrf = m2.group(1) if m2 else tok
    H = {'X-CSRFToken': csrf}

    # 页面可达
    rp = c.get(f'/grades/exams/{exam_id}/ai-import')
    ph = rp.get_data(as_text=True)
    check('AI 导入页 200', rp.status_code == 200, str(rp.status_code))
    check('页面含场次预览容器 #schedBox', 'id="schedBox"' in ph)
    check('页面 accept 含图片后缀', '.jpg' in ph and '.png' in ph)
    check('页面含「覆盖本场已有场次」开关', 'id="schedReplace"' in ph)

    # analyze：上传监考表
    data = {'files': (io.BytesIO(build_invig_xlsx()), '监考表.xlsx')}
    r = c.post(f'/grades/exams/{exam_id}/ai-import/analyze', data=data,
               content_type='multipart/form-data', headers=H)
    d = r.get_json() or {}
    check('analyze 200 且 success', r.status_code == 200 and d.get('success'), str(d.get('message')))
    token = d.get('token')
    sh0 = (d.get('sheets') or [{}])[0]
    check('analyze 识别 kind=exam_schedule', sh0.get('kind') == 'exam_schedule',
          str(sh0.get('kind')))
    check('analyze 返回场次预览 5 场',
          len(((sh0.get('schedule') or {}).get('sessions')) or []) == 5,
          str(sh0.get('schedule')))

    # apply：按 exam_schedule 写入
    body = {'token': token, 'schedule_mode': 'replace',
            'sheets': [{'name': sh0.get('name'), 'score_kind': 'exam_schedule'}]}
    r = c.post(f'/grades/exams/{exam_id}/ai-import/apply', json=body, headers=H)
    d = r.get_json() or {}
    check('apply 200 且 success', r.status_code == 200 and d.get('success'), str(d))
    si = d.get('schedule') or {}
    check('apply 场次写入 5 个', si.get('sessions') == 5, str(si))
    check('apply 考场 2 个', si.get('rooms') == 2, str(si))
    check('apply 监考 10 格', si.get('invigs') == 10, str(si))

    with app.app_context():
        af = ExamAffair.query.filter_by(exam_id=exam_id).first()
        check('考务批次已自动创建', af is not None)
        n_s = AffairSchedule.query.filter_by(affair_id=af.id).count()
        n_r = AffairRoom.query.filter_by(affair_id=af.id).count()
        n_i = AffairInvigilation.query.filter_by(affair_id=af.id).count()
        check('DB：场次 5 条', n_s == 5, str(n_s))
        check('DB：考场 2 条', n_r == 2, str(n_r))
        check('DB：监考 10 条', n_i == 10, str(n_i))
        row0 = (AffairSchedule.query.filter_by(affair_id=af.id)
                .order_by(AffairSchedule.sort_order).first())
        check('DB：首场 08:05-10:35 语文',
              row0.start_time == '08:05' and row0.end_time == '10:35' and row0.subject == '语文',
              f'{row0.start_time}-{row0.end_time} {row0.subject}')

    # 覆盖式：再导一次不翻倍（每次重建文件流，因为流用后即关）
    r2 = c.post(f'/grades/exams/{exam_id}/ai-import/analyze',
                data={'files': (io.BytesIO(build_invig_xlsx()), '监考表.xlsx')},
                content_type='multipart/form-data', headers=H)
    t2 = (r2.get_json() or {}).get('token')
    c.post(f'/grades/exams/{exam_id}/ai-import/apply',
           json={'token': t2, 'schedule_mode': 'replace',
                 'sheets': [{'name': sh0.get('name'), 'score_kind': 'exam_schedule'}]},
           headers=H)
    with app.app_context():
        af = ExamAffair.query.filter_by(exam_id=exam_id).first()
        n_s2 = AffairSchedule.query.filter_by(affair_id=af.id).count()
        check('覆盖模式重导不翻倍（仍 5 条）', n_s2 == 5, str(n_s2))
        n_i2 = AffairInvigilation.query.filter_by(affair_id=af.id).count()
        check('覆盖模式重导监考仍 10 条', n_i2 == 10, str(n_i2))

    # 前端回传场次为权威来源：用户给「四选二」场次补科目 → 落库应为补全后的科目
    rb = c.post(f'/grades/exams/{exam_id}/ai-import/analyze',
                data={'files': (io.BytesIO(build_invig_xlsx()), '监考表.xlsx')},
                content_type='multipart/form-data', headers=H)
    tb = (rb.get_json() or {}).get('token')
    edited = [
        {'date': '2026-10-08', 'start': '08:05', 'end': '10:35', 'subject': '语文'},
        {'date': '2026-10-08', 'start': '10:55', 'end': '12:10', 'subject': '物理/历史'},
        {'date': '2026-10-08', 'start': '14:30', 'end': '15:45', 'subject': '化学'},
        {'date': '2026-10-09', 'start': '07:50', 'end': '09:50', 'subject': '数学'},
        {'date': '2026-10-09', 'start': '10:10', 'end': '12:10', 'subject': '外语'},
    ]
    re_ = c.post(f'/grades/exams/{exam_id}/ai-import/apply',
                 json={'token': tb, 'schedule_mode': 'replace', 'schedules': edited,
                       'sheets': [{'name': sh0.get('name'), 'score_kind': 'exam_schedule'}]},
                 headers=H)
    check('回传场次（含补全科目）apply 成功', (re_.get_json() or {}).get('success'), str(re_.get_json()))
    with app.app_context():
        af = ExamAffair.query.filter_by(exam_id=exam_id).first()
        row1430 = (AffairSchedule.query.filter_by(affair_id=af.id, start_time='14:30')
                   .first())
        check('用户补全的科目落库（14:30 → 化学）',
              row1430 is not None and row1430.subject == '化学',
              str(row1430.subject if row1430 else None))
        n_i3 = AffairInvigilation.query.filter_by(affair_id=af.id).count()
        check('改科目后监考仍绑对上（10 格）', n_i3 == 10, str(n_i3))

    # 图片路径：无 AI Key 时应给出明确提示（不 500）
    d3 = c.post(f'/grades/exams/{exam_id}/ai-import/analyze',
                data={'files': (io.BytesIO(b'\x89PNG\r\n\x1a\n' + b'0' * 64), '考试时间.png')},
                content_type='multipart/form-data', headers=H)
    j3 = d3.get_json() or {}
    check('图片无 AI Key → 明确提示而非 500',
          d3.status_code == 200 and not j3.get('success')
          and 'AI Key' in (j3.get('message') or ''), str(j3.get('message')))

    # ============ 审查修复项回归（C / A / B / E / J③ + append / dry_run / 去重） ============

    # C：非法日期/时间必须被归一为 ''，不能抛异常（原会 500）
    check('C：非法日期 2026-13-45 → 空', aisvc.normalize_date('2026-13-45') == '')
    check('C：非法日期 2026-02-30 → 空', aisvc.normalize_date('2026-02-30') == '')
    check('C：合法日期正常', aisvc.normalize_date('2026-10-08') == '2026-10-08')
    check('C：非法时间 25:99 → 空', aisvc.normalize_time('25:99') == '')
    check('C：合法时间正常', aisvc.normalize_time('8:05') == '08:05')

    # J③：合并科目用「、」连接，环节⑥ 可拆开
    check('J③：物理/历史 → 物理、历史', aisvc.normalize_schedule_subject('物理/历史') == '物理、历史')
    check('J③：英语归一为外语', aisvc.normalize_schedule_subject('英语') == '外语')
    from app.modules.grades.routes.affair_invig import _split_subjects
    check('J③：物理、历史 可被 _split_subjects 拆开',
          _split_subjects('物理、历史') == {'物理', '历史'}, str(_split_subjects('物理、历史')))

    # A：混合上传（Excel + 图片）→ 图片不被静默忽略；无 Key 时应给出可见提示
    mix = c.post(f'/grades/exams/{exam_id}/ai-import/analyze',
                 data={'files': [(io.BytesIO(build_invig_xlsx()), '监考表.xlsx'),
                                 (io.BytesIO(b'\x89PNG\r\n\x1a\n' + b'0' * 64), '时间表.png')]},
                 content_type='multipart/form-data', headers=H)
    mj = mix.get_json() or {}
    check('A：混合上传返回 200 且 Excel 部分被识别',
          mix.status_code == 200 and mj.get('success') and len(mj.get('sheets') or []) >= 1,
          str(mj.get('message')))
    check('A：混合上传时图片无 Key 有可见提示（不再静默）',
          any('图片' in (e or '') for e in (mj.get('errors') or [])),
          str(mj.get('errors')))

    # 去重：同一 日期+开始+科目 两条只留一条
    rdd = c.post(f'/grades/exams/{exam_id}/ai-import/analyze',
                 data={'files': (io.BytesIO(build_invig_xlsx()), '监考表.xlsx')},
                 content_type='multipart/form-data', headers=H)
    tdd = (rdd.get_json() or {}).get('token')
    dup = [{'date': '2026-10-08', 'start': '08:05', 'end': '10:35', 'subject': '语文'},
           {'date': '2026-10-08', 'start': '08:05', 'end': '10:35', 'subject': '语文'}]
    rd2 = c.post(f'/grades/exams/{exam_id}/ai-import/apply',
                 json={'token': tdd, 'schedule_mode': 'replace', 'schedules': dup},
                 headers=H)
    check('去重：重复场次只留 1 条', (rd2.get_json() or {}).get('schedule', {}).get('sessions') == 1,
          str((rd2.get_json() or {}).get('schedule')))

    # dry_run：不落库、返回 preview.schedule
    rdr = c.post(f'/grades/exams/{exam_id}/ai-import/analyze',
                 data={'files': (io.BytesIO(build_invig_xlsx()), '监考表.xlsx')},
                 content_type='multipart/form-data', headers=H)
    tdr = (rdr.get_json() or {}).get('token')
    with app.app_context():
        af = ExamAffair.query.filter_by(exam_id=exam_id).first()
        n_before = AffairSchedule.query.filter_by(affair_id=af.id).count()
    rdy = c.post(f'/grades/exams/{exam_id}/ai-import/apply',
                 json={'token': tdr, 'dry_run': True, 'schedule_mode': 'replace',
                       'sheets': [{'name': sh0.get('name'), 'score_kind': 'exam_schedule'}]},
                 headers=H)
    dyj = rdy.get_json() or {}
    check('dry_run：返回 preview.schedule', bool((dyj.get('preview') or {}).get('schedule')), str(dyj))
    with app.app_context():
        af = ExamAffair.query.filter_by(exam_id=exam_id).first()
        n_after = AffairSchedule.query.filter_by(affair_id=af.id).count()
    check('dry_run：不落库（场次数不变）', n_before == n_after, f'{n_before}->{n_after}')

    # append：追加不覆盖（现有 1 条 → 追加后 >1）
    rap = c.post(f'/grades/exams/{exam_id}/ai-import/analyze',
                 data={'files': (io.BytesIO(build_invig_xlsx()), '监考表.xlsx')},
                 content_type='multipart/form-data', headers=H)
    tap = (rap.get_json() or {}).get('token')
    c.post(f'/grades/exams/{exam_id}/ai-import/apply',
           json={'token': tap, 'schedule_mode': 'replace', 'schedules': dup,
                 'sheets': [{'name': sh0.get('name'), 'score_kind': 'exam_schedule'}]},
           headers=H)                      # 先变成 1 条
    rap2 = c.post(f'/grades/exams/{exam_id}/ai-import/analyze',
                  data={'files': (io.BytesIO(build_invig_xlsx()), '监考表.xlsx')},
                  content_type='multipart/form-data', headers=H)
    tap2 = (rap2.get_json() or {}).get('token')
    c.post(f'/grades/exams/{exam_id}/ai-import/apply',
           json={'token': tap2, 'schedule_mode': 'append', 'schedules': dup,
                 'sheets': [{'name': sh0.get('name'), 'score_kind': 'exam_schedule'}]},
           headers=H)
    with app.app_context():
        af = ExamAffair.query.filter_by(exam_id=exam_id).first()
        n_ap = AffairSchedule.query.filter_by(affair_id=af.id).count()
    check('append：追加后场次增加（1 → 2）', n_ap == 2, str(n_ap))

    # E：同一天同一时段并行多科，监考不得错挂到别的场次
    rp = c.post(f'/grades/exams/{exam_id}/ai-import/analyze',
                data={'files': (io.BytesIO(build_parallel_xlsx()), '并行监考表.xlsx')},
                content_type='multipart/form-data', headers=H)
    tp = (rp.get_json() or {}).get('token')
    rpj = rp.get_json() or {}
    psh = (rpj.get('sheets') or [{}])[0]
    check('E：并行场次解析为 2 条',
          len(((psh.get('schedule') or {}).get('sessions')) or []) == 2,
          str(psh.get('schedule')))
    rpe = c.post(f'/grades/exams/{exam_id}/ai-import/apply',
                 json={'token': tp, 'schedule_mode': 'replace',
                       'sheets': [{'name': psh.get('name'), 'score_kind': 'exam_schedule'}]},
                 headers=H)
    check('E：并行场次 apply 成功（无唯一约束冲突）', (rpe.get_json() or {}).get('success'),
          str(rpe.get_json()))
    with app.app_context():
        af = ExamAffair.query.filter_by(exam_id=exam_id).first()
        rows = AffairSchedule.query.filter_by(affair_id=af.id).all()
        by_subj = {r.subject: r for r in rows}
        check('E：物理/历史 两条场次都在', set(by_subj) == {'物理', '历史'}, str(list(by_subj)))
        phy_names = {i.name for i in AffairInvigilation.query.filter_by(
            affair_id=af.id, schedule_id=by_subj['物理'].id).all()}
        his_names = {i.name for i in AffairInvigilation.query.filter_by(
            affair_id=af.id, schedule_id=by_subj['历史'].id).all()}
        check('E：物理场监考 = 徐卓/祝平平（未被错挂）',
              phy_names == {'徐卓', '祝平平'}, str(phy_names))
        check('E：历史场监考 = 吴若澜/郝文哲（未被错挂）',
              his_names == {'吴若澜', '郝文哲'}, str(his_names))

    # B：覆盖式写入回显「清空量」
    rbc = c.post(f'/grades/exams/{exam_id}/ai-import/analyze',
                 data={'files': (io.BytesIO(build_invig_xlsx()), '监考表.xlsx')},
                 content_type='multipart/form-data', headers=H)
    tbc = (rbc.get_json() or {}).get('token')
    rbc2 = c.post(f'/grades/exams/{exam_id}/ai-import/apply',
                  json={'token': tbc, 'schedule_mode': 'replace',
                        'sheets': [{'name': sh0.get('name'), 'score_kind': 'exam_schedule'}]},
                  headers=H)
    bj = (rbc2.get_json() or {}).get('schedule') or {}
    check('B：覆盖时返回 cleared_sessions/cleared_invigs',
          'cleared_sessions' in bj and bj.get('cleared_sessions') == 2,
          str(bj))

print(f'\n===== AI 场次导入冒烟：{ok} 通过 / {fail} 失败 =====')
sys.exit(1 if fail else 0)
