# -*- coding: utf-8 -*-
"""全局对比回归冒烟（脚本级）：
- 分科考试：物理方向 01(强基)/02(卓越) + 历史方向 03(卓越) → 去差/小计/总计口径
- 全科考试：无方向 3 班 9 科 → 单组 + 全年级行（无方向小计）
- 方向筛选、未划线报错、Excel 导出、页面与导航入口可达
口径断言按手算值逐项核对（去差剔卓越班总分末尾 2 人、上线按全体参考人数）。
用法：python scripts/smoke/smoke_global_compare.py
"""
import os as _os
import sys as _sys
_HERE = _os.path.dirname(_os.path.abspath(__file__))
_SCRIPTS = _os.path.dirname(_HERE)
_ROOT = _os.path.dirname(_SCRIPTS)
for _p in (_ROOT, _SCRIPTS):
    if _p not in _sys.path:
        _sys.path.insert(0, _p)

import io
import os
import re
import sys
import tempfile
from datetime import date

BASE = _ROOT
sys.path.insert(0, BASE)
_TMP = tempfile.mkdtemp(prefix='stulink_ld_check_')

import config as config_mod  # noqa: E402
config_mod.BASE_DIR = _TMP
config_mod.Config.SQLALCHEMY_DATABASE_URI = 'sqlite:///' + os.path.join(_TMP, 'system.db')
config_mod.Config.SQLALCHEMY_BINDS = {
    'dormitory': 'sqlite:///' + os.path.join(_TMP, 'dormitory.db'),
    'history': 'sqlite:///' + os.path.join(_TMP, 'history.db'),
    'grades': 'sqlite:///' + os.path.join(_TMP, 'grades.db'),
    'points': 'sqlite:///' + os.path.join(_TMP, 'points.db'),
    'academic': 'sqlite:///' + os.path.join(_TMP, 'academic.db'),
    # 2026-10-10：教务分库——查课 / 业绩 / 表单 各自独立库
    'inspection': 'sqlite:///' + os.path.join(_TMP, 'inspection.db'),
    'achievement': 'sqlite:///' + os.path.join(_TMP, 'achievement.db'),
    'forms': 'sqlite:///' + os.path.join(_TMP, 'forms.db'),
    # 2026-10-10：补 system 绑定 —— 启动建表会遍历 'system'，缺它会有 WARN 噪音
    'system': 'sqlite:///' + os.path.join(_TMP, 'system.db'),
    'portrait': 'sqlite:///' + os.path.join(_TMP, 'portrait.db'),
    'morality': 'sqlite:///' + os.path.join(_TMP, 'morality.db'),
    'timed': 'sqlite:///' + os.path.join(_TMP, 'timed.db'),
    'system': 'sqlite:///' + os.path.join(_TMP, 'system.db'),
    'timetable': 'sqlite:///' + os.path.join(_TMP, 'timetable.db'),
}
config_mod.Config.SECRET_KEY = 'ld-check-secret'

from app import create_app  # noqa: E402
from app.extensions import db  # noqa: E402
from app.models import Student, User  # noqa: E402
from app.models.grades import (Exam, ExamScore, ExamBand, ExamTeacherLink,  # noqa: E402
                               TOTAL_SUBJECT)

app = create_app()
ok = fail = 0


def check(name, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
        print(f'[PASS] {name}')
    else:
        fail += 1
        print(f'[FAIL] {name}  {extra}')


def mk_scores(exam_id, rows, grade):
    """rows: [(no, name, cls, ctype, dir_, sel, {sub: sc} 或 总分)}"""
    for no, name, cls, ctype, dir_, sel, subs in rows:
        for sub, sc in subs.items():
            db.session.add(ExamScore(
                exam_id=exam_id, student_no=no, student_name=name, grade=grade,
                class_name=cls, class_type=ctype, direction=dir_,
                subject_selection=sel, subject=sub, score=sc))


def mk_bands(exam_id, rows):
    """rows: [(direction, subject, [(seq, name, lower)])]"""
    for direction, subject, bands in rows:
        for seq, name, lower in bands:
            db.session.add(ExamBand(
                exam_id=exam_id, direction=direction, subject=subject, seq=seq,
                name=name, lower_mode='score', lower_value=lower))


with app.app_context():
    # 建表由 create_app() 内部初始化完成（与 smoke_grades 同一模式）
    adm = User(username='ldadm', real_name='全局对比管理员', role='admin',
               must_change_pwd=False)
    adm.set_password('LdAdm#2026')
    db.session.add(adm)

    # ================= 分科考试（2024 级） =================
    ex1 = Exam(name='2026学年第一学期期中考', exam_type='期中考', grade='2024级',
               exam_date=date(2026, 10, 1), term='2026-2027', status='imported')
    db.session.add(ex1)
    db.session.flush()

    # 学生（主库）
    for i, (no, nm, cls, sel) in enumerate([
            ('2401', '物一', '01班', '物化生'), ('2402', '物二', '01班', '物化生'),
            ('2403', '物三', '01班', '物化生'), ('2404', '物四', '01班', '物化生'),
            ('2405', '物五', '02班', '物化生'), ('2406', '物六', '02班', '物化生'),
            ('2407', '物七', '02班', '物化生'), ('2408', '物八', '02班', '物化生'),
            ('2409', '历一', '03班', '史政地'), ('2410', '历二', '03班', '史政地'),
            ('2411', '历三', '03班', '史政地')]):
        db.session.add(Student(student_number=no, name=nm, gender='男' if i % 2 else '女',
                               grade='2024级', class_name=cls, subject_selection=sel))

    # 成绩：01班强基（不去差）、02班卓越（去差剔总分末尾 2 人）、03班卓越（同样剔 2 人）
    mk_scores(ex1.id, [
        ('2401', '物一', '01班', '强基班', '物理', '物化生',
         {TOTAL_SUBJECT: 600, '语文': 110, '数学': 130, '物理': 90}),
        ('2402', '物二', '01班', '强基班', '物理', '物化生',
         {TOTAL_SUBJECT: 550, '语文': 100, '数学': 120, '物理': 85}),
        ('2403', '物三', '01班', '强基班', '物理', '物化生',
         {TOTAL_SUBJECT: 500, '语文': 90, '数学': 110, '物理': 80}),
        ('2404', '物四', '01班', '强基班', '物理', '物化生',
         {TOTAL_SUBJECT: 450, '语文': 80, '数学': 100, '物理': 75}),
        ('2405', '物五', '02班', '卓越班', '物理', '物化生',
         {TOTAL_SUBJECT: 620, '语文': 130, '数学': 140, '物理': 95}),
        ('2406', '物六', '02班', '卓越班', '物理', '物化生',
         {TOTAL_SUBJECT: 580, '语文': 120, '数学': 130, '物理': 90}),
        ('2407', '物七', '02班', '卓越班', '物理', '物化生',
         {TOTAL_SUBJECT: 520, '语文': 110, '数学': 120, '物理': 85}),
        ('2408', '物八', '02班', '卓越班', '物理', '物化生',
         {TOTAL_SUBJECT: 460, '语文': 100, '数学': 110, '物理': 80}),
        ('2409', '历一', '03班', '卓越班', '历史', '史政地',
         {TOTAL_SUBJECT: 560, '语文': 105, '数学': 115, '历史': 88}),
        ('2410', '历二', '03班', '卓越班', '历史', '史政地',
         {TOTAL_SUBJECT: 520, '语文': 95, '数学': 105, '历史': 82}),
        ('2411', '历三', '03班', '卓越班', '历史', '史政地',
         {TOTAL_SUBJECT: 480, '语文': 85, '数学': 95, '历史': 76}),
    ], '2024级')

    mk_bands(ex1.id, [
        ('物理', TOTAL_SUBJECT, [(1, '特控', 580), (2, '本科', 500)]),
        ('物理', '语文', [(1, '特控', 100), (2, '本科', 90)]),
        ('物理', '数学', [(1, '特控', 120), (2, '本科', 105)]),
        ('物理', '物理', [(1, '特控', 88), (2, '本科', 82)]),
        ('历史', TOTAL_SUBJECT, [(1, '特控', 540), (2, '本科', 470)]),
        ('历史', '语文', [(1, '特控', 100), (2, '本科', 90)]),
        ('历史', '数学', [(1, '特控', 110), (2, '本科', 100)]),
        ('历史', '历史', [(1, '特控', 84), (2, '本科', 78)]),
    ])

    # ================= 全科考试（2025 级，无方向，9 科） =================
    ex2 = Exam(name='2026学年第一学期月考', exam_type='月考', grade='2025级',
               exam_date=date(2026, 9, 20), term='2026-2027', status='imported')
    db.session.add(ex2)
    db.session.flush()

    base9 = {'语文': 100, '数学': 110, '外语': 120, '物理': 70, '化学': 80,
             '生物': 75, '政治': 65, '历史': 72, '地理': 78}
    for i, (no, nm, cls) in enumerate([('2501', '全一', '01班'), ('2502', '全二', '01班'),
                                       ('2503', '全三', '02班'), ('2504', '全四', '02班'),
                                       ('2505', '全五', '03班')]):
        db.session.add(Student(student_number=no, name=nm, gender='男', grade='2025级',
                               class_name=cls, subject_selection='全科'))
        add = 10 if i % 2 else 0
        row = {k: v + add for k, v in base9.items()}
        row[TOTAL_SUBJECT] = sum(row.values())
        mk_scores(ex2.id, [(no, nm, cls, '', '', '全科', row)], '2025级')

    mk_bands(ex2.id, [
        ('', TOTAL_SUBJECT, [(1, '特控', 850), (2, '本科', 700)]),
    ] + [( '', sub, [(1, '特控', v + 5), (2, '本科', v - 10)])
         for sub, v in base9.items()])

    # ================= 未划线考试 =================
    ex3 = Exam(name='未划线考试', exam_type='月考', grade='2025级',
               exam_date=date(2026, 9, 1), term='2026-2027', status='imported')
    db.session.add(ex3)
    db.session.flush()
    mk_scores(ex3.id, [('2501', '全一', '01班', '', '', '全科',
                        {TOTAL_SUBJECT: 700, '语文': 95})], '2025级')

    # ================= 小计/总计「按组合拆」验证（2026-10-16） =================
    # 用户原话「总分人数有分组合的显示，去差均分怎么没有呢」——多组合的方向小计 /
    # 全年级总计要逐个组合出行（人数、去差均分、上线都只算本组合）。
    ex4 = Exam(name='组合拆分验证考', exam_type='月考', grade='2024级',
               exam_date=date(2026, 9, 10), term='2026-2027', status='imported')
    db.session.add(ex4)
    db.session.flush()
    mk_scores(ex4.id, [
        ('2401', '物一', '01班', '卓越班', '物理', '物化生', {TOTAL_SUBJECT: 600, '语文': 110}),
        ('2402', '物二', '01班', '卓越班', '物理', '物化生', {TOTAL_SUBJECT: 560, '语文': 100}),
        ('2403', '物三', '01班', '卓越班', '物理', '物化生', {TOTAL_SUBJECT: 520, '语文': 90}),
        ('2404', '物四', '02班', '卓越班', '物理', '物化地', {TOTAL_SUBJECT: 500, '语文': 95}),
        ('2405', '物五', '02班', '卓越班', '物理', '物化地', {TOTAL_SUBJECT: 480, '语文': 85}),
        ('2406', '物六', '02班', '卓越班', '物理', '物化政', {TOTAL_SUBJECT: 400, '语文': 60}),
    ], '2024级')
    # ex4 故意配 6 层（2026-10-16 第九轮审计 UI-U04）：打印版层数封顶 5 层要能被断言
    mk_bands(ex4.id, [
        ('物理', TOTAL_SUBJECT, [(1, '清北', 620), (2, '985', 600), (3, '211', 590),
                                 (4, '双一流', 585), (5, '特控', 580), (6, '本科', 500)]),
        ('物理', '语文', [(1, '特控', 100), (2, '本科', 90)]),
    ])

    # ================= 任课教师快照（考试当时，source=import） =================
    # v1.18.9.0：全局对比在各科成绩下方显示该班该科教师，取本场快照（不回落当前映射）
    for _ex, _grade, _rows in ((ex1, '2024级', [
            ('01班', '语文', '张语文'), ('01班', '数学', '李数学'), ('01班', '物理', '王物理'),
            ('02班', '语文', '陈语文'), ('02班', '数学', '周数学'), ('02班', '物理', '吴物理'),
            ('03班', '语文', '赵语文'), ('03班', '数学', '钱数学'), ('03班', '历史', '孙历史'),
        ]), (ex2, '2025级', [
            ('01班', '语文', '高一语文'), ('01班', '数学', '高一数学'),
            ('02班', '语文', '高二语文'), ('03班', '语文', '高三语文'),
        # ex4（2026-10-16 第九轮审计 STC-S09）：混组合班也要有快照，否则
        # 「班主任只写在班级合计行、组合明细行不重复」这条渲染口径永远测不出来
        ]), (ex4, '2024级', [
            ('02班', '语文', '钱语文'), ('02班', '数学', '孙数学'),
            ('02班', '班主任', '周班主'),
        ])):
        for _cls, _sub, _nm in _rows:
            db.session.add(ExamTeacherLink(
                exam_id=_ex.id, grade=_grade, class_name=_cls, subject=_sub,
                teacher_name=_nm, source='import'))

    db.session.commit()
    e1, e2, e3, e4 = ex1.id, ex2.id, ex3.id, ex4.id
    print('== 数据就绪 ==', e1, e2, e3, e4)

    # ================= 服务层直算断言 =================
    from app.modules.grades.services import global_compare_service as ld

    d1 = ld.global_compare_report(e1)
    check('分科：无 error', 'error' not in d1, str(d1.get('error')))
    check('分科：mode=split', d1.get('mode') == 'split', str(d1.get('mode')))
    check('分科：列=总分+语文/数学/物理/历史并集',
          d1['subjects'] == ['总分', '语文', '数学', '物理', '历史'],
          str(d1['subjects']))
    check('分科：科目数（不含总分）=4', d1['subject_count'] == 4,
          str(d1['subject_count']))
    check('分科：组数=2（物理/历史）', len(d1['groups']) == 2,
          str([g['key'] for g in d1['groups']]))
    check('分科：可选层名（高→低）=特控/本科', d1['layers'] == ['特控', '本科'],
          str(d1.get('layers')))
    check('分科：默认勾选层=特控/本科（211/特控/本科口径）',
          d1['default_lines'] == ['特控', '本科'], str(d1.get('default_lines')))

    g_phy, g_his = d1['groups'][0], d1['groups'][1]
    check('物理组 2 个班', len(g_phy['rows']) == 2)
    # 2026-10-21 用户口径：同班型班级相邻（卓越班在前、强基班在后），
    # 所以物理组的行序是 02班(卓越) → 01班(强基)。断言一律按班名取行，不写下标。
    check('物理组行序：卓越班 02班 排在强基班 01班 之前（同班型相邻）',
          [x['class_name'] for x in g_phy['rows']] == ['02班', '01班'],
          str([x['class_name'] for x in g_phy['rows']]))
    r01 = next(x for x in g_phy['rows'] if x['class_name'] == '01班')
    r02 = next(x for x in g_phy['rows'] if x['class_name'] == '02班')
    check('01班：班型快照=强基班', r01['class_type'] == '强基班', r01['class_type'])
    check('01班：方向=物理', r01['direction'] == '物理')
    c = r01['cells']
    check('01班 语文 去差均分=95.0（强基不去差）', c['语文']['trim_avg'] == 95.0,
          str(c['语文']['trim_avg']))
    check('01班 语文 特控=2/50%',
          c['语文']['lines']['特控']['n'] == 2 and c['语文']['lines']['特控']['rate'] == 50,
          f"{c['语文']['lines']['特控']}")
    check('01班 语文 本科=3/75%',
          c['语文']['lines']['本科']['n'] == 3 and c['语文']['lines']['本科']['rate'] == 75,
          f"{c['语文']['lines']['本科']}")
    check('01班 总分 去差均分=525.0', c[TOTAL_SUBJECT]['trim_avg'] == 525.0,
          str(c[TOTAL_SUBJECT]['trim_avg']))
    check('01班 总分 特控=1/25%（600≥580）',
          c[TOTAL_SUBJECT]['lines']['特控']['n'] == 1
          and c[TOTAL_SUBJECT]['lines']['特控']['rate'] == 25,
          f"{c[TOTAL_SUBJECT]['lines']['特控']}")
    check('01班 历史科 无成绩=—（None）',
          c['历史']['trim_avg'] is None and c['历史']['lines']['特控'] is None,
          str(c['历史']))

    # ================= 组合列（2026-10-15：方向之后，读成绩行快照） =================
    check('01班 组合列=物化生×4（全班同组合）',
          r01['selections'] == [{'name': '物化生', 'count': 4}],
          str(r01.get('selections')))

    check('02班：班型快照=卓越班', r02['class_type'] == '卓越班')
    c2 = r02['cells']
    check('02班 语文 去差均分=125.0（剔总分末尾2人后）', c2['语文']['trim_avg'] == 125.0,
          str(c2['语文']['trim_avg']))
    check('02班 语文 特控=4/100%（上线按全体算）',
          c2['语文']['lines']['特控']['n'] == 4 and c2['语文']['lines']['特控']['rate'] == 100,
          f"{c2['语文']['lines']['特控']}")
    check('02班 总分 去差均分=600.0', c2[TOTAL_SUBJECT]['trim_avg'] == 600.0,
          str(c2[TOTAL_SUBJECT]['trim_avg']))
    check('02班 总分 本科=3/75%（620/580/520 ≥500）',
          c2[TOTAL_SUBJECT]['lines']['本科']['n'] == 3
          and c2[TOTAL_SUBJECT]['lines']['本科']['rate'] == 75,
          f"{c2[TOTAL_SUBJECT]['lines']['本科']}")

    sub_phy = g_phy['subtotal']
    # 2026-10-16 用户口径「物理全年级改成年级、历史全年级改成年级，全年级也改成年级」
    # —— 三类汇总行的「班级」列统一写「年级」，方向靠「方向」列区分
    check('物理小计：班级列=「年级」（方向列=物理）',
          sub_phy and sub_phy['class_name'] == '年级' and sub_phy['direction'] == '物理',
          str(sub_phy and (sub_phy['class_name'], sub_phy['direction'])))
    check('历史小计：班级列=「年级」（方向列=历史）',
          g_his['subtotal'] and g_his['subtotal']['class_name'] == '年级'
          and g_his['subtotal']['direction'] == '历史',
          str(g_his['subtotal'] and (g_his['subtotal']['class_name'],
                                     g_his['subtotal']['direction'])))
    check('物理小计：人数=8', sub_phy['count'] == 8, str(sub_phy['count']))
    check('物理小计 总分 去差均分=550.0（合并重算）',
          sub_phy['cells'][TOTAL_SUBJECT]['trim_avg'] == 550.0,
          str(sub_phy['cells'][TOTAL_SUBJECT]['trim_avg']))
    check('物理小计 总分 本科=6/75%（含 520）',
          sub_phy['cells'][TOTAL_SUBJECT]['lines']['本科']['n'] == 6
          and sub_phy['cells'][TOTAL_SUBJECT]['lines']['本科']['rate'] == 75,
          f"{sub_phy['cells'][TOTAL_SUBJECT]['lines']['本科']}")
    check('物理小计 组合列=物化生×8（01+02 班合并）',
          sub_phy['selections'] == [{'name': '物化生', 'count': 8}],
          str(sub_phy.get('selections')))

    check('历史组 1 个班', len(g_his['rows']) == 1 and g_his['rows'][0]['class_name'] == '03班')
    r03 = g_his['rows'][0]
    check('03班 历史科 去差均分=88.0（卓越班剔总分末尾2人，仅剩2409）',
          r03['cells']['历史']['trim_avg'] == 88.0,
          str(r03['cells']['历史']['trim_avg']))
    check('03班 历史科 本科=2/67%',
          r03['cells']['历史']['lines']['本科']['n'] == 2
          and r03['cells']['历史']['lines']['本科']['rate'] == 67,
          f"{r03['cells']['历史']['lines']['本科']}")
    check('03班 组合列=史地政×3（旧写法 史政地 自动归一）',
          r03['selections'] == [{'name': '史地政', 'count': 3}],
          str(r03.get('selections')))

    gr = d1['grand']
    check('全年级行存在（班级列=「年级」、方向列为空）',
          gr and gr['class_name'] == '年级' and not (gr['direction'] or ''),
          str(gr and (gr['class_name'], gr['direction'])))
    check('全年级：人数=11', gr['count'] == 11, str(gr['count']))
    check('全年级 总分 去差均分=551.4（跨方向合并）',
          gr['cells'][TOTAL_SUBJECT]['trim_avg'] == 551.4,
          str(gr['cells'][TOTAL_SUBJECT]['trim_avg']))
    check('全年级 总分 特控=4/36%（逐生按各自方向线，含卡线580）',
          gr['cells'][TOTAL_SUBJECT]['lines']['特控']['n'] == 4
          and gr['cells'][TOTAL_SUBJECT]['lines']['特控']['rate'] == 36,
          f"{gr['cells'][TOTAL_SUBJECT]['lines']['特控']}")
    check('全年级 总分 本科=9/82%',
          gr['cells'][TOTAL_SUBJECT]['lines']['本科']['n'] == 9
          and gr['cells'][TOTAL_SUBJECT]['lines']['本科']['rate'] == 82,
          f"{gr['cells'][TOTAL_SUBJECT]['lines']['本科']}")
    check('全年级 组合列=物化生×8 + 史地政×3（跨方向，人数降序）',
          gr['selections'] == [{'name': '物化生', 'count': 8}, {'name': '史地政', 'count': 3}],
          str(gr.get('selections')))

    # ============ 小计/总计「按组合拆」（2026-10-16 用户要求） ============
    # 组合列早就按组合显示人数，去差均分/上线也要按组合 —— 多组合的汇总行逐个组合出行。
    check('单组合小计不拆组合明细（ex1 物理方向全是物化生）',
          sub_phy.get('combo_rows') == [], str(sub_phy.get('combo_rows')))

    d4 = ld.global_compare_report(e4)
    check('ex4：无 error', 'error' not in d4, str(d4.get('error')))
    sub4 = d4['groups'][0]['subtotal']
    check('ex4 物理小计：合计行人数=6', sub4['count'] == 6, str(sub4['count']))
    cr = sub4.get('combo_rows') or []
    check('ex4 物理小计：按组合拆 3 行（物化生/物化地/物化政，人数降序）',
          [x['selections'][0]['name'] for x in cr] == ['物化生', '物化地', '物化政'],
          str([x.get('selections') for x in cr]))
    check('ex4 组合行：is_combo + is_subtotal 标记',
          all(x.get('is_combo') and x.get('is_subtotal') for x in cr))
    # 2026-10-20 用户口径变更：去差名额按**组合**分配（人数多的组合先拿），
    # 且「去差后 = min(参考人数, 班级人数 − 名额)」——缺考先顶替名额。
    # 01班 3 人参考（卓越班名额 2），02班 物化地 2 人 / 物化政 1 人。
    check('ex4 物化生行：人数=3 且给出三列（班级/参考/去差后）',
          cr[0]['count'] == 3
          and cr[0]['enrolled'] >= cr[0]['exam_count'] >= cr[0]['trim_count'],
          f"count={cr[0]['count']} enr={cr[0].get('enrolled')} "
          f"ref={cr[0].get('exam_count')} trim={cr[0].get('trim_count')} "
          f"n={cr[0].get('trim_n')} avg={cr[0]['cells'][TOTAL_SUBJECT]['trim_avg']}")
    check('ex4 物化地行：人数=2，去差后 ≤ 参考（名额先给人数多的组合）',
          cr[1]['count'] == 2
          and cr[1]['trim_count'] <= cr[1]['exam_count'],
          f"enr={cr[1].get('enrolled')} ref={cr[1].get('exam_count')} "
          f"trim={cr[1].get('trim_count')} avg={cr[1]['cells'][TOTAL_SUBJECT]['trim_avg']}")
    check('ex4 物化政行：仅 1 人的组合不再去差（名额已被大组合拿走）',
          cr[2]['trim_count'] == cr[2]['exam_count']
          and cr[2]['cells'][TOTAL_SUBJECT]['trim_avg'] == 400.0,
          f"ref={cr[2].get('exam_count')} trim={cr[2].get('trim_count')} "
          f"avg={cr[2]['cells'][TOTAL_SUBJECT]['trim_avg']}")
    # 2026-10-16 用户口径「数据有些重复了」：总计行的组合明细与方向小计**完全重复**
    # （每个组合只属于一个方向，同一批人同一套数字）→ 总计行不再重复拆组合
    check('ex4 全年级总计：不再重复拆组合（combo_rows 为空）',
          (d4['grand'].get('combo_rows') or []) == [],
          str(d4['grand'].get('combo_rows')))
    check('ex4 全年级总计：跨组合 → hide_avg=True（只给人数 + 上线率）',
          d4['grand'].get('hide_avg') is True, str(d4['grand'].get('hide_avg')))
    check('ex4 组合行不带教师/班主任（与小计行同口径）',
          all(((x.get('cells') or {}).get('语文') or {}).get('teacher') is None
              and not x.get('headteacher') for x in cr))

    # ============ 混组合班：班级合计行（2026-10-16 用户要求） ============
    # 原话「如果班里面有不同的组合的，应该有个上线率的合计的那种」——合计行给人数与
    # 上线率（逐生按各自组合取线后汇总），跨组合的均分无意义 → 展示层显示「—」。
    rows4 = d4['groups'][0]['rows']
    check('ex4 行数：01班 1 行 + 02班（合计+物化地+物化政）3 行 = 4 行',
          len(rows4) == 4, str([(r['class_name'], r['count']) for r in rows4]))
    tot4 = rows4[1]
    check('ex4 02班合计行：人数=3 且 hide_avg=True',
          tot4['class_name'] == '02班' and tot4['count'] == 3
          and tot4.get('hide_avg') is True,
          f"{tot4['class_name']} {tot4['count']} {tot4.get('hide_avg')}")
    check('ex4 02班合计行带班主任（本场快照）；组合明细行不重复',
          tot4['headteacher'] == '周班主'
          and all(not x['headteacher'] for x in rows4[2:]),
          f"{tot4['headteacher']} / {[x['headteacher'] for x in rows4[2:]]}")
    check('ex4 02班合计行：组合列=物化地×2 + 物化政×1（人数降序）',
          tot4['selections'] == [{'name': '物化地', 'count': 2},
                                 {'name': '物化政', 'count': 1}],
          str(tot4.get('selections')))
    check('ex4 02班合计行：本科 1/33%（逐生按各自组合判线后汇总）',
          tot4['cells'][TOTAL_SUBJECT]['lines']['本科'] == {'n': 1, 'rate': 33},
          str(tot4['cells'][TOTAL_SUBJECT]['lines']['本科']))
    check('ex4 02班合计行：数据层均分仍在（500+480+400)/3=460.0，只是展示层隐藏',
          tot4['cells'][TOTAL_SUBJECT]['avg'] == 460.0,
          str(tot4['cells'][TOTAL_SUBJECT]['avg']))
    det4 = rows4[2:]
    check('ex4 02班组合明细行：人数 2/1 + is_combo + 不重复班主任',
          [x['count'] for x in det4] == [2, 1]
          and all(x.get('is_combo') and not x.get('headteacher') for x in det4),
          str([(x['count'], x.get('is_combo')) for x in det4]))
    check('ex4 物化地明细行：本科 1/50%（500 过线）',
          det4[0]['cells'][TOTAL_SUBJECT]['lines']['本科'] == {'n': 1, 'rate': 50},
          str(det4[0]['cells'][TOTAL_SUBJECT]['lines']['本科']))
    check('ex4 纯组合班 01班：仍 1 行、不隐藏均分（无多余合计行）',
          rows4[0]['class_name'] == '01班' and not rows4[0].get('hide_avg')
          and rows4[0]['count'] == 3,
          str((rows4[0]['class_name'], rows4[0].get('hide_avg'), rows4[0]['count'])))

    # ============ 全体均分「不去差」（2026-10-16 用户要求） ============
    # 用户原话「没有一个合计的人数吗？不带去差均分的那种」——人数列是全体人数，
    # 均分补一个全体口径（含被剔学生），页面/导出/打印版与去差均分同格上下显示。
    check('ex1 02班 总分 全体均分=545.0（4 人平均；去差只留 620/580=600.0）',
          c2[TOTAL_SUBJECT].get('avg') == 545.0,
          f"avg={c2[TOTAL_SUBJECT].get('avg')} trim={c2[TOTAL_SUBJECT]['trim_avg']}")
    check('ex1 02班 语文 全体均分=115.0',
          c2['语文'].get('avg') == 115.0, str(c2['语文'].get('avg')))
    check('ex1 01班（强基不去差）：总分 avg 与 trim_avg 相同=525.0（页面不重复显示）',
          c[TOTAL_SUBJECT].get('avg') == c[TOTAL_SUBJECT]['trim_avg'] == 525.0,
          f"{c[TOTAL_SUBJECT].get('avg')} / {c[TOTAL_SUBJECT]['trim_avg']}")
    check('ex4 物化生组合行：全体均分=560.0（去差 600.0）',
          cr[0]['cells'][TOTAL_SUBJECT].get('avg') == 560.0,
          str(cr[0]['cells'][TOTAL_SUBJECT].get('avg')))

    # ============ 跨组合合计行不给均分（2026-10-16 用户口径） ============
    # 原话「总共的组合都不同，我们按组合给的，你搞这个全年级的去差均分还有啥意思，
    # 搞个上线率还差不多」——组合线不同，混算均分没有可比性；
    # 展示层隐藏（hide_avg），数据层保留 trim_avg/avg（供别处复用），上线率列不受影响。
    check('ex1 物理小计（单组合）：hide_avg=False（均分保留）',
          sub_phy.get('hide_avg') is False, str(sub_phy.get('hide_avg')))
    check('ex1 全年级（跨方向物化生+史地政）：hide_avg=True（不给混算均分）',
          gr.get('hide_avg') is True, str(gr.get('hide_avg')))
    check('ex4 物理小计（3 个组合）：hide_avg=True；组合明细行不受影响',
          sub4.get('hide_avg') is True
          and all(not x.get('hide_avg') for x in cr), str(sub4.get('hide_avg')))
    check('hide_avg 只影响展示：数据层均分仍保留（全年级总分 551.4）',
          gr['cells'][TOTAL_SUBJECT]['trim_avg'] == 551.4,
          str(gr['cells'][TOTAL_SUBJECT]['trim_avg']))
    check('ex1 全年级 上线率不受影响（本科 9/82%）',
          gr['cells'][TOTAL_SUBJECT]['lines']['本科']['n'] == 9
          and gr['cells'][TOTAL_SUBJECT]['lines']['本科']['rate'] == 82)

    # ================= 任课教师（v1.18.9.0：各科成绩下方，本场快照） =================
    check('01班 语文 教师=张语文（取本场快照）',
          c['语文'].get('teacher') == '张语文', str(c['语文'].get('teacher')))
    check('02班 物理 教师=吴物理', c2['物理'].get('teacher') == '吴物理',
          str(c2['物理'].get('teacher')))
    check('03班 历史 教师=孙历史', r03['cells']['历史'].get('teacher') == '孙历史',
          str(r03['cells']['历史'].get('teacher')))
    check('03班 物理 无教师（该班不开物理课）',
          r03['cells']['物理'].get('teacher') is None, str(r03['cells']['物理']))
    check('总分列不带教师（班主任改在总分格下方由前端展示）',
          c[TOTAL_SUBJECT].get('teacher') is None, str(list(c[TOTAL_SUBJECT].keys())))
    check('物理小计行不带教师', sub_phy['cells']['语文'].get('teacher') is None)
    check('全年级行不带教师', gr['cells']['语文'].get('teacher') is None)

    # 方向筛选
    dp = ld.global_compare_report(e1, '物理')
    check('筛选物理：仅 1 组 2 班', len(dp['groups']) == 1 and len(dp['groups'][0]['rows']) == 2)
    check('筛选物理：不出全年级行', dp['grand'] is None)
    # v1.19.62.64 回归：筛选后 directions 仍返回**完整**方向集（否则前端下拉会丢另一个方向）
    check('筛选物理：方向下拉仍含 物理+历史（不丢另一方向）',
          dp['directions'] == ['物理', '历史'], str(dp.get('directions')))
    check('未筛选：方向下拉含 物理+历史',
          d1['directions'] == ['物理', '历史'], str(d1.get('directions')))

    # 全科考试
    d2 = ld.global_compare_report(e2)
    check('全科：mode=all', d2.get('mode') == 'all', str(d2.get('mode')))
    check('全科：1 组 3 班', len(d2['groups']) == 1 and len(d2['groups'][0]['rows']) == 3)
    check('全科：不出方向小计', d2['groups'][0]['subtotal'] is None)
    check('全科：9 科列 + 总分列', len(d2['subjects']) == 10 and d2['subject_count'] == 9,
          str(d2['subjects']))
    check('全科：有全年级行', d2['grand'] is not None)
    rr01 = d2['groups'][0]['rows'][0]
    check('全科 01班 语文 去差均分=105.0（无卓越去差）',
          rr01['cells']['语文']['trim_avg'] == 105.0,
          str(rr01['cells']['语文']['trim_avg']))
    check('全科 01班 语文 特控=1/50%（105 线，110 过）',
          rr01['cells']['语文']['lines']['特控']['n'] == 1
          and rr01['cells']['语文']['lines']['特控']['rate'] == 50,
          f"{rr01['cells']['语文']['lines']['特控']}")
    check('全科 全年级 人数=5', d2['grand']['count'] == 5, str(d2['grand']['count']))
    check('全科 01班 语文 教师=高一语文',
          rr01['cells']['语文'].get('teacher') == '高一语文',
          str(rr01['cells']['语文'].get('teacher')))
    check('全科 03班（无快照的科目）教师为空',
          d2['groups'][0]['rows'][2]['cells']['数学'].get('teacher') is None,
          str(d2['groups'][0]['rows'][2]['cells']['数学']))
    check('全科考试：组合列全空（未分科，页面显示「—」）',
          all(r['selections'] == [] for r in d2['groups'][0]['rows'])
          and d2['grand']['selections'] == [],
          str([r.get('selections') for r in d2['groups'][0]['rows']]))

    # 未划线
    d3 = ld.global_compare_report(e3)
    check('未划线：返回 error 提示', 'error' in d3 and '划线' in d3['error'],
          str(d3.get('error')))

    # ============ 学科分析各班对比表：班型分组 + 组内均分降序（2026-10-21）============
    # 用户口径「学科分析也改成班型分组 + 组内按均分降序」——卓越班在前、
    # 强基班在后，组内仍是排名视图（均分降序）。
    from app.modules.grades.services.tab_service import subject_tab
    _t1 = subject_tab(e1, '语文')['tables']['class_compare']['rows']
    _names = [r['class_name'] for r in _t1]
    _avgs = [r['avg'] for r in _t1]
    check('学科分析：卓越班（02/03）排在强基班（01）之前',
          _names[-1] == '01班' and set(_names[:2]) == {'02班', '03班'},
          str(_names))
    check('学科分析：同班型组内仍按均分降序',
          _avgs[0] is not None and _avgs[1] is not None and _avgs[0] >= _avgs[1],
          str(_avgs))
    _hm = subject_tab(e1, '语文')['charts']['heatmap']
    check('学科分析热力图横轴：同班型相邻（02/03 在前、01 在后）',
          _hm['xAxis'][-1] == '01班', str(_hm['xAxis']))

# ================= HTTP 层断言 =================
with app.test_client() as c:
    html = c.get('/login').get_data(as_text=True)
    m = re.search(r'name="csrf_token"[^>]*value="([^"]*)"', html)
    tok = m.group(1) if m else ''
    r = c.post('/login', data={'csrf_token': tok, 'username': 'ldadm',
                               'password': 'LdAdm#2026'}, follow_redirects=False)
    check('管理员登录成功', r.status_code in (301, 302), str(r.status_code))

    rp = c.get('/grades/global-compare')
    check('页面 /grades/global-compare 200', rp.status_code == 200, str(rp.status_code))
    check('页面含标题「全局对比」', '全局对比' in rp.get_data(as_text=True))

    ra = c.get(f'/grades/api/global-compare/report?exam_id={e1}')
    j = ra.get_json() or {}
    check('接口 200 且 success', ra.status_code == 200 and j.get('success') is True,
          f'{ra.status_code} {str(j)[:120]}')
    check('接口返回 groups/grand', bool(j.get('data', {}).get('groups'))
          and bool(j.get('data', {}).get('grand')))

    rn = c.get(f'/grades/api/global-compare/report?exam_id={e3}')
    jn = rn.get_json() or {}
    check('未划线考试接口 success=false 且有提示',
          jn.get('success') is False and '划线' in (jn.get('message') or ''),
          str(jn)[:140])

    # 导航入口
    rh = c.get('/')
    check('侧边栏含「全局对比」入口', '全局对比' in rh.get_data(as_text=True))

    # Excel 导出
    rx = c.get(f'/grades/api/global-compare/export?exam_id={e1}')
    check('导出 xlsx 200', rx.status_code == 200, str(rx.status_code))
    check('导出 Content-Type xlsx',
          'spreadsheetml' in (rx.headers.get('Content-Type') or ''),
          rx.headers.get('Content-Type') or '')
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(rx.data))
        ws = wb.active
        heads = [ws.cell(row=2, column=i).value for i in range(1, 13)]
        # 2026-10-20 用户口径：新增「参考人数 / 去差后」两列（固定列 5 列）
        check('导出表头正确（年级/班级/班级人数/参考人数/去差后/…/组合/总分组）',
              heads[:8] == ['年级', '班级', '班级人数', '参考人数', '去差后',
                            '方向', '组合', '属性']
              and heads[8] == '总分', str(heads))
        # 第二行指标名：每科「去差均分 + 特控线 + 本科线」（默认勾选层）
        heads2 = [ws.cell(row=3, column=i).value for i in range(1, 13)]
        check('导出第二行指标名（去差均分/特控线/本科线）',
              heads2[8:11] == ['去差均分', '特控线', '本科线'], str(heads2))
        # 2026-10-20：三列人数都要有值（班级人数 ≥ 参考人数 ≥ 去差后）
        _r4 = [ws.cell(row=4, column=i).value for i in (3, 4, 5)]
        check('导出人数三列齐全且 班级人数 ≥ 参考人数 ≥ 去差后',
              all(isinstance(v, int) for v in _r4)
              and _r4[0] >= _r4[1] >= _r4[2], str(_r4))
        # 2026-10-16 用户口径「excel 导出时没有注释解释」→ 第 1 行是口径注释（合并 + 换行）
        note = ws.cell(row=1, column=1).value or ''
        check('导出首行口径注释（合并单元格 + 去差/全体均分/上线率口径）',
              '口径' in note and '去差均分' in note and '全体均分' in note
              and '上线率' in note and 'A1' in str(ws.merged_cells.ranges),
              str(note)[:60])
        check('导出数据从第 4 行开始（注释 + 两行表头）',
              ws.cell(row=4, column=2).value not in (None, ''), str(ws.cell(row=4, column=2).value))
        # 三类汇总行的班级列都是「年级」（方向列区分：物理/历史/空）
        check('导出含方向小计行（班级=年级 + 方向=物理）',
              any(ws.cell(row=r, column=2).value == '年级'
                  and ws.cell(row=r, column=6).value == '物理'
                  for r in range(4, ws.max_row + 1)))
        check('导出含全年级总计行（班级=年级 + 方向列空）',
              any(ws.cell(row=r, column=2).value == '年级'
                  and not (ws.cell(row=r, column=6).value or '')
                  for r in range(4, ws.max_row + 1)))
        all_vals = [ws.cell(row=r, column=cl).value
                    for r in range(4, ws.max_row + 1)
                    for cl in range(1, ws.max_column + 1)]
        check('导出成绩格含任课教师（换行）',
              any(isinstance(v, str) and '张语文' in v and '\n' in v
                  for v in all_vals),
              '未找到「均分\\n教师」单元格')
        check('导出组合列含「物化生 4」（组合 + 人数，与页面一致）',
              any(isinstance(v, str) and '物化生 4' in v for v in all_vals),
              '未找到组合列单元格')
        check('导出均分格含全体均分「(545.0)」（去差 600.0 下一行）',
              any(isinstance(v, str) and '(545.0)' in v for v in all_vals),
              '未找到全体均分单元格')
    except Exception as e:  # noqa: BLE001
        check('导出文件可被 openpyxl 打开', False, str(e))

    # 导出按页面勾选层（lines 参数）：只勾特控 → 每科仅 去差均分 + 特控线
    rl = c.get(f'/grades/api/global-compare/export?exam_id={e1}&lines=%E7%89%B9%E6%8E%A7')
    try:
        import openpyxl
        wb2 = openpyxl.load_workbook(io.BytesIO(rl.data))
        ws2 = wb2.active
        h2 = [ws2.cell(row=3, column=i).value for i in range(1, 13)]
        check('导出仅勾特控：指标行=去差均分/特控线（无本科线）',
              h2[8:10] == ['去差均分', '特控线'] and h2[10] != '本科线', str(h2))
    except Exception as e:  # noqa: BLE001
        check('导出 lines 参数生效', False, str(e))

    # 导出：小计/总计也按组合拆（2026-10-16）——ex4 = 合计行 + 3 个组合明细行
    # 2026-10-16 第九轮审计（STC-S03）：定位行一律按（班级列, 方向列）匹配，
    # 不再写死行号（插一行 combo_rows 就会让旧断言变成移动靶）
    r4 = c.get(f'/grades/api/global-compare/export?exam_id={e4}')
    try:
        import openpyxl
        wb4 = openpyxl.load_workbook(io.BytesIO(r4.data))
        ws4 = wb4.active

        def _find(cls, direction=''):
            for r in range(4, ws4.max_row + 1):
                if (ws4.cell(row=r, column=2).value == cls
                        and (ws4.cell(row=r, column=6).value or '') == direction):
                    return r
            return None

        n_sub = sum(1 for r in range(4, ws4.max_row + 1)
                    if ws4.cell(row=r, column=2).value == '年级'
                    and ws4.cell(row=r, column=6).value == '物理')
        n_grand = sum(1 for r in range(4, ws4.max_row + 1)
                      if ws4.cell(row=r, column=2).value == '年级'
                      and not (ws4.cell(row=r, column=6).value or ''))
        check('导出：物理小计 4 行（合计 + 3 组合明细，班级列都是「年级」）',
              n_sub == 4, str(n_sub))
        check('导出：全年级总计 1 行（不再重复列组合明细）', n_grand == 1, str(n_grand))
        # 混组合班：合计行 + 组合明细行，合计行总分格为「—」
        n_02 = sum(1 for r in range(4, ws4.max_row + 1)
                   if ws4.cell(row=r, column=2).value == '02班')
        check('导出：ex4 02班 3 行（合计 + 物化地 + 物化政）', n_02 == 3, str(n_02))
        rr = _find('02班', '物理')
        check('导出：混组合班合计行总分格=「—」（下方带班主任）',
              rr and str(ws4.cell(row=rr, column=9).value).startswith('—'),
              repr(rr and ws4.cell(row=rr, column=9).value))
        check('导出：其下组合明细行有均分（500.0 起，含全体均分换行）',
              rr and str(ws4.cell(row=rr + 1, column=9).value).startswith('500.0'),
              repr(rr and ws4.cell(row=rr + 1, column=9).value))
        check('导出：班主任只写在班级合计行（周班主），组合明细行不重复',
              rr and '周班主' in str(ws4.cell(row=rr, column=9).value)
              and '周班主' not in str(ws4.cell(row=rr + 1, column=9).value),
              repr(rr and ws4.cell(row=rr, column=9).value))
        # 跨组合合计行的总分格 = 「—」（不给混算均分）；组合明细行仍有均分
        rr = _find('年级', '物理')
        check('导出：跨组合合计行总分格=「—」',
              rr and ws4.cell(row=rr, column=9).value == '—',
              str(rr and ws4.cell(row=rr, column=9).value))
        check('导出：组合明细行仍有均分（含全体均分换行）',
              rr and isinstance(ws4.cell(row=rr + 1, column=9).value, str)
              and '(560.0)' in ws4.cell(row=rr + 1, column=9).value,
              repr(rr and ws4.cell(row=rr + 1, column=9).value))
        check('导出：小计行加粗 / 其组合明细行不加粗',
              rr and ws4.cell(row=rr, column=1).font.bold is True
              and ws4.cell(row=rr + 1, column=1).font.bold is False,
              f"{rr and ws4.cell(row=rr, column=1).font.bold} / "
              f"{rr and ws4.cell(row=rr + 1, column=1).font.bold}")
        # 「分组的统计不用加粗，其他的都加粗」（2026-10-16 用户口径）：
        # 班级行 / 班级合计行加粗，班级的组合明细行不加粗 —— 与页面 tr.combo 同款
        r01 = _find('01班', '物理')
        r02 = _find('02班', '物理')
        rg = _find('年级', '')
        check('导出：班级行 / 班级合计行 / 全年级行都加粗',
              r01 and r02 and rg
              and ws4.cell(row=r01, column=1).font.bold is True
              and ws4.cell(row=r02, column=1).font.bold is True
              and ws4.cell(row=rg, column=1).font.bold is True,
              f'{r01}/{r02}/{rg}')
        check('导出：班级的组合明细行不加粗',
              ws4.cell(row=r02 + 1, column=1).font.bold is False
              and ws4.cell(row=r02 + 2, column=1).font.bold is False,
              f'{ws4.cell(row=r02 + 1, column=1).font.bold}')
        check('导出：全年级总计行加粗',
              all(ws4.cell(row=r, column=1).font.bold is True
                  for r in range(4, ws4.max_row + 1)
                  if ws4.cell(row=r, column=2).value == '年级'
                  and not (ws4.cell(row=r, column=6).value or '')),
              '全年级行非加粗')
    except Exception as e:  # noqa: BLE001
        check('导出组合明细可被 openpyxl 打开', False, str(e))

    # ================= 打印版（2026-10-16 第九轮审计 STC-S02：此前零覆盖） =================
    # 打印版是第三个消费方，页面/导出的口径它都要有一份；CSS 规则顺序类缺陷
    # （UI-U01：tr.sub 覆盖 tr.combo → 小计的组合明细行被打成 700）只有这里能抓到。
    rp = c.get(f'/grades/global-compare/print?exam={e4}')
    check('打印版 200', rp.status_code == 200, str(rp.status_code))
    html = rp.get_data(as_text=True)
    # 2026-10-20：班级行额外带斑马纹 class（zb-a / zb-b），故用「combo 开头」匹配
    check('打印版：班级的组合明细行带 class="combo"（不加粗）',
          'class="combo' in html, '')
    check('打印版：方向小计的组合明细行带 class="sub combo"',
          'class="sub combo' in html, '')
    check('打印版：班级行带斑马纹 class（zb-a/zb-b 交替）',
          'zb-a' in html and 'zb-b' in html, '')
    check('打印版：班级列统一为「年级」', '>年级</td>' in html, '')
    check('打印版：hide_avg 行均分位为灰色「—」（avg na）', 'avg na' in html, '')
    check('打印版：层数缺省用 default_lines（211/特控/本科，不是全部 6 层）',
          '清北' not in html.split('gc-logic')[0], '')
    # 层数封顶 5（UI-U04：真实库 8 场限时练是 12 层名次分段 → 134 列不可读）
    rp6 = c.get(f'/grades/global-compare/print?exam={e4}'
                '&lines=清北,985,211,双一流,特控,本科')
    m = re.search(r'科　·　([^<]+)</div>', rp6.get_data(as_text=True))
    check('打印版：请求 6 层只印前 5 层（封顶）',
          bool(m) and '本科' not in (m.group(1) if m else '')
          and '特控' in (m.group(1) if m else ''),
          str(m and m.group(1)))

print(f'\n===== 全局对比验证：{ok} 通过 / {fail} 失败 =====')
sys.exit(1 if fail else 0)
