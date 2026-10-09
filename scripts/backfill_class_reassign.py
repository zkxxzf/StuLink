# -*- coding: utf-8 -*-
# StuLink v1.18.7.1 2026-09-30
# Copyright (c) 2026 zkxxzf. Apache License 2.0
"""历史「高一选科分班」记录回填（幂等，默认 dry-run）

用户口径（2026-09-30）：
- 2024级 与 2025级 各发生过一次大规模调班（高一选科分班）：
    2024级 → 2025-02-13（依据 20250107 分班名单）
    2025级 → 2026-03-05（依据 20251219 核对名单）
- 两份旧名单给出的是**分班前的班级**；当前系统班级即为分班后（准确，不动）
- 分班前全体「全科、不分方向」→ 记录内容：`2024级01班·全科` → `2024级03班·卓越班·物理·物化生`
- **只对旧名单里查得到的学号回填**；查不到的（借读/后来入校）不写记录

安全措施：
- 姓名一致性闸门：旧名单姓名与库中姓名不符即中止（防学号错配串人）
- 幂等：同一学生同日期已有 reassign 记录则跳过，可重复执行

用法：
    python -m scripts.backfill_class_reassign            # dry-run 预览
    python -m scripts.backfill_class_reassign --apply    # 实际写库
"""
from __future__ import annotations

import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl                                                  # noqa: E402

from app import create_app                                       # noqa: E402
from app.models import Student                                   # noqa: E402
from app.utils.helpers import (write_change_log, location_desc,   # noqa: E402
                               student_location_desc)

# 旧名单 → 年级 / 调班日期 / 来源标记
_D24 = os.environ.get('STULINK_EXCEL_DIR') or r'd:\Users\lenovo\Desktop\2024级历次考试成绩'
_ROST = os.environ.get('STULINK_ROSTER_DIR') or r'd:\Users\lenovo\Desktop'

SOURCES = [
    {'path': os.path.join(_ROST, '20250107分班2024级-借读一人.xlsx'),
     'sheet': '111', 'grade': '2024级',
     'date': '2025-02-13 08:00:00', 'src': '20250107分班名单'},
    {'path': os.path.join(_ROST, '2025级20251219核对人员1人未入班.xlsx'),
     'sheet': None, 'grade': '2025级',
     'date': '2026-03-05 08:00:00', 'src': '20251219核对名单'},
]
FULL_SUBJECT = '全科'
OPERATOR = '系统回填'


def norm_no(v):
    s = str(v).strip() if v is not None else ''
    return s[:-2] if s.endswith('.0') else s


def read_roster(path, sheet):
    """读旧名单 → {学号: (姓名, 分班)}"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    hdr = [str(c.value).strip() if c.value else '' for c in ws[1]]
    i_no = hdr.index('学号')
    i_name = hdr.index('姓名') if '姓名' in hdr else None
    i_cls = next((i for i, h in enumerate(hdr) if h.strip() == '分班'), None)
    if i_cls is None:
        wb.close()
        raise RuntimeError(f'{path} 缺少「分班」列')
    data = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        no = norm_no(row[i_no]) if i_no < len(row) else ''
        nm = str(row[i_name]).strip() if i_name is not None and i_name < len(row) and row[i_name] else ''
        cls = str(row[i_cls]).strip() if i_cls < len(row) and row[i_cls] else ''
        if no and no.isdigit():
            data[no] = (nm, cls)
    wb.close()
    return data


def main():
    ap = argparse.ArgumentParser(description='历史高一选科分班记录回填（幂等）')
    ap.add_argument('--apply', action='store_true', help='实际写库；缺省为 dry-run')
    args = ap.parse_args()
    tag = 'APPLIED' if args.apply else 'DRY-RUN（未写库）'

    app = create_app()
    with app.app_context():
        from app.extensions import db
        from sqlalchemy import text

        # history.db 独立引擎（student_change_log 在 history 库，不能用默认 session）
        hist_engine = db.engines.get('history')
        if hist_engine is None:
            print('[中止] 未找到 history 数据库引擎')
            return 1

        print(f'模式: {tag}\n')
        total_plan = total_skip = 0

        for src in SOURCES:
            if not os.path.exists(src['path']):
                print(f'[跳过] 找不到文件: {src["path"]}')
                continue
            roster = read_roster(src['path'], src['sheet'])
            stu = {norm_no(s.student_number): s for s in
                   Student.query.filter_by(grade=src['grade']).all() if s.student_number}

            matched = [(no, roster[no][0], roster[no][1], stu[no])
                       for no in roster if no in stu]
            only_file = [no for no in roster if no not in stu]
            uncovered = [no for no in stu if no not in roster]

            # 姓名一致性闸门
            bad_names = [(no, fn, s.name) for no, fn, _c, s in matched if fn and fn != s.name]
            print(f'=== {src["grade"]}（调班日 {src["date"][:10]}，据 {src["src"]}）===')
            print(f'  名单 {len(roster)} 条；可匹配 {len(matched)} 人；'
                  f'仅名单有 {len(only_file)}；库内无原班级 {len(uncovered)}')
            if bad_names:
                print(f'  [中止] 姓名不一致 {len(bad_names)} 处，疑似学号错配：')
                for no, fn, dn in bad_names[:10]:
                    print(f'    {no} 名单={fn} 库={dn}')
                print('  请先核对名单后再执行。')
                return 1
            print('  姓名校验: 全部一致 OK')

            # 幂等：查已回填（走 history 引擎）
            ids = [s.id for _n, _f, _c, s in matched]
            done = set()
            if ids:
                CH = 800
                with hist_engine.connect() as conn:
                    for i in range(0, len(ids), CH):
                        chunk = ids[i:i + CH]
                        ph = ','.join(f':p{k}' for k in range(len(chunk)))
                        params = {f'p{k}': v for k, v in enumerate(chunk)}
                        params['d'] = src['date'][:10] + '%'
                        rows = conn.execute(text(
                            "SELECT student_id FROM student_change_log "
                            f"WHERE change_type='reassign' AND changed_at LIKE :d "
                            f"AND student_id IN ({ph})"), params).fetchall()
                        done.update(r[0] for r in rows)

            plan = []
            noop = 0
            for no, _fn, fcls, s in matched:
                if s.id in done:
                    total_skip += 1
                    continue
                old = location_desc(src['grade'], fcls, None, None, FULL_SUBJECT)
                new = student_location_desc(s)
                if old == new:
                    # 前后位置完全一致（如均为「不分班」）：无信息量，不写记录
                    noop += 1
                    continue
                plan.append((s, old, new))
            print(f'  待写 {len(plan)} 条；已存在跳过 {len(done)} 条；'
                  f'前后无变化跳过 {noop} 条')

            for s, old, new in plan[:3]:
                print(f'    {s.student_number} {old} → {new}')
            if len(plan) > 3:
                print(f'    ...（共 {len(plan)} 条）')

            if args.apply and plan:
                for s, old, new in plan:
                    write_change_log(
                        'reassign',
                        [{'id': s.id, 'student_number': s.student_number or '',
                          'name': s.name}],
                        old_value=old, new_value=new,
                        detail=f'高一选科分班（历史回填，据 {src["src"]}）',
                        operator_name=OPERATOR, changed_at=src['date'])
                print(f'  已写库 {len(plan)} 条')
            total_plan += len(plan)
            print('')

        print(f'合计：待写/已写 {total_plan} 条；幂等跳过 {total_skip} 条')
        if not args.apply:
            print('\n确认无误后执行：python -m scripts.backfill_class_reassign --apply')
    return 0


if __name__ == '__main__':
    sys.exit(main())
