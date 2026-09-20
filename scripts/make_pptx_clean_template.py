# -*- coding: utf-8 -*-
"""制备 PPT 导出用的干净母版模板（v1.13.2）。

从原始郑开学校母版 report_template.pptx 删除自带的 8 张示例 slide，
完整保留 slideLayout（Logo/背景装饰）、slideMaster、media、theme，
输出 report_template_clean.pptx 供 report_export_service.build_pptx 加载。

用法（在项目根目录）：
    python scripts/make_pptx_clean_template.py

原始母版更新后重跑一次本脚本即可同步干净模板。
"""
import os
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TPL_DIR = os.path.join(BASE, 'app', 'templates', 'pptx')
SRC = os.path.join(TPL_DIR, 'report_template.pptx')
DST = os.path.join(TPL_DIR, 'report_template_clean.pptx')


def main():
    if not os.path.exists(SRC):
        print(f'原始母版不存在: {SRC}')
        sys.exit(1)

    prs = Presentation(SRC)
    print(f'原始母版示例 slide 数: {len(prs.slides)}')

    # python-pptx 标准删除姿势：drop_rel + 移除 sldIdLst 引用
    # （save 时自动清理 Content_Types 与 rels，不产生悬空引用）
    xml_slides = prs.slides._sldIdLst
    for sldId in list(xml_slides):
        prs.part.drop_rel(sldId.get(qn('r:id')))
        xml_slides.remove(sldId)

    prs.save(DST)
    print(f'干净模板已生成: {DST} ({os.path.getsize(DST)} bytes)')

    # 自检：可加载、0 slide、layout 装饰保留
    prs2 = Presentation(DST)
    assert len(prs2.slides) == 0, '干净模板应无示例 slide'
    assert len(prs2.slide_layouts) > 0, 'layout 应保留'
    print(f'自检通过：slides=0, layouts={len(prs2.slide_layouts)}')


if __name__ == '__main__':
    main()
