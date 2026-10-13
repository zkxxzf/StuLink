# -*- coding: utf-8 -*-
"""考务环节④⑤⑥⑦ 冒烟（脚本级）：任课核对 / 考试时间表 / 监考编排 / 监考表导出

覆盖：
- 环节④ 首次自动带入当前任课映射；保存/同步矩阵
- 环节⑤ 时间表新增/保存；导入 Excel 表头容错
- 环节⑥ 教师池（含无账号的校外监考员）；自动编排**同科回避**、每人每场次只监一个考场
- 环节⑦ 导出监考表 xlsx / pdf
- 页面 /grades/affairs/<aid> 200 且含 7 环节

用法：python scripts/smoke_invig.py
"""
import io
import os
import re
import sys
import tempfile
from datetime import date

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)
_TMP = tempfile.mkdtemp(prefix='stulink_invig_check_')

import config as config_mod  # noqa: E402
config_mod.BASE_DIR = _TMP
config_mod.Config.SQLALCHEMY_DATABASE_URI = 'sqlite:///' + os.path.join(_TMP, 'system.db')
config_mod.Config.SQLALCHEMY_BINDS = {
    k: 'sqlite:///' + os.path.join(_TMP, k + '.db')
    for k in ('dormitory', 'history', 'grades', 'points', 'academic', 'inspection',
              'achievement', 'forms', 'system', 'portrait', 'timetable')
}
config_mod.Config.SECRET_KEY = 'invig-check-secret'

from app import create_app  # noqa: E402
from app.extensions import db  # noqa: E402
from app.models import Student, User  # noqa: E402
from app.models.grades import (ExamAffair, AffairRoom, AffairStudent, TeacherSubjectLink,  # noqa: E402
                               AffairTeacher, AffairSchedule, AffairInvigilator,
                               AffairInvigilation, Exam)

app = create_app()
ok = fail = 0


def check(name, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
        print(f'[PASS] {name}')
    else:
        fail += 1
        print(f'[FAIL] {name}  {extra}')


with app.app_context():
    adm = User(username='invadm', real_name='考务管理员', role='admin', must_change_pwd=False)
    adm.set_password('InvAdm#2026')
    db.session.add(adm)

    # 教师（含学科）
    teachers = {}
    for uname, nm in [('t_chinese', '张语文'), ('t_math', '李数学'),
                      ('t_chem', '王化学'), ('t_phys', '赵物理')]:
        u = User(username=uname, real_name=nm, role='teacher', is_active=True,
                 must_change_pwd=False)
        u.set_password('x')
        db.session.add(u)
        teachers[uname] = u
    db.session.flush()

    # 学生（01/02 班）
    for i in range(8):
        db.session.add(Student(student_number=f'250{i}', name=f'学生{i}',
                               gender='男', grade='2025级',
                               class_name='01班' if i < 4 else '02班',
                               subject_selection='物化生'))

    # 任课映射：01班 语文=张、数学=李、化学=王；02班 语文=张、物理=赵
    links = [
        ('01班', '语文', teachers['t_chinese'].id),
        ('01班', '数学', teachers['t_math'].id),
        ('01班', '化学', teachers['t_chem'].id),
        ('02班', '语文', teachers['t_chinese'].id),
        ('02班', '物理', teachers['t_phys'].id),
    ]
    for cls, sub, uid in links:
        db.session.add(TeacherSubjectLink(grade='2025级', class_name=cls,
                                          subject=sub, user_id=uid, active=True))

    # v1.19.47.3 页头「本场考试」信息：绑一场考试（名字 / 性质 / 阶段 / 时间）
    # v1.19.47.3 侧栏 7 环节也依赖 affair.exam_id 反查考试 → 必须先建考试再建批次
    _exam = Exam(name='2025级2026-2027学年第一学期期中考（冒烟）', exam_kind='联考',
                 exam_type='期中考', grade='2025级', exam_date=date(2026, 10, 8),
                 term='2026-2027', status='imported')
    db.session.add(_exam)
    db.session.flush()

    affair = ExamAffair(name='冒烟考务', grade='2025级', exam_date=date(2026, 10, 8),
                        exam_id=_exam.id, status='draft')
    db.session.add(affair)
    db.session.flush()
    aid = affair.id

    # 考场 ×2
    for rn, loc in [('01', '教学楼A-101'), ('02', '教学楼A-102')]:
        db.session.add(AffairRoom(affair_id=aid, room_no=rn, location=loc,
                                  subject='默认', capacity=30))
    # 名单（参加考试）
    for i in range(8):
        db.session.add(AffairStudent(affair_id=aid, student_no=f'250{i}',
                                     name=f'学生{i}', class_name='01班' if i < 4 else '02班',
                                     subject_selection='物化生', is_attend=True))
    db.session.commit()
    print('== 数据就绪 affair_id=%d ==' % aid)

# ================= HTTP 层 =================
with app.test_client() as c:
    html = c.get('/login').get_data(as_text=True)
    m = re.search(r'name="csrf_token"[^>]*value="([^"]*)"', html)
    tok = m.group(1) if m else ''
    r = c.post('/login', data={'csrf_token': tok, 'username': 'invadm',
                               'password': 'InvAdm#2026'}, follow_redirects=False)
    check('管理员登录成功', r.status_code in (301, 302), str(r.status_code))
    # 登录后会话重建，CSRF 需重新取
    m2 = re.search(r'var\s+csrfToken\s*=\s*"([^"]+)"', c.get('/').get_data(as_text=True))
    csrf = m2.group(1) if m2 else tok
    H = {'X-CSRFToken': csrf}

    # ---- 页面：7 环节 ----
    rp = c.get(f'/grades/affairs/{aid}')
    ph = rp.get_data(as_text=True)
    check('考务页面 200', rp.status_code == 200, str(rp.status_code))
    check('页面含 7 个环节标签',
          all(x in ph for x in ['① 学生名单', '② 考场设置', '③ 编排考号',
                                '④ 核对任课教师', '⑤ 考试时间表', '⑥ 编排监考表', '⑦ 导出']))
    check('页面含监考表导出入口', '导出监考表 Excel' in ph and '导出监考表 PDF' in ph)
    # v1.19.47.3 页头与去重
    check('重复步骤条已删除（无 step-pill）', 'step-pill' not in ph)
    check('页头显示本场考试名字', '2025级2026-2027学年第一学期期中考（冒烟）' in ph)
    check('页头显示考试类型（性质/阶段）', '联考' in ph and '期中考' in ph)
    check('页头显示考试时间', '2026-10-08' in ph)
    check('侧栏标题=考试名字（非「年级·性质·阶段·日期」串）',
          '2025级2026-2027学年第一学期期中考（冒烟）' in ph
          and '2025级·联考·期中考' not in ph)
    # v1.19.47.3 点环节标签不滚动：改用 history.replaceState 写 hash（不再直接赋 location.hash）
    check('点环节标签不滚动（用 history.replaceState）', 'history.replaceState' in ph)
    # v1.19.47.3 侧栏 7 环节：断言**渲染产物**（affair_steps 是 Jinja 变量名，
    # 不会出现在 HTML 里，故改数圆点 / 数「环节」文案）。
    # 默认不带 hash 打开 → 当前环节 = 第 1 步：恰好 1 个圆点带 on，其余 6 个既非 on 也非 done
    #（此刻 4/5/6 环节无数据），以此证明「只亮当前、不把前面的全亮」。
    _n_dot = ph.count('class="efn-dot')
    check('侧栏渲染 7 个环节圆点', _n_dot == 7, f'efn-dot 数={_n_dot}')
    check('侧栏圆点只亮当前 1 个（其余未完成不亮）',
          ph.count('efn-dot on"') == 1 and ph.count('efn-dot on done"') == 0,
          f'on={ph.count("efn-dot on"+chr(34))}')
    check('侧栏标注 7 环节', '/7 环节' in ph)
    check('侧栏圆点带环节标题（鼠标悬浮可见第几步叫什么）', '环节1 学生名单' in ph)
    check('选科明细默认展开（collapse show）', 'collapse show' in ph)
    check('改容量实时刷新选科明细（refreshStats）', 'refreshStats' in ph)
    check('侧栏圆点只亮当前环节（syncSidebarSteps）', 'syncSidebarSteps' in ph)

    # v1.19.47.3 实时统计接口：改容量后选科明细即时反映
    r = c.get(f'/grades/affairs/{aid}/stats')
    js = r.get_json() or {}
    check('实时统计接口 200', r.status_code == 200, str(r.status_code))
    check('实时统计：2 个考场 / 容量 60', js.get('rooms') == 2 and js.get('capacity') == 60,
          str({k: js.get(k) for k in ('rooms', 'capacity')}))
    check('实时统计：选科明细含 物化生 8 人', (js.get('per_sel') or {}).get('物化生', {}).get('students') == 8,
          str(js.get('per_sel')))

    # ---- 环节④ ----
    r = c.get(f'/grades/affairs/{aid}/teachers')
    d = r.get_json()
    check('环节④ 自动带入任课映射（首次自动 sync）',
          d.get('success') and len(d['data']['cells']) == 5,
          str(d.get('data', {}).get('cells', {}).keys()))
    check('环节④ 班级=01/02班', d['data']['classes'] == ['01班', '02班'],
          str(d['data']['classes']))

    # 保存：01班 化学 改派为 李数学（模拟临时换老师）
    with app.app_context():
        math_uid = User.query.filter_by(username='t_math').first().id
    r = c.post(f'/grades/affairs/{aid}/teachers/save', json={'cells': {'01班|化学': math_uid}},
               headers=H)
    check('环节④ 保存矩阵 ok', (r.get_json() or {}).get('ok') is True, str(r.get_json()))
    with app.app_context():
        t = AffairTeacher.query.filter_by(affair_id=aid, class_name='01班', subject='化学').first()
        check('环节④ 化学改由李数学（source=manual）',
              t and t.user_id == math_uid and t.source == 'manual')

    # 同步回填（覆盖）
    r = c.post(f'/grades/affairs/{aid}/teachers/sync', headers=H)
    check('环节④ 重新同步覆盖', (r.get_json() or {}).get('ok') is True)
    with app.app_context():
        t = AffairTeacher.query.filter_by(affair_id=aid, class_name='01班', subject='化学').first()
        check('环节④ 同步后化学回到王化学', t and t.teacher_name == '王化学', str(t and t.teacher_name))

    # ---- 环节⑤ ----
    r = c.post(f'/grades/affairs/{aid}/schedules/save',
               data={'subject': '语文', 'exam_date': '2026-10-08',
                     'start_time': '09:00', 'end_time': '10:30'}, headers=H)
    check('环节⑤ 新增场次 语文', (r.get_json() or {}).get('ok') is True)
    r = c.post(f'/grades/affairs/{aid}/schedules/save',
               data={'subject': '化学', 'exam_date': '2026-10-08',
                     'start_time': '14:00', 'end_time': '15:30'}, headers=H)
    check('环节⑤ 新增场次 化学', (r.get_json() or {}).get('ok') is True)
    # 导入 Excel（表头容错：日期/开始/结束/科目）
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['日期', '开始时间', '结束时间', '科目', '备注'])
    ws.append(['2026-10-09', '09:00', '10:30', '数学', ''])
    ws.append([None, None, None, '', '无科目应跳过'])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    r = c.post(f'/grades/affairs/{aid}/schedules/import',
               data={'file': (buf, 'sched.xlsx')},
               content_type='multipart/form-data', headers=H)
    jd = r.get_json() or {}
    check('环节⑤ 导入 Excel（容错表头，跳过空科目）',
          jd.get('ok') and jd.get('added') == 1 and jd.get('skipped') == 1, str(jd))
    with app.app_context():
        check('环节⑤ 场次数=3', AffairSchedule.query.filter_by(affair_id=aid).count() == 3)

    # ---- 环节⑥ 其他监考员（无账号） ----
    r = c.post(f'/grades/affairs/{aid}/invig/others/add',
               data={'name': '校外陈老师', 'subject': '外语'}, headers=H)
    check('环节⑥ 添加无账号校外监考员', (r.get_json() or {}).get('ok') is True, str(r.get_json()))
    with app.app_context():
        o = AffairInvigilator.query.filter_by(affair_id=aid, name='校外陈老师').first()
        # 「其他监考员」表刻意不含 user_id：校外老师无账号，仅本次考试使用
        check('环节⑥ 校外监考员落库（无账号字段）',
              o is not None and not hasattr(o, 'user_id'))

    # 设置每考场 2 人
    r = c.post(f'/grades/affairs/{aid}/invig/settings', data={'slots': 2}, headers=H)
    check('环节⑥ 设置每考场 2 人', (r.get_json() or {}).get('slots') == 2)

    # 自动编排
    r = c.post(f'/grades/affairs/{aid}/invig/auto', headers=H)
    jd = r.get_json() or {}
    check('环节⑥ 自动编排成功', jd.get('ok') is True, str(jd))
    check('环节⑥ 编排规模=3场×2考场×2人=12格', jd.get('rows') == 12, str(jd))

    with app.app_context():
        scheds = {s.id: s for s in AffairSchedule.query.filter_by(affair_id=aid).all()}
        ivs = AffairInvigilation.query.filter_by(affair_id=aid).all()
        # 同科回避：化学场次不得出现 王化学
        chem_ids = [s.id for s in scheds.values() if s.subject == '化学']
        chem_names = {iv.name for iv in ivs if iv.schedule_id in chem_ids}
        check('环节⑥ 化学场次回避化学老师（王化学不排）',
              '王化学' not in chem_names, str(chem_names))
        wang = User.query.filter_by(username='t_chem').first()
        # 语文场次不得出现 张语文（教语文）
        ch_ids = [s.id for s in scheds.values() if s.subject == '语文']
        ch_names = {iv.name for iv in ivs if iv.schedule_id in ch_ids}
        check('环节⑥ 语文场次回避语文老师（张语文不排）',
              '张语文' not in ch_names, str(ch_names))
        # 每人每场次只监一个考场
        dup = False
        for sid in scheds:
            names = [iv.name for iv in ivs if iv.schedule_id == sid]
            if len(names) != len(set(names)):
                dup = True
        check('环节⑥ 同一场次一人只监一个考场', not dup)
        # 每考场恰好 2 位（本场监考位填满）
        check('环节⑥ 每考场填满 2 位',
              all(len([iv for iv in ivs if iv.schedule_id == s.id and iv.room_no == rn]) == 2
                  for s in scheds.values() for rn in ('01', '02')))

    # 手动指定：化学场次给 王化学 → 应被回避拒绝
    with app.app_context():
        chem_sched = AffairSchedule.query.filter_by(affair_id=aid, subject='化学').first()
        csid = chem_sched.id
    r = c.post(f'/grades/affairs/{aid}/invig/set',
               data={'schedule_id': csid, 'room_no': '01', 'slot': 1, 'name': '王化学'},
               headers=H)
    check('环节⑥ 手动指派回避教师被拒', (r.get_json() or {}).get('ok') is False,
          str(r.get_json()))
    # 手动指定：派一个从未被编排的校外监考员 → 允许
    # 2026-10-13 审计修复：v1.19.47.29 起手动指派与自动编排共用
    # 「一人一场次一考场」硬约束，而自动编排已把 12 格填满，
    # 已入池的老师（含校外陈老师）都已被排到某场某考场 —— 再派他们都会被拒
    # （这是产品的正确行为，不是缺陷）。
    # 因此新增一个未入池的校外监考员，先清一格再派，验证手动指派链路可用。
    r = c.post(f'/grades/affairs/{aid}/invig/others/add',
               data={'name': '校外刘老师', 'subject': '地理'}, headers=H)
    check('环节⑥ 追加第二位校外监考员', (r.get_json() or {}).get('ok') is True,
          str(r.get_json()))
    with app.app_context():
        chem_sched2 = AffairSchedule.query.filter_by(affair_id=aid, subject='化学').first()
        chsid = chem_sched2.id
    r = c.post(f'/grades/affairs/{aid}/invig/set',
               data={'schedule_id': chsid, 'room_no': '02', 'slot': 2, 'name': ''},
               headers=H)
    check('环节⑥ 手动清空格成功', (r.get_json() or {}).get('ok') is True,
          str(r.get_json()))
    r = c.post(f'/grades/affairs/{aid}/invig/set',
               data={'schedule_id': chsid, 'room_no': '02', 'slot': 2, 'name': '校外刘老师'},
               headers=H)
    check('环节⑥ 手动指派校外监考员成功', (r.get_json() or {}).get('ok') is True,
          str(r.get_json()))

    # ---- 环节⑦ 导出 ----
    r = c.get(f'/grades/affairs/{aid}/export/invig')
    check('环节⑦ 监考表 xlsx 200', r.status_code == 200, str(r.status_code))
    try:
        wb = openpyxl.load_workbook(io.BytesIO(r.data))
        ws = wb.active
        heads = [ws.cell(row=1, column=i).value for i in range(1, 7)]
        check('环节⑦ 监考表表头（考试时间/科目/考场号/考场位置/监考1/监考2）',
              heads == ['考试时间', '考试科目', '考场号', '考场位置',
                        '监考老师1', '监考老师2'], str(heads))
        check('环节⑦ 监考表行数=3场×2考场=6', ws.max_row - 1 == 6, str(ws.max_row - 1))
    except Exception as e:  # noqa: BLE001
        check('环节⑦ 监考表 xlsx 可打开', False, str(e))
    r = c.get(f'/grades/affairs/{aid}/export/invig/pdf')
    check('环节⑦ 监考表 PDF 200', r.status_code == 200
          and 'pdf' in (r.headers.get('Content-Type') or ''), str(r.status_code))

print(f'\n===== 考务监考冒烟：{ok} 通过 / {fail} 失败 =====')
sys.exit(1 if fail else 0)
