# -*- coding: utf-8 -*-
"""读取「年级成绩分析」页里的特控线 / 本科线。

该页结构（每块 6 行：清北/985/211/特控/本科/专科）：
    r01  分数线
    r02  科别 | 批次 | 总赋 | 语文 | 数学 | 外语 | 物理 | 历史 | 化赋 | 生赋 | 政赋 | 地赋
    r03  全科 | 清北 | …
    r06        | 特控 | …        ← 只要这两行
    r07        | 本科 | …
高一下的场次有两块：物理类、历史类（第二块前面再出现一次「科别」表头）。

返回 {方向: {科目: {'特控线': x, '本科线': y}}}，
方向取值 ''（全科）/ '物理' / '历史'，科目已映射成系统标准名。
"""
import openpyxl

DIR_MAP = {'全科': '', '物理类': '物理', '历史类': '历史'}
SUB_MAP = {'总赋': '总分', '化赋': '化学', '生赋': '生物',
           '政赋': '政治', '地赋': '地理'}
SEQ_NAME = {'特控': '特控线', '本科': '本科线'}


def _num(value):
    """单元格转 float；空、非数值返回 None。"""
    if value is None or value == '':
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _text(value):
    """单元格转去空白字符串。"""
    return '' if value is None else str(value).strip()


def _map_subject(header):
    """表头 → 系统科目名；'科别'/'批次' 等返回空串。"""
    text = _text(header)
    if not text or text in ('科别', '批次'):
        return ''
    return SUB_MAP.get(text, text)


def _take_line(row, header, direction, out):
    """把一行「特控/本科」按表头写进 out。"""
    band = SEQ_NAME.get(_text(row[1]))
    if not band:
        return
    for idx, head in enumerate(header):
        subject = _map_subject(head)
        value = _num(row[idx]) if idx < len(row) else None
        if not subject or value is None:
            continue
        out.setdefault(direction, {}).setdefault(subject, {})[band] = value


def read_lines(path, sheet='年级成绩分析'):
    """读取一场考试的分数线，返回 {方向: {科目: {层名: 分数}}}。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    if sheet in wb.sheetnames:
        header, direction, active = None, '', False
        for row in wb[sheet].iter_rows(max_col=12, values_only=True):
            first = _text(row[0]) if row else ''
            if first == '科别':
                header, active = [_text(c) for c in row], True
                continue
            if not active or header is None:
                continue
            if first:
                if first not in DIR_MAP:
                    break
                direction = DIR_MAP[first]
            _take_line(row, header, direction, out)
    wb.close()
    return out
